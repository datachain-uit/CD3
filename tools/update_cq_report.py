from docx import Document


REPORT = "CQ/docs/REPORT_CQ_V1.docx"


def add_text(document: Document, text: str, *, heading: bool = False) -> None:
    paragraph = document.add_paragraph(text)
    if heading:
        style = next(style for style in document.styles if style.name == "Heading 1")
        paragraph._p.get_or_add_pPr().get_or_add_pStyle().val = style.style_id


document = Document(REPORT)
for paragraph in document.paragraphs:
    if paragraph.text.startswith("Cập nhật:"):
        paragraph.text = "Cập nhật: 26/09/2026"
        break

add_text(document, "14. RNN V0 và sanity components (cập nhật 26/09/2026)", heading=True)
add_text(document, "Phạm vi đã hoàn thành: CQ, pipeline V0 (RAW_CONSTANT_FILL_WITH_MASKS), mô hình RNN, seed model 42, GPU L4 và AMP bfloat16. W1 và W2 đã hoàn thành một shared-checkpoint run; W3 chưa được ghi nhận hoàn thành trong báo cáo này.")
add_text(document, "Giao thức mô hình: mỗi window huấn luyện đúng một checkpoint trên TRAIN thực với sequence P1–P4 đầy đủ. Mỗi epoch chọn checkpoint bằng mean validation Macro-F1 trên P1–P4, tie-break bằng mean cross-entropy. Checkpoint đóng băng sau selection và được suy luận một lần trên TEST/P1, TEST/P2, TEST/P3, TEST/P4; không retrain, fine-tune, refit scaler hoặc refit imputer theo phase test.")
add_text(document, "Input V0 không phải chỉ là null mask. Tensor gồm feature numeric đã transform train-only, missingness masks và availability/modality masks. V0 neutral-fill observed missing values nhưng giữ missing__*, phase_available_*, video_observed_mask_*, problem_observed_mask_* và comment_observed_mask_* để phân biệt zero quan sát, missing đủ điều kiện và structural absence. Raw ID/high-cardinality ID code không được đưa vào tensor.")
add_text(document, "Lưu trữ model theo immutable naming contract: task={task}/feature_regime={feature_regime}/window_id={window_id}/pipeline_id=V0/model_name=RNN/seed=42/run_id={sha256 canonical config}/attempt_id={run_id}-NNNN. L1 lưu checkpoint/model.pt, train_history, prediction tables, probability matrices, confusion matrices, timing và manifest. L2 phân vùng theo phase_id, stage=S3_MODEL, pipeline_id, model_name, run_id và attempt_id; cost_shared_across_phases=true để không nhân bốn chi phí train.")
add_text(document, "Mỗi VALIDATION/TEST phase P1–P4 lưu prediction table, probability matrix c0/c1/c2 và confusion matrix. L2 performance/calibration lưu Macro-F1, balanced accuracy, MCC, Cohen kappa, G-mean, per-class precision/recall/F1/support, NLL, multiclass Brier, ECE-15 và cross-entropy. W1 có 8 prediction và 8 probability artifacts; mỗi file xấp xỉ 24.4–24.7 MiB trên Modal.")
add_text(document, "Sanity hậu kiểm chạy riêng, không train lại RNN, cho VALIDATION và TEST ở từng P1–P4. Công thức hiện hành acctempo_m3_v1_jsd_base2_six_components gồm s_nan, s_maj_jsd, s_ent, s_drift, s_eff và s_leak; S_san_plus là geometric mean cố định của sáu component. s_maj_jsd = 1 − JSD_base2(predicted class distribution, true class distribution); không dùng code legacy tách s_maj và s_jsd. s_drift dùng class distribution TRAIN thực trước balancing làm tham chiếu.")
add_text(document, "Ví dụ CQ/W1 TEST/P1: s_nan=1.000000, s_maj_jsd=0.992046, s_ent=0.076604, s_drift=0.999927, s_eff=0.697147, s_leak=0.452014, S_san_plus=0.536872, S_perf=0.586045 và AccTEMPO_M3=0.565857. Probability-sum max absolute error=1.45e-07; leakage probe AUC=0.773993 trên TRAIN pooled P1–P4 (12,874,208 phase rows). Các score thấp là diagnostic cần phân tích, không phải lỗi artifact.")
add_text(document, "Trạng thái CQ: W1 và W2 có đủ 8 sanity facts mỗi window (VALIDATION/TEST × P1–P4) gắn với RNN-V0 run/attempt tương ứng. Cần hoàn thành RNN-V0 và model_sanity W3 trước khi tổng hợp CQ W1–W3 hoặc so sánh pipeline/model chính thức.")
document.save(REPORT)
