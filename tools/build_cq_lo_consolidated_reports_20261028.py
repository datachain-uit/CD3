#!/usr/bin/env python3
"""Build the current, provenance-aware CQ and LO evidence reports.

The reports deliberately distinguish completed seed-42 evidence from pilots and
from protocol steps that have been registered but not yet executed.  They are
generated only from downloaded local facts and the integrity-verified LO
closure archive; no model artefact is changed by this script.
"""
from __future__ import annotations

import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
import sys as _sys
if str(ROOT / "tools") not in _sys.path:
    _sys.path.insert(0, str(ROOT / "tools"))
CQ_FACTS = ROOT / "CQ/resuilt/cq_recurrent_allfamilies_grid_r4_20261023"
CQ_D23 = ROOT / "CQ/resuilt/cq_patchtst_d23_stability_pilot_20261026"
CQ_REPLAY_RECEIPT = ROOT / "CQ/resuilt/cq_rnn_equivalence_20261028/canonical_facts_comparison.json"
LO_ROOT = ROOT / "LO/resuilt/lo_static_v2_seed42_20261026"
LO_CLOSURE_AUDIT = LO_ROOT / "_closure_audit_temp/audits_seed42/LO/BILSTM/seed=42"
LO_CLOSURE = LO_ROOT / "LO_static_v2_closure_20261028.tar.gz"


PIPE_ORDER = {f"V{i}": i for i in range(17)}
FAMILY_ORDER = {name: index for index, name in enumerate(("RNN", "LSTM", "GRU", "BILSTM"))}
DISPLAY_FAMILY_ORDER = ["RNN", "LSTM", "GRU", "BILSTM"]


def normalise_split(value: object) -> str:
    return str(value).upper()


def setup_document(title: str, subtitle: str) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.62)
    section.right_margin = Inches(0.62)
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)
    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run(title)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(18)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run(subtitle)
    sr.italic = True
    sr.font.name = "Arial"
    sr.font.size = Pt(10.5)
    return doc


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(9 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt({1: 14, 2: 12, 3: 11}.get(level, 10.5))


def add_text(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(10)


def _format(value: object) -> str:
    if value is None or pd.isna(value):
        return ""
    if isinstance(value, float):
        if abs(value) >= 1000:
            return f"{value:,.0f}"
        return f"{value:.4f}"
    return str(value)


def add_table(doc: Document, frame: pd.DataFrame, columns: list[tuple[str, str]], *, font_size: float = 9.5) -> None:
    table = doc.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    table.autofit = True
    for cell, (_, label) in zip(table.rows[0].cells, columns):
        cell.text = label
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
                run.font.name = "Arial"
                run.font.size = Pt(font_size + 0.5)
    for _, row in frame.iterrows():
        cells = table.add_row().cells
        for cell, (name, _) in zip(cells, columns):
            cell.text = _format(row.get(name))
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(font_size)
    doc.add_paragraph()


def sort_matrix(frame: pd.DataFrame) -> pd.DataFrame:
    result = frame.copy()
    result["_family"] = result["model_name"].map(FAMILY_ORDER).fillna(99)
    result["_pipeline"] = result["pipeline_id"].map(PIPE_ORDER).fillna(99)
    return result.sort_values(["_family", "_pipeline"]).drop(columns=["_family", "_pipeline"])


def summary_table(metrics: pd.DataFrame, *, metric_map: dict[str, str]) -> pd.DataFrame:
    group_cols = ["eval_split", "model_name", "pipeline_id"]
    out = metrics.groupby(group_cols, as_index=False)[list(metric_map)].mean()
    return sort_matrix(out)


def coverage_table(manifests: list[dict], *, task: str) -> pd.DataFrame:
    rows = []
    for manifest in manifests:
        rows.append({
            "Family": manifest["model_name"],
            "Runs": manifest["runs"],
            "Sanity": manifest["rows"]["sanity_components"],
            "S_leak": manifest["rows"]["canonical_s_leak"],
            "Slices": manifest["rows"]["evaluation_slices"],
            "Per-class": manifest["rows"]["metrics_per_class"],
            "Bootstrap B=2,000": manifest["rows"]["bootstrap_metrics_per_class"],
            "Checkpoint": manifest["rows"]["checkpoint_selection"],
            "Release": manifest["release_id"],
        })
    return pd.DataFrame(rows).sort_values("Family", key=lambda s: s.map(FAMILY_ORDER))


def load_lo_metrics() -> tuple[pd.DataFrame, list[dict]]:
    roots = {
        "RNN": LO_ROOT / "RNN/audits/LO/RNN/seed=42",
        "LSTM": LO_ROOT / "LSTM/audits/LO/LSTM/seed=42",
        "GRU": LO_ROOT / "GRU/audits/LO/GRU/seed=42",
        "BILSTM": LO_CLOSURE_AUDIT,
    }
    frames, manifests = [], []
    for family, path in roots.items():
        manifest = json.loads((path / "audit_manifest.json").read_text(encoding="utf-8"))
        if manifest["model_name"] != family or manifest["runs"] != 51:
            raise ValueError(f"Unexpected LO audit contract: {path}")
        sanity = pd.read_parquet(path / "sanity_components.parquet")
        if len(sanity) != 408:
            raise ValueError(f"{family}: expected 408 sanity rows, got {len(sanity)}")
        frames.append(sanity)
        manifests.append(manifest)
    metrics = pd.concat(frames, ignore_index=True)
    if len(metrics) != 1632:
        raise ValueError(f"LO expected 1,632 evaluation facts, got {len(metrics)}")
    return metrics, manifests


def static_contract() -> list[str]:
    layout = json.loads((LO_ROOT / "RNN/LO/W1/V0/RNN/seed=42/feature_layout.json").read_text(encoding="utf-8"))
    return layout["static_columns"]


# The shared training loop changed in 044d6c7 (optimizer_step: unscale, gradient division, optional warm-up and
# clipping).  POST_MATERIALIZE_CHECKS[27] closes only when a cell trained BEFORE that commit and its replay trained
# AFTER it agree; two runs from checkouts on the same side of 044d6c7 prove determinism, not loop equivalence.
TRAINING_LOOP_CHANGE_COMMIT = "044d6c7"
COMMITS_BEFORE_LOOP_CHANGE = ("5b14052", "7886a85", "585343e", "eb80516", "83c5056", "75eae0c", "c225961", "8d105db", "0851aed")
COMMITS_AFTER_LOOP_CHANGE = ("044d6c7", "b0e5ae1", "bc20089", "af079d8", "651b458", "29da3b2", "e0ae433")


def _loop_side(commit: str | None) -> str:
    """'before' / 'after' 044d6c7 by git ancestry when the checkout knows both commits, else by the reviewed list."""
    short = (commit or "").strip()[:7]
    if not short:
        return "unknown"
    try:
        import subprocess
        probe = subprocess.run(["git", "-C", str(ROOT), "merge-base", "--is-ancestor", TRAINING_LOOP_CHANGE_COMMIT, short],
                               capture_output=True, text=True, timeout=10)
        if probe.returncode == 0:
            return "after"
        if probe.returncode == 1 and subprocess.run(["git", "-C", str(ROOT), "cat-file", "-e", f"{short}^{{commit}}"],
                                                     capture_output=True, timeout=10).returncode == 0:
            return "before"
    except (OSError, subprocess.SubprocessError):
        pass
    if any(short.startswith(c) or c.startswith(short) for c in COMMITS_BEFORE_LOOP_CHANGE):
        return "before"
    if any(short.startswith(c) or c.startswith(short) for c in COMMITS_AFTER_LOOP_CHANGE):
        return "after"
    return "unknown"


def replay_provenance(receipt: dict, *, reference_key: str = "reference", current_key: str = "current") -> dict:
    """Commits of both sides of a paired replay and what the delta-0 result proves.

    Receipts written before this field existed carry no commits: the report then says so instead of
    claiming that the training-loop check is closed.
    """
    def field_of(side: str, keys: tuple[str, ...], nested_keys: tuple[str, ...]) -> str | None:
        for key in keys:
            if receipt.get(key):
                return str(receipt[key])
        nested = receipt.get(side) or {}
        if isinstance(nested, dict):
            for key in nested_keys:
                if nested.get(key):
                    return str(nested[key])
        return None

    def commit_of(side: str) -> str | None:
        return field_of(side, (f"{side}_runner_commit", f"{side}_commit", f"{side}_runner_git_commit"),
                        ("runner_git_commit", "runner_commit", "commit"))

    def attributed_commit(side: str) -> str | None:
        """A side without a recorded commit is attributed through its runner_sha256 (tools/identify_runner_commit.py)."""
        digest = field_of(side, (f"{side}_runner_sha256",), ("runner_sha256",))
        if not digest:
            return None
        try:
            import identify_runner_commit as attribution
            found = attribution.identify(digest, root=ROOT)
        except Exception:                     # no git history in this checkout: the hash stays unattributed
            return None
        return found["introduced_in"] if found["status"] == "matched" and found["side_of_loop_change"] != "unknown" else None

    reference, current = commit_of(reference_key), commit_of(current_key)
    attributed = []
    if reference is None and attributed_commit(reference_key):
        reference = attributed_commit(reference_key)
        attributed.append(f"reference ({reference}, theo runner_sha256)")
    if current is None and attributed_commit(current_key):
        current = attributed_commit(current_key)
        attributed.append(f"current ({current}, theo runner_sha256)")
    sides = (_loop_side(reference), _loop_side(current))
    if reference is None or current is None:
        verdict = "Delta 0; receipt chưa ghi commit hai phía — kiểm tương đương vòng lặp huấn luyện (044d6c7) chưa đóng"
        sentence = (f"Receipt không ghi runner commit của phía reference và phía current, nên chưa xác định được phía reference "
                    f"có được huấn luyện trước {TRAINING_LOOP_CHANGE_COMMIT} hay không; delta 0 vì thế chưa đóng kiểm tương đương vòng lặp huấn luyện.")
    elif sides == ("before", "after"):
        verdict = f"Tương đương chính xác giữa vòng lặp trước và sau {TRAINING_LOOP_CHANGE_COMMIT}"
        sentence = (f"Phía reference ({reference}) được huấn luyện trước {TRAINING_LOOP_CHANGE_COMMIT} và phía current ({current}) sau đó, "
                    f"nên delta 0 đóng kiểm tương đương vòng lặp huấn luyện cho họ recurrent.")
    elif "unknown" in sides:
        verdict = "Delta 0; chưa xếp được commit hai phía so với 044d6c7"
        sentence = (f"Commit {reference} và {current} chưa nằm trong danh sách commit đã rà soát; "
                    f"cần ghi rõ hai phía so với {TRAINING_LOOP_CHANGE_COMMIT} trước khi đọc delta 0 là tương đương vòng lặp.")
    else:
        verdict = f"Delta 0 giữa hai checkout cùng phía {TRAINING_LOOP_CHANGE_COMMIT}: chứng minh tất định, chưa chứng minh tương đương vòng lặp"
        sentence = (f"Cả hai phía ({reference}, {current}) nằm cùng phía {TRAINING_LOOP_CHANGE_COMMIT}, nên delta 0 chứng minh tính tất định "
                    f"của backend, không phải tương đương giữa vòng lặp cũ và mới.")
    if attributed:
        sentence += " Commit suy ra từ runner_sha256 bằng tools/identify_runner_commit.py: " + "; ".join(attributed) + "."
    return {"reference_commit": reference or "UNRECORDED", "current_commit": current or "UNRECORDED",
            "sides": sides, "verdict": verdict, "sentence": sentence, "attributed_by_runner_hash": attributed}


def cq_report() -> Path:
    target = ROOT / "CQ/docs/REPORT_CQ_CONSOLIDATED_EVIDENCE_2026-10-28_R20.docx"
    summary = json.loads((CQ_FACTS / "summary.json").read_text(encoding="utf-8"))
    metrics = pd.read_parquet(CQ_FACTS / "cq_family_pipeline_window_phase_metrics.parquet")
    resources = pd.read_parquet(CQ_FACTS / "cq_resource_per_run.parquet")
    d23 = json.loads((CQ_D23 / "audits/CQ/PATCHTST/seed=42/audit_manifest.json").read_text(encoding="utf-8"))
    replay_receipt = json.loads(CQ_REPLAY_RECEIPT.read_text(encoding="utf-8"))
    if len(metrics) != 1632 or len(resources) != 204:
        raise ValueError("CQ complete-grid coverage is not 1,632 metrics / 204 resources")
    if not bool(summary["all_energy_nonnull"]):
        raise ValueError("CQ grid has missing energy values")
    if not replay_receipt.get("keys_align") or not replay_receipt.get("acctempo_equivalent_exact"):
        raise ValueError("CQ paired replay receipt does not establish exact canonical equivalence")
    if any(float(value) != 0.0 for value in replay_receipt["max_abs_delta"].values()):
        raise ValueError("CQ paired replay receipt has non-zero canonical metric deltas")

    doc = setup_document(
        "BÁO CÁO CQ — BẰNG CHỨNG THỰC NGHIỆM HIỆN HÀNH",
        "Clean-grain V2.2 · seed 42 · F1 recurrent grid batch 2.048 · cập nhật 28/10/2026",
    )
    add_heading(doc, "1. Kết luận điều hành")
    add_text(doc, "Lưới F1 CQ đã hoàn tất và có facts hậu kiểm đầy đủ: 4 họ mô hình × 17 pipeline × 3 window × 4 phase × 2 tập đánh giá = 1.632 evaluation cells; 204 run có NVML GPU-only energy estimate. AccTEMPO là độ đo hiệu năng chính. Macro-F1, balanced accuracy, recall từng lớp, calibration và các thành phần sanity là bằng chứng bổ trợ.")
    add_text(doc, "PatchTST D23 chỉ là sensitivity pilot: W2/V0 chọn epoch 4 nhưng W3/V0 chọn epoch 1 với lịch sử validation không đơn đỉnh. D23 không được dùng để đổi giao thức. Nhóm đã khóa cấu hình F2 v2; paired replay RNN đạt delta 0 nhưng cần xác minh commit hai phía trước khi kết luận tương đương training loop.")

    add_heading(doc, "2. Release, thiết kế và độ bao phủ")
    scope = pd.DataFrame([{
        "Release": "CQ_V2_2 clean-grain",
        "Split / phase": "v2_2 / wide_prefix_v2_2",
        "Seed": 42,
        "Pipelines": "V0–V16",
        "Families": ", ".join(DISPLAY_FAMILY_ORDER),
        "Windows / phases": "W1–W3 / P1–P4",
        "Evaluation": "VALIDATION + TEST",
        "Evaluation facts": int(summary["evaluation_cells"]),
        "Resource runs": int(summary["resource_runs"]),
    }])
    add_table(doc, scope, [(x, x) for x in scope.columns])
    add_text(doc, "Mọi run F1 dùng NVIDIA L4, bfloat16 logits/loss float32, deterministic algorithms, batch 2.048 và NVML_POWER_SAMPLING_ESTIMATE. CI bootstrap dùng B=2.000. Validation chỉ chọn checkpoint; TEST không được dùng để chọn hyperparameter.")

    add_heading(doc, "3. Hậu kiểm, grain và integrity")
    audit = pd.DataFrame([{
        "Grain / lineage": "Clean-grain; mỗi input audit có rows = distinct enrollment_id",
        "Per-class facts": int(summary["per_class_facts"]),
        "Bootstrap facts": int(summary["bootstrap_per_class_facts"]),
        "Energy": summary["energy_measurement_types"][0],
        "Energy missing": "0/204 run",
        "Sanity definition": "S_leak từ parent imputer thực, probe fit một lần",
        "Slices": "ALL, STRICT, OVERLAP; validation/test frozen",
    }])
    add_table(doc, audit, [(x, x) for x in audit.columns])
    add_text(doc, "Các kết quả cải thiện do augmentation phải được đọc kèm s_ent, S_perf và S_san+; S_san+ v1 là metric chất lượng dữ liệu chính, còn s_cal/S_san+ v2 là sensitivity. Energy chỉ là ước lượng GPU NVML, không tự suy diễn CO₂.")

    metric_map = {"acctempo_m3": "AccTEMPO", "f1_macro": "Macro-F1", "balanced_accuracy": "Balanced acc.", "s_ent": "s_ent", "s_san_plus": "S_san+", "s_leak": "s_leak", "c0_recall": "Recall C0", "c1_recall": "Recall C1", "c2_recall": "Recall C2"}
    aggregate = summary_table(metrics, metric_map=metric_map)
    add_heading(doc, "4. Kết quả đầy đủ theo family × pipeline")
    add_text(doc, "Mỗi hàng là trung bình của 12 cell (W1–W3 × P1–P4) trong cùng split. Facts chi tiết window × phase được giữ trong CQ/resuilt/cq_recurrent_allfamilies_grid_r4_20261023.")
    cols = [("model_name", "Family"), ("pipeline_id", "Pipeline"), *[(k, v) for k, v in metric_map.items()]]
    for split in ("TEST", "VALIDATION"):
        add_heading(doc, f"4.{1 if split == 'TEST' else 2}. {split} — 68 cấu hình", 2)
        add_table(doc, aggregate[aggregate.eval_split.map(normalise_split).eq(split)], cols, font_size=8.2)

    add_heading(doc, "5. Hạt W1/P1 — kiểm tra diễn giải sớm nhất")
    add_text(doc, "Bảng này không gộp phase/window, giúp đối chiếu trực tiếp bối cảnh dự đoán sớm nhất. Không dùng để thay thế bảng primary đầy đủ ở mục 4.")
    w1p1 = sort_matrix(metrics[(metrics.window_id.eq("W1")) & (metrics.phase_id.eq("P1"))].copy())
    for split in ("TEST", "VALIDATION"):
        add_heading(doc, f"5.{1 if split == 'TEST' else 2}. {split} W1/P1", 2)
        add_table(doc, w1p1[w1p1.eval_split.map(normalise_split).eq(split)], cols, font_size=8.2)

    add_heading(doc, "6. Thời gian train và năng lượng GPU")
    resource_summary = (resources.groupby("model_name", as_index=False)
                        .agg(runs=("pipeline_id", "count"), train_seconds_mean=("train_seconds", "mean"),
                             train_seconds_sum=("train_seconds", "sum"), energy_kwh_sum=("energy_gpu_estimated_kwh", "sum"),
                             energy_kwh_mean=("energy_gpu_estimated_kwh", "mean"), missing_rate_max=("energy_missing_sample_rate", "max")))
    resource_summary = resource_summary.sort_values("model_name", key=lambda s: s.map(FAMILY_ORDER))
    add_table(doc, resource_summary, [("model_name", "Family"), ("runs", "Runs"), ("train_seconds_mean", "Train s/run TB"), ("train_seconds_sum", "Train s tổng"), ("energy_kwh_mean", "kWh/run TB"), ("energy_kwh_sum", "kWh tổng"), ("missing_rate_max", "NVML missing max")])
    add_text(doc, "So sánh hiệu năng thời gian dùng train_seconds. Năng lượng chỉ phản ánh công suất GPU lấy mẫu mỗi giây; không có carbon factor/region khóa, nên không báo cáo CO₂.")

    add_heading(doc, "7. PatchTST D23 và giao thức F2")
    d23_rows = []
    for path in sorted((CQ_D23 / "CQ").glob("*/V0/PATCHTST/seed=42/run_manifest.json")):
        item = json.loads(path.read_text(encoding="utf-8"))
        d23_rows.append({
            "Window": item["window_id"], "Best epoch": item["best_epoch"], "LR": item["learning_rate"],
            "Warm-up": item["warmup_epochs"], "Clip": item["gradient_clip_norm"],
            "Batch / micro / accum.": f"{item['batch_size']} / {item['micro_batch_size']} / {item['accumulation_steps']}",
            "Energy kWh": item["energy_gpu_estimated_kwh"], "Purpose": item["run_purpose"],
        })
    add_table(doc, pd.DataFrame(d23_rows), [(x, x) for x in ["Window", "Best epoch", "LR", "Warm-up", "Clip", "Batch / micro / accum.", "Energy kWh", "Purpose"]])
    add_text(doc, f"D23 audit: {d23['runs']} run, {d23['bootstrap_repetitions']} bootstrap repetitions, {d23['rows']['sanity_components']} sanity rows. Vì W3 dừng ở epoch 1 và validation trajectory không đủ ổn định, D23 là sensitivity pilot. Protocol F2 v2 đã được nhóm khóa (LR 1e-3, warm-up 0, không clipping); replay delta 0 chưa tự chứng minh tương đương trước–sau thay đổi training loop.")

    add_heading(doc, "8. Replay và việc còn mở")
    provenance = replay_provenance(replay_receipt)
    replay = pd.DataFrame([{
        "Cell": "CQ/W1/V0/RNN/seed42",
        "Backend": "VastAI L4, batch 2.048, deterministic",
        "Reference (commit)": provenance["reference_commit"],
        "Current (commit)": provenance["current_commit"],
        "Canonical facts mỗi phía": int(replay_receipt["rows_current"]),
        "Delta tuyệt đối lớn nhất": max(float(value) for value in replay_receipt["max_abs_delta"].values()),
        "Kết luận": provenance["verdict"],
    }])
    add_table(doc, replay, [(x, x) for x in replay.columns], font_size=8.6)
    add_text(doc, "Paired replay CQ/W1/V0/RNN/seed 42 so sánh 8 canonical facts theo split × phase, với delta bằng 0. Sai khác khoảng 0,0144 khi so với export lịch sử chưa có receipt chứng minh nguyên nhân; không được kết luận nguyên nhân từ phép so này. " + provenance["sentence"] + " Không dùng TEST để chỉnh model; seed 43/44 và MOMENT chưa thuộc claim seed-42 này.")

    add_heading(doc, "9. Định danh artefact")
    add_text(doc, f"Facts chính: {CQ_FACTS.relative_to(ROOT)}. Pilot D23: {CQ_D23.relative_to(ROOT)}. Paired-replay receipt: {CQ_REPLAY_RECEIPT.relative_to(ROOT)}. Tài liệu sinh lúc {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}.")
    doc.save(target)
    return target


def lo_report() -> Path:
    target = ROOT / "LO/docs/REPORT_LO_STATIC_V2_CONSOLIDATED_EVIDENCE_2026-10-28_R20.docx"
    metrics, manifests = load_lo_metrics()
    static_cols = static_contract()
    prefix_path = ROOT / "LO/resuilt/lo_prefix_source_audit_20261028/summary_static13_full_scan.json"
    prefix_audit = json.loads(prefix_path.read_text(encoding="utf-8"))
    if not prefix_audit["parity_all_ok"] or not all(row["future_mask_ok_sample"] for row in prefix_audit["mask"]):
        raise ValueError("LO static-13 prefix audit failed")
    cohort_root = ROOT / "LO/resuilt/lo_static_v2_cohort_audit_20261028"
    lo_replay = json.loads((cohort_root / "paired_replay_equivalence.json").read_text(encoding="utf-8"))
    if not lo_replay["paired_exact"] or len(lo_replay["cells"]) != 8:
        raise ValueError("LO paired replay equivalence not established")
    for family in ("RNN", "LSTM", "GRU", "BILSTM"):
        for index in range(17):
            cohort = json.loads((cohort_root / family / f"V{index}" / "summary.json").read_text(encoding="utf-8"))
            if cohort["cells"] != 96 or cohort["prediction_files"] != 24 or cohort["duplicate_hashes"] != 0 or cohort["invalid_strict"] != 0:
                raise ValueError(f"LO cohort audit failed: {family}/V{index}")
    closure_sha = hashlib.sha256(LO_CLOSURE.read_bytes()).hexdigest()
    receipt = json.loads((LO_ROOT / "LO_auto_download_stop_receipt.json").read_text(encoding="utf-8"))
    doc = setup_document(
        "BÁO CÁO LO — BẰNG CHỨNG STATIC-V2 HIỆN HÀNH",
        "LO V3.1 clean-grain · 13 static fields · seed 42 · F1 recurrent grid batch 2.048 · cập nhật 28/10/2026",
    )
    add_heading(doc, "1. Kết luận điều hành")
    add_text(doc, "Lưới LO static-v2 13 cột là lưới F1 seed 42 chính cho báo cáo này, gồm bốn họ recurrent: 4 family × 17 pipeline × 3 window × 4 phase × VALIDATION/TEST = 1.632 evaluation facts. Mỗi family có 51 run và audit hậu kiểm B=2.000. BiLSTM được lưu trong closure archive đã tải về và xác minh SHA-256; RNN/LSTM/GRU có facts audit trực tiếp tại local. Lưới 15 cột chỉ là bằng chứng lịch sử và không trộn vào các bảng chính.")
    add_text(doc, "AccTEMPO là metric hiệu năng chính. Accuracy không được dùng làm kết luận chính do bất cân bằng nhãn LO mạnh. S_san+ v1 là độ đo chất lượng dữ liệu chính; s_cal và S_san+ v2 là sensitivity. Toàn bộ kết quả dưới đây là static-v2 current layout 13 cột, không trộn historical layout 15 cột.")

    add_heading(doc, "2. Release và static contract")
    scope = pd.DataFrame([{
        "Release": manifests[0]["release_id"],
        "Data release": manifests[0]["data_release_id"],
        "Split / phase": "v3_1 / wide_prefix_v3_1",
        "Label": "CATALOG_NORMALIZED_PRIMARY_V3_1 — exclude no scored signal",
        "Seed": 42,
        "Windows / phases": "W1–W3 / P1–P4",
        "Pipelines / families": "V0–V16 / RNN, LSTM, GRU, BiLSTM",
        "Static policy": "current; 13 fields; split_rule_version và timeline_source loại trừ",
    }])
    add_table(doc, scope, [(x, x) for x in scope.columns])
    add_text(doc, "13 static fields: " + ", ".join(static_cols) + ". Dynamic P1–P4 và các missing/phase masks vẫn được giữ theo prefix. Hợp đồng này được dùng đồng nhất cho bốn họ trong report này.")
    add_heading(doc, "2.1. Audit prefix và static trên source V3.1", 2)
    add_text(doc, "Full-scan tất cả file và row group của TEST prefix V3.1: mỗi phase có 409.508 dòng; P1–P3 kiểm tra 189.192.696 ô thuộc phase tương lai và không có ô nào khác null. So 13 cột static của layout RNN/W1/V0 theo enrollment giữa P1 và P2–P4: 39/39 phép kiểm tra khớp, không có giá trị lệch. Kết quả chứng minh source prefix masking và static parity cho layout được kiểm tra; không tự xác nhận thời điểm sẵn có thực tế của feature catalogue hay tính đúng của toàn bộ prediction.")
    add_text(doc, f"Bằng chứng machine-readable: {prefix_path.relative_to(ROOT)}. View V3.1 chứa các cột nguồn; hợp đồng 13 static columns được áp dụng khi model chọn tensor, không phải tên một release view riêng.")

    add_heading(doc, "3. Độ bao phủ audit và provenance")
    add_table(doc, coverage_table(manifests, task="LO"), [(x, x) for x in ["Family", "Runs", "Sanity", "S_leak", "Slices", "Per-class", "Bootstrap B=2,000", "Checkpoint", "Release"]])
    add_text(doc, "S_leak được tính từ parent imputation train thực và frozen evaluation prefix; không refit riêng theo augmented pipeline. Audit lưu support/precision/recall/F1 theo lớp, CI bootstrap, checkpoint selection, STRICT/OVERLAP và probability invariant.")
    add_text(doc, f"Audit cohort current static-v2 hoàn tất 68/68 tổ hợp family × pipeline (204 run, 1.632 prediction files): mỗi pipeline có 24 file dự đoán và 96 cell SEEN/UNSEEN × STRICT/OVERLAP; không có hash enrollment trùng hoặc cờ STRICT không hợp lệ. Kết quả: {cohort_root.relative_to(ROOT)}. Phân nhóm SEEN xác định theo course_id có trong TRAIN cùng window; các lát là phân tích mô tả, không thay đổi tập TEST.")
    lo_prov = replay_provenance({"reference_commit": lo_replay.get("baseline_commit"), "current_commit": lo_replay.get("replay_commit")})
    lo_delta = max(float(cell["max_probability_delta"]) for cell in lo_replay["cells"])
    add_text(doc, f"Replay equivalence LO/W1/V0/RNN/seed 42: baseline {lo_prov['reference_commit']} và replay {lo_prov['current_commit']} cùng contract static-v2 13 cột, batch 2.048 và cùng chọn epoch {lo_replay['best_epoch']}. Ghép từng enrollment nguồn qua salt riêng từng run: cả {len(lo_replay['cells'])} split × phase khớp nhãn thật, nhãn dự đoán và xác suất, delta tối đa {lo_delta:g}. {lo_prov['sentence']} Receipt: LO/resuilt/lo_static_v2_cohort_audit_20261028/paired_replay_equivalence.json.")
    cq_match_path = ROOT / 'CQ/resuilt/cq_rnn_matched_historical_20261029/paired_replay_equivalence.json'
    cq_match = json.loads(cq_match_path.read_text(encoding='utf-8'))
    if not cq_match.get('paired_exact') or len(cq_match.get('cells', [])) != 8:
        raise ValueError('CQ matched replay evidence is incomplete')
    add_text(doc, 'Bằng chứng hỗ trợ loop chung từ CQ/W1/V0/RNN: matched replay giữ layout lịch sử và ghép theo enrollment, cả 8 split × phase có nhãn, prediction và xác suất khớp tuyệt đối; cùng chọn epoch 2. Baseline CQ có timestamp 20/10 trước 044d6c7 nhưng không ghi Git commit chính xác, nên không tự gán revision. Phép so này hỗ trợ kiểm tra loop chung, không phải phép before–after riêng trên LO. Việc không có baseline LO 13 cột trước 044d6c7 không làm lưới static-v2 hiện tại trở thành lưới lịch sử hay bắt buộc train lại.')
    archive = pd.DataFrame([{
        "Closure archive": LO_CLOSURE.name,
        "SHA-256": f"sha256:{closure_sha}",
        "Receipt UTC": receipt["completed_at_utc"],
        "BILSTM audit": "51 run / 408 sanity rows / B=2.000",
        "Status": "Downloaded and verified before remote stop",
    }])
    add_table(doc, archive, [(x, x) for x in archive.columns], font_size=8.6)

    metric_map = {"acctempo_m3": "AccTEMPO", "macro_f1": "Macro-F1", "balanced_acc": "Balanced acc.", "s_ent": "s_ent", "s_san_plus": "S_san+", "s_leak": "s_leak", "c0_recall": "Recall I/D", "c1_recall": "Recall G", "c2_recall": "Recall E"}
    aggregate = summary_table(metrics, metric_map=metric_map)
    add_heading(doc, "4. Kết quả đầy đủ theo family × pipeline")
    add_text(doc, "Mỗi hàng là trung bình của 12 cell W1–W3 × P1–P4 trong split tương ứng. Đây là bảng tổng hợp; facts cấp window × phase nằm trong từng family audit và được giữ nguyên để tái xuất.")
    cols = [("model_name", "Family"), ("pipeline_id", "Pipeline"), *[(k, v) for k, v in metric_map.items()]]
    for split in ("TEST", "VALIDATION"):
        add_heading(doc, f"4.{1 if split == 'TEST' else 2}. {split} — 68 cấu hình", 2)
        add_table(doc, aggregate[aggregate.eval_split.map(normalise_split).eq(split)], cols, font_size=8.2)

    add_heading(doc, "5. Hạt W1/P1 — dự đoán sớm nhất")
    add_text(doc, "W1/P1 được đưa riêng để tránh việc trung bình 12 cell che khuất bối cảnh khó nhất. Đây là mô tả theo config, không dùng TEST để chọn pipeline.")
    w1p1 = sort_matrix(metrics[(metrics.window_id.eq("W1")) & (metrics.phase_id.eq("P1"))].copy())
    for split in ("TEST", "VALIDATION"):
        add_heading(doc, f"5.{1 if split == 'TEST' else 2}. {split} W1/P1", 2)
        add_table(doc, w1p1[w1p1.eval_split.map(normalise_split).eq(split)], cols, font_size=8.2)

    add_heading(doc, "6. Runtime, năng lượng và giới hạn so sánh")
    add_text(doc, "Mọi static-v2 run dùng NVIDIA L4, batch 2.048, bfloat16 logits/loss float32, deterministic algorithms và NVML_POWER_SAMPLING_ESTIMATE. Train_seconds là chỉ số so sánh thời gian; energy GPU chỉ là ước lượng NVML (không tự suy diễn CO₂). Mỗi run manifest, train_history và checkpoint selection được lưu trong archive family/closure archive.")
    add_text(doc, "Historical LO 15-static-column evidence được quarantined: có thể dùng để truy vết lịch sử nhưng không trộn vào bảng static-v2. Salt cohort historical hiện được giữ tại local ngoài Git và có thể dùng để tái audit/đối chiếu aggregate SEEN/UNSEEN × STRICT/OVERLAP mà không retrain. Replay W1/V0/RNN theo patch provenance đã được archive/audit cùng closure. Nội dung salt không được đưa vào report, manifest hay repository.")

    add_heading(doc, "7. Trạng thái protocol và việc còn mở")
    status = pd.DataFrame([
        {"Hạng mục": "F1 recurrent static-v2 seed 42", "Trạng thái": "Đủ 204 run và 1.632 evaluation facts; 4 family audit xong"},
        {"Hạng mục": "Grain/prefix/static", "Trạng thái": "Clean-grain và static-v2 contract đã khóa; 13 cột loại trừ split_rule_version, timeline_source"},
        {"Hạng mục": "SEEN/UNSEEN × STRICT/OVERLAP current static-v2", "Trạng thái": "Đủ 204 run / 68 family × pipeline; không trùng hash và không có cờ STRICT lỗi"},
        {"Hạng mục": "F2 PatchTST/MOMENT", "Trạng thái": "Chưa thuộc claim primary của báo cáo này; chỉ mở sau gate protocol/replay tương ứng"},
        {"Hạng mục": "Seed stability", "Trạng thái": "Seed 43/44 là pha tiếp theo, không được suy diễn từ seed 42"},
    ])
    add_table(doc, status, [(x, x) for x in status.columns])
    add_heading(doc, "8. Định danh artefact")
    add_text(doc, f"Audit local: {LO_ROOT.relative_to(ROOT)}. Closure archive: {LO_CLOSURE.relative_to(ROOT)}. Báo cáo sinh lúc {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}.")
    doc.save(target)
    return target


def main() -> None:
    cq = cq_report()
    lo = lo_report()
    print(json.dumps({"cq_report": str(cq), "lo_report": str(lo)}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
