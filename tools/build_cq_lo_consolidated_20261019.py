"""Create reader-facing consolidated CQ and LO reports from audited sources."""
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

ROOT=Path(__file__).resolve().parents[1]

def prepend(doc, title, paragraphs):
    anchor=doc.element.body[0]
    heading=doc.add_heading(title,level=1)
    anchor.addprevious(heading._p)
    for value in paragraphs:
        p=doc.add_paragraph(value)
        anchor.addprevious(p._p)
    p=doc.add_paragraph('Nội dung chi tiết và các bảng theo từng pipeline, window, phase nằm ở các phần bên dưới.')
    anchor.addprevious(p._p)

def improve_font(doc):
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        if run.font.size is None or run.font.size.pt < 10:
                            run.font.size=Pt(10)

cq=Document(ROOT/'CQ/docs/REPORT_CQ_FULL_E2E_SEED42_CLOSURE_AUDIT_2026-10-19.docx')
prepend(cq,'TỔNG QUAN HIỆN HÀNH CQ — 19/10/2026',[
    'Phạm vi: CQ V2.2 clean-grain, V0–V16 (17 pipeline), W1–W3, P1–P4. AccTEMPO là độ đo chính; các thành phần sanity, per-class và uncertainty phải được đọc cùng kết quả dự đoán.',
    'LSTM/GRU/BiLSTM: bộ batch 2.048 seed 42 mới gồm 153 run và facts slice r2 đã ingest L2. Bảng batch 512 ở phần đầu là lịch sử, không trộn với bộ mới. RNN seed 42 là bộ riêng; không suy kết quả của model khác hoặc seed 43/44.',
    'Bằng chứng mới: 153/153 manifest batch 2.048 có GPU-only NVML estimate; 153 epoch được chọn khớp history. Replay checkpoint là mẫu W1/V0 trên CPU, chưa xác minh identity bằng hash. CQ TEST prefix đã full-scan; thời điểm có sẵn của static vẫn cần duyệt.'])
for para in cq.paragraphs:
    if para.text.startswith('Per-class facts giữ support, precision, recall, F1; bootstrap là B=2,000.'):
        para.text='[Mô tả bộ lịch sử; batch 2.048 mới có NVML GPU-only, xem phụ lục telemetry] '+para.text
improve_font(cq)
cq_out=ROOT/'CQ/docs/REPORT_CQ_TONG_HOP_DAY_DU_SEED42_2026-10-19.docx'; cq.save(cq_out)

lo=Document(ROOT/'LO/docs/REPORT_LO_FULL_E2E_SEED42_STATIC_HISTORY_2026-10-19.docx')
prepend(lo,'TỔNG QUAN HIỆN HÀNH LO — 19/10/2026',[
    'Phạm vi: LO V3.1 clean-grain, V0–V16 (17 pipeline), W1–W3, P1–P4, seed 42. Bốn family RNN/LSTM/GRU/BiLSTM có 204 run; các kết quả và audit của từng family được trình bày trong báo cáo.',
    'TEST prefix đã được quét đầy đủ các file/row group; không tìm thấy ô future-phase không null. Static-only probe đã chạy W1/V0, nhưng không chứng minh nguyên nhân chênh lệch validation–test.',
    'Cả 204 layout lịch sử chứa split_rule_version; toàn bộ 126 input đã kiểm tra chỉ chứa mã 1. Code hiện hành loại cột; kết quả cũ giữ nguyên với lineage lịch sử. Điều này không chứng minh checkpoint huấn luyện lại sẽ giống hệt.',
    'Còn mở: ablation hai cell không static; xác minh point-in-time của catalog/lịch, đối chiếu cohort và quyết định disposition timeline_source/window_duration_seconds. Không suy kết luận stability từ một seed hoặc suy CO₂ từ telemetry không phù hợp.'])
improve_font(lo)
lo_out=ROOT/'LO/docs/REPORT_LO_TONG_HOP_DAY_DU_SEED42_2026-10-19.docx'; lo.save(lo_out)
print(cq_out); print(lo_out)
