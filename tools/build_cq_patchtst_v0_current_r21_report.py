"""Build a traceable CQ PatchTST V0 report from the three downloaded runs."""
from __future__ import annotations

import hashlib
import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
RUNS = ROOT / 'CQ/resuilt/cq_patchtst_v0_current_seed42_r21'
AUDIT = ROOT / 'CQ/resuilt/cq_patchtst_v0_current_seed42_r21_audit/CQ/PATCHTST/seed=42'
PILOT = ROOT / 'CQ/resuilt/cq_patchtst_v0_seed42_20261023'
R4 = Path('G:/My Drive/Jobs/CQ/organized_moved_20261029/L2_facts/fact_exports/cq_recurrent_allfamilies_grid_r4_20261023')
OUT = ROOT / 'CQ/resuilt/cq_patchtst_v0_current_seed42_r21_report'
DOC = ROOT / 'CQ/docs/REPORT_CQ_PATCHTST_V0_CURRENT_2026-10-31.docx'
WINDOWS = ('W1', 'W2', 'W3')
KEYS = ['window_id', 'eval_split', 'phase_id']
CLASSES = ('warning', 'average', 'good')
EXPECTED_STATIC = ['ex_counts', 'gender', 'primary_risk_set', 'resource_count',
    'school_avg_bio_length_chars', 'school_avg_motto_length_chars', 'school_count',
    'teacher_avg_bio_length_chars', 'teacher_count', 'teacher_org_count', 'window_duration_seconds']


def sha(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


def fmt(v, digits=4):
    return '—' if pd.isna(v) else f'{float(v):.{digits}f}'


def table(doc, headers, records, fontsize=10):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = str(h)
        tcpr = cell._tc.get_or_add_tcPr()
        shading = OxmlElement('w:shd'); shading.set(qn('w:fill'), '1F4E78'); tcpr.append(shading)
        for run in cell.paragraphs[0].runs:
            run.bold = True; run.font.color.rgb = RGBColor(255,255,255)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for row in records:
        cells = t.add_row().cells
        for i, value in enumerate(row): cells[i].text = str(value)
    for row in t.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(0)
                for run in para.runs:
                    run.font.name = 'Calibri'; run.font.size = Pt(fontsize)
    doc.add_paragraph()


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    audit_manifest = json.loads((AUDIT/'audit_manifest.json').read_text(encoding='utf-8'))
    receipt = json.loads((RUNS/'download_verification.json').read_text(encoding='utf-8'))
    risk_audit_path=ROOT/'CQ/resuilt/cq_primary_risk_set_audit_20261031/audit.json'
    risk_audit=json.loads(risk_audit_path.read_text(encoding='utf-8'))
    assert risk_audit['source_constant_one'] and risk_audit['v0_input_constant_zero'] and len(risk_audit['files'])==18
    replay_current=json.loads((OUT/'current_epoch1_replay_receipt.json').read_text(encoding='utf-8'))
    assert replay_current['pass']
    assert audit_manifest['runs'] == 3 and audit_manifest['bootstrap_repetitions'] == 2000
    assert audit_manifest['release_id'] == 'CQ_V2_2' and len(receipt['runs']) == 3
    run_records=[]; histories=[]; raw_metrics=[]; pilot_metrics=[]; pilot_sanity=[]
    pilot_layouts={}; current_layouts={}
    for w in WINDOWS:
        p=RUNS/'CQ'/w/'V0/PATCHTST/seed=42'
        m=json.loads((p/'run_manifest.json').read_text(encoding='utf-8'))
        layout=json.loads((p/'feature_layout.json').read_text(encoding='utf-8'))
        recorded = next(x['run_manifest_sha256'] for x in receipt['runs'] if x['window']==w)
        assert sha(p/'run_manifest.json') == recorded
        assert (m['task'],m['window_id'],m['pipeline_id'],m['model_name'],m['seed']) == ('CQ',w,'V0','PATCHTST',42)
        assert (m['release_id'],m['split_version'],m['phase_version']) == ('CQ_V2_2','v2_2','wide_prefix_v2_2')
        assert (m['batch_size'],m['micro_batch_size'],m['accumulation_steps']) == (2048,512,4)
        assert m['static_layout_policy']=='current' and m['run_purpose']=='PRIMARY_GRID'
        assert layout['static_columns']==EXPECTED_STATIC and 'timeline_source' not in layout['static_columns']
        current_layouts[w]=sha(p/'feature_layout.json')
        h=pd.read_parquet(p/'train_history.parquet').copy(); h.insert(0,'window_id',w); histories.append(h)
        metrics=pd.read_parquet(p/'metrics_overall.parquet'); assert len(metrics)==8; raw_metrics.append(metrics)
        old=PILOT/w
        old_layout=json.loads((old/'feature_layout.json').read_text(encoding='utf-8'))
        pilot_layouts[w]=sha(old/'feature_layout.json')
        assert len(old_layout['static_columns'])==12 and 'timeline_source' in old_layout['static_columns']
        pm=pd.read_parquet(old/'metrics_overall.parquet').copy(); pilot_metrics.append(pm)
        run_records.append({'window':w,'best_epoch':m['best_epoch'],'epochs_run':len(h),
          'best_validation_macro_f1':float(h.validation_macro_f1_mean_p1_p4.max()),
          'elapsed_seconds':m['elapsed_seconds'],'train_seconds':h.train_seconds.sum(),
          'energy_gpu_estimated_kwh':m.get('energy_gpu_estimated_kwh'),
          'energy_measurement_type':m.get('energy_measurement_type'),
          'runner_git_commit':m.get('runner_git_commit'),'layout_sha256':current_layouts[w],
          'pilot_layout_sha256':pilot_layouts[w], 'data_release_id':m.get('data_release_id'),
          'train_rows':m.get('train_rows'),'validation_rows':m.get('validation_rows'),
          'test_rows':m.get('test_rows',{}).get('P1')})
    assert len({r['data_release_id'] for r in run_records}) == 1
    assert run_records[0]['data_release_id']==audit_manifest['data_release_id']
    run_df=pd.DataFrame(run_records); history=pd.concat(histories,ignore_index=True)
    selection=pd.read_parquet(AUDIT/'checkpoint_selection.parquet')
    sanity=pd.read_parquet(AUDIT/'sanity_components.parquet')
    slices=pd.read_parquet(AUDIT/'evaluation_slices.parquet')
    cls=pd.read_parquet(AUDIT/'metrics_per_class.parquet')
    boot=pd.read_parquet(AUDIT/'bootstrap_metrics_per_class.parquet')
    assert (len(selection),len(sanity),len(slices),len(cls),len(boot))==(3,24,72,72,216)
    for row in run_records:
        selected=selection[selection.window_id==row['window']].iloc[0]
        assert int(selected.best_epoch)==int(row['best_epoch'])
        assert abs(float(selected.best_score)-row['best_validation_macro_f1']) < 1e-12
    assert not sanity.duplicated(KEYS).any() and int(sanity.prediction_nan_inf_count.sum())==0
    assert not sanity[['acctempo_m3','macro_f1','balanced_acc','s_perf','s_san_plus','s_nan','s_maj_jsd','s_ent','s_drift','s_eff','s_leak']].isna().any().any()
    assert set(boot.bootstrap_repetitions)=={2000} and set(boot.bootstrap_method)=={'stratified_binomial_equivalent_v1'}
    assert set(sanity.s_leak_probe_fit_once)=={True} and all(sanity.s_leak_probe_n_eval==sanity.n_eval_samples)
    ci=boot[boot.metric_name=='recall'][KEYS+['class_code','ci_lower','ci_upper','bootstrap_method','bootstrap_repetitions']]
    class_ci=cls.merge(ci,on=KEYS+['class_code'],validate='one_to_one'); assert len(class_ci)==72
    rec=pd.read_parquet(R4/'cq_family_pipeline_window_phase_metrics.parquet')
    assert len(rec)==1632
    rec=rec[(rec.pipeline_id=='V0') & (rec.seed==42)].copy()
    assert len(rec)==96 and set(rec.model_name)=={'RNN','LSTM','GRU','BILSTM'}
    # Identical enrollment counts and class supports are necessary checks, not proof of identical cohort membership.
    sup=cls.pivot(index=KEYS,columns='class_index',values='n_support')
    for row in rec.itertuples(index=False):
        k=(row.window_id,row.eval_split,row.phase_id)
        assert tuple(int(z) for z in row.per_class_support)==tuple(int(sup.loc[k,i]) for i in range(3))
    paired=sanity[KEYS+['n_eval_samples','acctempo_m3','macro_f1','balanced_acc']].rename(columns={'acctempo_m3':'PATCHTST_current'})
    for family in ('RNN','LSTM','GRU','BILSTM'):
        q=rec[rec.model_name==family][KEYS+['acctempo_m3']].rename(columns={'acctempo_m3':family})
        paired=paired.merge(q,on=KEYS,validate='one_to_one')
    old_s=pd.read_parquet(PILOT/'audits_v0_seed42/CQ/PATCHTST/seed=42/sanity_components.parquet')
    old_s=old_s[KEYS+['acctempo_m3','macro_f1','balanced_acc']].rename(columns={'acctempo_m3':'pilot_acctempo','macro_f1':'pilot_macro_f1','balanced_acc':'pilot_balanced_acc'})
    pilot_compare=sanity[KEYS+['n_eval_samples','acctempo_m3','macro_f1','balanced_acc']].merge(old_s,on=KEYS,validate='one_to_one')
    assert len(pilot_compare)==24 and len(paired)==24
    for name,frame in [('run_contracts',run_df),('training_history',history),('checkpoint_selection',selection),
        ('metrics_and_sanity',sanity),('per_class_recall_ci',class_ci),('bootstrap_all_metrics',boot),
        ('per_class_metrics',cls),('evaluation_slices',slices),
        ('comparison_recurrent_v0',paired),('comparison_old_patchtst_v0',pilot_compare)]:
        frame.to_csv(OUT/f'{name}.csv',index=False,encoding='utf-8-sig')

    doc=Document(); sec=doc.sections[0]; sec.orientation=WD_ORIENT.LANDSCAPE
    sec.page_width=Inches(11.7); sec.page_height=Inches(8.3)
    sec.left_margin=sec.right_margin=Inches(.55); sec.top_margin=sec.bottom_margin=Inches(.55)
    normal=doc.styles['Normal']; normal.font.name='Calibri'; normal.font.size=Pt(9.5)
    title=doc.add_heading('CQ — ĐÁNH GIÁ PATCHTST V0 LAYOUT HIỆN HÀNH',0); title.alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph('CQ V2.2 · seed 42 · W1–W3 · P1–P4 · kết quả và audit tải về 31/10/2026')
    doc.add_heading('1. Kết luận và phạm vi',1)
    doc.add_paragraph('Ba run V0 đã hoàn tất và có audit hậu kiểm; đây là kết quả seed 42, không đại diện cho độ ổn định giữa các seed. AccTEMPO M3 là độ đo chính; Macro-F1, balanced accuracy và recall theo lớp giúp giải thích hiệu năng. Kết quả TEST mang tính báo cáo, không dùng để chọn hyperparameter.')
    primary=paired[(paired.window_id=='W1') & (paired.eval_split=='TEST') & (paired.phase_id=='P1')].iloc[0]
    doc.add_paragraph(f"Ở W1/TEST/P1, PatchTST V0 đạt AccTEMPO {fmt(primary.PATCHTST_current)}; RNN/LSTM/GRU/BiLSTM V0 lần lượt đạt {fmt(primary.RNN)}/{fmt(primary.LSTM)}/{fmt(primary.GRU)}/{fmt(primary.BILSTM)}. Chênh lệch nhỏ trong một ô không đủ kết luận model nào tốt hơn trên toàn lưới.")
    w2p4=sanity[(sanity.window_id=='W2') & (sanity.eval_split=='TEST') & (sanity.phase_id=='P4')].iloc[0]
    doc.add_paragraph(f"W2/TEST/P4 có AccTEMPO {fmt(w2p4.acctempo_m3)}, Macro-F1 {fmt(w2p4.macro_f1)} và s_ent {fmt(w2p4.s_ent)}. Cần đọc hiệu năng và thành phần sanity cùng nhau thay vì chỉ nhìn một điểm tổng hợp.")
    doc.add_paragraph('Cả ba run dùng layout CQ hiện hành 11 cột static. Pilot PatchTST cũ dùng 12 cột, có timeline_source; mẫu recurrent lịch sử W1/V0 của LSTM/GRU/BiLSTM dùng 13 cột, có timeline_source và split_rule_version. So sánh giữa các layout mang tính mô tả; không quy chênh lệch cho riêng kiến trúc. Counts và support khớp R4, nhưng riêng số lượng không chứng minh từng enrollment trùng nhau.')
    doc.add_paragraph('Kiểm tra CQ primary_risk_set: materializer gán hằng số 1; full scan ở view xác nhận giá trị này trên TRAIN/VALIDATION và TEST P1–P4. Trong cả 18 model-input V0 thuộc ba window, sau preprocessing cột này và mask đều bằng 0; bản thân cột hằng này không mã hóa thông tin thay đổi theo enrollment. Kiểm tra này không thay thế duyệt các static khác hoặc kiểm tra tiêu chí chọn cohort.')
    doc.add_paragraph('Trước khi mở 48 ô còn lại: kiểm tra review và contract của các static còn lại, xem kết quả V0 và thống nhất mở lưới. Replay một epoch dùng layout mới không khớp pilot cũ, vì vậy không dùng để xác nhận tái sử dụng ba pilot cũ.')
    doc.add_heading('2. Contract, tài nguyên và checkpoint',1)
    doc.add_paragraph('CQ_V2_2 / v2_2 / wide_prefix_v2_2; runner 29da3b2; PRIMARY_GRID; batch 2.048 = 4 × 512; LR 0,001; warm-up 0; không clip; tối đa 50 epoch, patience 5; GPU L4, BF16. Năng lượng là ước lượng GPU bằng NVML, không suy ra CO₂.')
    sel=selection.set_index('window_id')
    table(doc,['W','TRAIN','VALID','TEST/P','Epoch chọn','Epoch chạy','Score chọn','Margin','Trong margin','Train giờ','GPU kWh'],[
      [r['window'],f"{r['train_rows']:,}",f"{r['validation_rows']:,}",f"{r['test_rows']:,}",r['best_epoch'],r['epochs_run'],fmt(sel.loc[r['window'],'best_score']),fmt(sel.loc[r['window'],'selection_margin']),int(sel.loc[r['window'],'n_checkpoints_within_margin']),fmt(r['train_seconds']/3600,2),fmt(r['energy_gpu_estimated_kwh'],4)] for r in run_records])
    doc.add_paragraph('W3 chọn epoch 1 và dừng sau epoch 6; validation Macro-F1 trung bình giảm ở epoch 2–3 rồi phục hồi ở epoch 5. Đây là dao động trong một run, chưa đủ để quy nguyên nhân cho optimizer hay khẳng định model kém ổn định trên mọi seed.')
    doc.add_heading('2.1 Kiểm tra primary_risk_set cho CQ V0',2)
    doc.add_paragraph('Cột view CQ được gán F.lit(1), không được tính từ nhãn hay cutoff. Wide imputer thực tế giữ cột không có hậu tố này; layout V0 dùng cả cột và missing mask. Audit quét toàn bộ giá trị ở 18 file input và đối chiếu SHA-256 từng file với receipt upload. Giá trị trong tensor đầu vào là 0, mask là 0 trên mọi dòng được quét.')
    table(doc,['Window','Input','Số dòng','Giá trị sau preprocessing','Missing mask'],[
      [x['window'],x['split'],f"{x['primary_risk_set']['rows']:,}",str(x['primary_risk_set']['values']),str(x['missing_mask']['values'])] for x in risk_audit['files']])
    doc.add_paragraph('Kết luận chỉ áp dụng CQ V0 của release và input đã kiểm tra; không suy rộng sang các biến risk-set theo phase của LO hoặc V1–V16. Việc chọn sample sau join nhãn và tính sẵn có của các cột catalog/lịch là các câu hỏi riêng.')
    doc.add_heading('3. Metrics chính — riêng từng window, split và phase',1)
    table(doc,['W','Tập','P','N','AccTEMPO M3','Macro-F1','Balanced acc.','S_perf','S_san+ v1'],[
      [r.window_id,r.eval_split,r.phase_id,f'{int(r.n_eval_samples):,}',fmt(r.acctempo_m3),fmt(r.macro_f1),fmt(r.balanced_acc),fmt(r.s_perf),fmt(r.s_san_plus)] for r in sanity.itertuples(index=False)])
    doc.add_heading('4. Sáu thành phần sanity',1)
    doc.add_paragraph('s_leak được fit một lần trên TRAIN thật của parent V0 theo từng window và đánh giá trên tập thật; s_eff dùng model input của parent. S_san+ v1 là thước đo chất lượng dữ liệu; s_cal / S_san+ v2 là sensitivity.')
    table(doc,['W','Tập','P','s_nan','s_maj_jsd','s_ent','s_drift','s_eff','s_leak','S_san+'],[
      [r.window_id,r.eval_split,r.phase_id,fmt(r.s_nan),fmt(r.s_maj_jsd),fmt(r.s_ent),fmt(r.s_drift),fmt(r.s_eff),fmt(r.s_leak),fmt(r.s_san_plus)] for r in sanity.itertuples(index=False)])
    doc.add_heading('5. Recall, support và khoảng tin cậy từng lớp',1)
    doc.add_paragraph('CI 95% của recall sử dụng 2.000 lần theo phương pháp stratified_binomial_equivalent_v1 được ghi trong audit. Cờ support nhỏ theo ngưỡng 400. Thứ tự lớp: warning, average, good.')
    table(doc,['W','Tập','P','Lớp','Support','Recall','CI thấp','CI cao','Precision','F1','Nhỏ'],[
      [r.window_id,r.eval_split,r.phase_id,r.class_code,f'{int(r.n_support):,}',fmt(r.recall),fmt(r.ci_lower),fmt(r.ci_upper),fmt(r.precision),fmt(r.f1),int(r.small_support_flag)] for r in class_ci.itertuples(index=False)])
    doc.add_heading('6. So với recurrent V0 — AccTEMPO M3',1)
    doc.add_paragraph('Bốn họ recurrent là facts R4 seed 42; so sánh cùng W/split/P và V0. Đây là mô tả giữa các kết quả có layout static và runner khác nhau, không phải ablation kiểm soát.')
    table(doc,['W','Tập','P','PatchTST 11','RNN','LSTM','GRU','BiLSTM'],[
      [r.window_id,r.eval_split,r.phase_id,fmt(r.PATCHTST_current),fmt(r.RNN),fmt(r.LSTM),fmt(r.GRU),fmt(r.BILSTM)] for r in paired.itertuples(index=False)])
    doc.add_heading('7. So với PatchTST V0 cũ',1)
    doc.add_paragraph('Pilot cũ có timeline_source và runner khác; ba ô cũ không được gộp vào lưới layout mới. Replay một epoch W1 đã cho layout hash và ba metric epoch 1 khác nhau khi so với pilot cũ, do đó không chứng minh tương đương hai layout. Khi so cùng run V0 mới với replay một epoch, layout hash, contract và ba metric epoch 1 khớp bit-for-bit; kiểm tra này không thay thế replay toàn bộ prediction hoặc phép so loop cũ–mới.')
    table(doc,['W','Tập','P','AccTEMPO mới','AccTEMPO cũ','Macro-F1 mới','Macro-F1 cũ','BA mới','BA cũ'],[
      [r.window_id,r.eval_split,r.phase_id,fmt(r.acctempo_m3),fmt(r.pilot_acctempo),fmt(r.macro_f1),fmt(r.pilot_macro_f1),fmt(r.balanced_acc),fmt(r.pilot_balanced_acc)] for r in pilot_compare.itertuples(index=False)])
    doc.add_heading('8. STRICT và OVERLAP',1)
    doc.add_paragraph('Hai lát chỉ mô tả các tập đánh giá; không thay cohort ALL hoặc quy tắc chọn checkpoint. N và support từng lớp trong audit cần được xem khi diễn giải lát nhỏ.')
    filtered=slices[slices['slice'].isin(['STRICT','OVERLAP'])]
    table(doc,['W','Tập','P','Lát','N','Số lớp','Macro-F1','Balanced acc.','AccTEMPO'],[
      [r.window_id,r.eval_split,r.phase_id,r.slice,f'{int(r.n_eval_samples):,}',int(r.ev_n_classes_present),fmt(r.macro_f1),fmt(r.balanced_acc),fmt(r.acctempo_m3)] for r in filtered.itertuples(index=False)])
    doc.add_heading('9. Lịch sử validation từng epoch',1)
    table(doc,['W','Epoch','Train CE','VALID Macro-F1 TB','VALID CE TB','Train giây','LR cuối'],[
      [r.window_id,int(r.epoch),fmt(r.train_cross_entropy_mean_p1_p4),fmt(r.validation_macro_f1_mean_p1_p4),fmt(r.validation_cross_entropy_mean_p1_p4),fmt(r.train_seconds,1),fmt(r.learning_rate_end,4)] for r in history.itertuples(index=False)])
    doc.add_heading('10. Nguồn và giới hạn',1)
    doc.add_paragraph('Nguồn run: CQ/resuilt/cq_patchtst_v0_current_seed42_r21; audit: CQ/resuilt/cq_patchtst_v0_current_seed42_r21_audit; recurrent R4: Jobs/CQ/organized_moved_20261029/L2_facts/fact_exports/cq_recurrent_allfamilies_grid_r4_20261023. Bảng CSV đi kèm được tạo từ cùng dữ liệu; receipt đã kiểm tra hash ba manifest tải về.')
    doc.add_paragraph('Audit có 24 dòng sanity, 72 dòng per-class, 216 khoảng bootstrap, 72 lát và ba dòng checkpoint selection; không có prediction NaN/Inf trong audit. Bằng chứng primary_risk_set: CQ/resuilt/cq_primary_risk_set_audit_20261031/audit.json. Cần duyệt availability của các static khác và provenance trước khi mở lưới; hash runner baseline CQ lịch sử hiện chưa khớp lịch sử Git, nên không dùng để tự kết luận tương đương loop trước–sau.')
    doc.save(DOC)
    print(json.dumps({'report':str(DOC),'tables':str(OUT),'runs':3,'sanity':len(sanity),'per_class':len(class_ci),'bootstrap':len(boot),'recurrent_comparisons':len(paired),'slices':len(filtered),'max_probability_error':float(sanity.probability_sum_max_abs_error.max())},ensure_ascii=False,indent=2))

if __name__ == '__main__': main()
