"""Create a standalone CQ--LO Split V2.2 reference report."""
from __future__ import annotations

from pathlib import Path

import polars as pl
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "SPLIT_V2_2_REFERENCE_QA_REPORT_2026-09-29.docx"
NAVY, BLUE, GREY, GREEN = "17365D", "1F4E78", "F2F2F2", "E2F0D9"


def read(task: str, name: str) -> pl.DataFrame:
    files = list((ROOT / task / "resuilt" / "split_registry_v2_2_reference" / name).glob("*.parquet"))
    if len(files) != 1:
        raise RuntimeError(f"{task}/{name}: expected one parquet, found {len(files)}")
    return pl.read_parquet(files[0])


def shade(cell, colour: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), colour)
    props.append(node)


def paragraph_shade(paragraph, colour: str) -> None:
    props = paragraph._p.get_or_add_pPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), colour)
    props.append(node)


def table(doc: Document, headers: list[str], rows: list[list[object]], widths: list[float] | None = None) -> None:
    result = doc.add_table(rows=1, cols=len(headers))
    result.style = "Table Grid"; result.autofit = False
    for i, header in enumerate(headers):
        cell = result.rows[0].cells[i]; cell.text = ""
        run = cell.paragraphs[0].add_run(str(header))
        run.bold = True; run.font.size = Pt(11.5); run.font.color.rgb = RGBColor(255, 255, 255)
        shade(cell, BLUE)
        if widths: cell.width = Inches(widths[i])
    for j, row in enumerate(rows):
        cells = result.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            cells[i].paragraphs[0].add_run(str(value)).font.size = Pt(11.5)
            if j % 2: shade(cells[i], GREY)
            if widths: cells[i].width = Inches(widths[i])
    doc.add_paragraph()


def heading(doc: Document, value: str) -> None:
    p = doc.add_heading(value, 1)
    for run in p.runs: run.font.color.rgb = RGBColor.from_string(NAVY)


def pct(value: float) -> str:
    return f"{value * 100:.2f}%"


def main() -> None:
    qa = {task: read(task, "qa22_v2_2").sort("window_id") for task in ("CQ", "LO")}
    actual = {task: read(task, "split_actual_ratios") for task in ("CQ", "LO")}
    labels = {task: read(task, "split_label_distribution") for task in ("CQ", "LO")}
    strict = {task: read(task, "temporal_strict_share") for task in ("CQ", "LO")}
    calendar = {task: read(task, "calendar_boundaries").row(0, named=True) for task in ("CQ", "LO")}
    keys = ("a_boundary", "slice_1_boundary", "slice_2_boundary", "slice_3_boundary")
    if any(calendar["CQ"][key] != calendar["LO"][key] for key in keys):
        raise RuntimeError("CQ and LO do not share the V2.2 reference calendar")
    if not all(qa[task].select(pl.col("qa22_ok").all()).item() for task in qa):
        raise RuntimeError("V2.2 reference QA22 did not pass")

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.7); section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.8); section.right_margin = Cm(1.8)
    doc.styles["Normal"].font.name = "Aptos"; doc.styles["Normal"].font.size = Pt(10.5)
    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("CQ–LO · Split V2.2 reference · 2026-09-29 · Registry-derived").font.size = Pt(8)
    title = doc.add_paragraph(); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("BÁO CÁO CHUNG CQ–LO: SPLIT V2.2 REFERENCE")
    run.bold = True; run.font.size = Pt(21); run.font.color.rgb = RGBColor.from_string(NAVY)
    subtitle = doc.add_paragraph(); subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("scenario = hybrid · shared temporal calendar · reference protocol").italic = True
    verdict = doc.add_paragraph(); verdict.add_run("Kết luận: ").bold = True
    verdict.add_run("V2.2 reference đạt QA22 ở W1–W3 cho cả CQ và LO. Đây là baseline immutable dùng để tái hiện protocol V2.2 gốc.")
    paragraph_shade(verdict, GREEN)

    heading(doc, "1. Định nghĩa V2.2 reference")
    doc.add_paragraph(
        "V2.2 reference là bản chạy đúng contract V2.2 gốc theo tempo_fix. Nó dùng calendar chung CQ–LO, assignment ở cấp offering, "
        "và chỉ khóa QA22 khi: (1) label-end order đúng, nghĩa là offering TRAIN kết thúc không muộn hơn offering evaluation sớm nhất; "
        "và (2) validation, test của từng window đều không rỗng. Nó không áp ngưỡng tối thiểu theo tỷ lệ hoặc 50,000 dòng."
    )
    table(doc, ["Thuộc tính", "Quy ước V2.2 reference"], [
        ["scenario", "hybrid"], ["split_version", "v2_2"], ["Calendar anchor", "CQ; LO dùng cùng boundary"],
        ["Đơn vị gán", "offering + timeline_source; không split offering"], ["QA label-end order", "Bắt buộc"],
        ["QA quy mô arm", "validation ≥1 và test ≥1 enrollment"], ["purge_mode", "NONE ở reference registry; strict context được lưu riêng"],
        ["Nguồn immutable", "CQ/LO/resuilt/split_registry_v2_2_reference"],
    ], [2.4, 4.6])

    heading(doc, "2. Cơ chế chia dữ liệu")
    doc.add_paragraph(
        "A gồm 60% enrollment sớm nhất theo offering_end_date. Future pool được sắp theo thời gian và tách thành T1, T2, T3. "
        "Mỗi lát được tách thành hai block validation/test: T1→B/C, T2→D/E, T3→F/G. Do đó W1=A|B|C; "
        "W2=A+B+C|D|E; W3=A–E|F|G. CQ ghép balanced, còn LO ghép rarest_first để không bỏ rơi nhãn hiếm trong từng cặp."
    )
    cal = calendar["CQ"]
    table(doc, ["Boundary", "UTC"], [["A", cal["a_boundary"]], ["T1 (B/C)", cal["slice_1_boundary"]], ["T2 (D/E)", cal["slice_2_boundary"]], ["T3 (F/G)", cal["slice_3_boundary"]]], [2.3, 3.6])

    heading(doc, "3. Kết quả QA22")
    rows = []
    for task in ("CQ", "LO"):
        for row in qa[task].iter_rows(named=True):
            rows.append([task, row["window_id"], f"{row['n_enrollments_validation']:,}", f"{row['n_enrollments_test']:,}", f"{row['n_enrollments_evaluation']:,}", "Đạt" if row["end_order_ok"] else "Không đạt", "Đạt" if row["qa22_ok"] else "Không đạt"])
    table(doc, ["Task", "W", "Validation", "Test", "Eval gộp", "End order", "QA22"], rows, [.55, .35, 1., 1., 1., .75, .55])
    doc.add_paragraph("Mọi arm đều có dữ liệu và mọi window thỏa label-end order. Vì vậy V2.2 reference đạt đầy đủ điều kiện của chính protocol V2.2.")

    heading(doc, "4. Tỷ lệ split thực tế sau join nhãn")
    rows = []
    for task in ("CQ", "LO"):
        for window in ("W1", "W2", "W3"):
            by_split = {row["split"]: row for row in actual[task].filter(pl.col("window_id") == window).iter_rows(named=True)}
            rows.append([task, window, f"{by_split['train']['enrollment_count']:,}", pct(by_split['train']['split_enrollment_share']), f"{by_split['validation']['enrollment_count']:,}", f"{by_split['test']['enrollment_count']:,}"])
    table(doc, ["Task", "W", "Train", "Train share", "Validation", "Test"], rows, [.55, .4, 1.2, .9, 1.1, 1.1])

    heading(doc, "5. Phân phối nhãn ở validation và test")
    doc.add_paragraph(
        "Các tỷ lệ dưới đây được tính sau join nhãn của từng task, trong từng split và window; vì vậy không dùng một tỷ lệ chung cho CQ và LO. "
        "Chênh lệch giữa các window phản ánh temporal label drift, không phải lỗi cần cân bằng xuyên các lát thời gian."
    )
    for task in ("CQ", "LO"):
        doc.add_paragraph(f"{task}").runs[0].bold = True
        rows = []
        subset = labels[task].filter(pl.col("split").is_in(["validation", "test"]))
        for row in subset.sort(["window_id", "split", "label"]).iter_rows(named=True):
            rows.append([row["window_id"], row["split"], row["label"], f"{row['label_enrollment_count']:,}", pct(row["label_ratio"])])
        table(doc, ["W", "Split", "Nhãn", "Số enrollment", "Tỷ lệ"], rows, [.45, 1.1, 1.7, 1.3, .9])

    heading(doc, "6. Temporal strict và overlap")
    rows = []
    for task in ("CQ", "LO"):
        p1_test = strict[task].filter((pl.col("split") == "test") & (pl.col("phase_id") == "P1"))
        for row in qa[task].join(p1_test, on="window_id").iter_rows(named=True):
            rows.append([task, row["window_id"], pct(row["ev_train_overlap_share"]), pct(row["temporal_strict_share"]), f"{row['ev_train_overlap_share'] + row['temporal_strict_share']:.12f}"])
    table(doc, ["Task", "W", "TEST P1 overlap", "TEST P1 strict", "Tổng"], rows, [.55, .4, 1.35, 1.25, .9])
    doc.add_paragraph("ev_train_overlap_share là tỷ lệ TEST có cutoff P1 sớm hơn tau_w. Nó được báo riêng theo TEST; tổng với temporal_strict_share(P1, TEST) bằng 1.")

    heading(doc, "7. Cấu trúc lưu trữ và giới hạn sử dụng")
    doc.add_paragraph("split_registry_v2_2_reference/\n├── manifest/\n├── qa22_v2_2/\n├── split_registry/\n├── split_actual_ratios/\n├── split_label_distribution/\n├── temporal_strict_context/\n├── temporal_strict_share/\n├── duration_summary/\n├── calendar_boundaries/\n└── block_label_audit/")
    doc.add_paragraph("V2.2 reference là baseline tái hiện và đối chiếu. Khi chạy downstream theo protocol hiện hành, dùng manifest V2.3; không dùng các artifact split_registry_v2_2 legacy đã bị loại về integrity/semantics.")
    OUT.parent.mkdir(exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
