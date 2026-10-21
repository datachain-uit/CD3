from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "CQ" / "docs" / "REPORT_CQ_V2.2.docx"
OUTPUT = ROOT / "CQ" / "docs" / "REPORT_CQ_V2.2_CLEAN_GRAIN_S1_S2_2026-10-07.docx"


def clear_document(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    tc_pr.append(shade)


def set_cell_text(cell, text: str, *, bold: bool = False, size: float = 10.5) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)


def table(document: Document, headers: list[str], rows: list[list[object]]) -> None:
    item = document.add_table(rows=1, cols=len(headers))
    item.style = "Table Grid"
    for cell, text in zip(item.rows[0].cells, headers):
        set_cell_text(cell, text, bold=True)
        set_cell_shading(cell, "D9EAF7")
    for row in rows:
        cells = item.add_row().cells
        for cell, value in zip(cells, row):
            set_cell_text(cell, value)
    document.add_paragraph()


def para(document: Document, text: str, *, bold: bool = False) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = bold
    return paragraph


def bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.add_run("• " + text)


def heading(document: Document, text: str, level: int = 1) -> None:
    # The historical template has Word heading styles whose display names are
    # visible but not addressable through python-docx.  Use explicit formatting
    # while retaining the template's page/header/footer settings.
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(14 if level == 1 else 12)


def main() -> None:
    document = Document(TEMPLATE)
    clear_document(document)
    styles = document.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    styles["Normal"].font.size = Pt(11)
    for style_name in ("Title",):
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("BÁO CÁO CQ — RELEASE V2.2 CLEAN-GRAIN\nS0/S1/S2 HOÀN TẤT, S3 ĐANG KHỞI ĐỘNG")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Phạm vi: learner–course enrollment | scenario=hybrid | cập nhật 07/10/2026")
    para(document, "Trạng thái báo cáo: đây là bản mới, độc lập với REPORT_CQ_V2.2.docx. Các kết quả RNN seed 42 cũ không được mang sang, vì S3 được chạy lại trên input clean-grain và release identity CQ_V2_2 đã khóa.")

    heading(document, "Tóm tắt điều hành")
    table(document, ["Hạng mục", "Trạng thái", "Kết luận"], [
        ["Release input", "LOCKED", "CQ_V2_2; hybrid split v2_2; wide_prefix_v2_2; one row per enrollment."],
        ["S0 natural-state profile", "SUCCESS", "ALL_ROWS; dùng làm baseline data-quality chính thức."],
        ["S1 imputation", "SUCCESS 15/15", "V0, V1, V5, V9, V13 × W1–W3; input/model artifact đầy đủ."],
        ["S2 augmentation", "SUCCESS 36/36", "CDSMOTE, SASMOTE, RadiusSMOTE × 4 parent imputers × W1–W3; seed augmentation=42."],
        ["S3 model", "PENDING", "RNN seed 42 được khởi động lại; cần secret tempo-prediction-salt trước khi worker train được tạo."],
        ["Outcome chính", "LOCKED", "AccTEMPO-M3 trên TEST; Macro-F1/balanced accuracy/MCC/calibration là chỉ số hỗ trợ."],
    ])

    heading(document, "1. Mục tiêu, đơn vị phân tích và phạm vi")
    para(document, "Mục tiêu của CQ là dự báo nhãn vận hành canonical ở cấp learner–course enrollment. Đơn vị phân tích duy nhất là enrollment_id; user_id và course_id chỉ là context/categorical codes được quản lý theo TRAIN, không phải khóa để nhân dòng.")
    bullet(document, "Không random split giữa enrollment hoặc giữa các phase của cùng enrollment.")
    bullet(document, "Không dùng feature/event sau prediction cutoff của P1–P4.")
    bullet(document, "Validation và TEST luôn là dữ liệu thật; augmentation chỉ materialize TRAIN.")
    bullet(document, "Không dùng output từ release V1/label_stratified_v1 hay wide_prefix_v1 trong báo cáo này.")

    heading(document, "2. Khóa release và grain dữ liệu")
    para(document, "Release thí nghiệm là CQ_V2_2, feature_regime=CQ_RAW_EARLY, scenario=hybrid, split_version=v2_2 và phase_version=wide_prefix_v2_2. Materializer hiện dùng cumulative_phase_features_v2 sạch-grain; nếu nguồn feature còn exact duplicate thì materializer dừng cứng thay vì âm thầm khử trùng trong S1.")
    table(document, ["Window", "TRAIN rows", "Ý nghĩa"], [
        ["W1", "1,609,276", "Cohort lịch sử A làm TRAIN; validation/test thuộc lát đánh giá W1."],
        ["W2", "1,963,060", "TRAIN tích lũy các block trước W2 theo split registry khóa."],
        ["W3", "2,324,027", "TRAIN tích lũy lịch sử đến trước lát đánh giá W3."],
    ])
    para(document, "GRAIN_OK được kiểm tra trước S1 trên TRAIN, VALIDATION và từng TEST prefix: enrollment_id không null và không lặp. Đây là điều kiện chặn vì duplicate sẽ làm sai support, k-NN/SMOTE, số bước train, confidence interval và sanity facts.")
    para(document, "Bằng chứng materialization hiện hành ghi nguồn feature 11,807,090 enrollment sau clean grain. Nhãn, enrollment windows và joined base không có duplicate tại grain yêu cầu. Release input trên Modal không bị thay đổi giữa S0 và S1/S2.")

    heading(document, "3. Split V2.2 và phase-safe evaluation")
    para(document, "CQ dùng hybrid temporal split V2.2 với CQ là calendar anchor. Offering_id × timeline_source là unit split; một offering không được xuất hiện ở hai block. A là lịch sử, còn các lát T1/T2/T3 được dùng để tạo validation/test W1/W2/W3 theo split registry đã khóa.")
    bullet(document, "TEST P1 chỉ quan sát P1; P2 quan sát P1–P2; P3 quan sát P1–P3; P4 quan sát P1–P4.")
    bullet(document, "Giá trị future phase giữ unavailable/NULL; không được impute hoặc augment thành thông tin tương lai.")
    bullet(document, "Purge, embargo, label drift và split ratio theo task là metadata của registry, không suy diễn từ tỷ lệ chung giữa task.")

    heading(document, "4. S0/S1: data quality và imputation")
    para(document, "S0 và S1 đều profile ALL_ROWS. ‘Ô thiếu đủ điều kiện’ là numeric value missing khi modality tương ứng đã quan sát; đó là đối tượng imputation. ‘Ô vắng mặt cấu trúc’ là modality không tồn tại trong dữ liệu và được bảo toàn bằng mask, không đưa vào mẫu số imputation rate.")
    table(document, ["Pipeline", "Variant", "Fit/transform contract", "Trạng thái W1–W3"], [
        ["V0", "zero-fill + masks", "Baseline; masks không impute/scaled.", "SUCCESS"],
        ["V1", "median", "Fit TRAIN pooled P1–P4; transform toàn bộ TRAIN/VALIDATION/TEST prefix.", "SUCCESS"],
        ["V5", "mean", "Cùng leakage contract.", "SUCCESS"],
        ["V9", "ExtraTrees light", "Fit TRAIN only; estimator nhẹ theo runtime contract.", "SUCCESS"],
        ["V13", "MICE", "3 iterations; fit sample tối đa 1,000,000 TRAIN rows; transform full rows.", "SUCCESS"],
    ])
    para(document, "QA S1 hiện hành đạt 15/15 cell: có đủ train, validation, test_P1..test_P4, fitted preprocessor, không có unfittable static/dynamic column, và future_imputed_cells=0. Numeric nulls ở TEST P1–P3 chỉ là future-phase unavailable được bảo toàn; TRAIN, VALIDATION và TEST P4 không còn residual eligible numeric null.")

    heading(document, "5. S2: augmentation TRAIN-only")
    para(document, "S2 không thay validation/test, không thay label release và không refit imputer. Mỗi artifact balanced TRAIN kế thừa parent imputation manifest, hash manifest parent, release/split/phase/label identity và synthetic ledger có parent_a, parent_b, alpha cùng phương pháp sinh.")
    table(document, ["Parent imputation", "CDSMOTE", "SASMOTE", "RadiusSMOTE"], [
        ["V1 median", "V2", "V3", "V4"],
        ["V5 mean", "V6", "V7", "V8"],
        ["V9 ExtraTrees", "V10", "V11", "V12"],
        ["V13 MICE", "V14", "V15", "V16"],
    ])
    table(document, ["Window", "Original TRAIN", "Balanced TRAIN", "Synthetic rows", "S2 status"], [
        ["W1", "1,609,276", "1,866,048", "256,772", "12/12 methods × parents SUCCESS"],
        ["W2", "1,963,060", "2,269,260", "306,200", "12/12 methods × parents SUCCESS"],
        ["W3", "2,324,027", "2,688,368", "364,341", "12/12 methods × parents SUCCESS"],
    ])
    para(document, "S2 seed được cố định là 42. Đây là một fixed-balancer design: stability seeds 42/43/44 sẽ đo variation của model initialization và dataloader, không nhân thêm biến thiên của balancer. Mỗi balanced artifact có QA quality, mask integrity, phase-boundary integrity, availability integrity và row-count reconciliation.")
    para(document, "RadiusSMOTE chạy chậm hơn vì phải tìm radius-neighbors trong không gian numeric đã scale. fallback_nearest_neighbour_rate phải được xuất theo window; ngưỡng >0.50 là cảnh báo radius thưa, không phải runtime failure. Không dùng fallback rate để xóa artifact sau khi đã hoàn tất; nó là biến diễn giải/quality fact.")

    heading(document, "6. S3: model protocol đang thực thi")
    para(document, "Trình tự hiện hành là chạy RNN seed 42 trên V0–V16 và W1–W3 trước. Sau đó cùng input/model contract được mở rộng cho LSTM, GRU và BiLSTM; các seed model 43/44 là replication. GPU là NVIDIA L4, AMP bfloat16; model seed và dataloader seed được sinh deterministic từ master seed, role, task, window, pipeline và model.")
    table(document, ["Thành phần", "Quy tắc khóa"], [
        ["TRAIN", "model_inputs/train.parquet với V0/V1/V5/V9/V13; balanced_train.parquet với V2–V4/V6–V8/V10–V12/V14–V16."],
        ["VALIDATION", "Dữ liệu thật, không augmentation; chọn checkpoint bằng validation aggregate P1–P4."],
        ["TEST", "Dữ liệu thật P1–P4 prefix, chỉ inference bằng checkpoint/scaler đã khóa."],
        ["Seed", "Imputation=20260922; augmentation=42 fixed; model/dataloader master seeds=42,43,44."],
        ["Prediction privacy", "TEMPO_PREDICTION_SALT từ Modal secret tempo-prediction-salt; không ghi raw sensitive identifier vào prediction facts."],
    ])
    para(document, "Tại thời điểm viết report, S3 chưa có kết quả mới để diễn giải: worker bị chặn trước khi start do account Modal thiếu secret tempo-prediction-salt/TEMPO_PREDICTION_SALT. Đây là lỗi deploy configuration, không phải lỗi data, imputation hay augmentation. Sau khi tạo secret, rerun từ W1; các lần thất bại không tạo model artifact.")

    heading(document, "7. Outcome, sanity và nguyên tắc đọc kết quả")
    para(document, "AccTEMPO-M3 là primary outcome trên TEST đã khóa. Macro-F1, balanced accuracy, MCC, per-class support/recall, calibration (NLL/Brier/ECE), probability matrices và confusion matrices là secondary/supporting outcomes. Accuracy đơn lẻ không được dùng để chọn winner khi nhãn mất cân bằng.")
    bullet(document, "S_san+ v1 là data-quality/sanity component chính; s_cal v2 chỉ là sensitivity report. Nó không thay AccTEMPO-M3 làm performance outcome.")
    bullet(document, "s_leak, s_eff và s_drift tính trên real rows theo pipeline contract; không nhân lại chỉ vì thay kiến trúc RNN/LSTM/GRU/BiLSTM.")
    bullet(document, "Selection dùng VALIDATION; TEST chỉ báo cáo sau khi recipe được khóa. Không chọn pipeline từ TEST mean P1–P4.")
    bullet(document, "Báo cáo tách P1/P2/P3/P4 và W1/W2/W3; không gộp temporal drift thành một con số duy nhất.")

    heading(document, "8. Cấu trúc lưu trữ và truy vết")
    para(document, "Mọi retry là immutable attempt mới; không ghi đè attempt cũ. Model/augmentation resolver chỉ chọn attempt SUCCESS khớp đầy đủ release identity và parent manifest hash.")
    tree = """tempo-data-v1/
├── input/CQ_v2_2/
│   ├── release_manifest.json
│   ├── phase_views_v2_2/                 # TRAIN + VALIDATION
│   └── test_prefix_views_v2_2/P1..P4/    # TEST phase prefixes
└── meta_release=imputation-v1/
    ├── L0_registry/                       # state, environment, pipeline registry
    ├── L1_runs/task=CQ/feature_regime=CQ_RAW_EARLY/window_id=W*/
    │   ├── pipeline_id=V0,V1,V5,V9,V13/model_name=IMPUTATION_ONLY/
    │   │   └── seed=20260922/run_id=*/attempt_id=*/model_inputs/
    │   ├── pipeline_id=V2..V16/model_name=BALANCED_TRAIN_ONLY/
    │   │   └── seed=42/run_id=*/attempt_id=*/{balanced_train,synthetic_ledger}.parquet
    │   └── pipeline_id=V0..V16/model_name=RNN/seed=42|43|44/
    │       └── run_id=*/attempt_id=*/{checkpoint,predictions,probabilities,confusions}
    ├── L2_facts/                          # DQ, fidelity, augmentation quality, metrics, sanity, resource
    └── L3_meta/                           # aggregate meta-observations after paired grid is complete"""
    paragraph = document.add_paragraph()
    run = paragraph.add_run(tree)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)

    heading(document, "9. Quyết định và bước kế tiếp")
    table(document, ["Ưu tiên", "Việc cần làm", "Điều kiện hoàn thành"], [
        ["1", "Tạo Modal secret tempo-prediction-salt với key TEMPO_PREDICTION_SALT.", "Model app deploy được; worker RNN bắt đầu."],
        ["2", "Chạy RNN seed 42 cho V0–V16, W1–W3; mỗi cell kèm model sanity.", "51 model cells và sanity facts SUCCESS."],
        ["3", "Xuất facts validation/TEST, probability/confusion matrices và resource usage.", "Đủ P1–P4 × W1–W3 cho mỗi pipeline."],
        ["4", "Chọn shortlist bằng validation; mới mở seed 43/44 và LSTM/GRU/BiLSTM.", "Không test-select; recipe khóa trước TEST synthesis."],
        ["5", "Báo cáo Radius fallback rate theo window/pipeline.", "Cảnh báo >0.50 được diễn giải cùng hiệu năng."],
    ])
    para(document, "Kết luận: CQ V2.2 clean-grain đã qua toàn bộ bước dữ liệu trước model. S0/S1 đủ 15 imputation cells và S2 đủ 36 balanced TRAIN artifacts, không leakage vào future phase hoặc validation/test. S3 chỉ đang chờ cấu hình secret prediction salt; chưa có kết quả model mới để xếp hạng pipeline.")

    document.core_properties.title = "CQ V2.2 clean-grain — S0/S1/S2 report"
    document.core_properties.subject = "CQ release lock, imputation, augmentation, and S3 protocol"
    document.core_properties.author = "CD3 experimental pipeline"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
