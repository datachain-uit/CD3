"""Refresh the legacy CQ V1 report with the active V2.2 experiment contract."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "CQ" / "docs" / "REPORT_CQ_V1.docx"
FALLBACK = ROOT / "CQ" / "docs" / "REPORT_CQ_V1_CAP_NHAT_V2_2_2026-09-30.docx"
NAVY, BLUE, GREY = "17365D", "1F4E78", "F2F2F2"


def shade(cell, color: str) -> None:
    props = cell._tc.get_or_add_tcPr(); node = OxmlElement("w:shd")
    node.set(qn("w:fill"), color); props.append(node)


def heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph(); run = p.add_run(text)
    run.bold = True; run.font.size = Pt(14.5); run.font.color.rgb = RGBColor.from_string(NAVY)


def table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    value = doc.add_table(rows=1, cols=len(headers)); value.style = "Table Grid"
    for i, item in enumerate(headers):
        cell = value.rows[0].cells[i]; cell.text = ""
        run = cell.paragraphs[0].add_run(item); run.bold = True; run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(255, 255, 255); shade(cell, BLUE)
    for r, row in enumerate(rows):
        cells = value.add_row().cells
        for i, item in enumerate(row):
            cells[i].text = ""; cells[i].paragraphs[0].add_run(str(item)).font.size = Pt(10.5)
            if r % 2: shade(cells[i], GREY)
    doc.add_paragraph()


def replace(p, text: str) -> None:
    p.clear(); p.add_run(text)


def delete_from(doc: Document, marker: str) -> None:
    target = next((p for p in doc.paragraphs if p.text.strip().startswith(marker)), None)
    if target is None: raise RuntimeError(f"Không tìm thấy marker {marker}")
    body = doc._element.body; children = list(body); start = children.index(target._element)
    for child in children[start:]:
        if not child.tag.endswith("}sectPr"): body.remove(child)


def main() -> None:
    doc = Document(REPORT)
    doc.core_properties.title = "Báo cáo quy trình CQ — trạng thái V2.2 hiện hành"
    replace(doc.paragraphs[0], "BÁO CÁO QUY TRÌNH CQ — TRẠNG THÁI V2.2 HIỆN HÀNH")
    replace(doc.paragraphs[1], "Phạm vi: learner–course enrollment | feature/label/split V2.2 (hybrid) | cập nhật 30/09/2026")
    semantic = next(p for p in doc.paragraphs if p.text.startswith("Quy ước dự án:"))
    replace(semantic, "Quy ước hiện hành: CQ canonical là nhãn vận hành chính, cố định ở cấp enrollment từ vector [COELO, AFELO, ACELO]. Nhãn không phải completion hoặc final grade quan sát trực tiếp; mọi component cấu tạo nhãn phải theo proxy/leakage contract.")
    for old, new in {
        "6. Tách phase P1–P4": "6. Split V2.2 hiện hành và phase P1–P4",
        "7. Chia temporal block": "7. Cohort CQ, nhãn và temporal drift",
        "8. Các bước tiếp theo": "8. Integrity, support và lộ trình downstream",
    }.items():
        p = next((p for p in doc.paragraphs if p.text.strip().startswith(old)), None)
        if p: replace(p, new)
    delete_from(doc, "6. Phase P1–P4 và Method 5")

    heading(doc, "6. Split V2.2 hiện hành và phase P1–P4")
    doc.add_paragraph("Split V1/V2.1 chỉ là provenance; release thực nghiệm là scenario=hybrid, split_version=v2.2, CQ là calendar anchor. Đơn vị gán offering_id × timeline_source không được tách giữa split. Test Pj chỉ phơi lộ prefix P1…Pj.")
    table(doc, ["Window", "TRAIN", "VALIDATION", "TEST", "Lát thời gian"], [
        ["W1", "A", "B", "C", "T1"], ["W2", "A+B+C", "D", "E", "T2"], ["W3", "A+B+C+D+E", "F", "G", "T3"],
    ])
    table(doc, ["Mốc calendar chung", "UTC"], [
        ["Kết thúc A", "2020-07-31"], ["Kết thúc T1 (B/C)", "2020-09-13"],
        ["Kết thúc T2 (D/E)", "2020-12-31"], ["Kết thúc T3 (F/G)", "2022-01-18"],
    ])
    table(doc, ["W", "Validation", "Test", "Eval gộp", "End-order", "QA22"], [
        ["W1", "175,689", "178,095", "353,784", "Đạt", "Đạt"],
        ["W2", "180,485", "180,482", "360,967", "Đạt", "Đạt"],
        ["W3", "180,032", "180,031", "360,063", "Đạt", "Đạt"],
    ])

    heading(doc, "7. Cohort CQ, nhãn và temporal drift")
    table(doc, ["Hạng mục", "Giá trị"], [
        ["Cohort hybrid canonical", "2,684,090 enrollment; 3,716 offering"],
        ["Nhãn", "CQ_label_final: warning / average / good"],
        ["Nhãn cố định", "Một label cuối offering dùng chung P1–P4; chỉ feature prefix thay đổi."],
        ["Phân tầng block", "CQ pairing = balanced; label drift giữa T1/T2/T3 được giữ nguyên để đánh giá đúng thời gian."],
    ])
    table(doc, ["Block", "Offering", "Enrollment", "warning", "average", "good"], [
        ["A", "1,596", "1,609,276", "1,576,716", "19,394", "13,166"],
        ["B", "367", "175,689", "157,642", "10,099", "7,948"],
        ["C", "379", "178,095", "156,692", "10,096", "11,307"],
        ["D", "281", "180,485", "174,628", "2,462", "3,395"],
        ["E", "282", "180,482", "174,628", "2,462", "3,392"],
        ["F", "404", "180,032", "171,054", "5,808", "3,170"],
        ["G", "407", "180,031", "171,054", "5,808", "3,169"],
    ])
    table(doc, ["W", "Split", "Enrollment", "warning", "average", "good"], [
        ["W1", "validation", "175,689", "89.728%", "5.748%", "4.524%"], ["W1", "test", "178,095", "87.982%", "5.669%", "6.349%"],
        ["W2", "validation", "180,485", "96.755%", "1.364%", "1.881%"], ["W2", "test", "180,482", "96.756%", "1.364%", "1.879%"],
        ["W3", "validation", "180,032", "95.013%", "3.226%", "1.761%"], ["W3", "test", "180,031", "95.013%", "3.226%", "1.760%"],
    ])

    heading(doc, "8. Integrity assignment, strictness và offering dài")
    table(doc, ["Assignment QA", "Kết quả"], [
        ["Distinct enrollment", "2,684,090"], ["Enrollment đa rolling block", "0"],
        ["Enrollment đa split trong cùng window", "0"], ["Unit đa block", "0"], ["Kết luận", "Đạt"],
    ])
    table(doc, ["W", "TEST strict P1", "P1 cutoff-overlap"], [
        ["W1", "12.50%", "87.50%"], ["W2", "84.89%", "15.11%"], ["W3", "1.46%", "98.54%"],
    ])
    doc.add_paragraph("Offering dài (>270 ngày) được giữ nguyên đơn vị split. Prediction-level output phải mang timeline_source và long_offering_flag để phân tầng metric; không tạo subgroup bắt buộc mới chỉ vì duration dài.")

    heading(doc, "9. Data quality, imputation và model contract")
    table(doc, ["Hạng mục", "Quy tắc hiện hành"], [
        ["Feature dictionary", "Khóa trước imputation; label/proxy, ID, binary mask và future-phase availability không được làm predictor sai quy tắc."],
        ["S0", "Chạy lại trên V2.2 W1–W3; báo eligible missing, structural absence và event sparsity theo temporal slice × timeline_source."],
        ["Imputation", "Fit train-only; validation/test transform-only; observed missing only; không impute structural absence/future phase."],
        ["Augmentation", "TRAIN-only của từng window; validation/test không augment."],
        ["Model", "Seed 42, L4, bf16 AMP; selection = MEAN_PHASE_MACRO_F1 trên VALIDATION, tiebreak mean cross-entropy thấp nhất."],
    ])
    doc.add_paragraph("Kết quả S0/imputation/model V1 là provenance, không được trộn với V2.2. CQ không cần rebuild label hay registry; cần materialize view V2.2 rồi chạy lại P0/S0/downstream trên release mới.")

    heading(doc, "10. Support, metric và điều kiện báo cáo")
    table(doc, ["Contract", "Quy tắc"], [
        ["Metric chính", "Macro-F1, balanced accuracy, MCC, Cohen kappa, G-mean; không diễn giải Accuracy một mình do imbalance."],
        ["Per-class", "Precision/recall/F1 + n_support, probability matrix và confusion matrix cho mỗi window × phase × eval split."],
        ["Support nhỏ", "small_support_flag = n_support <400; CQ evaluation support nhỏ nhất hiện là 2,462, nên không có cell class chính nào bị cờ trước khi phân tầng STRICT/context."],
        ["STRICT/context", "Báo metric tổng thể và tách STRICT/OVERLAP; khi có prediction context, tách thêm timeline_source và long_offering_flag."],
        ["Sanity", "Báo s_cal v2 cạnh s_ent v1, cùng S_san+, leakage, calibration và drift facts."],
    ])

    heading(doc, "11. Lộ trình thực nghiệm CQ")
    table(doc, ["Bước", "Trạng thái", "Điều kiện"], [
        ["Registry V2.2", "Đạt", "CQ calendar anchor; ghi scenario, split_version, parent manifest/audit, config/code hash và seed."],
        ["Materialize V2.2", "Cần kiểm tra/chạy", "phase_views_v2_2 và test_prefix_views_v2_2/P1…P4 phải có _SUCCESS."],
        ["P0 + S0", "Chạy sau materialize", "Khóa dictionary, QA proxy/availability rồi đọc DQ theo lát."],
        ["V0 + imputation", "Chạy lại trên V2.2", "Không tái dùng model inputs hay facts V1."],
        ["Augmentation + model", "Được đi tiếp sau input V2.2", "Augment train-only; metric có support/CI/context theo contract."],
        ["L3 meta-dataset", "Chưa materialize", "Chỉ build sau run grid paired đủ V0/imputation/balancing/model."],
    ])
    doc.add_paragraph("Kết luận: CQ đã đủ điều kiện đi tiếp toàn bộ experimental pipeline V2.2. Không có điều kiện label tương tự LO; blocker thực tế duy nhất là materialize đúng release và không trộn artifact V1 vào run V2.2.")

    footer = doc.sections[0].footer.paragraphs[0]; footer.clear(); footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("CQ · Báo cáo quy trình cập nhật V2.2 · 30/09/2026").font.size = Pt(8.5)
    try: doc.save(REPORT); print(REPORT)
    except PermissionError: doc.save(FALLBACK); print(FALLBACK)


if __name__ == "__main__":
    main()
