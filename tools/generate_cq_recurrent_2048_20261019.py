"""Build a separate CQ recurrent batch-2048 audit report from local facts."""
from pathlib import Path
import json
import pandas as pd
from docx import Document
from docx.shared import Pt, Inches

ROOT = Path(__file__).resolve().parents[1]
REC = ROOT / 'CQ/resuilt/cq_recurrent_batch2048_slice_r2_20261019/audits/CQ'
OUT = ROOT / 'CQ/docs/REPORT_CQ_RECURRENT_BATCH2048_SEED42_SLICE_R2_2026-10-19.docx'
paths = {'LSTM': REC / 'LSTM/seed=42', 'GRU': REC / 'GRU/seed=42', 'BiLSTM': REC / 'BILSTM/seed=42'}
doc = Document()
style = doc.styles['Normal']; style.font.name = 'Arial'; style.font.size = Pt(11)
section = doc.sections[0]; section.top_margin = Inches(.65); section.bottom_margin = Inches(.65)
doc.add_heading('BÁO CÁO CQ — RECURRENT BATCH 2.048, SEED 42', 0)
doc.add_paragraph('Ngày tổng hợp: 19/10/2026. Báo cáo bổ sung từ facts audit VastAI; không thay thế báo cáo protocol/split và không đại diện cho kết quả nhiều seed.')
doc.add_heading('1. Phạm vi và cách đọc', 1)
doc.add_paragraph('Release CQ_V2_2, clean-grain; LSTM, GRU và BiLSTM, mỗi model 3 window × 17 pipeline × 4 phase; seed model 42, batch 2.048. AccTEMPO là độ đo chính; S_perf, S_san+ v1 và các thành phần được trình bày tách biệt. Bảng TEST chỉ mô tả, không dùng để chọn model hay checkpoint. Chưa có kết luận độ ổn định giữa các seed.')
doc.add_paragraph('S_leak được tính trên parent imputer thực, fit TRAIN và đánh giá theo split/phase; các pipeline tăng cường dùng chung parent. S_san+ v2/s_cal là phân tích nhạy. STRICT/OVERLAP là lát mô tả, không thay tập TEST chính. Facts đã ingest vào L2 CQ với ingest_id vastai_cq_f1_batch2048_seed42_slice_r2_20261019; receipt và số partition đã được kiểm tra trên Modal.')

def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = 'Table Grid'
    for c, v in zip(t.rows[0].cells, headers): c.text = str(v)
    for row in rows:
        for c, v in zip(t.add_row().cells, row): c.text = str(v)
    for r in t.rows:
        for c in r.cells:
            for p in c.paragraphs:
                for run in p.runs: run.font.size = Pt(10)
    doc.add_paragraph('')

frames = {}; manifest_rows = []
for model, base in paths.items():
    info = json.loads((base / 'audit_manifest.json').read_text(encoding='utf-8'))
    if info['runs'] != 51 or info['bootstrap_repetitions'] != 2000 or info['release_id'] != 'CQ_V2_2': raise ValueError(f'Unexpected audit: {model}')
    fact = pd.read_parquet(base / 'sanity_components.parquet')
    if len(fact) != 408 or fact.duplicated(['window_id','pipeline_id','eval_split','phase_id']).any(): raise ValueError(f'Unexpected fact coverage: {model}')
    frames[model] = fact
    manifest_rows.append([model, info['runs'], info['rows']['sanity_components'], info['rows']['metrics_per_class'], info['rows']['bootstrap_metrics_per_class'], info['rows']['checkpoint_selection']])

doc.add_heading('2. Bao phủ audit', 1)
table(['Model', 'Run', 'Sanity', 'Per-class', 'Bootstrap', 'Checkpoint'], manifest_rows)
doc.add_paragraph('Mỗi model còn có 120 fact canonical S_leak và 1.224 fact đánh giá ALL/STRICT/OVERLAP. Bootstrap 2.000 lần; bootstrap được triển khai theo phương pháp được ghi trong từng fact, không mặc định là bootstrap theo enrollment trực tiếp.')
doc.add_paragraph('Các artifact gốc LSTM (766 file), GRU (714 file) và BiLSTM (714 file) đã được đối chiếu SHA-256 với VastAI, không phát hiện file thiếu, thừa hoặc sai hash. Audit slice r2 được tính lại từ prediction đã xác minh; toàn bộ 408 dòng ALL của mỗi model khớp bảng sanity cho các thành phần và AccTEMPO.')
doc.add_heading('3. Trung bình mô tả 12 cell TEST mỗi pipeline', 1)
doc.add_paragraph('Các cột dưới đây là trung bình không trọng số của 3 window × 4 phase. Chỉ dùng để định hướng; không thay thế bảng riêng W/P và không chọn model theo TEST.')
for model, f in frames.items():
    doc.add_heading(model, 2)
    x = f[f.eval_split.eq('TEST')].groupby('pipeline_id')[['acctempo_m3','s_perf','s_san_plus','macro_f1','balanced_acc']].mean()
    table(['V','AccTEMPO','S_perf','S_san+','Macro-F1','Balanced acc.'], [[f'V{i}']+[f'{x.loc[f"V{i}",c]:.4f}' for c in x.columns] for i in range(17)])

doc.add_heading('4. AccTEMPO TEST chi tiết từng window và phase', 1)
for model, f in frames.items():
    doc.add_heading(model, 2)
    for w in ('W1','W2','W3'):
        doc.add_heading(w, 3)
        q = f[(f.eval_split=='TEST') & (f.window_id==w)].pivot(index='pipeline_id', columns='phase_id', values='acctempo_m3')
        table(['Pipeline','P1','P2','P3','P4'], [[f'V{i}']+[f'{q.loc[f"V{i}",p]:.4f}' for p in ('P1','P2','P3','P4')] for i in range(17)])

doc.add_heading('5. VALIDATION và checkpoint', 1)
doc.add_paragraph('Checkpoint được lựa chọn theo validation macro-F1 trung bình P1–P4; không lựa chọn theo TEST. Bảng dưới là thống kê mô tả, không thay thế so sánh ghép cặp hay khoảng tin cậy giữa model.')
rows=[]
for model, f in frames.items():
    v=f[f.eval_split=='VALIDATION']; c=pd.read_parquet(paths[model]/'checkpoint_selection.parquet')
    rows.append([model, f"{v.acctempo_m3.mean():.4f}", f"{v.macro_f1.mean():.4f}", f"{c.best_epoch.min()}–{c.best_epoch.max()}",f"{c.n_checkpoints_within_margin.min()}–{c.n_checkpoints_within_margin.max()}"])
table(['Model','AccTEMPO TB','Macro-F1 TB','Best epoch min–max','Số checkpoint trong margin'],rows)
doc.add_heading('6. Thành phần sanity TEST theo pipeline', 1)
doc.add_paragraph('Trung bình không trọng số trên W1–W3 và P1–P4; s_nan là trường hiện có trong fact. Diễn giải sáu thành phần riêng, đặc biệt s_ent, thay vì chỉ nhìn AccTEMPO.')
components = ['s_nan', 's_maj_jsd', 's_ent', 's_drift', 's_eff', 's_leak']
for model, f in frames.items():
    doc.add_heading(model, 2)
    x = f[f.eval_split.eq('TEST')].groupby('pipeline_id')[components].mean()
    table(['V','s_nan','s_maj_jsd','s_ent','s_drift','s_eff','s_leak'], [[f'V{i}']+[f'{x.loc[f"V{i}",c]:.4f}' for c in components] for i in range(17)])
doc.add_heading('7. Recall từng lớp: TEST W1/P1', 1)
doc.add_paragraph('Ba cột tương ứng thứ tự nhãn warning, average, good. Đây là kết quả mô tả theo một lát cố định; support và CI chi tiết được lưu trong fact per-class và bootstrap.')
for model in paths:
    doc.add_heading(model, 2)
    p = pd.read_parquet(paths[model] / 'metrics_per_class.parquet')
    p = p[(p.eval_split=='TEST') & (p.window_id=='W1') & (p.phase_id=='P1')]
    q = p.pivot(index='pipeline_id', columns='class_code', values='recall')
    table(['V','warning','average','good'], [[f'V{i}']+[f'{q.loc[f"V{i}",c]:.4f}' for c in ('warning','average','good')] for i in range(17)])
doc.add_heading('8. STRICT/OVERLAP theo model', 1)
doc.add_paragraph('Giá trị trung bình không trọng số theo 3 window × 17 pipeline × 4 phase. s_eff và s_leak dùng parent thực của toàn eval split/phase; không fit lại probe theo lát. s_drift và các metric dự đoán được tính trên từng lát.')
rows=[]
for model, path in paths.items():
    f = pd.read_parquet(path/'evaluation_slices.parquet')
    f = f[f.eval_split.eq('TEST')]
    for sl in ('ALL','STRICT','OVERLAP'):
        x=f[f['slice'].eq(sl)]
        rows.append([model,sl,len(x),f'{x.acctempo_m3.mean():.4f}',f'{x.macro_f1.mean():.4f}',f'{x.balanced_acc.mean():.4f}'])
table(['Model','Slice','Cell','AccTEMPO','Macro-F1','Balanced acc.'], rows)
doc.add_heading('9. Giới hạn và bước tiếp theo', 1)
doc.add_paragraph('Đây là seed 42; chưa có bằng chứng đầy đủ cho seed 43/44, F2/F3, hoặc toàn bộ đối chiếu cross-backend. Các metric audit được tính từ artifacts VastAI đã đối chiếu hash và ingest vào L2; việc ingest không thay thế các kiểm định khoa học còn mở. Không diễn giải năng lượng ước lượng NVML GPU-only là năng lượng toàn hệ thống hoặc CO₂. Không suy ra nhân quả từ bảng trung bình.')
doc.add_paragraph('Nguồn: CQ/resuilt/cq_recurrent_batch2048_slice_r2_20261019/audits; mỗi thư mục audit có manifest, hash manifest nguồn và các bảng Parquet tương ứng. R1 cũ giữ riêng để truy vết; không dùng composite metric từ bảng slice R1.')
doc.save(OUT)
print(OUT)
