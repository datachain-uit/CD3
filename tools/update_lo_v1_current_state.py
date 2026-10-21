"""Refresh REPORT_LO_V1.docx in place with the current LO V2.2 status.

The early V1 process description is retained for provenance.  Content from
section 6 onward is replaced because its V1 split, imputation and downstream
run status is no longer the active experimental contract.
"""
from __future__ import annotations

from pathlib import Path

import polars as pl
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "LO" / "docs" / "REPORT_LO_V1.docx"
UPDATED_REPORT = ROOT / "LO" / "docs" / "REPORT_LO_V1_CAP_NHAT_V2_2_2026-09-30.docx"
R2 = ROOT / "LO" / "resuilt" / "split_registry_v2_2_overlap_audit_v1_r2"
LONG = ROOT / "LO" / "resuilt" / "long_offering_assignment_audit_v1"
SCORE = ROOT / "LO" / "resuilt" / "final_score_contract_audit_v1"
NAVY, BLUE, GREY = "17365D", "1F4E78", "F2F2F2"


def read(root: Path, name: str) -> pl.DataFrame:
    files = list((root / name).glob("*.parquet"))
    if not files:
        raise RuntimeError(f"Không tìm thấy parquet: {root / name}")
    return pl.read_parquet(max(files, key=lambda p: p.stat().st_mtime))


def shade(cell, color: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), color)
    props.append(node)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(14.5 if level == 1 else 12)
    run.font.color.rgb = RGBColor.from_string(NAVY)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, value in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(value)
        run.bold = True
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(255, 255, 255)
        shade(cell, BLUE)
    for r_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            cells[i].paragraphs[0].add_run(str(value)).font.size = Pt(10.5)
            if r_idx % 2:
                shade(cells[i], GREY)
    doc.add_paragraph()


def replace_paragraph(paragraph, text: str) -> None:
    paragraph.clear()
    paragraph.add_run(text)


def remove_from_paragraph(doc: Document, startswith: str) -> None:
    target = next((p for p in doc.paragraphs if p.text.strip().startswith(startswith)), None)
    if target is None:
        raise RuntimeError(f"Không tìm thấy mốc cập nhật: {startswith}")
    body = doc._element.body
    children = list(body)
    start = children.index(target._element)
    # Keep sectPr, which must remain the final body element.
    for child in children[start:]:
        if child.tag.endswith("}sectPr"):
            continue
        body.remove(child)


def pct(value: float) -> str:
    return f"{value * 100:.3f}%"


def slice_of(block: str) -> str:
    return "A" if block == "A" else "T1" if block in {"B", "C"} else "T2" if block in {"D", "E"} else "T3"


def main() -> None:
    doc = Document(REPORT)
    doc.core_properties.title = "Báo cáo quy trình LO — trạng thái V2.2 hiện hành"
    doc.core_properties.subject = "LO final_score, split V2.2 và điều kiện downstream"

    # Correct the title and the semantic statement in the preserved V1 provenance.
    replace_paragraph(doc.paragraphs[0], "BÁO CÁO QUY TRÌNH LO — TRẠNG THÁI V2.2 HIỆN HÀNH")
    replace_paragraph(doc.paragraphs[1], "Phạm vi: learner–course enrollment | feature/label/split V2.2 (hybrid) | cập nhật 30/09/2026")
    first_semantic = next(p for p in doc.paragraphs if p.text.startswith("Quy ước của dự án:"))
    replace_paragraph(
        first_semantic,
        "Quy ước hiện hành: nhãn LO là final_score — điểm tổng hợp từ các hoạt động được chấm theo trọng số course. "
        "Nó không phải final grade quan sát độc lập, completion hay dropout. Mọi thành phần tạo nhãn phải bị loại khỏi feature/imputer để tránh leakage.",
    )
    label_heading = next(p for p in doc.paragraphs if p.text.startswith("5. Gán nhãn LO"))
    replace_paragraph(label_heading, "5. Nhãn LO: final_score từ hoạt động được chấm")
    label_desc = next(p for p in doc.paragraphs if p.text.startswith("LO là operational behavioural"))
    replace_paragraph(
        label_desc,
        "LO dùng final_score hoạt động có trọng số, không phải final grade độc lập. Exam problem là problem ở chapter phân cấp cuối khi course có exam_weight; problem còn lại là assignment. "
        "Nhãn chỉ được dùng cho course có course_limit weights hợp lệ; không downsample cohort.",
    )
    # Refresh the plain-text contents list retained from the original report.
    contents = {
        "6. Tách phase P1–P4": "6. Split V2.2 hiện hành và phase P1–P4",
        "7. Chia temporal block": "7. Cohort LO, coverage và quy ước final_score",
        "8. Các bước tiếp theo": "8. Audit final_score, leakage contract và điều kiện chưa khóa",
    }
    for prefix, revised in contents.items():
        paragraph = next((p for p in doc.paragraphs if p.text.strip().startswith(prefix)), None)
        if paragraph is not None:
            replace_paragraph(paragraph, revised)

    # The former V1 section 6+ is obsolete as an active contract.
    remove_from_paragraph(doc, "6. Phase P1–P4 và Method 5")

    contract = read(SCORE, "label_semantics_contract").row(0, named=True)
    version = read(SCORE, "label_version_summary").row(0, named=True)
    exclusions = read(SCORE, "exclusion_reason_summary")
    profiles = read(SCORE, "score_range_by_profile")
    block = read(R2, "block_label_audit")
    registry = read(R2, "split_registry")
    qa = read(R2, "qa22_v2_2")
    label = read(R2, "split_label_distribution")
    actual = read(R2, "split_actual_ratios")
    duration = read(LONG, "duration_by_timeline")
    assignment = read(LONG, "enrollment_assignment_qa").row(0, named=True)
    strict = read(LONG, "strict_by_window_split_phase")

    add_heading(doc, "6. Split V2.2 hiện hành và phase P1–P4")
    doc.add_paragraph(
        "Split V1 và V2.1 chỉ còn là provenance, không được ghép vào meta-release hiện hành. Split chính là "
        "scenario=hybrid, split_version=v2.2, với CQ làm calendar anchor. Đơn vị gán là offering_id × timeline_source; "
        "một offering không được nằm ở nhiều split trong cùng window. P1/P2/P3/P4 là các prefix tích lũy; test Pj chỉ phơi lộ P1…Pj."
    )
    add_table(doc, ["Window", "TRAIN", "VALIDATION", "TEST", "Quy tắc"], [
        ["W1", "A", "B", "C", "A là lịch sử chung; B/C nằm trong lát tương lai T1."],
        ["W2", "A+B+C", "D", "E", "D/E nằm trong lát tương lai T2."],
        ["W3", "A+B+C+D+E", "F", "G", "F/G nằm trong lát tương lai T3."],
    ])
    qa_rows = [[r["window_id"], f"{r['n_enrollments_validation']:,}", f"{r['n_enrollments_test']:,}",
                "Đạt" if r["end_order_ok"] else "Không đạt", "Đạt" if r["qa22_ok"] else "Không đạt"]
               for r in qa.sort("window_id").iter_rows(named=True)]
    add_table(doc, ["Window", "Validation", "Test", "Label-end order", "QA22 V2.2"], qa_rows)

    add_heading(doc, "7. Cohort LO, coverage và quy ước final_score")
    total = int(version["enrollment_count"])
    eligible = int(version["eligible_enrollment_count"])
    add_table(doc, ["Hạng mục", "Giá trị"], [
        ["Raw enrollment", f"{total:,}"],
        ["Cohort LO eligible", f"{eligible:,} ({eligible / total * 100:.3f}% raw enrollment)"],
        ["Field nhãn", contract["label_field"]],
        ["Field score đang lưu", contract["persisted_score_field"]],
        ["Công thức", contract["formula"]],
        ["Ngưỡng lớp", contract["class_thresholds"]],
        ["Label rule version", version["label_rule_version"]],
    ])
    ex_rows = [[r["proxy_status"], f"{r['enrollment_count']:,}", f"{r['course_count']:,}", pct(r["enrollment_share"])]
               for r in exclusions.sort("enrollment_count", descending=True).iter_rows(named=True)]
    add_table(doc, ["Trạng thái raw", "Enrollment", "Course", "Tỷ lệ raw"], ex_rows)
    doc.add_paragraph(
        "Coverage 98.7–100% ở các slice A/T1/T2/T3 là coverage trong cohort LO sau điều kiện weights, không phải coverage toàn raw enrollment. "
        "Phần lớn raw enrollment bị loại vì thiếu course_limit weights."
    )

    add_heading(doc, "8. Audit final_score, leakage contract và điều kiện chưa khóa")
    score_gt_100 = int(profiles.select(pl.sum("n_final_score_gt_100")).item())
    profile_rows = [[r["assessment_modality_signature"], f"{r['eligible_enrollment_count']:,}",
                     f"{r['final_score_p99']:.2f}", f"{r['final_score_max']:.2f}", f"{r['n_final_score_gt_100']:,}"]
                    for r in profiles.sort("eligible_enrollment_count", descending=True).iter_rows(named=True)]
    add_table(doc, ["Assessment profile", "Eligible", "P99", "Max", "Score >100"], profile_rows)
    doc.add_paragraph(
        f"Cảnh báo cần xác nhận nghiệp vụ: {score_gt_100:,} enrollment ({score_gt_100 / eligible * 100:.3f}% cohort) có final_score >100; "
        "nguồn nằm ở profile có EXAM vì công thức đang dùng total_correct_ratio_exam. Nếu đây là weighted activity aggregate uncapped có chủ đích, "
        "giữ nhãn và nêu rõ quy ước. Nếu final_score phải thuộc 0–100, phải chuẩn hóa component exam, rebuild label và chạy lại split LO trước downstream."
    )
    add_table(doc, ["Proxy-label contract", "Quy tắc"], [
        ["Danh sách canonical", "final_score/performance_score, 3 weights, watch_percent, watched_videos/video_counts, assignment/exam attempt và ratios; kể cả derived fallback."],
        ["Feature dictionary", "Tất cả canonical proxy và biến thể _P1…_P4 phải is_label_proxy=1; loại khỏi imputer predictor và model input."],
        ["Artifact kiểm chứng", "final_score_contract_audit_v1/{label_semantics_contract, proxy_feature_contract, exclusion_reason_summary, score_range_by_profile}."],
    ])

    add_heading(doc, "9. Phân bố nhãn, drift và support")
    slices = []
    for name in ("A", "T1", "T2", "T3"):
        data = block.filter(pl.col("rolling_block").map_elements(slice_of, return_dtype=pl.String) == name)
        n = int(data["enrollment_count"].sum())
        vals = {key: int(data[key].sum()) for key in ("I/D_count", "G_count", "E_count")}
        slices.append([name, f"{n:,}", f"{vals['I/D_count']/n*100:.3f}%", f"{vals['G_count']/n*100:.3f}%", f"{vals['E_count']/n*100:.3f}%"])
    add_table(doc, ["Slice", "Enrollment", "I/D", "G", "E"], slices)
    split_rows = []
    for window in ("W1", "W2", "W3"):
        for arm in ("train", "validation", "test"):
            n = actual.filter((pl.col("window_id") == window) & (pl.col("split") == arm)).row(0, named=True)["enrollment_count"]
            values = {r["label"]: r for r in label.filter((pl.col("window_id") == window) & (pl.col("split") == arm)).iter_rows(named=True)}
            split_rows.append([window, arm, f"{n:,}", f"{values['I/D']['label_ratio']*100:.3f}%", f"{values['G']['label_ratio']*100:.3f}%", f"{values['E']['label_ratio']*100:.3f}%"])
    add_table(doc, ["W", "Split", "Enrollment", "I/D", "G", "E"], split_rows)
    doc.add_paragraph(
        "Label drift giữa T1/T2/T3 là thay đổi thật theo thời gian, không được san bằng giữa các lát. Metric per-class phải luôn kèm support. "
        "small_support_flag=1 khi n_support<400 cho từng window × phase × split × subgroup × class; cell nhỏ cần CI bootstrap 95%."
    )

    add_heading(doc, "10. Integrity assignment, strictness và offering dài")
    add_table(doc, ["Assignment QA", "Giá trị"], [
        ["Distinct enrollment", f"{assignment['distinct_enrollment_count']:,}"],
        ["Enrollment ở nhiều rolling block", str(assignment["n_enrollments_multiple_rolling_blocks"])],
        ["Enrollment ở nhiều split trong cùng window", str(assignment["n_enrollments_multiple_splits_same_window"])],
        ["Context không map rolling block", str(assignment["n_context_rows_without_block"])],
        ["Kết luận", "Đạt" if assignment["assignment_qa_pass"] == 1 else "Không đạt"],
    ])
    dur_rows = [[r["timeline_source"], str(r["duration_days_p50"]), str(r["duration_days_p90"]), str(r["duration_days_max"]), pct(r["long_offering_enrollment_share"])]
                for r in duration.sort("timeline_source").iter_rows(named=True)]
    add_table(doc, ["Timeline source", "Duration p50", "p90", "Max", "Enrollment >270 ngày"], dur_rows)
    p1_test = strict.filter((pl.col("split") == "test") & (pl.col("phase_id") == "P1")).sort("window_id")
    add_table(doc, ["TEST strict P1", "W1", "W2", "W3"], [["Tỷ lệ", *[pct(r["temporal_strict_share"]) for r in p1_test.iter_rows(named=True)]]])
    doc.add_paragraph(
        "timeline_source và long_offering_flag phải đi vào prediction-level context để báo cáo metric phân tầng. Không loại offering dài mặc định và không tạo subgroup bắt buộc chỉ vì duration dài."
    )

    add_heading(doc, "11. Trạng thái downstream và lộ trình thực nghiệm LO")
    add_table(doc, ["Hạng mục", "Trạng thái hiện tại", "Điều kiện/ghi chú"], [
        ["Split registry V2.2", "Đạt", "Primary split; run phải lưu scenario, split_version, parent manifest/audit, code/config hash và seed."],
        ["Materialize V2.2 + test prefixes", "Được chạy", "Đọc registry V2.2; giữ timeline_source, duration_days, long_offering_flag."],
        ["S0 data quality", "Được chạy", "Báo eligible missing, structural absence, event sparsity theo slice × timeline source; không điền structural/future cells."],
        ["Feature dictionary/imputation", "Chờ quyết định final_score", "Khóa proxy-label contract trước; fit train-only, validation/test transform-only."],
        ["Augmentation/model", "Chờ dictionary LO", "Không augment validation/test; metric tổng thể + strict/context/per-class cells cùng CI/support."],
    ])
    doc.add_paragraph(
        "Kết luận: split V2.2, assignment integrity và audit cohort đã đủ để LO materialize và S0. Cổng duy nhất chưa thể bỏ qua trước khi khóa feature/imputation/model là xác nhận cách diễn giải và miền giá trị final_score, đặc biệt component exam."
    )

    # Footer should identify that this formerly V1-named file now contains current status.
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.clear()
    footer.add_run("LO · Báo cáo quy trình cập nhật V2.2 · 30/09/2026").font.size = Pt(8.5)
    try:
        doc.save(REPORT)
        print(REPORT)
    except PermissionError:
        doc.save(UPDATED_REPORT)
        print(UPDATED_REPORT)


if __name__ == "__main__":
    main()
