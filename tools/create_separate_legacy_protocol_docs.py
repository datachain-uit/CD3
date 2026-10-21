"""Create separate, non-current CQ V1 and LO V1 protocol snapshots."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]


def setup(title: str, subtitle: str) -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.65)
    sec.bottom_margin = Inches(0.65)
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(10.5)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title); r.bold = True; r.font.size = Pt(18)
    doc.add_paragraph(subtitle)
    return doc


def heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = True; r.font.size = Pt(13)


def table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    out = doc.add_table(rows=1, cols=len(headers)); out.style = "Table Grid"
    for cell, value in zip(out.rows[0].cells, headers): cell.text = value
    for values in rows:
        for cell, value in zip(out.add_row().cells, values): cell.text = value


# CQ V1
cq = setup(
    "GIAO THỨC LEGACY — CQ V1",
    "Snapshot historical/reference. Không dùng để chạy thực nghiệm hoặc gộp meta-release hiện hành; CQ hiện hành dùng split V2.2.",
)
heading(cq, "1. Định danh và artifact")
table(cq, ["Trường", "CQ V1"], [
    ["Task / label", "CQ / CQ_label_final"],
    ["Label rule", "cq_vector_proximity_v1"],
    ["Label artifact", "labels/cq_labels_v1"],
    ["Feature regime", "CQ_RAW_EARLY"],
    ["Scenario", "hybrid"],
    ["Phase", "P1=25%, P2=50%, P3=75%, P4=90%"],
    ["Split / phase version", "label_stratified_v1 / wide_prefix_v1"],
    ["Views", "tasks/CQ/hybrid/phase_views_v1; test_prefix_views_v1"],
])
heading(cq, "2. Cách chia Method-5 cũ")
cq.add_paragraph("A gồm các offering sớm nhất theo offering_end_date ở ranh giới whole-offering, đạt xấp xỉ 60% enrollment. Các offering còn lại phân bổ B–G bằng allocator deterministic, count-aware để cân bằng số enrollment và nhãn; B–G không được xếp theo thứ tự thời gian.")
table(cq, ["Window", "TRAIN", "VALIDATION", "TEST"], [
    ["W1", "A", "B", "C"], ["W2", "A+B+C", "D", "E"], ["W3", "A+B+C+D+E", "F", "G"],
])
heading(cq, "3. Hạn chế và quyết định thay thế")
table(cq, ["Điểm", "Kết luận historical"], [
    ["W1", "Thiết kế A historical, B/C future-pool; phù hợp thứ tự lịch ở mức mô tả."],
    ["W2–W3", "B–G không phải lát thời gian liên tiếp; không chứng minh strict temporal ordering giữa TRAIN và evaluation."],
    ["Registry", "Chưa có đầy đủ QA22: boundaries, purge/embargo, overlap và drift metadata."],
    ["Trạng thái", "Retired; thay bằng CQ split_version=v2_2. Run V1 chỉ giữ provenance."],
])
cq.add_paragraph("Nguồn: CQ/task.yaml legacy, CQ Feature_extraction config V1 và materialize_label_stratified_views.py.")
cq.save(ROOT / "CQ" / "docs" / "REPORT_CQ_PROTOCOL_V1_LEGACY.docx")

# LO V1
lo = setup(
    "GIAO THỨC LEGACY — LO V1",
    "Snapshot historical/reference. Không dùng cho model/meta-release hiện hành; LO hiện hành dùng nhãn catalog-normalized và split/view V3.1.",
)
heading(lo, "1. Định danh và artifact")
table(lo, ["Trường", "LO V1"], [
    ["Task / label", "LO operational / LO_performance_label_3"],
    ["Label rule", "lo_operational_performance_v1"],
    ["Label artifact", "labels/lo_labels_v1"],
    ["Ý nghĩa label", "behavioural performance proxy; không phải final grade, completion hoặc dropout quan sát trực tiếp"],
    ["Feature regime", "LO_FULL_EARLY"],
    ["Scenario", "hybrid"],
    ["Phase", "P1=25%, P2=50%, P3=75%, P4=90%"],
    ["Split / phase version", "label_stratified_v1 / wide_prefix_v1"],
    ["Views", "tasks/LO/hybrid/phase_views_v1; test_prefix_views_v1"],
])
heading(lo, "2. Cách chia Method-5 cũ")
lo.add_paragraph("A lấy khoảng 60% enrollment đầu theo offering_end_date tại whole-offering boundaries. Later offerings được allocator deterministic phân vào B–G theo số enrollment và số nhãn G/E để giữ cân bằng; các block B–G không là những lát lịch liên tiếp.")
table(lo, ["Window", "TRAIN", "VALIDATION", "TEST"], [
    ["W1", "A", "B", "C"], ["W2", "A+B+C", "D", "E"], ["W3", "A+B+C+D+E", "F", "G"],
])
heading(lo, "3. Lịch sử tài liệu LO còn lưu")
table(lo, ["Tài liệu", "Vai trò"], [
    ["REPORT_LO_V1_CAP_NHAT_V2_2_2026-09-30.docx", "Bản report LO V1 được cập nhật tại thời điểm chuyển V2.2"],
    ["REPORT_LO_V2.2.docx", "Đánh giá chuyển split V2.2"],
    ["REPORT_LO_V3.docx", "Đánh giá LO catalog-normalized V3"],
    ["REPORT_LO_V3_1.docx", "Report chính thức hiện hành, có V3.1 và model facts"],
])
heading(lo, "4. Hạn chế và quyết định thay thế")
table(lo, ["Điểm", "Kết luận historical"], [
    ["Temporal", "W2/W3 không có strict temporal separation vì B–G không chronological."],
    ["Label", "Proxy behavioural V1 bị thay bởi final-score catalog-normalized V3.1."],
    ["QA registry", "Thiếu calendar/overlap/purge/label-drift contract đầy đủ."],
    ["Trạng thái", "Retired; giữ để giải thích provenance, không gộp với V3.1."],
])
lo.add_paragraph("Nguồn: protocol config V1, materialize_label_stratified_views.py và các report LO lịch sử còn lưu trong LO/docs.")
lo.save(ROOT / "LO" / "docs" / "REPORT_LO_PROTOCOL_V1_LEGACY.docx")

print(ROOT / "CQ" / "docs" / "REPORT_CQ_PROTOCOL_V1_LEGACY.docx")
print(ROOT / "LO" / "docs" / "REPORT_LO_PROTOCOL_V1_LEGACY.docx")
