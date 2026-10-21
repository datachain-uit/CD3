"""Keep only the canonical current storage tree in CQ/LO reports (no gap notes)."""
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]


def remove_from_marker_to_end(doc: Document, marker: str) -> None:
    body = doc.element.body
    children = list(body)
    start = None
    for index, child in enumerate(children):
        if child.tag == qn("w:p") and marker in "".join(child.itertext()):
            start = index
            break
    if start is not None:
        # The compliance block was appended last; remove it without touching
        # any earlier report content or the terminal section-properties node.
        for child in children[start:]:
            if child.tag != qn("w:sectPr"):
                body.remove(child)


def tree(task: str, regime: str, input_root: str, views: str, prefixes: str, label: str) -> str:
    return f'''tempo-data-v1/
├── input/                                          # immutable experiment source
│   └── {input_root}/
│       ├── {views}/                                # TRAIN + VALIDATION; full P1–P4
│       ├── {prefixes}/P1 ... P4/                   # TEST prefix views
│       └── release_manifest.json                   # data/split/phase/schema/label provenance
│
└── meta_release=imputation-v1/
    ├── L0_registry/                                # definitions, locks and provenance
    │   ├── data_release_{task}.json
    │   ├── split_registry.json  pipeline_registry.json  metric_definitions.json
    │   ├── state_manifests/s0_{task}.json
    │   ├── feature_dictionary/task={task}/window_id=W{{1..3}}/feature_dictionary.parquet
    │   ├── run_registry/run_id={{run_id}}.json
    │   ├── model_registry/model_name=RNN/model_revision={{revision}}.json
    │   └── environments/{{s0_{task},attempt_id={{attempt_id}}}}.json
    │
    ├── L1_runs/                                    # immutable executable artifacts
    │   └── task={task}/feature_regime={regime}/window_id=W1/
    │       ├── pipeline_id=V0/model_name=IMPUTATION_ONLY/seed=20260922/
    │       │   └── run_id={{imputation_sha}}/attempt_id={{attempt}}-NNNN/
    │       │       ├── run_manifest.json  fitted_preprocessors/
    │       │       └── model_inputs/{{train,validation,test_P1..test_P4}}.parquet
    │       ├── phase_id=P4/pipeline_id=V2/model_name=BALANCED_TRAIN_ONLY/seed=42/
    │       │   └── run_id={{augmentation_sha}}/attempt_id={{attempt}}-NNNN/
    │       │       ├── run_manifest.json  balanced_train.parquet
    │       │       ├── synthetic_ledger.parquet  augmentation_quality.json
    │       │       └── resource/timing artifact
    │       └── pipeline_id=V0/model_name=RNN/seed=42/
    │           └── run_id={{model_sha}}/attempt_id={{attempt}}-NNNN/
    │               ├── run_manifest.json  checkpoint/model.pt  train_history.parquet
    │               ├── feature_layout.json  timing.parquet
    │               ├── predictions/eval_split={{VALIDATION,TEST}}/phase_id=P{{1..4}}/
    │               ├── probability_matrices/eval_split={{VALIDATION,TEST}}/phase_id=P{{1..4}}/
    │               └── confusion_matrices/eval_split={{VALIDATION,TEST}}/phase_id=P{{1..4}}/
    │
    ├── L2_facts/                                   # queryable facts, partitioned by run/phase
    │   ├── bin_scheme/  dq_feature_profile/  dataset_quality_summary/
    │   ├── distribution_sketch/  class_distribution/  drift_measures/
    │   ├── imputation_fidelity/  augmentation_quality/  resource_usage/
    │   ├── metrics_overall/  metrics_per_class/  calibration/  confusion_matrix/
    │   ├── checkpoint_selection/  sanity_components/  bootstrap_metrics_per_class/
    │   ├── report_exports/task={task}/report_id={{report_id}}/
    │   └── .../task={task}/feature_regime={regime}/window_id=W1/
    │       └── phase_id=P1/.../run_id={{...}}/attempt_id={{...}}/part-00000.parquet
    │
    └── L3_meta/                                    # aggregated meta-dataset layer
        └── task={task}/label={label}/...'''


def update(path: Path, *, task: str, regime: str, input_root: str, views: str, prefixes: str, label: str) -> None:
    doc = Document(path)
    remove_from_marker_to_end(doc, "ĐỐI CHIẾU CẤU TRÚC VỚI META-DATASET V1.6")
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("tempo-data-v1/"):
            paragraph.clear()
            run = paragraph.add_run(tree(task, regime, input_root, views, prefixes, label))
            run.font.name = "Consolas"; run.font.size = Pt(8.5)
            break
    else:
        raise RuntimeError(f"Current storage tree not found in {path}")
    doc.save(path)


update(ROOT / "CQ" / "docs" / "REPORT_CQ_V2.2.docx", task="CQ", regime="CQ_RAW_EARLY",
       input_root="CQ_v2_2", views="phase_views_v2_2", prefixes="test_prefix_views_v2_2", label="CQ_label_final")
update(ROOT / "LO" / "docs" / "REPORT_LO_V3_1.docx", task="LO", regime="LO_FULL_EARLY",
       input_root="LO_v3_1", views="phase_views_v3_1_scored_signal_excluded",
       prefixes="test_prefix_views_v3_1_scored_signal_excluded", label="LO_performance_label_3")
print("current report storage trees normalized")
