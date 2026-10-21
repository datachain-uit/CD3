"""Append the evidence-backed Split V2.1 status to the CQ and LO reports."""
from __future__ import annotations

from pathlib import Path

import polars as pl
from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
MARKER = "Cập nhật Split V2.1 — registry khóa ngày 28/09/2026"


def parquet(root: Path, name: str) -> pl.DataFrame:
    return pl.read_parquet(next((root / name).glob("*.parquet")))


def remove_previous_update(doc: Document) -> None:
    start = None
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(MARKER):
            start = paragraph._element
            break
    if start is None:
        return
    deleting = False
    for element in list(doc._element.body):
        if element is start:
            deleting = True
        if deleting and element.tag.rsplit("}", 1)[-1] != "sectPr":
            doc._element.body.remove(element)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(9)
    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)


def add_heading(doc: Document, text: str, level: int) -> None:
    """Use direct formatting because the supplied reports omit Heading styles."""
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(16 if level == 1 else 12)
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(5)


def pct(value: float | None) -> str:
    return "—" if value is None else f"{value * 100:.2f}%"


def update(task: str, report_path: Path, registry_root: Path) -> None:
    qa = parquet(registry_root, "qa22_v2_1")
    reg = parquet(registry_root, "split_registry")
    ratios = parquet(registry_root, "split_actual_ratios")
    labels = parquet(registry_root, "split_label_distribution")
    context = parquet(registry_root, "temporal_strict_context")
    durations = parquet(registry_root, "duration_summary")

    doc = Document(report_path)
    remove_previous_update(doc)
    doc.add_page_break()
    add_heading(doc, MARKER, level=1)
    doc.add_paragraph(
        "Phạm vi cập nhật: registry Split V2.1 của " + task +
        ". Đây là bằng chứng cho thiết kế, phân tầng và kiểm định split; "
        "phase views V2.1, S0, imputation, augmentation và model chưa được chạy lại trên release này."
    )

    add_heading(doc, "1. Quy tắc Split V2.1", level=2)
    doc.add_paragraph(
        "Đơn vị gán là offering (kèm timeline_source), do đó một offering không bị chia "
        "giữa TRAIN, VALIDATION và TEST trong cùng window. Offering được sắp theo "
        "offering_end_date: block A là lịch sử sớm; phần sau được cắt thành ba lát thời gian "
        "liên tiếp và chỉ VALIDATION/TEST trong cùng lát được phân tầng theo nhãn và quy mô."
    )
    doc.add_paragraph("W1: TRAIN=A; VALIDATION=B; TEST=C.  W2: TRAIN=A+B+C; VALIDATION=D; TEST=E.  W3: TRAIN=A+B+C+D+E; VALIDATION=F; TEST=G.")
    doc.add_paragraph(
        "Khác với V2 strict-purge đã loại quá nhiều evaluation rows, V2.1 dùng purge_mode=NONE. "
        "Overlap được đo và báo cáo, không bị loại. QA22 chỉ khóa registry khi thứ tự end-date "
        "giữa TRAIN và evaluation đúng, evaluation không rỗng và vượt ngưỡng tối thiểu."
    )

    add_heading(doc, "2. QA22 và quy mô thực tế", level=2)
    qrows = []
    for row in qa.sort("window_id").to_dicts():
        qrows.append([
            row["window_id"], f"{row['n_enrollments_validation']:,}",
            f"{row['n_enrollments_test']:,}", f"{row['min_eval_enrollments']:,}",
            "PASS" if row["qa22_ok"] else "FAIL", pct(row["ev_train_overlap_share"]),
            str(row["days_train_end_to_test_start"]), str(row["train_history_span_days"]),
        ])
    add_table(doc,
        ["Window", "Validation", "Test", "Ngưỡng eval", "QA22", "Overlap eval–train", "Gap (ngày)", "History (ngày)"],
        qrows)
    doc.add_paragraph(
        "Cả ba window đạt end_order_ok và eval_size_ok. Gap âm phản ánh các khóa học diễn ra "
        "đồng thời; theo V2.1 đây là thống kê overlap, không phải điều kiện loại dữ liệu."
    )

    pivot = ratios.pivot(on="split", index="window_id", values="split_enrollment_share", aggregate_function="first").sort("window_id")
    share_rows = [[r["window_id"], pct(r.get("train")), pct(r.get("validation")), pct(r.get("test"))] for r in pivot.to_dicts()]
    add_table(doc, ["Window", "TRAIN", "VALIDATION", "TEST"], share_rows)

    add_heading(doc, "3. Registry metadata, purge và drift", level=2)
    rrows = []
    for row in reg.sort("window_id").to_dicts():
        drift = row["label_drift_window_prev_jsd"]
        rrows.append([
            row["window_id"], row["purge_mode"], str(row["n_purged"]),
            pct(row["purged_enrollment_share"]), "Có" if row["purge_gt_half_eval"] else "Không",
            "—" if drift is None else f"{drift:.4f}",
        ])
    add_table(doc, ["Window", "Purge mode", "n_purged", "Tỷ lệ purge", "Purge >50% eval", "Label drift JSD"], rrows)
    doc.add_paragraph(
        "Metadata khóa: scenario=hybrid; split_version=v2.1. Label drift là thay đổi phân phối "
        "nhãn thật giữa các lát thời gian và không bị san bằng qua các lát."
    )

    add_heading(doc, "4. Phân tầng temporal-strict và overlap", level=2)
    strict_rows = []
    for split in ("validation", "test"):
        frame = context.filter(pl.col("split") == split)
        strict_rows.append([split.upper(), *[pct(frame.select(pl.col(f"temporal_strict_P{i}").mean()).item()) for i in range(1, 5)]])
    add_table(doc, ["Evaluation split", "P1 strict", "P2 strict", "P3 strict", "P4 strict"], strict_rows)
    doc.add_paragraph(
        "temporal_strict_Pj=1 khi cutoff của pha Pj không sớm hơn tau_w (thời điểm nhãn "
        "TRAIN cuối cùng đã biết); các hàng còn lại là OVERLAP. Khi có model V2.1, mọi metric "
        "phải được báo cáo chung và tách riêng STRICT/OVERLAP."
    )

    add_heading(doc, "5. Hỗ trợ theo lớp và phạm vi diễn giải", level=2)
    minimum = (labels.filter(pl.col("split").is_in(["validation", "test"]))
        .group_by("label").agg(pl.col("label_enrollment_count").min().alias("minimum"))
        .sort("label"))
    add_table(doc, ["Nhãn", "Số enrollment nhỏ nhất trong một validation/test"],
              [[row["label"], f"{row['minimum']:,}"] for row in minimum.to_dicts()])
    if task == "LO":
        doc.add_paragraph(
            "Cảnh báo: LO có support rất nhỏ cho lớp E và G trong một số evaluation split. "
            "Có thể báo cáo metric tổng thể, nhưng per-class precision/recall/F1 của E/G phải "
            "kèm support và khoảng tin cậy; không được diễn giải là ước lượng ổn định."
        )
    else:
        doc.add_paragraph(
            "CQ có đủ support tuyệt đối cho các lớp evaluation, nhưng vẫn rất mất cân bằng. "
            "Accuracy không phải metric chính; cần macro-F1, balanced accuracy và metric per-class."
        )

    add_heading(doc, "6. Timeline coverage và trạng thái downstream", level=2)
    drows = [[r["timeline_source"], f"{r['offering_count']:,}", f"{r['enrollment_count']:,}",
              str(r["duration_days_p50"]), str(r["duration_days_p90"]), pct(r["end_at_dataset_last_date_share"])]
             for r in durations.to_dicts()]
    add_table(doc, ["Timeline source", "Offerings", "Enrollments", "Duration P50", "Duration P90", "End at data boundary"], drows)
    doc.add_paragraph(
        "Registry V2.1 đã khóa thành công. Tuy nhiên chưa có phase_views_v2_1/test_prefix_views_v2_1, "
        "nên mọi data-quality, imputation, augmentation và model result đã có từ V1 không được gộp "
        "hay diễn giải là kết quả V2.1. Những thí nghiệm đó chỉ bắt đầu sau khi materialize view thành công."
    )

    doc.save(report_path)


if __name__ == "__main__":
    update("CQ", ROOT / "CQ/docs/REPORT_CQ_V1.docx", ROOT / "CQ/resuilt/split_registry_v2_1")
    update("LO", ROOT / "LO/docs/REPORT_LO_V1.docx", ROOT / "LO/resuilt/split_registry_v2_1")
