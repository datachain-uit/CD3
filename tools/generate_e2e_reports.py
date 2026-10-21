"""Generate evidence-bound end-to-end Vietnamese reports for CQ and LO.

The script reads only local immutable/parquet evidence.  It deliberately keeps
execution batches separate when their runner provenance differs.
"""
from __future__ import annotations

import json
from datetime import date
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
TODAY = "2026-10-16"
PIPES = [f"V{i}" for i in range(17)]
PIPELINE_LABELS = {
    "V0": "zero-fill + masks",
    "V1": "mean", "V5": "median", "V9": "ExtraTrees", "V13": "MICE",
    "V2": "mean + CDSMOTE", "V3": "mean + RadiusSMOTE", "V4": "mean + SASMOTE",
    "V6": "median + CDSMOTE", "V7": "median + RadiusSMOTE", "V8": "median + SASMOTE",
    "V10": "ExtraTrees + CDSMOTE", "V11": "ExtraTrees + RadiusSMOTE", "V12": "ExtraTrees + SASMOTE",
    "V14": "MICE + CDSMOTE", "V15": "MICE + RadiusSMOTE", "V16": "MICE + SASMOTE",
}


def read_many(paths: list[Path]) -> pd.DataFrame:
    return pd.concat([pd.read_parquet(p) for p in paths], ignore_index=True) if paths else pd.DataFrame()


def fmt(x, digits=4):
    if x is None or pd.isna(x):
        return "—"
    if isinstance(x, (float, np.floating)):
        return f"{x:.{digits}f}"
    return str(x)


def md_table(df: pd.DataFrame) -> str:
    if df.empty:
        return "*Không có dòng dữ liệu.*\n"
    cols = list(df.columns)
    lines = ["| " + " | ".join(cols) + " |", "|" + "|".join(["---"] * len(cols)) + "|"]
    for row in df.itertuples(index=False, name=None):
        lines.append("| " + " | ".join(str(x).replace("|", "\\|") for x in row) + " |")
    return "\n".join(lines) + "\n"


def add_table(doc: Document, df: pd.DataFrame):
    if df.empty:
        doc.add_paragraph("Không có dòng dữ liệu.")
        return
    table = doc.add_table(rows=1, cols=len(df.columns))
    table.style = "Table Grid"
    table.autofit = True
    for cell, col in zip(table.rows[0].cells, df.columns):
        cell.text = str(col)
    for row in df.itertuples(index=False, name=None):
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = str(value)
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9.5)
    for cell in table.rows[0].cells:
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.bold = True
                run.font.size = Pt(10)


def setup_doc(title: str) -> Document:
    d = Document()
    sec = d.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Inches(11.69), Inches(8.27)
    sec.left_margin = sec.right_margin = Inches(0.45)
    style = d.styles["Normal"]
    style.font.name = "Aptos"
    style.font.size = Pt(9)
    p = d.add_heading(title, 0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.add_paragraph(f"Ngày tổng hợp: {TODAY}. Báo cáo được sinh trực tiếp từ artefact cục bộ và facts parquet; không suy diễn số liệu thiếu.")
    return d


def render(out: Path, title: str, sections: list[tuple[str, str, list[pd.DataFrame]]]):
    md = [f"# {title}", "", f"Ngày tổng hợp: {TODAY}.", ""]
    doc = setup_doc(title)
    for heading, text, tables in sections:
        md += [f"## {heading}", "", text, ""]
        doc.add_heading(heading, 1)
        doc.add_paragraph(text)
        for table in tables:
            md.append(md_table(table)); md.append("")
            add_table(doc, table)
            doc.add_paragraph("")
    out.with_suffix(".md").write_text("\n".join(md), encoding="utf-8")
    doc.save(out.with_suffix(".docx"))


def performance(metrics: pd.DataFrame, sanity: pd.DataFrame, split: str) -> pd.DataFrame:
    m = metrics.copy()
    s = sanity.copy()
    if "eval_split" in m: m = m[m.eval_split.eq(split)]
    if "eval_split" in s: s = s[s.eval_split.eq(split)]
    keys = ["model_name", "pipeline_id"]
    mcols = ["accuracy", "balanced_accuracy", "f1_macro", "f1_weighted", "pr_auc_macro", "roc_auc_macro_ovr", "multiclass_nll", "multiclass_brier", "top_label_ece_15"]
    scols = ["acctempo_m3", "s_perf", "s_san_plus", "s_nan", "s_maj_jsd", "s_ent", "s_drift", "s_eff", "s_leak", "s_cal", "s_san_plus_v2"]
    mg = m.groupby(keys, as_index=False)[[c for c in mcols if c in m]].mean(numeric_only=True)
    sg = s.groupby(keys, as_index=False)[[c for c in scols if c in s]].mean(numeric_only=True)
    out = mg.merge(sg, on=keys, how="outer")
    out["pipeline"] = out.pipeline_id.map(lambda x: f"{x} — {PIPELINE_LABELS.get(x, x)}")
    ordered = ["model_name", "pipeline"] + [c for c in ["acctempo_m3", "f1_macro", "balanced_accuracy", "accuracy", "pr_auc_macro", "roc_auc_macro_ovr", "multiclass_nll", "multiclass_brier", "top_label_ece_15", "s_perf", "s_san_plus", "s_nan", "s_maj_jsd", "s_ent", "s_drift", "s_eff", "s_leak", "s_cal", "s_san_plus_v2"] if c in out]
    out = out[ordered].sort_values(["model_name", "pipeline"])
    for c in out.columns[2:]: out[c] = out[c].map(fmt)
    return out


def family_summary(metrics: pd.DataFrame, sanity: pd.DataFrame, split: str) -> pd.DataFrame:
    m = metrics[metrics.eval_split.eq(split)]
    s = sanity[sanity.eval_split.eq(split)]
    a = m.groupby("model_name", as_index=False)[[c for c in ["accuracy", "balanced_accuracy", "f1_macro", "pr_auc_macro", "roc_auc_macro_ovr"] if c in m]].mean(numeric_only=True)
    b = s.groupby("model_name", as_index=False)[[c for c in ["acctempo_m3", "s_perf", "s_san_plus", "s_san_plus_v2"] if c in s]].mean(numeric_only=True)
    r = a.merge(b, on="model_name", how="outer")
    for c in r.columns[1:]: r[c] = r[c].map(fmt)
    return r


def per_class_summary(per_class: pd.DataFrame, split: str = "TEST") -> pd.DataFrame:
    d = per_class[per_class.eval_split.eq(split)].copy()
    cols = [c for c in ["n_support", "precision", "recall", "f1", "average_precision", "small_support_flag"] if c in d]
    out = d.groupby(["model_name", "class_code"], as_index=False)[cols].mean(numeric_only=True)
    if "n_support" in out: out["n_support"] = out["n_support"].round().astype("Int64")
    if "small_support_flag" in out: out["small_support_flag"] = (out["small_support_flag"] * 100).map(lambda x: fmt(x, 1) + "%")
    for c in ["precision", "recall", "f1", "average_precision"]:
        if c in out: out[c] = out[c].map(fmt)
    return out


def slice_summary(slices: pd.DataFrame) -> pd.DataFrame:
    d = slices.copy()
    cols = [c for c in ["n_eval_samples", "macro_f1", "balanced_acc", "acctempo_m3", "s_perf", "s_san_plus"] if c in d]
    out = d.groupby(["model_name", "eval_split", "slice"], as_index=False)[cols].mean(numeric_only=True)
    if "n_eval_samples" in out: out["n_eval_samples"] = out["n_eval_samples"].round().astype("Int64")
    for c in ["macro_f1", "balanced_acc", "acctempo_m3", "s_perf", "s_san_plus"]:
        if c in out: out[c] = out[c].map(fmt)
    return out


def cq_data():
    root = ROOT / "CQ/resuilt/vastai_f1_l2_exports_20261016"
    folders = [root / "lstm_gru", root / "bilstm"]
    frames = {}
    for n in ["metrics_test_long", "metrics_validation_long", "sanity_test_long", "sanity_validation_long", "metrics_per_class_long", "bootstrap_metrics_per_class_long", "strict_overlap_metrics_long", "checkpoint_selection_long", "canonical_s_leak_long", "resource_long"]:
        frames[n] = read_many([f / f"{n}.parquet" for f in folders])
    return frames


def legacy_cq_rnn_summary() -> pd.DataFrame:
    """Read the preserved provenance-strict RNN report rather than retype values."""
    p = ROOT / "CQ/docs/REPORT_CQ_V2.2_CLEAN_GRAIN_RNN_SEED42_V2_2026-10-07.docx"
    if not p.exists():
        return pd.DataFrame()
    d = Document(p)
    # Table 4 in the frozen report is the full 17-pipeline RNN TEST table.
    t = d.tables[4]
    rows = [[c.text.strip() for c in r.cells] for r in t.rows]
    return pd.DataFrame(rows[1:], columns=rows[0])


def lo_data():
    root = ROOT / "LO/resuilt/vastai_f2_batch2048_seed42_20261013"
    metrics = read_many(list(root.glob("**/metrics_overall.parquet")))
    audit_root = root / "audits"
    data = {"metrics": metrics}
    for name in ["sanity_components", "metrics_per_class", "bootstrap_metrics_per_class", "evaluation_slices", "checkpoint_selection", "canonical_s_leak"]:
        data[name] = read_many(list(audit_root.glob(f"**/{name}.parquet")))
    manifests = []
    for p in root.glob("**/run_manifest.json"):
        d = json.loads(p.read_text(encoding="utf-8"))
        manifests.append({k: d.get(k) for k in ["task", "window_id", "pipeline_id", "model_name", "seed", "execution_backend", "gpu", "batch_size", "max_epochs", "best_epoch", "elapsed_seconds", "peak_ram_mb", "precision_policy", "energy_measurement"]})
    data["resource"] = pd.DataFrame(manifests)
    eq = root / "modal_vast_equivalence_LO_W1_V0_RNN_seed42.json"
    data["equivalence"] = json.loads(eq.read_text(encoding="utf-8")) if eq.exists() else {}
    audit_root = ROOT / "LO/resuilt/audit_prefix_validation_alignment_v1-20261016T105149Z-1-001/audit_prefix_validation_alignment_v1"
    if audit_root.is_dir():
        data["prefix_audit_summary"] = json.loads((audit_root / "summary.json").read_text(encoding="utf-8"))
        for name in ["w1_validation_test_p1_p4_summary", "w1_p1_seen_unseen_strict_overlap_metrics", "w1_p1_course_cohort_label_support"]:
            data[name] = pd.read_csv(audit_root / f"{name}.csv")
    else:
        data["prefix_audit_summary"] = {}
    return data


def fact_counts(items: dict[str, pd.DataFrame], labels: dict[str, str]) -> pd.DataFrame:
    rows=[]
    for k, label in labels.items():
        d=items.get(k, pd.DataFrame())
        rows.append({"Fact":label,"Số dòng":len(d),"Khóa/tính bao phủ":"—"})
    return pd.DataFrame(rows)


def main():
    cq=cq_data()
    cq_metrics=pd.concat([cq["metrics_test_long"], cq["metrics_validation_long"]], ignore_index=True)
    cq_sanity=pd.concat([cq["sanity_test_long"], cq["sanity_validation_long"]], ignore_index=True)
    cq_counts=fact_counts(cq, {"metrics_test_long":"Performance TEST", "metrics_validation_long":"Performance VALIDATION", "sanity_test_long":"Sanity TEST", "sanity_validation_long":"Sanity VALIDATION", "metrics_per_class_long":"Per-class", "bootstrap_metrics_per_class_long":"Bootstrap per-class (B=2,000)", "strict_overlap_metrics_long":"STRICT/OVERLAP", "checkpoint_selection_long":"Checkpoint selection", "canonical_s_leak_long":"Canonical S_leak", "resource_long":"Resource"})
    legacy_rnn = legacy_cq_rnn_summary()
    cq_resource=cq["resource_long"].groupby(["model_name","gpu","batch_size"],dropna=False).agg(runs=("run_id","nunique"), epochs_mean=("epochs_run","mean"), wall_s=("elapsed_seconds","mean"), ram_peak_mb=("peak_ram_mb","max"), energy_records=("energy_gpu_estimated_kwh",lambda x:x.notna().sum())).reset_index()
    for c in ["epochs_mean","wall_s","ram_peak_mb"]: cq_resource[c]=cq_resource[c].map(lambda x:fmt(x,1))
    cq_runprov=cq["resource_long"].groupby(["model_name","attempt_id","runner_sha256","execution_backend"],dropna=False).agg(runs=("run_id","nunique")).reset_index()
    cq_sections=[
      ("1. Phạm vi, release và quy tắc đọc kết quả", "CQ dùng release CQ_V2_2, split v2_2 và feature phase prefix v2_2. Ba window dùng A–G: W1 TRAIN=A, VALIDATION=B, TEST=C; W2 TRAIN=A+B+C, VALIDATION=D, TEST=E; W3 TRAIN=A+B+C+D+E, VALIDATION=F, TEST=G. A là 60% enrollment sớm nhất theo offering, phần còn lại theo lát thời gian; validation/test chỉ dùng các prefix P1–P4 tương ứng. AccTEMPO là metric hiệu năng chính; S_san+ v1 là metric chất lượng dữ liệu chính, còn S_san+ v2/s_cal là sensitivity. Accuracy không được dùng để kết luận khi nhãn mất cân bằng.", []),
      ("2. Pipeline và kế hoạch thực nghiệm", "V0 là zero-fill + mask. V1/V5/V9/V13 lần lượt là mean/median/ExtraTrees/MICE; các pipeline V2–V4, V6–V8, V10–V12, V14–V16 là CDSMOTE/RadiusSMOTE/SASMOTE trên parent tương ứng. S2 giữ validation/test nguyên vẹn. Seed augmentation của batch hiện có là 42; seed model là 42.", [pd.DataFrame({"Pipeline":PIPES,"Cấu hình":[PIPELINE_LABELS[p] for p in PIPES]})]),
      ("3. Kiểm soát grain, lineage và S2", "Nguồn cumulative_phase_features v1 từng có exact duplicate 23,614,180 → 11,807,090; release clean-grain v2 là đầu vào hiện hành và materializer dừng cứng nếu còn duplicate. Grain audit downstream của V0/V2 xác nhận mỗi enrollment có đúng một hàng. S2 có 36 plan immutable: 4 parent imputer × 3 window × 3 method, seed 42. RadiusSMOTE có fallback nearest-neighbour cao (>50% trong 24/24 cell lớp–window), nên chỉ là sensitivity comparator, không phải bằng chứng nội suy local-radius ổn định.", [pd.DataFrame([{"Window":"W1","V0 TRAIN": "1,609,276", "V2 balanced TRAIN":"1,866,048"},{"Window":"W2","V0 TRAIN":"1,963,060","V2 balanced TRAIN":"2,269,260"},{"Window":"W3","V0 TRAIN":"2,324,027","V2 balanced TRAIN":"2,688,368"}])]),
      ("4. Bao phủ facts mới và provenance", "Các facts dưới đây là export L2 đã ingest từ các compact contract VastAI. LSTM và GRU thuộc cùng batch provenance; BiLSTM thuộc batch runner khác, nên bảng dùng để mô tả và so sánh trong batch. Không gộp chúng thành một tuyên bố reproducibility đồng nhất.", [cq_counts, cq_runprov]),
      ("4.1 RNN Modal L4/BF16 đã nghiệm thu", "RNN seed 42 là batch Modal L4/BF16 provenance-strict trước đó: 51/51 cell, release hash CQ_V2_2 hiện hành, model revision recurrent_shared_phase_v2_4_prefixsafe_epoch50. Bảng được đọc lại từ report provenance-strict ngày 07/10; giữ tách khỏi batch VastAI bên dưới vì backend/runner khác. Nó vẫn là evidence thực nghiệm hợp lệ cho RNN riêng lẻ.", [legacy_rnn]),
      ("5. Kết quả TEST đầy đủ theo family × pipeline", "Trung bình đều trên 3 window × 4 phase. Đây là bảng đầy đủ các độ đo hiệu năng, calibration và sáu thành phần sanity có trong facts. So sánh cross-family cần ghi nhớ provenance runner ở Mục 4.", [family_summary(cq_metrics,cq_sanity,"TEST"), performance(cq_metrics,cq_sanity,"TEST")]),
      ("6. Kết quả VALIDATION đầy đủ theo family × pipeline", "Validation được dùng cho chọn checkpoint theo mean Macro-F1 P1–P4 với margin đã lưu. AccTEMPO validation được materialize riêng và không thay thế kết quả TEST.", [family_summary(cq_metrics,cq_sanity,"VALIDATION"), performance(cq_metrics,cq_sanity,"VALIDATION")]),
      ("7. Kiểm định mở rộng và tài nguyên", "Per-class facts giữ support, precision, recall, F1; bootstrap là B=2,000. STRICT/OVERLAP là lát mô tả, không thay cohort TEST. Canonical S_leak fit trên real parent imputer TRAIN pooled P1–P4, không refit theo augmented variants. Năng lượng không được kết luận vì batch này không có telemetry năng lượng hợp lệ; chỉ so sánh thời gian bằng elapsed_seconds/train_seconds khi có trường đó.", [cq_resource, cq["checkpoint_selection_long"].groupby("model_name").agg(runs=("run_id","nunique"), checkpoints_min=("checkpoint_count_evaluated","min"), checkpoints_max=("checkpoint_count_evaluated","max"), within_margin_min=("n_checkpoints_within_margin","min"), within_margin_max=("n_checkpoints_within_margin","max")).reset_index()]),
      ("7.1 Kết quả theo lớp và lát đánh giá", "Bảng theo lớp thể hiện rõ support/recall của từng nhãn; đây là bằng chứng chống diễn giải sai accuracy khi phân bố nhãn không đều. Các lát ALL/OVERLAP/STRICT chỉ mô tả độ ổn định theo lát, không thay all-row TEST primary.", [per_class_summary(cq["metrics_per_class_long"]), slice_summary(cq["strict_overlap_metrics_long"])]),
      ("8. Trạng thái và việc tiếp theo", "Đã đầy đủ facts cho 153 recurrent cell VastAI (LSTM, GRU, BiLSTM), gồm performance, sanity, canonical leakage, per-class, bootstrap, STRICT/OVERLAP, checkpoint và resource. RNN Modal là batch lịch sử khác và không được nhập vào bảng current-Vast này. Cần chạy/đưa vào evidence audit prefix-alignment cho CQ trước khi coi selection validation là không thiên lệch: bảng W1 VALIDATION/TEST P1–P4, static/mask-static parity và SEEN/UNSEEN. PatchTST/MOMENT, seed 43/44 và replication backend cùng revision vẫn là các pha tiếp theo; không được suy ra stability từ seed 42 duy nhất.", [])]
    render(ROOT/"CQ/docs/REPORT_CQ_FULL_E2E_SEED42_2026-10-16", "BÁO CÁO ĐẦY ĐỦ CQ — CLEAN-GRAIN, S2 VÀ RECURRENT SEED 42", cq_sections)

    lo=lo_data(); lo_metrics=lo["metrics"]; lo_sanity=lo["sanity_components"]
    lo_counts=fact_counts(lo,{"metrics":"Performance overall", "sanity_components":"Sanity TEST + VALIDATION", "metrics_per_class":"Per-class", "bootstrap_metrics_per_class":"Bootstrap per-class (B=2,000)", "evaluation_slices":"STRICT/OVERLAP", "checkpoint_selection":"Checkpoint selection", "canonical_s_leak":"Canonical S_leak", "resource":"Run manifests/resources"})
    lr=lo["resource"].groupby(["model_name","gpu","batch_size","precision_policy"],dropna=False).agg(runs=("pipeline_id","count"), epochs_config=("max_epochs","max"), wall_s=("elapsed_seconds","mean"), ram_peak_mb=("peak_ram_mb","max")).reset_index()
    for c in ["wall_s","ram_peak_mb"]: lr[c]=lr[c].map(lambda x:fmt(x,1))
    eq=lo["equivalence"]
    eq_text = "Không tìm thấy artefact equivalence cục bộ."
    if eq: eq_text = "Modal–Vast equivalence LO/W1/V0/RNN/seed42 đã có artefact riêng. Các trường hợp salted ID khác nhau được diễn giải theo contract; đối chiếu prediction/metric được giữ trong JSON audit, không thay thế provenance từng run."
    pa=lo.get("prefix_audit_summary", {})
    prefix_phase=pd.DataFrame()
    prefix_slice=pd.DataFrame()
    prefix_labels=pd.DataFrame()
    if pa:
        d=lo["w1_validation_test_p1_p4_summary"]
        prefix_phase=d.groupby(["eval_split","phase_id"],as_index=False)[["macro_f1","balanced_accuracy"]].mean()
        prefix_phase["n_rows"] = d.groupby(["eval_split","phase_id"])["n_rows"].first().to_numpy()
        prefix_phase["macro_f1"]=prefix_phase["macro_f1"].map(fmt);prefix_phase["balanced_accuracy"]=prefix_phase["balanced_accuracy"].map(fmt)
        d=lo["w1_p1_seen_unseen_strict_overlap_metrics"]
        prefix_slice=d.groupby(["axis","slice"],as_index=False)[["n_rows","macro_f1","balanced_accuracy"]].mean()
        prefix_slice["n_rows"]=prefix_slice["n_rows"].round().astype(int);prefix_slice["macro_f1"]=prefix_slice["macro_f1"].map(fmt);prefix_slice["balanced_accuracy"]=prefix_slice["balanced_accuracy"].map(fmt)
        prefix_labels=lo["w1_p1_course_cohort_label_support"].copy()
        prefix_labels["share_within_cohort"]=prefix_labels["share_within_cohort"].map(lambda x:fmt(x*100,2)+"%")
    lo_sections=[
      ("1. Phạm vi, release và split", "LO dùng nhãn final_score source-faithful/catalog-normalized, release V3.1 clean-grain và split chronological hybrid. A là 60% enrollment lịch sử theo whole-offering boundary; future pool chia ba lát T1–T3 và trong mỗi lát pairing validation/test. W1: A/B/C; W2: A+B+C/D/E; W3: A+B+C+D+E/F/G. Test-side purge được ghi registry; label drift giữa lát được giữ lại, không san phẳng. AccTEMPO là metric chính; S_san+ v1 là chất lượng dữ liệu chính và v2/s_cal sensitivity.", []),
      ("2. Nhãn, mất cân bằng và diễn giải", "LO có các lớp I/D, G, E. Mất cân bằng rất mạnh khiến accuracy cao có thể chỉ phản ánh lớp I/D; vì vậy quyết định không dựa trên accuracy. Bảng chi tiết per-class/CI nằm trong facts và phần kết quả dưới đây dùng Macro-F1, balanced accuracy, PR-AUC và AccTEMPO song song.", []),
      ("3. Pipeline và chất lượng đầu vào", "Cùng ma trận 17 pipeline với CQ: V0 zero-fill + masks; bốn imputer và ba balancer trên từng parent. Nguồn feature cũ từng duplicate chính xác; clean-grain materializer đã khử/khóa grain trước khi sinh input. V0 và S2 chỉ sử dụng TRAIN để fit; validation/test giữ nguyên.", [pd.DataFrame({"Pipeline":PIPES,"Cấu hình":[PIPELINE_LABELS[p] for p in PIPES]})]),
      ("4. Bao phủ facts và đối chiếu mới", "Bốn family RNN/LSTM/GRU/BiLSTM đều đủ 51 cell (17 pipeline × 3 window) và audit hậu kiểm. Các facts có performance, sanity, canonical S_leak, per-class, bootstrap B=2,000, STRICT/OVERLAP và checkpoint selection.", [lo_counts, pd.DataFrame([{"Đối chiếu":"Modal–Vast LO/W1/V0/RNN/seed42","Trạng thái":"Có artefact equivalence" if eq else "Không tìm thấy","Diễn giải":eq_text}])]),
      ("4.1 Chẩn đoán W1/P1: validation-prefix và cohort", "Audit read-only prefix-validation đã chạy trên release LO V3.1 clean-grain và 204 layout recurrent. Có một signature layout chung, 15 static fields; không có future dynamic cell không-null trong TEST prefix (P1 kiểm 231 field, P2 154, P3 77), và mọi static value theo enrollment giống nhau P1–P4. Do đó bác bỏ H1 dưới dạng lỗi future-dynamic prefix/masking. TEST W1/P1 gồm 86.77% seen course; nhưng unseen course có Macro-F1/balanced accuracy cao hơn. Vì vậy H2 đơn giản (unseen course gây suy giảm) cũng không được ủng hộ; chênh lệch bị gắn với temporal slice và support nhãn hiếm, đặc biệt E chỉ có trong seen/OVERLAP. Không cần chọn lại checkpoint nếu static contract được giữ nguyên; 15 static field vẫn được lưu danh sách để protocol owner review định nghĩa.", [prefix_phase, prefix_slice, prefix_labels]),
      ("5. Kết quả TEST đầy đủ theo family × pipeline", "Trung bình đều trên W1–W3 và P1–P4. Bảng đầy đủ giữ riêng tất cả metric có thể tổng hợp trực tiếp từ facts.", [family_summary(lo_metrics,lo_sanity,"TEST"),performance(lo_metrics,lo_sanity,"TEST")]),
      ("6. Kết quả VALIDATION đầy đủ theo family × pipeline", "Validation hỗ trợ chọn checkpoint; rule và margin được lưu trong checkpoint facts. Không dùng validation để thay thế claim TEST.", [family_summary(lo_metrics,lo_sanity,"VALIDATION"),performance(lo_metrics,lo_sanity,"VALIDATION")]),
      ("7. Sanity, CI, lát STRICT/OVERLAP và tài nguyên", "S_leak là probe canonical theo parent imputer real TRAIN pooled P1–P4. Nó không được tái fit trên data augmented. Per-class facts, bootstrap B=2,000 và STRICT/OVERLAP có đầy đủ audit rows. Năng lượng/CO₂ chỉ được báo cáo khi có telemetry thích hợp; artefact hiện tại không đủ điều kiện suy ra năng lượng hoặc carbon.", [lr, lo["checkpoint_selection"].groupby("model_name").agg(runs=("pipeline_id","count"), checkpoints_min=("checkpoint_count_evaluated","min"),checkpoints_max=("checkpoint_count_evaluated","max"),within_margin_min=("n_checkpoints_within_margin","min"),within_margin_max=("n_checkpoints_within_margin","max")).reset_index()]),
      ("7.1 Kết quả theo lớp và lát đánh giá", "Các chỉ số theo lớp là cần thiết vì I/D chiếm đa số. Bảng ALL/OVERLAP/STRICT là descriptive; kết luận chính luôn là all-row TEST AccTEMPO cùng Macro-F1/balanced accuracy.", [per_class_summary(lo["metrics_per_class"]), slice_summary(lo["evaluation_slices"])]),
      ("8. Kết luận và các pha còn mở", "Audit prefix-alignment đã đóng gate W1/P1: không có lỗi future-prefix/masking và không có bằng chứng rằng unseen course một mình giải thích TEST thấp; không cần S3 rerun hoặc checkpoint reselection khi static contract giữ nguyên. LO recurrent seed 42 vì thế có evidence để báo cáo theo window/phase/pipeline/family, với điều kiện luôn trình bày W1/P1 cùng per-class, STRICT/OVERLAP và course cohort thay vì diễn giải một giá trị gộp. PatchTST/MOMENT và seed replication 43/44 vẫn là pha tiếp theo. Nếu bổ sung backend mới, cần ghi provenance và đối chiếu equivalence trước khi trộn với batch hiện hữu.", [])]
    render(ROOT/"LO/docs/REPORT_LO_FULL_E2E_SEED42_2026-10-16", "BÁO CÁO ĐẦY ĐỦ LO — V3.1 CLEAN-GRAIN, S2 VÀ RECURRENT SEED 42", lo_sections)


if __name__ == "__main__": main()
