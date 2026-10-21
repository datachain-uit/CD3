"""Generate the LO V3 catalog-normalized release report from downloaded audits."""
from __future__ import annotations

from datetime import date
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
RESULT = ROOT / "LO" / "resuilt"
OUT = ROOT / "LO" / "docs" / "REPORT_LO_V3.docx"


def parquet(root: Path, name: str) -> pd.DataFrame:
    matches = list((root / name).glob("*.parquet"))
    if len(matches) != 1:
        raise FileNotFoundError(f"Expected one parquet in {root / name}, found {len(matches)}")
    return pd.read_parquet(matches[0])


CONTRACT = RESULT / "final_score_contract_audit_v3_catalog_normalized"
REGISTRY = RESULT / "split_registry_v3_catalog_normalized_overlap_audit_v1"
COVERAGE = RESULT / "label_coverage_audit_v3_catalog_normalized"
LONG = RESULT / "long_offering_assignment_audit_v3_catalog_normalized"
CEILING = RESULT / "course_score_ceiling_audit_v3_catalog_normalized" / "course_score_ceiling_audit_v3_catalog_normalized"


def shade(cell, value: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    item = OxmlElement("w:shd")
    item.set(qn("w:fill"), value)
    props.append(item)


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def table(doc: Document, headers: list[str], rows: list[list[object]], widths: list[float] | None = None) -> None:
    tab = doc.add_table(rows=1, cols=len(headers))
    tab.style = "Table Grid"
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_repeat_header(tab.rows[0])
    for cell, text in zip(tab.rows[0].cells, headers):
        cell.text = str(text)
        shade(cell, "1F4E78")
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)
    for values in rows:
        cells = tab.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = str(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
    if widths:
        for row in tab.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Inches(width)
    doc.add_paragraph()


def heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    if level == 1:
        p.paragraph_format.space_before = Pt(12)


def pct(value: float | int | None, digits: int = 3) -> str:
    if value is None or pd.isna(value):
        return "—"
    return f"{100 * float(value):.{digits}f}%"


def n(value: float | int | None) -> str:
    if value is None or pd.isna(value):
        return "—"
    return f"{int(value):,}"


def day(value) -> str:
    """Format either pandas Timestamp, datetime, or datetime.date."""
    return str(value.date() if hasattr(value, "date") else value)


def main() -> None:
    semantics = parquet(CONTRACT, "label_semantics_contract").iloc[0]
    label_version = parquet(CONTRACT, "label_version_summary").iloc[0]
    score = parquet(CONTRACT, "score_range_by_profile")
    overflow = parquet(CONTRACT, "legacy_overflow_mechanism")
    exclusions = parquet(CONTRACT, "exclusion_reason_summary")
    proxy = parquet(CONTRACT, "proxy_feature_contract")
    blocks = parquet(REGISTRY, "block_label_audit")
    split = parquet(REGISTRY, "split_registry")
    split_ratios = parquet(REGISTRY, "split_actual_ratios")
    distribution = parquet(REGISTRY, "split_label_distribution")
    calendar = parquet(REGISTRY, "calendar_boundaries").iloc[0]
    concentration = parquet(COVERAGE, "rare_label_concentration")
    coverage_summary = parquet(COVERAGE, "slice_timeline_summary")
    assignment = parquet(LONG, "enrollment_assignment_qa").iloc[0]
    duration = parquet(LONG, "duration_by_timeline")
    strict = parquet(LONG, "strict_by_window_split_phase")
    ceiling_summary = parquet(CEILING, "course_score_ceiling_summary").iloc[0]
    slice_exclusions = parquet(COVERAGE, "slice_exclusion_reason_summary")

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.62)
    section.right_margin = Inches(0.62)
    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(22)

    title = document.add_heading("BÁO CÁO LO — CATALOG-NORMALIZED V3", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = document.add_paragraph("Hybrid scenario | CQ shared calendar | cập nhật từ artifact V3 ngày 30/09/2026")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].font.color.rgb = RGBColor(89, 89, 89)

    document.add_paragraph(
        "Kết luận điều hành: V3 thay thế nhãn source-faithful V2.2 làm PRIMARY. "
        "Điểm hoạt động được chuẩn hóa theo catalog nên bị chặn cấu trúc trong [0,100]. "
        "Split V3 đạt QA22 theo định nghĩa đã khóa của giao thức; phân phối E/G rất hiếm và tập trung theo offering "
        "được giữ là limitation bắt buộc báo cáo, không phải lý do chặn materialize."
    )

    heading(document, "1. Phạm vi, đơn vị và nguyên tắc", 1)
    document.add_paragraph(
        "Đơn vị phân tích là enrollment_id (một user–course enrollment). Feature ở grain enrollment × phase P1–P4, "
        "được tổng hợp tích lũy từ enroll_time tới cutoff tương ứng. Nhãn là activity-derived final_score, không phải "
        "final grade quan sát độc lập, completion hay dropout."
    )
    table(document, ["Thành phần", "Quy ước V3"], [
        ["Scenario", "hybrid; timeline_source = course_specific_timeline / enrollment_anchored_proxy / global_template_fallback"],
        ["Prediction", "P1–P4 cumulative prefix; test Pj chỉ nhìn P1…Pj"],
        ["Split unit", "offering_id × timeline_source; không tách một unit qua nhiều rolling block"],
        ["Label", str(semantics["label_field"]) + "; persisted score = " + str(semantics["persisted_score_field"])],
        ["Primary version", str(label_version["label_rule_version"]) + " | " + str(label_version["label_threshold_set"])],
    ])

    heading(document, "2. Nhãn V3 và lý do thay thế V2.2", 1)
    document.add_paragraph(
        "V2.2 cộng correct_ratio qua nhiều bài exam nên final_score có thể vượt 100; nhánh TEMPO cũ khi đó gán score >100 vào I/D, "
        "tạo mapping không đơn điệu theo hiệu năng. V3 dùng một best automatic score cho mỗi catalog problem; problem không nộp đóng góp 0."
    )
    table(document, ["Hạng mục", "Contract PRIMARY V3"], [
        ["Công thức", str(semantics["formula"])],
        ["Assignment / exam", "sum(best correct ratio per catalog problem) / catalog problem count; unattempted = 0"],
        ["Bands", str(semantics["class_thresholds"])],
        ["Sensitivity", "TEMPO_LEGACY_CLIP: legacy attempt-level score clip tại 100; không dùng làm PRIMARY"],
        ["Final-grade claim", str(semantics["final_grade_claim"])],
    ])

    heading(document, "3. Cohort và kiểm chứng score", 1)
    raw_total = int(exclusions["enrollment_count"].sum())
    eligible = int(label_version["eligible_enrollment_count"])
    eligible_courses = int(exclusions.loc[exclusions["proxy_status"] == "ELIGIBLE", "course_count"].iloc[0])
    table(document, ["Trạng thái", "Enrollment", "Course", "Tỷ lệ raw"], [
        [r.proxy_status, n(r.enrollment_count), n(r.course_count), pct(r.enrollment_share)]
        for r in exclusions.sort_values("enrollment_count", ascending=False).itertuples(index=False)
    ])
    document.add_paragraph(
        f"Cohort label-eligible là {n(eligible)} / {n(raw_total)} raw enrollment ({pct(eligible/raw_total)}), từ "
        f"{n(eligible_courses)} course eligible (artifact bao phủ {n(label_version.course_count)} course trước loại trừ). "
        "Course thiếu course_limit weights là nguyên nhân loại trừ lớn nhất; "
        "đây là giới hạn coverage dữ liệu nguồn, không phải missing được impute."
    )
    heading(document, "3.1. Audit trần điểm cấp course và quyết định cohort", 2)
    table(document, ["Chỉ số", "Course", "Enrollment eligible", "Tỷ lệ cohort eligible"], [
        ["Trần điểm quan sát <60", n(ceiling_summary.course_count_ceiling_lt_60),
         n(ceiling_summary.enrollment_count_ceiling_lt_60),
         pct(ceiling_summary.enrollment_count_ceiling_lt_60 / ceiling_summary.eligible_enrollment_count)],
        ["Trần điểm quan sát <85", n(ceiling_summary.course_count_ceiling_lt_85),
         n(ceiling_summary.enrollment_count_ceiling_lt_85),
         pct(ceiling_summary.enrollment_count_ceiling_lt_85 / ceiling_summary.eligible_enrollment_count)],
        ["Không có attempt ở mọi thành phần assignment/exam có trọng số — QUYẾT ĐỊNH LOẠI", n(ceiling_summary.course_count_proposed_exclude),
         n(ceiling_summary.enrollment_count_proposed_exclude),
         pct(ceiling_summary.share_eligible_enrollments_proposed_exclude)],
    ])
    document.add_paragraph(
        "course_score_ceiling_observed là performance_score PRIMARY lớn nhất quan sát được trong từng course, "
        "không phải trần lý thuyết. Quyết định protocol: loại 102 course / 172,003 enrollment với "
        "proxy_exclusion_reason = no_scored_signal_in_course. Điều kiện là course có assignment hoặc exam mang trọng số "
        "dương nhưng không có attempt quan sát được ở mọi thành phần chấm điểm có trọng số. Vì cohort thay đổi, các artifact "
        "V3 hiện tại là baseline-audit; phải tái tạo label, split, views và S0 LO trước model."
    )
    score_rows = []
    for r in score.itertuples(index=False):
        score_rows.append([r.assessment_modality_signature, n(r.eligible_enrollment_count), f"{r.final_score_min:.3f}",
                           f"{r.final_score_max:.3f}", n(r.n_final_score_gt_100), n(r.n_legacy_score_gt_100)])
    table(document, ["Assessment profile", "Eligible", "V3 min", "V3 max", "V3 >100", "Legacy >100"], score_rows)
    table(document, ["Nhóm TEMPO legacy", "Enrollment", "Exam catalog p50", "Exam attempted p50", "Legacy score p50", "Legacy max"], [
        [r.legacy_overflow_group, n(r.enrollment_count), n(r.exam_catalog_count_p50),
         n(r.exam_attempted_problem_count_p50), f"{r.legacy_score_p50:.3f}", f"{r.legacy_score_max:.3f}"]
        for r in overflow.itertuples(index=False)
    ])

    heading(document, "4. Leakage contract", 1)
    document.add_paragraph(
        "Các cột tạo trực tiếp final_score bị đánh dấu label proxy và phải loại khỏi predictor của imputer/model; "
        "điều này gồm score, weights, tử/mẫu số catalog, ratio, legacy score và coverage ratio."
    )
    table(document, ["Nhóm proxy", "Canonical columns"], [
        ["Score / weights", "performance_score; video_weight; assignment_weight; exam_weight; course_weight_total"],
        ["Video", "watched_videos; video_counts; watch_percent; video_catalog_coverage"],
        ["Problem", "assignment/exam catalog counts, best-correct sums, ratios; problem_catalog_coverage; problem_score_catalog_coverage"],
        ["Sensitivity", "tempo_legacy_unclipped_score; tempo_legacy_score_clipped"],
        ["Audit artifact", f"Snapshot đã refresh: {len(proxy)} canonical fields, bao gồm ba coverage fields."],
    ])

    heading(document, "4.1. Current-grade regime", 2)
    document.add_paragraph(
        "Khóa nhóm current-grade cho LO_FULL_EARLY (và loại toàn bộ khỏi LO_NO_CURRENT_GRADE): "
        "score_sum, score_mean, score_std, score_fraction_mean, score_fraction_std, correct_problem_count, "
        "incorrect_problem_count, correct_ratio, correct_ratio_std; áp dụng cùng chính sách cho mọi hậu tố _P1…_P4. "
        "problem_score_catalog_coverage là label_proxy nên bị cấm hoàn toàn khỏi imputer/model."
    )

    heading(document, "5. Split V3 và QA22", 1)
    document.add_paragraph(
        "Calendar được neo theo CQ. A là history; future pool được chia theo offering_end_date thành T1=B/C, T2=D/E, T3=F/G. "
        "Trong từng lát, rarest_first ghép unit nguyên vẹn để cân bằng enrollment và lớp hiếm."
    )
    table(document, ["Calendar boundary", "Ngày"], [
        ["A boundary", day(calendar.a_boundary)],
        ["T1 boundary", day(calendar.slice_1_boundary)],
        ["T2 boundary", day(calendar.slice_2_boundary)],
        ["T3 boundary", day(calendar.slice_3_boundary)],
    ])
    document.add_paragraph(
        "QA22 được giữ theo giao thức: tau_w ≤ min_evaluation_end_utc và hai evaluation arm thỏa size policy. "
        "P1–P4 strictness/overlap được lưu như mô tả bối cảnh, không là gate chặn split hoặc materialize."
    )
    train_counts = (split_ratios[split_ratios["split"] == "train"]
                    .set_index("window_id")["enrollment_count"].to_dict())
    table(document, ["W", "TRAIN / VALID / TEST", "tau_w", "Min evaluation end", "QA22", "Purge"], [
        [r.window_id, f"{n(train_counts[r.window_id])} / {n(r.n_enrollments_validation)} / {n(r.n_enrollments_test)}",
         day(r.tau_w), day(r.min_evaluation_end_utc), "Đạt" if r.qa22_ok else "Không đạt", r.purge_mode]
        for r in split.itertuples(index=False)
    ])

    heading(document, "6. Phân bố nhãn và support", 1)
    # Namedtuple cannot expose I/D-derived column; use records safely.
    block_rows = []
    for _, r in blocks.sort_values("rolling_block").iterrows():
        total = int(r.enrollment_count)
        block_rows.append([r.rolling_block, n(total), f"{n(r['I/D_count'])} ({pct(r['I/D_count']/total)})",
                           f"{n(r.G_count)} ({pct(r.G_count/total)})", f"{n(r.E_count)} ({pct(r.E_count/total)})"])
    table(document, ["Block", "Enrollment", "I/D", "G", "E"], block_rows)
    pivot = distribution.pivot(index=["window_id", "split", "enrollment_count"], columns="label", values="label_enrollment_count").reset_index().fillna(0)
    split_rows = []
    for r in pivot.sort_values(["window_id", "split"]).itertuples(index=False):
        total = int(r.enrollment_count)
        split_rows.append([r.window_id, r.split, n(total), f"{n(getattr(r, '_3'))} ({pct(getattr(r, '_3')/total)})" if False else "", "", ""])
    # DataFrame lookup avoids invalid identifier I/D.
    split_rows = []
    for _, r in pivot.sort_values(["window_id", "split"]).iterrows():
        total = int(r.enrollment_count)
        split_rows.append([r.window_id, r.split, n(total), f"{n(r['I/D'])} ({pct(r['I/D']/total)})",
                           f"{n(r['G'])} ({pct(r['G']/total)})", f"{n(r['E'])} ({pct(r['E']/total)})"])
    table(document, ["W", "Split", "Enrollment", "I/D", "G", "E"], split_rows)
    document.add_paragraph(
        "E/G có support nhỏ ở một số cell W2/W3. Không dùng resampling để diễn giải validation/test: augmentation chỉ thuộc TRAIN."
    )
    table(document, ["Policy support/selection", "Quy ước phải áp dụng downstream"], [
        ["small_support_flag", "1 khi support thực của ô (window, phase, eval_split, subgroup, class) < 400."],
        ["small_subgroup_flag", "Giữ riêng: n_g < 30 hoặc một class support < 5."],
        ["Class absent", "Ghi NULL với reason = CLASS_ABSENT; không thay bằng 0 hoặc nội suy."],
        ["Uncertainty", "Per-class recall/F1, confusion matrix, bootstrap CI 95%; paired bootstrap B=2,000 cho LO W2/W3."],
        ["Checkpoint", "MEAN_PHASE_MACRO_F1 trên VALIDATION; đọc selection_margin W1 thận trọng vì concentration theo offering."],
        ["Sanity", "Báo cáo S_san+ v1 và s_cal v2 cạnh nhau cho LO W2/W3."],
        ["Sampling PRIMARY", "IR10_K10: oversample tối đa 10× mỗi lớp hiếm; synthetic share không quá 90%. W1 là kết luận chính; W2/W3 exploratory."],
    ])

    heading(document, "7. Coverage nhãn theo lát", 1)
    coverage_agg = (coverage_summary.groupby("temporal_slice", dropna=False)[
        ["n_enrollments_total", "n_proxy_eligible", "n_performance_score_observed", "n_E", "n_G", "n_I_D"]
    ].sum(min_count=1).reset_index())
    order = {"A": 0, "T1": 1, "T2": 2, "T3": 3, "UNASSIGNED": 4}
    coverage_agg["_order"] = coverage_agg["temporal_slice"].map(order).fillna(99)
    coverage_agg = coverage_agg.sort_values("_order")
    table(document, ["Slice", "All label rows", "Eligible", "Eligible share", "I/D", "G", "E"], [
        [r.temporal_slice, n(r.n_enrollments_total), n(r.n_proxy_eligible), pct(r.n_proxy_eligible / r.n_enrollments_total),
         n(r.n_I_D), n(r.n_G), n(r.n_E)]
        for r in coverage_agg.itertuples(index=False)
    ])
    document.add_paragraph(
        "Mẫu số split là enrollment label-eligible đã có rolling block; bảng coverage dùng toàn bộ label artifact, nên có "
        "UNASSIGNED và có thể lớn hơn population split. proxy_exclusion_reason phải luôn được báo cáo cùng mẫu số này; "
        "không diễn giải giảm E/G là log drift trước khi S0 đọc missing_rate và event_sparsity theo lát/timeline_source."
    )
    exclusion_assigned = slice_exclusions[
        (slice_exclusions["temporal_slice"].isin(["A", "T1", "T2", "T3"]))
        & (slice_exclusions["proxy_status"] != "ELIGIBLE")
    ].groupby(["temporal_slice", "proxy_status"], as_index=False)["n_enrollments"].sum()
    ex_order = {"A": 0, "T1": 1, "T2": 2, "T3": 3}
    exclusion_assigned["_order"] = exclusion_assigned["temporal_slice"].map(ex_order)
    exclusion_assigned = exclusion_assigned.sort_values(["_order", "proxy_status"])
    table(document, ["Slice", "Lý do loại", "Enrollment"], [
        [r.temporal_slice, r.proxy_status, n(r.n_enrollments)]
        for r in exclusion_assigned.itertuples(index=False)
    ])
    document.add_paragraph(
        f"Có {n(exclusion_assigned.n_enrollments.sum())} enrollment không đủ điều kiện trong A/T1/T2/T3; "
        "T3 chiếm phần lớn, chủ yếu thiếu assignment catalog hoặc có tổng trọng số course khác 100. "
        "Đây là coverage/source artifact, phải được báo cáo cùng drift theo lát."
    )

    heading(document, "8. Concentration, assignment và offering dài", 1)
    conc = concentration.sort_values(["temporal_slice", "timeline_source"])
    table(document, ["Slice", "Timeline", "E", "G", "Top-E share", "Top-G share"], [
        [r.temporal_slice, r.timeline_source, n(r.n_E), n(r.n_G), pct(r.top_E_offering_share), pct(r.top_G_offering_share)]
        for r in conc.itertuples(index=False)
    ])
    document.add_paragraph(
        "Concentration theo offering là limitation: chẳng hạn T1 course-specific có một offering chiếm phần lớn E, "
        "và T3 course-specific có một offering chiếm phần lớn G. Báo cáo model phải phân tầng theo timeline_source và nêu support."
    )
    table(document, ["Timeline source", "Units", "Enrollment", "Duration p50/p90/max", "Enrollment >270 days"], [
        [r.timeline_source, n(r.offering_count), n(r.enrollment_count),
         f"{n(r.duration_days_p50)} / {n(r.duration_days_p90)} / {n(r.duration_days_max)}", pct(r.long_offering_enrollment_share)]
        for r in duration.itertuples(index=False)
    ])
    table(document, ["Assignment integrity", "Giá trị"], [
        ["Distinct enrollment", n(assignment.distinct_enrollment_count)],
        ["Enrollment ở nhiều rolling block", n(assignment.n_enrollments_multiple_rolling_blocks)],
        ["Enrollment ở nhiều split trong cùng window", n(assignment.n_enrollments_multiple_splits_same_window)],
        ["Context không map block", n(assignment.n_context_rows_without_block)],
        ["Kết luận", "Đạt" if assignment.assignment_qa_pass == 1 else "Không đạt"],
    ])
    document.add_paragraph(
        "Quyết định đơn vị đã khóa: giữ nguyên offering dài, không tách pseudo-run và không cắt lại calendar. "
        "Hai cột context bắt buộc ở prediction-level là timeline_source và long_offering_flag (duration >270 ngày); "
        "không tạo subgroup bắt buộc mới chỉ vì duration dài."
    )

    heading(document, "9. Strictness context và trạng thái downstream", 1)
    strict_eval = strict[strict["split"] != "train"].copy()
    strict_wide = (strict_eval.pivot(index=["window_id", "split"], columns="phase_id", values="temporal_strict_share")
                   .reset_index().fillna(0).sort_values(["window_id", "split"]))
    table(document, ["W", "Split", "P1 strict", "P2 strict", "P3 strict", "P4 strict"], [
        [r.window_id, r.split, pct(r.P1), pct(r.P2), pct(r.P3), pct(r.P4)]
        for r in strict_wide.itertuples(index=False)
    ])
    document.add_paragraph(
        "Các tỷ lệ trên là context audit theo giao thức đã khóa; QA22 không được thay đổi bởi các tỷ lệ này. "
        "V3 baseline đạt audit split; tuy nhiên quyết định loại course không có tín hiệu chấm thay đổi cohort. Vì vậy phải rebuild label → split → views → S0 trước imputation, augmentation và model."
    )
    table(document, ["Gate", "Trạng thái V3", "Điều kiện ghi nhận"], [
        ["Course ceiling decision", "Đã chốt loại", "102 course / 172,003 enrollment sẽ nhận reason no_scored_signal_in_course; rebuild downstream là bắt buộc."],
        ["Registry / QA22 / assignment", "Baseline đạt", "Sau khi cohort đổi phải chạy lại LO split QA22 cùng CQ calendar anchor; pairing=rarest_first."],
        ["Materialize views + prefixes", "Chờ rebuild label/split", "Giữ timeline_source, duration_days, long_offering_flag và label_threshold_set."],
        ["P0 feature dictionary", "Đã khóa policy", "19 label proxy + suffix _P1…_P4 bị cấm; current-grade chỉ thuộc LO_FULL_EARLY."],
        ["S0", "Được chạy sau materialize", "Đọc missing_rate/event_sparsity theo A/T1/T2/T3 × timeline_source trước khi diễn giải drift."],
        ["Imputation / augmentation", "Sau P0", "Fit TRAIN-only; validation/test transform-only; augmentation chỉ TRAIN."],
        ["Model", "Sau S0/imputation", "Áp dụng policy support cấp ô, CI và metric S_san+/s_cal đã đăng ký."],
    ])
    table(document, ["Artifact", "Đường dẫn release"], [
        ["Label contract", "final_score_contract_audit_v3_catalog_normalized/"],
        ["Coverage", "label_coverage_audit_v3_catalog_normalized/"],
        ["Long-offering / assignment", "long_offering_assignment_audit_v3_catalog_normalized/"],
        ["Split registry", "split_registry_v3_catalog_normalized_overlap_audit_v1/"],
        ["Course ceiling", "course_score_ceiling_audit_v3_catalog_normalized/"],
    ])

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("LO catalog-normalized V3 | immutable audit snapshot | ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
