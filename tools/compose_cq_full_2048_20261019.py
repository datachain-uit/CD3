"""Append the validated batch-2048 findings to the existing CQ full report."""
from copy import deepcopy
from pathlib import Path
from docx import Document

root = Path(__file__).resolve().parents[1]
original = root / 'CQ/docs/REPORT_CQ_FULL_E2E_SEED42_2026-10-16.docx'
update = root / 'CQ/docs/REPORT_CQ_RECURRENT_BATCH2048_SEED42_SLICE_R2_2026-10-19.docx'
output = root / 'CQ/docs/REPORT_CQ_FULL_E2E_BATCH2048_SEED42_2026-10-19.docx'
doc = Document(original)
doc.add_page_break()
doc.add_heading('Cập nhật kết quả recurrent batch 2.048 — 19/10/2026', 0)
doc.add_paragraph('Phần trước giữ nguyên lịch sử và bối cảnh của báo cáo gốc. Với LSTM, GRU và BiLSTM, số liệu batch 2.048 và audit r2 ở phần sau thay thế các kết quả batch 512 và composite slice r1 trước đây; không trộn hai bộ khi so sánh. RNN không thuộc đợt ingest này.')
for element in Document(update).element.body:
    if not element.tag.endswith('}sectPr'):
        doc.element.body.insert(-1, deepcopy(element))
doc.save(output)
print(output)
