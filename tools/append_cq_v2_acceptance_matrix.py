"""Append the explicit GĐ1 clean-grain acceptance matrix to the CQ V2 report.

This is a report-only step.  It validates downloaded r2 facts and never
changes a Modal Volume artifact or an experimental result.
"""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "CQ" / "docs" / "REPORT_CQ_V2.2_CLEAN_GRAIN_RNN_SEED42_V2_2026-10-07.docx"
FACTS = ROOT / "_tmp_cq_rnn_r3_facts" / "report_id=RNN_seed42_v2_2_clean_20261007_r2"
RADIUS = ROOT / "_tmp_cq_radius_audit"
MARKER = "10. Ma tran nghiem thu GĐ1 clean-grain (r2)"


def _paragraph(document: Document, value: str, size: int = 10, bold: bool = False) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(value)
    run.bold = bold
    run.font.size = Pt(size)


def _table(document: Document, headers: tuple[str, ...], rows: list[tuple[object, ...]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = str(value)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(9)
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = str(value)
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(9)


def _read(name: str) -> pd.DataFrame:
    return pd.read_parquet(FACTS / name)


def main() -> None:
    document = Document(REPORT)
    if any(paragraph.text == MARKER for paragraph in document.paragraphs):
        raise RuntimeError("Acceptance matrix already present; refusing to append twice.")

    inventory = _read("run_inventory.parquet")
    test_sanity = _read("sanity_test_long.parquet")
    validation_sanity = _read("sanity_validation_long.parquet")
    validation_metrics = _read("metrics_validation_long.parquet")
    selection = _read("checkpoint_selection_long.parquet")
    per_class = _read("metrics_per_class_long.parquet")
    bootstrap = _read("bootstrap_metrics_per_class_long.parquet")
    slices = _read("strict_overlap_metrics_long.parquet")
    resources = _read("resource_long.parquet")
    radius = pd.concat([pd.read_parquet(path) for path in sorted(RADIUS.rglob("*.parquet"))], ignore_index=True)

    for frame, expected, name in (
        (inventory, 51, "run inventory"),
        (test_sanity, 204, "TEST sanity"),
        (validation_sanity, 204, "VALIDATION sanity"),
        (validation_metrics, 204, "VALIDATION metrics"),
        (selection, 51, "checkpoint selection"),
        (per_class, 1224, "per-class"),
        (bootstrap, 3672, "bootstrap"),
        (slices, 3672, "STRICT–OVERLAP"),
        (resources, 51, "resource"),
    ):
        if len(frame) != expected:
            raise ValueError(f"Expected {expected} {name} rows, got {len(frame)}")
    if radius["fallback_nearest_neighbour_rate"].isna().any() or len(radius) != 24:
        raise ValueError("Radius fallback audit must contain 24 non-null class records")

    sanity = pd.concat([test_sanity, validation_sanity], ignore_index=True)
    groups = sanity.groupby(["eval_split", "window_id", "phase_id"])
    if groups["s_eff"].nunique().max() != 1 or groups["s_leak"].nunique().max() != 1:
        raise ValueError("Canonical real-row sanity invariance failed")
    if sanity["s_leak_probe_n_eval"].isna().any():
        raise ValueError("Canonical sanity lacks s_leak_probe_n_eval")

    _paragraph(document, MARKER, size=14, bold=True)
    _paragraph(
        document,
        "Bang nay doi chieu truc tiep voi checklist trong Danh_gia_GD1_clean-grain_2026-10-07. "
        "Tat ca so lieu la facts r2 cua 51 run RNN clean-grain seed 42; khong co buoc train, impute "
        "hay augmentation nao duoc chay lai khi lap bang.",
    )
    created_min, created_max = inventory["manifest_mtime_utc"].min(), inventory["manifest_mtime_utc"].max()
    _table(document, ("Muc nghiem thu", "Bang chung hien co", "Trang thai"), [
        (
            "R1 — provenance cua 51 run",
            f"run_inventory=51; model_revision v2_4; data_release_id, input_root, parent IDs va seed bundle day du. "
            f"Thoi diem UTC cua cac manifest: {created_min} den {created_max}.",
            "DAT",
        ),
        (
            "R2 — grain va lineage",
            "Lineage sidecar labels/features/windows/joined_base: 0 exact duplicate; grain audit V0 + V2 cho W1–W3 "
            "co rows = distinct enrollment va histogram {1:n}.",
            "DAT",
        ),
        (
            "R3 — sanity real-parent",
            f"TEST=204 va VALIDATION=204 row; s_eff va s_leak deu co 1 gia tri trong moi eval_split × window × phase; "
            "co s_leak_probe_fit_rows, s_leak_probe_n_eval, probe_fit_scope/probe_score_scope, s_cal va S_san+ v2.",
            "DAT",
        ),
        (
            "R4 — validation va checkpoint",
            f"Validation AccTEMPO + 6 sanity components: 204 row. Checkpoint selection: 51 row; "
            f"{int(selection.checkpoint_count_evaluated.min())}–{int(selection.checkpoint_count_evaluated.max())} checkpoint/run; "
            f"n_checkpoints_within_margin={int(selection.n_checkpoints_within_margin.min())}.",
            "DAT",
        ),
        (
            "R5 — per-class, CI, STRICT–OVERLAP",
            f"Per-class={len(per_class):,}; bootstrap B=2,000={len(bootstrap):,}; STRICT–OVERLAP={len(slices):,}. "
            "Cac fact giu support, recall, small_support_flag va khoang CI.",
            "DAT",
        ),
        (
            "R6 — S2 / RadiusSMOTE",
            f"Sampling plan 36 S2 run da bao phu trong muc 9.1; Radius audit=24 window × class record; "
            f"fallback range {radius.fallback_nearest_neighbour_rate.min():.3f}–{radius.fallback_nearest_neighbour_rate.max():.3f}; canh bao >0.50.",
            "DAT (co canh bao)",
        ),
        (
            "R7 — resource run moi",
            f"resource=51; wall-clock {resources.time_run_total_s.sum()/3600:.2f} gio; peak RAM median/max "
            f"{resources.peak_ram_mb.median():,.0f}/{resources.peak_ram_mb.max():,.0f} MB; GPU L4/BF16 trong manifest.",
            "DAT",
        ),
        (
            "Stability F1",
            "Seed 43/44 va RNN/LSTM/GRU/BiLSTM full grid la pha stability tiep theo theo GRID_POLICY, khong phai dieu kien tu choi pilot seed 42.",
            "KE HOACH",
        ),
    ])
    _paragraph(
        document,
        "Luu y provenance: 51 run seed-42 da hoan thanh truoc khi manifest co truong created_at_utc; "
        "inventory vi vay luu manifest_mtime_utc nhu timestamp provenance minh bach. Model runner hien ghi created_at_utc "
        "cho moi run moi va exporter ghi created_at_source de phan biet timestamp goc voi fallback mtime.",
    )
    document.save(REPORT)
    print(REPORT)


if __name__ == "__main__":
    main()
