"""Append current task-specific Modal Volume architecture to the main CQ/LO reports."""
from pathlib import Path

from docx import Document
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]


def append_architecture(path: Path, *, task: str, regime: str, release: str,
                        input_root: str, phase_views: str, test_views: str, label: str) -> None:
    doc = Document(path)
    marker = "CẤU TRÚC LƯU TRỮ HIỆN HÀNH TRÊN MODAL VOLUME"
    if any(marker in paragraph.text for paragraph in doc.paragraphs):
        return
    doc.add_page_break()
    p = doc.add_paragraph()
    r = p.add_run(marker); r.bold = True; r.font.size = Pt(14)
    doc.add_paragraph(
        f"Cấu trúc dưới đây là contract vật lý cho {task} release {release}. Artifact L1/L2 là bất biến; "
        "retry tạo attempt_id mới, không ghi đè run/attempt cũ. L3 chỉ là lớp tổng hợp và không thay thế provenance L1/L2."
    )
    tree = f'''tempo-data-v1/
├── input/                                          # immutable source supplied to experiments
│   └── {input_root}/
│       ├── {phase_views}/                          # TRAIN + VALIDATION, full P1–P4
│       ├── {test_views}/
│       │   ├── P1/                                 # TEST observes P1 only
│       │   ├── P2/                                 # TEST observes P1–P2
│       │   ├── P3/                                 # TEST observes P1–P3
│       │   └── P4/                                 # TEST observes P1–P4
│       └── release_manifest.json                   # split/phase/schema/label hash and provenance
│
└── meta_release=imputation-v1/
    ├── L0_registry/                                # registry, state, schema/input hash, environment
    │   ├── state_manifests/s0_{task}.json
    │   ├── run_registry/run_id={{run_id}}.json
    │   ├── model_registry/model_name=RNN/model_revision={{revision}}.json
    │   └── environments/attempt_id={{attempt_id}}.json
    │
    ├── L1_runs/                                    # immutable executable artifacts
    │   └── task={task}/feature_regime={regime}/window_id=W1/
    │       ├── pipeline_id=V0/
    │       │   ├── model_name=IMPUTATION_ONLY/seed=20260922/
    │       │   │   └── run_id={{imputation_sha}}/attempt_id={{attempt}}-NNNN/
    │       │   │       ├── run_manifest.json
    │       │   │       ├── fitted_preprocessors/
    │       │   │       └── model_inputs/{{train, validation, test_P1..test_P4}}.parquet
    │       │   └── model_name=RNN/seed=42/
    │       │       └── run_id={{model_sha}}/attempt_id={{attempt}}-NNNN/
    │       │           ├── run_manifest.json
    │       │           ├── checkpoint/model.pt + train_history.parquet
    │       │           ├── feature_layout.json + timing.parquet
    │       │           ├── predictions/{{eval_split=VALIDATION, eval_split=TEST}}/phase_id=P{{1..4}}/
    │       │           ├── probability_matrices/{{VALIDATION, TEST}}/phase_id=P{{1..4}}/
    │       │           └── confusion_matrices/{{VALIDATION, TEST}}/phase_id=P{{1..4}}/
    │       └── phase_id=P4/pipeline_id=V2/         # augmentation: TRAIN only
    │           └── model_name=BALANCED_TRAIN_ONLY/seed=42/
    │               └── run_id={{augmentation_sha}}/attempt_id={{attempt}}-NNNN/
    │                   ├── run_manifest.json
    │                   ├── balanced_train.parquet
    │                   ├── synthetic_ledger.parquet
    │                   └── augmentation_quality.json
    │
    ├── L2_facts/                                   # queryable facts; partitioned by task/regime/window/phase/run
    │   ├── dataset_quality_summary/  dq_feature_profile/  imputation_fidelity/
    │   ├── augmentation_quality/  resource_usage/  drift_measures/
    │   ├── metrics_overall/  metrics_per_class/  calibration/  confusion_matrix/
    │   ├── checkpoint_selection/  sanity_components/  bootstrap_metrics_per_class/
    │   ├── report_exports/task={task}/report_id={{report_id}}/
    │   └── .../task={task}/feature_regime={regime}/window_id=W1/
    │       └── phase_id=P1/.../run_id={{...}}/attempt_id={{...}}/
    │
    └── L3_meta/                                    # aggregated meta-dataset; never replaces L1/L2
        └── task={task}/label={label}/...'''
    p = doc.add_paragraph()
    run = p.add_run(tree)
    run.font.name = "Consolas"; run.font.size = Pt(8.5)
    doc.add_paragraph(
        "Nguyên tắc thực thi: TRAIN chỉ đọc model_inputs/train (hoặc balanced_train của augmentation); "
        "VALIDATION không được augment/refit; TEST prefix chỉ inference với checkpoint/scaler/imputer đã khóa. "
        "Các probability/confusion matrices được tách theo VALIDATION và TEST để tái tính metric mà không đọc raw identifier."
    )
    doc.save(path)


append_architecture(
    ROOT / "CQ" / "docs" / "REPORT_CQ_V2.2.docx",
    task="CQ", regime="CQ_RAW_EARLY", release="V2.2", input_root="CQ_v2_2",
    phase_views="phase_views_v2_2", test_views="test_prefix_views_v2_2", label="CQ_label_final",
)
append_architecture(
    ROOT / "LO" / "docs" / "REPORT_LO_V3_1.docx",
    task="LO", regime="LO_FULL_EARLY", release="V3.1", input_root="LO_v3_1",
    phase_views="phase_views_v3_1_scored_signal_excluded",
    test_views="test_prefix_views_v3_1_scored_signal_excluded", label="LO_performance_label_3",
)
print("current storage architecture appended to CQ and LO reports")
