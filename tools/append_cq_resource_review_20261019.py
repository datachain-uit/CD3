"""Export CQ batch-2048 GPU energy and train-only time and append evidence."""
import json
from pathlib import Path
import pandas as pd
from docx import Document

root = Path(__file__).resolve().parents[1]
artifacts = root / 'CQ/resuilt/cq_recurrent_batch2048_slice_r2_20261019/artifacts'
out = root / 'CQ/resuilt/cq_batch2048_closure_audit_20261019'
records = []
for path in sorted(artifacts.rglob('run_manifest.json')):
    m = json.loads(path.read_text(encoding='utf-8'))
    c = json.loads((path.parent / 'metadata_contract.json').read_text(encoding='utf-8'))
    h = pd.read_parquet(path.parent / 'train_history.parquet')
    records.append({'family': m['model_name'], 'pipeline': m['pipeline_id'], 'window': m['window_id'],
                    'train_seconds': float(h.train_seconds.sum()), 'n_epochs': len(h),
                    'energy_gpu_estimated_kwh': m.get('energy_gpu_estimated_kwh'),
                    'energy_type': m.get('energy_measurement_type'),
                    'runner_sha256': c['lineage']['runner_sha256']})
data = pd.DataFrame(records)
if len(data) != 153 or data.duplicated(['family', 'pipeline', 'window']).any():
    raise ValueError('CQ run inventory incomplete or duplicated')
if data.runner_sha256.nunique() != 1 or data.energy_gpu_estimated_kwh.isna().any():
    raise ValueError('Inconsistent runner or missing energy values')
if not data.energy_type.eq('NVML_POWER_SAMPLING_ESTIMATE').all():
    raise ValueError('Unexpected energy telemetry type')
data.to_csv(out / 'resource_153_runs_train_time_energy.csv', index=False)
group = data.groupby(['family', 'pipeline'], as_index=False).agg(
    mean_train_seconds=('train_seconds', 'mean'), mean_gpu_energy_estimated_kwh=('energy_gpu_estimated_kwh', 'mean'),
    n_windows=('window', 'nunique'))
group['averaging_rule'] = 'arithmetic mean over W1 W2 W3, one family and pipeline; not pooled over phases'
if not group.n_windows.eq(3).all():
    raise ValueError('Missing window in family/pipeline')
group.to_csv(out / 'resource_family_pipeline_mean_3_windows.csv', index=False)
doc = Document(root / 'CQ/docs/REPORT_CQ_TONG_HOP_DAY_DU_SEED42_2026-10-19_REVIEW.docx')
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            if 'Một manifest batch mới có NVML' in cell.text:
                cell.text = '153/153 manifest recurrent batch 2.048 có NVML GPU-only estimate; không suy CO₂ hoặc năng lượng toàn hệ thống.'
doc.add_paragraph('Bổ sung provenance và resource CQ batch 2.048')
doc.add_paragraph('Runner SHA-256: ' + data.runner_sha256.iloc[0] + '. Tất cả 153 contract ghi cùng hash; không suy commit Git nếu chưa có bằng chứng byte khớp. Train_seconds là tổng thời gian train từng epoch trong history; năng lượng là ước lượng GPU-only bằng NVML, đo trên cửa sổ ghi trong manifest, không phải năng lượng toàn hệ thống hay CO₂. Bảng sau lấy trung bình 3 window, không gộp pha; chi tiết từng run có trong resource_153_runs_train_time_energy.csv.')
t = doc.add_table(rows=1, cols=5); t.style = 'Table Grid'
for cell, value in zip(t.rows[0].cells, ['Họ', 'V', 'Số W', 'Train TB (giây)', 'GPU TB (kWh, ước lượng)']): cell.text = value
for _, r in group.sort_values(['family','pipeline'], key=lambda x: x.map(lambda p: int(p[1:])) if x.name == 'pipeline' else x).iterrows():
    for cell, value in zip(t.add_row().cells, [r.family, r.pipeline, str(r.n_windows), f'{r.mean_train_seconds:.2f}', f'{r.mean_gpu_energy_estimated_kwh:.6f}']): cell.text = value
target = root / 'CQ/docs/REPORT_CQ_TONG_HOP_DAY_DU_SEED42_2026-10-19_RESOURCE.docx'
doc.save(target)
print(target, 'runs=', len(data), 'groups=', len(group))
