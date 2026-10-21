"""Replace local-path table in section 12 with a Modal L0--L3 tree."""
from pathlib import Path

from docx import Document
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[2]
TREE = """tempo-data-v1/
└── meta_release=imputation-v1/
    ├── L0_registry/
    │   ├── state_manifests/                 # S0_RAW, release inventory, schema/hash
    │   └── environments/                    # runtime environment theo run_id
    ├── L1_runs/
    │   └── task={CQ|LO}/regime={...}/window=Wk/
    │       └── model=IMPUTATION_ONLY/pipeline=V*/seed=20260922/
    │           ├── run_manifest.json        # release, split, label, seed, code/config hash
    │           ├── model_inputs/            # train, validation, test_P1…P4 đã transform
    │           └── fitted_preprocessors/    # pipeline .joblib và SHA-256
    ├── L2_facts/
    │   ├── dq_feature_profile/              # DQ trước/sau xử lý
    │   ├── dataset_quality_summary/          # validity, uniqueness, label/mask integrity
    │   ├── bin_scheme/                      # bins khóa theo train
    │   ├── distribution_sketch/             # phân bố feature
    │   ├── drift_measures/                  # drift giữa split/phase
    │   ├── imputation_fidelity/             # masked-validation MAE/RMSE/MedAE
    │   └── resource_usage/                  # wall time, CPU, RAM, GPU/cost telemetry
    └── L3_meta/
        └── imputation_diagnostic_seed_level/ # diagnostic; chưa phải prediction metric
"""


def remove_table(table) -> None:
    element = table._element
    element.getparent().remove(element)


def update(path: Path) -> None:
    doc = Document(path)
    for table in list(doc.tables):
        header = [cell.text.strip() for cell in table.rows[0].cells]
        if "Vị trí local" in header:
            remove_table(table)

    title = next((p for p in doc.paragraphs if p.text.startswith(("Artifact release V1 được khóa", "Cấu trúc lưu trữ manifest"))), None)
    if title is None:
        raise RuntimeError(f"Artifact-storage title not found: {path}")
    title.text = "Cấu trúc lưu trữ manifest và artifact trên Modal Volume:"
    for run in title.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(13)
        run.bold = True

    storage = next((p for p in doc.paragraphs if p.text.startswith(("Trên Modal Volume tempo-data-v1", "tempo-data-v1/"))), None)
    if storage is None:
        raise RuntimeError(f"Modal storage paragraph not found: {path}")
    storage.text = TREE
    for run in storage.runs:
        run.font.name = "Consolas"
        run.font.size = Pt(11)

    temporary = path.with_name(path.stem + "_updating.docx")
    doc.save(temporary)
    temporary.replace(path)


for report in (ROOT / "CQ/docs/REPORT_CQ_V1.docx", ROOT / "LO/docs/REPORT_LO_V1.docx"):
    update(report)
    print(f"updated {report}")
