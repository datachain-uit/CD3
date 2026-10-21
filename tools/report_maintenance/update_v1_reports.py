"""Refresh CQ/LO V1 Word reports from the frozen label-stratified audits."""
from __future__ import annotations

import os
from pathlib import Path

import polars as pl
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[2]


def only_parquet(directory: Path) -> Path:
    files = list(directory.glob("*.parquet"))
    if len(files) != 1:
        raise FileNotFoundError(f"expected exactly one parquet in {directory}, found {len(files)}")
    return files[0]
REPORTS = {
    "CQ": {
        "report": ROOT / "CQ/docs/REPORT_CQ_V1.docx",
        "audit": only_parquet(ROOT / "CQ/resuilt/CQ/hybrid/phase_views_v1_audit"),
        "blocks": only_parquet(ROOT / "CQ/resuilt/CQ/hybrid/split_manifest_v1_block_audit"),
        "label": "CQ_label_final",
        "labels": ("warning", "average", "good"),
        "regime": "CQ_RAW_EARLY",
        "table_indexes": (7, 8, 9, 10),
    },
    "LO": {
        "report": ROOT / "LO/docs/REPORT_LO_V1.docx",
        "audit": only_parquet(ROOT / "LO/resuilt/LO/hybrid/phase_views_v1_audit"),
        "blocks": only_parquet(ROOT / "LO/resuilt/LO/hybrid/split_manifest_v1_block_audit"),
        "label": "LO_performance_label_3",
        "labels": ("I/D", "E", "G"),
        "regime": "LO_FULL_EARLY",
        "table_indexes": (6, 7, 8, 9),
    },
}


def set_cell(cell, value: str, size: int = 11, bold: bool = False) -> None:
    cell.text = value
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(size)
            run.bold = bold


def set_table(table, rows: list[list[str]]) -> None:
    if len(table.rows) != len(rows) or len(table.columns) != len(rows[0]):
        raise ValueError("Existing report table shape differs from expected V1 template")
    for r, values in enumerate(rows):
        for c, value in enumerate(values):
            set_cell(table.cell(r, c), value, bold=(r == 0))


def append_heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)


def append_text(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(13)


def append_table(doc: Document, rows: list[list[str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for r, values in enumerate(rows):
        for c, value in enumerate(values):
            set_cell(table.cell(r, c), value, size=11, bold=(r == 0))


def fmt_count_ratio(count: int, ratio: float) -> str:
    return f"{count:,} ({ratio * 100:.3f}%)"


def update(task: str, cfg: dict) -> None:
    audit = pl.read_parquet(cfg["audit"])
    blocks = pl.read_parquet(cfg["blocks"]).sort("rolling_block")
    label = cfg["label"]
    labels = cfg["labels"]
    split_order = ("train", "validation", "test")
    windows = ("W1", "W2", "W3")
    p1 = audit.filter(pl.col("phase") == "P1")

    doc = Document(cfg["report"])
    title = "BÁO CÁO QUY TRÌNH CQ V1" if task == "CQ" else "BÁO CÁO QUY TRÌNH LO V1"
    doc.paragraphs[0].text = title
    doc.paragraphs[1].text = "Phạm vi: learner–course enrollment | release feature/label/split V1\nCập nhật: 23/09/2026"
    if task == "CQ":
        doc.paragraphs[3].text = (
            "Quy ước dự án: CQ canonical là nhãn vận hành chính cho CQ. Nhãn được cố định "
            "ở cấp enrollment từ vector [COELO, AFELO, ACELO]; không phải completion hoặc grade quan sát trực tiếp."
        )

    map_index, label_index, phase_index, next_index = cfg["table_indexes"]
    split_map = [
        ["Window", "Train (quá khứ)", "Validation / Test và phần không dùng"],
        ["W1", "A: 60% enrollment P1 sớm nhất", "Validation B; Test C; D–G không dùng"],
        ["W2", "A+B+C", "Validation D; Test E; F–G không dùng"],
        ["W3", "A+B+C+D+E", "Validation F; Test G"],
    ]
    set_table(doc.tables[map_index], split_map)

    label_rows = [["W", "Split", *labels]]
    phase_rows = [["W", "Split", "P1", "P2", "P3", "P4"]]
    for window in windows:
        for split in split_order:
            sub = p1.filter((pl.col("window") == window) & (pl.col("split") == split))
            total = int(sub["split_enrollment_count"][0])
            by_label = {row[label]: row for row in sub.to_dicts()}
            label_rows.append([
                window,
                split,
                *[fmt_count_ratio(int(by_label[name]["enrollment_count"]), float(by_label[name]["label_ratio"])) for name in labels],
            ])
            phase_counts = []
            for phase in ("P1", "P2", "P3", "P4"):
                phase_subset = audit.filter(
                    (pl.col("window") == window)
                    & (pl.col("split") == split)
                    & (pl.col("phase") == phase)
                )
                phase_counts.append(f"{int(phase_subset['split_enrollment_count'][0]):,}")
            phase_rows.append([window, split, *phase_counts])
    set_table(doc.tables[label_index], label_rows)
    set_table(doc.tables[phase_index], phase_rows)

    impute_table = [["Bước", "Quy tắc"]]
    impute_table += [
        ["Impute", "Fit chỉ trên train(W); observed missing only; validation/test chỉ transform; future phase không được impute."],
        ["Scaling", "log phù hợp → impute → RobustScaler. Scaler fit chỉ từ observed numeric train; áp dụng chung cho V0/Median/Mean/ExtraTrees/MICE."],
        ["Identifier", "user/course/teacher/school: bỏ prefix, map categorical theo train; không numeric-scale hoặc impute."],
        ["Checkpoint", "Tune bằng validation; test sealed tới khi khóa quyết định."],
        ["Audit", "L0/L1/L2/L3 lưu release, schema, split, missingness, fidelity, thời gian CPU và peak RAM."],
    ]
    set_table(doc.tables[next_index], impute_table)

    total_enrollment = int(blocks["enrollment_count"].sum())
    block_rows = [["Block", "Offerings", "Enrollments", *labels]]
    for row in blocks.to_dicts():
        if task == "CQ":
            count_keys = {"warning": "warning_count", "average": "average_count", "good": "good_count"}
            ratio_keys = {"warning": "warning_ratio", "average": "average_ratio", "good": "good_ratio"}
        else:
            count_keys = {"I/D": "id_count", "E": "e_count", "G": "g_count"}
            ratio_keys = {"I/D": "id_ratio", "E": "e_ratio", "G": "g_ratio"}
        block_rows.append([
            row["rolling_block"], str(row["offering_count"]), f"{row['enrollment_count']:,}",
            *[fmt_count_ratio(int(row[count_keys[name]]), float(row[ratio_keys[name]])) for name in labels],
        ])
    count_keys_total = (
        {"warning": "warning_count", "average": "average_count", "good": "good_count"}
        if task == "CQ"
        else {"I/D": "id_count", "E": "e_count", "G": "g_count"}
    )
    totals = {
        name: sum(int(row[count_keys_total[name]]) for row in blocks.to_dicts())
        for name in labels
    }
    block_rows.append(["Tổng", str(int(blocks["offering_count"].sum())), f"{total_enrollment:,}", *[fmt_count_ratio(totals[name], totals[name] / total_enrollment) for name in labels]])

    append_heading(doc, "9. Release V1 đã khóa và trạng thái thực thi")
    append_text(doc, f"Population hybrid đã khóa: {total_enrollment:,} enrollment thuộc {int(blocks['offering_count'].sum()):,} offering. Bảng dưới là phân tầng nhãn theo block A–G dùng để tạo split manifest V1; một offering chỉ thuộc một block.")
    append_table(doc, block_rows)
    append_text(doc, "Cơ chế split chính thức: A là 60% enrollment P1 sớm nhất theo offering_end_date. Future pool B–G được phân tầng theo nhãn hiếm và quy mô offering; validation/test đều nằm sau train theo thời gian. W1 dùng A/B/C, W2 dùng A+B+C/D/E, W3 dùng A+B+C+D+E/F/G. Không random enrollment và không chia một offering giữa nhiều split.")
    append_text(doc, "Artifact đã khóa: phase_views_v1, split_manifest_v1, phase_views_v1_audit và split_manifest_v1_block_audit. Train/validation có đủ P1–P4; test materialize thành prefix P1, P2, P3, P4 và không sử dụng feature tương lai.")
    append_text(doc, "Tiến độ imputation (Modal): S0_RAW đã materialize cho CQ và LO, W1–W3. Pipeline hiện dùng log phù hợp → impute → RobustScaler; Median/Mean fit toàn bộ training values; ExtraTrees fit 500,000 training rows; MICE giữ budget bảo thủ. L2 ghi data-quality, distribution/drift, masked-validation fidelity, wall time, CPU và peak RAM. Chưa khóa imputer cuối cùng hoặc chạy mô hình dự báo; V0/Median/Mean/ExtraTrees/MICE cần được so sánh trên cùng release trước khi chọn.")

    for p in doc.paragraphs:
        for run in p.runs:
            if run.font.size is None:
                run.font.name = "Times New Roman"

    temporary = cfg["report"].with_name(cfg["report"].stem + "_updating.docx")
    doc.save(temporary)
    os.replace(temporary, cfg["report"])


for task_name, config in REPORTS.items():
    update(task_name, config)
    print(f"updated {config['report']}")
