"""Append frozen W1 imputation benchmark results to both V1 Word reports."""
from __future__ import annotations

from pathlib import Path

import polars as pl
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[2]
BASE = ROOT / "downloads/W1_compare"
MARKER = "11. Kết quả benchmark điền khuyết W1"
RUNS = (
    ("V0", "v0"),
    ("Median", "median"),
    ("Mean", "mean"),
    ("ExtraTrees full train", "extra_trees_full"),
)
TASKS = {
    "CQ": ROOT / "CQ/docs/REPORT_CQ_V1.docx",
    "LO": ROOT / "LO/docs/REPORT_LO_V1.docx",
}


def add_paragraph(doc: Document, text: str, *, heading: bool = False) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.bold = heading
    run.font.name = "Times New Roman"
    run.font.size = Pt(16 if heading else 13)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for r, values in enumerate(rows):
        for c, value in enumerate(values):
            cell = table.cell(r, c)
            cell.text = value
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(13)
                    run.bold = r == 0


def facts(task_lower: str, tag: str) -> tuple[int, float, float, float, float]:
    fidelity = pl.read_parquet(BASE / f"{task_lower}_{tag}_fidelity.parquet")
    resource = pl.read_parquet(BASE / f"{task_lower}_{tag}_resource.parquet").row(0, named=True)
    n = int(fidelity["n_masked_cells"].sum())
    mae = float((fidelity["mae"] * fidelity["n_masked_cells"]).sum() / n)
    return n, mae, float(resource["time_run_total_s"]), float(resource["cpu_process_s"]), float(resource["peak_ram_mb"])


def update(task: str, report: Path) -> None:
    doc = Document(report)
    if any(MARKER in p.text for p in doc.paragraphs):
        raise RuntimeError(f"{report} already contains W1 benchmark results")
    task_lower = task.lower()
    data = [(name, *facts(task_lower, tag)) for name, tag in RUNS]
    n_probe = data[0][1]
    median_mae = next(row[2] for row in data if row[0] == "Median")
    et_mae = next(row[2] for row in data if row[0] == "ExtraTrees full train")
    improvement = (median_mae - et_mae) / median_mae * 100

    add_paragraph(doc, MARKER, heading=True)
    add_paragraph(
        doc,
        f"Phạm vi: {task} W1, masked-validation fidelity trên cùng {n_probe:,} ô numeric được che có kiểm soát. "
        "Weighted MAE là trung bình MAE có trọng số theo số ô bị che ở từng feature family; đây là chất lượng tái tạo dữ liệu, "
        "không phải metric dự báo nhãn.",
    )
    rows = [["Biến thể", "Weighted MAE ↓", "Wall time", "CPU process", "Peak RAM"]]
    for name, _, mae, wall, cpu, ram in data:
        rows.append([name, f"{mae:.4f}", f"{wall:.1f}s", f"{cpu:.1f}s", f"{ram / 1024:.1f}GB"])
    add_table(doc, rows)
    add_paragraph(
        doc,
        f"Kết quả tạm thời: ExtraTrees full train tốt nhất ở W1 ({et_mae:.4f}), giảm {improvement:.1f}% weighted MAE so với Median ({median_mae:.4f}). "
        "Mean gần Median nhưng nhanh hơn; Median vẫn là baseline đơn giản, ổn định. ExtraTrees dùng nhiều thời gian và RAM hơn, "
        "nên chỉ được khóa làm lựa chọn cuối sau khi tái lập kết quả ở W2 và W3.",
    )
    add_paragraph(
        doc,
        "MICE full-train chưa có benchmark hợp lệ trong bảng: run hiện tại gặp NaN ở predictor thưa của sklearn IterativeImputer. "
        "Pipeline đã bổ sung robust scaling nội bộ và train-median fallback cho predictor thưa; chỉ thêm MICE vào so sánh sau khi run hoàn tất thành công. "
        "Tất cả run trong bảng không dùng GPU; resource facts gồm CPU process time và peak RAM.",
    )
    temporary = report.with_name(report.stem + "_updating.docx")
    doc.save(temporary)
    temporary.replace(report)


for task_name, report_path in TASKS.items():
    update(task_name, report_path)
    print(f"updated {report_path}")
