"""Append split-stratification and manifest-storage specification to V1 reports."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[2]
MARKER = "12. Phân tầng split và lưu trữ manifest"
CONFIG = {
    "CQ": {
        "report": ROOT / "CQ/docs/REPORT_CQ_V1.docx",
        "labels": "Warning / Average / Good",
        "label_rule": "cq_vector_proximity_v1",
        "local": "CQ/resuilt/CQ/hybrid",
        "regime": "CQ_RAW_EARLY",
    },
    "LO": {
        "report": ROOT / "LO/docs/REPORT_LO_V1.docx",
        "labels": "I/D / E / G",
        "label_rule": "lo_performance_proxy_v1",
        "local": "LO/resuilt/LO/hybrid",
        "regime": "LO_FULL_EARLY",
    },
}


def paragraph(doc: Document, text: str, *, heading: bool = False) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.bold = heading
    run.font.name = "Times New Roman"
    run.font.size = Pt(16 if heading else 13)


def table(doc: Document, rows: list[list[str]]) -> None:
    result = doc.add_table(rows=len(rows), cols=len(rows[0]))
    result.style = "Table Grid"
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            cell = result.cell(row_index, column_index)
            cell.text = value
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(13)
                    run.bold = row_index == 0


def update(task: str, cfg: dict) -> None:
    doc = Document(cfg["report"])
    if any(MARKER in p.text for p in doc.paragraphs):
        raise RuntimeError(f"{cfg['report']} already includes manifest section")

    paragraph(doc, MARKER, heading=True)
    paragraph(
        doc,
        f"Nhãn dùng để phân tầng là {cfg['labels']}. Đơn vị phân tầng là offering, không phải từng enrollment: "
        "toàn bộ enrollment của một offering được gán duy nhất vào một block và một split. Nhờ đó tránh leakage "
        "cùng offering giữa train/validation/test, đồng thời số enrollment không làm phóng đại cỡ mẫu hiệu dụng của lớp hiếm.",
    )
    table(doc, [
        ["Window", "Train", "Validation", "Test", "Block không dùng"],
        ["W1", "A: 60% enrollment P1 sớm nhất", "B", "C", "D–G"],
        ["W2", "A + B + C", "D", "E", "F–G"],
        ["W3", "A + B + C + D + E", "F", "G", "Không có"],
    ])
    paragraph(
        doc,
        "A được xác định theo offering_end_date để tạo mốc train quá khứ. Future pool được chia B–G bằng phân tầng theo "
        "nhãn hiếm và quy mô offering; mục tiêu là giữ phân bố nhãn của từng block gần population tương ứng. Trong một window, "
        "validation/test ở sau train theo thời gian; giữa validation và test, offering được phân tầng để giảm nguy cơ lớp hiếm bằng 0. "
        "Test Pj chỉ phơi lộ P1…Pj; train và validation có P1–P4.",
    )

    paragraph(doc, "Artifact release V1 được khóa trước imputation:")
    table(doc, [
        ["Artifact", "Vị trí local", "Vai trò"],
        ["phase views", f"{cfg['local']}/phase_views_v1", "Train/validation wide P1–P4"],
        ["test prefixes", f"{cfg['local']}/test_prefix_views_v1/P1…P4", "Test phase-safe prefix"],
        ["split manifest", f"{cfg['local']}/split_manifest_v1", "Membership enrollment/offering, block và split"],
        ["phase audit", f"{cfg['local']}/phase_views_v1_audit", "Số sample/nhãn theo W, split, phase"],
        ["block audit", f"{cfg['local']}/split_manifest_v1_block_audit", "Số offering/enrollment và nhãn theo A–G"],
    ])
    paragraph(
        doc,
        "Trên Modal Volume tempo-data-v1, mỗi release được ghi dưới meta_release=imputation-v1. L0_registry/state_manifests "
        "lưu S0_RAW, inventory, schema/release hash và môi trường. L1_runs lưu run_manifest.json, model inputs và fitted preprocessor "
        f"theo task={task}, regime={cfg['regime']}, window, pipeline và seed. L2_facts lưu DQ profile, bin scheme, distribution/drift, "
        "masked-validation fidelity và resource usage; các fact này partition thêm run_id nên các run cũ không bị ghi đè.",
    )
    paragraph(
        doc,
        f"Mỗi run_manifest phải mang data_release_id, split_version=label_stratified_v1, phase_version=wide_prefix_v1, "
        f"label_rule_version={cfg['label_rule']}, feature_dictionary_version, seed, fit/impute scope, scaling contract, "
        "code_version, run_id, attempt_id, status và hash preprocessor. Vì vậy có thể đối chiếu hoặc tái chạy chính xác một variant "
        "mà không dựa vào tên file thủ công.",
    )
    temporary = cfg["report"].with_name(cfg["report"].stem + "_updating.docx")
    doc.save(temporary)
    temporary.replace(cfg["report"])


for task_name, config in CONFIG.items():
    update(task_name, config)
    print(f"updated {config['report']}")
