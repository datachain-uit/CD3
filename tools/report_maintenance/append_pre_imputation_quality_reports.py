"""Append the frozen pre-imputation quality baseline to CQ/LO V1 reports.

The split audits in the Median manifests are measured before values are filled;
they are therefore valid baseline DQ facts irrespective of the imputer used.
"""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[2]
MARKER = "10. Chất lượng dữ liệu trước điền khuyết (baseline S0/W1)"
REPORTS = {
    "CQ": {
        "report": ROOT / "CQ/docs/REPORT_CQ_V1.docx",
        "manifest": ROOT / "downloads/W1_compare/cq_median_manifest.json",
        "s0": ROOT / "s0_CQ.json",
    },
    "LO": {
        "report": ROOT / "LO/docs/REPORT_LO_V1.docx",
        "manifest": ROOT / "downloads/W1_compare/lo_median_manifest.json",
        "s0": ROOT / "s0_LO.json",
    },
}


def paragraph(doc: Document, text: str, *, bold: bool = False, size: int = 13) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)


def heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)


def table(doc: Document, rows: list[list[str]]) -> None:
    item = doc.add_table(rows=len(rows), cols=len(rows[0]))
    item.style = "Table Grid"
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            cell = item.cell(row_index, column_index)
            cell.text = value
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(13)
                    run.bold = row_index == 0


def source_bytes(s0: dict) -> int:
    return sum(
        int(item["bytes"])
        for item in s0["release"]["inventory"]["files"]
        if item["path"].endswith(".parquet")
    )


def run(task: str, config: dict) -> None:
    manifest = json.loads(config["manifest"].read_text(encoding="utf-8"))
    s0 = json.loads(config["s0"].read_text(encoding="utf-8"))
    doc = Document(config["report"])
    if any(MARKER in p.text for p in doc.paragraphs):
        raise RuntimeError(f"{config['report']} already contains this baseline section.")

    splits = manifest["splits"]
    rows = [["Tập", "Số dòng", "Phase được quan sát", "Ô thiếu đủ điều kiện", "Ô vắng mặt cấu trúc"]]
    for name in ("train", "validation", "test_P1", "test_P2", "test_P3", "test_P4"):
        item = splits[name]
        rows.append([
            name.replace("test_", "Test "),
            f"{int(item['rows']):,}",
            ", ".join(item["observed_phases"]),
            f"{int(item['eligible_missing_cells']):,}",
            f"{int(item['structural_cells']):,}",
        ])

    heading(doc, MARKER)
    paragraph(
        doc,
        "Đây là baseline chất lượng dữ liệu trước điền khuyết của window W1, "
        "đọc từ manifest audit của release V1 đã khóa. Các số ô là số đếm chính xác "
        "trên toàn bộ split, không phải ước lượng từ sample.",
    )
    table(doc, rows)
    paragraph(
        doc,
        "Định nghĩa: “ô thiếu đủ điều kiện” là giá trị numeric bị thiếu khi modality "
        "tương ứng đã được quan sát; đây mới là đối tượng của imputation. “Ô vắng mặt cấu trúc” "
        "là modality không có trong dữ liệu (ví dụ course không có video/problem/comment); "
        "nó không phải missing ngẫu nhiên, không được điền khuyết và được bảo toàn bằng mask. "
        "Vì hai đại lượng có mẫu số khác nhau, báo cáo không lấy ô thiếu chia cho ô cấu trúc để tạo tỉ lệ thiếu giả.",
    )
    paragraph(
        doc,
        f"S0_RAW đã hoàn tất cho {task}, W1–W3 với profile mẫu tối đa "
        f"{s0['profile_sample_rows_max']:,} dòng mỗi trạng thái. Release đầu vào gồm "
        f"{source_bytes(s0) / 1024**3:.2f} GiB Parquet (phase views và test prefixes); "
        f"data_release_id={s0['release']['data_release_id']}. ",
    )
    paragraph(
        doc,
        "Quality gates trước khi fit: nhãn đầu ra đầy đủ ở train/validation/test; "
        "không có dynamic hoặc static numeric column nào không thể fit trong W1; "
        "test Pj chỉ phơi lộ P1…Pj. Các phase tương lai bị khóa khỏi imputer, nên số "
        "future_imputed_cells luôn phải bằng 0.",
    )
    paragraph(
        doc,
        "Các số baseline này là đầu vào để so sánh V0, Median, Mean, ExtraTrees và MICE. "
        "Sau mỗi run, L2 lưu DQ sau xử lý, masked-validation fidelity, wall time, CPU time và peak RAM; "
        "không chọn imputer chỉ từ số null còn lại.",
    )

    # The older progress paragraph mentioned the superseded 500k ET budget.
    for p in doc.paragraphs:
        if "ExtraTrees fit 500,000" in p.text:
            for run_item in p.runs:
                if "ExtraTrees fit 500,000" in run_item.text:
                    run_item.text = run_item.text.replace("ExtraTrees fit 500,000", "ExtraTrees fit 1,000,000")

    temporary = config["report"].with_name(config["report"].stem + "_updating.docx")
    doc.save(temporary)
    temporary.replace(config["report"])


for task_name, task_config in REPORTS.items():
    run(task_name, task_config)
    print(f"updated {task_config['report']}")
