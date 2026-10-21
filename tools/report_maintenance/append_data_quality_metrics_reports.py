"""Document the expanded data-quality contract in CQ/LO V1 reports."""
from pathlib import Path

from docx import Document
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[2]
MARKER = "13. Bộ thước đo chất lượng dữ liệu"


def paragraph(doc: Document, text: str, heading: bool = False) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.bold = heading
    run.font.name = "Times New Roman"
    run.font.size = Pt(16 if heading else 13)


def table(doc: Document, rows: list[list[str]]) -> None:
    result = doc.add_table(rows=len(rows), cols=len(rows[0]))
    result.style = "Table Grid"
    for r, values in enumerate(rows):
        for c, value in enumerate(values):
            cell = result.cell(r, c)
            cell.text = value
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(13)
                    run.bold = r == 0


def update(path: Path) -> None:
    doc = Document(path)
    if any(MARKER in p.text for p in doc.paragraphs):
        raise RuntimeError(f"{path} already has DQ metric section")
    paragraph(doc, MARKER, True)
    paragraph(doc, "S0_RAW và S1_IMPUTED đều là kiểm kê toàn bộ dòng của từng trạng thái W/split/phase, nên số đếm và tỷ lệ chất lượng trước/sau điền khuyết là chính xác trên release. Các metric luôn lưu số đếm và mẫu số để không suy diễn tỷ lệ sai từ structural absence.")
    table(doc, [
        ["Nhóm", "Thước đo / quality gate", "Thời điểm và nơi lưu"],
        ["Completeness", "Eligible missing, completeness observed, zero rate, residual NaN, imputed rate", "S0/S1 · L2/dq_feature_profile"],
        ["Availability", "Structural absence, phase future unavailable, không impute structural/future", "S0/S1 · L2/dq_feature_profile"],
        ["Validity", "Numeric non-finite; mask phải thuộc {0,1}; giá trị âm ở count/duration để review", "S0/S1 · L2/dataset_quality_summary"],
        ["Integrity", "Label null; enrollment_id null/duplicate; số numeric column toàn thiếu", "S0/S1 · L2/dataset_quality_summary"],
        ["Consistency", "qa_label_complete, qa_enrollment_id_unique, qa_numeric_finite, qa_binary_masks_valid", "S0/S1 · L2/dataset_quality_summary"],
        ["Distribution", "p05/p50/p95, histogram khóa theo train, JSD observed–imputed và drift split/phase", "S0/S1 · L2/distribution_sketch, drift_measures"],
        ["Imputation fidelity", "Masked-validation MAE, RMSE, MedAE theo feature family", "S1 · L2/imputation_fidelity"],
        ["Efficiency", "Wall time, CPU process, allocated cores, peak RAM, GPU/cost telemetry availability", "S1 · L2/resource_usage"],
    ])
    paragraph(doc, "Gate tối thiểu để một run được dùng tiếp: label complete; enrollment_id không trùng trong split; numeric source không có non-finite; binary availability masks hợp lệ; residual numeric null = 0 ở output; future_imputed_cells = 0; và không có imputation vào structural absence. Các chỉ số negative count/duration là cờ review dữ liệu, không tự động xóa record vì một số feature có thể dùng sentinel/derived signed value.")
    paragraph(doc, "S0 cần chạy lại một lần sau bản code này để materialize dataset_quality_summary cho CQ và LO W1–W3. Việc này chỉ bổ sung L2 facts/registry mới, không thay đổi phase views, labels hoặc split manifest đã khóa.")
    temporary = path.with_name(path.stem + "_updating.docx")
    doc.save(temporary)
    temporary.replace(path)


for report in (ROOT / "CQ/docs/REPORT_CQ_V1.docx", ROOT / "LO/docs/REPORT_LO_V1.docx"):
    update(report)
    print(f"updated {report}")
