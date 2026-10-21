"""Record historic LO static contract mismatch without modifying old results."""
import json
from pathlib import Path
from docx import Document

root=Path(__file__).resolve().parents[1]
audit=json.loads((root/'LO/resuilt/static_contract_history_20261019/history_audit.json').read_text(encoding='utf-8'))
doc=Document(root/'LO/docs/REPORT_LO_FULL_E2E_SEED42_UPDATED_2026-10-19.docx')
doc.add_page_break()
doc.add_heading('Đối chiếu cột split_rule_version lịch sử — 19/10/2026',0)
doc.add_paragraph(f"Đã kiểm tra {audit['historical_layouts_total']} feature_layout.json của 204 run RNN/LSTM/GRU/BiLSTM seed 42; {audit['historical_layouts_with_split_rule_version']} layout chứa split_rule_version. Đây là sai lệch so với registry và fit_layout hiện hành, vốn xếp cột bookkeeping vào EXCLUDE. Không sửa checkpoint, prediction hay ghi lại lịch sử các run này.")
values={tuple(r['value_counts'].items()) for r in audit['files']}
if values != {(('1',sum(r['value_counts'].values())),) for r in audit['files']}:
    raise ValueError('Unexpected split_rule_version values; do not use fixed report wording')
doc.add_paragraph(f"Quét toàn bộ {len(audit['files'])} file input local: TRAIN/VALIDATION/TEST P1–P4 của 5 parent × 3 window và 36 balanced TRAIN. Mọi dòng trong mỗi file đều mang mã số 1, không null. Trong phạm vi các input được quét, cột không phân biệt enrollment, nhãn, train/validation/test hoặc phase. Điều này không chứng minh các checkpoint sẽ giống hệt nếu huấn luyện lại không có cột: cột hằng vẫn có thể tác động bias và quá trình tối ưu hóa.")
doc.add_paragraph('Audit với fit_layout hiện tại trên 10.000 dòng đầu W1/V0 chọn 14 static fields, không chọn split_rule_version; kết quả strict theo allowlist: 0 DENIED, 0 UNLISTED. Phép kiểm tra layout này là mẫu W1/V0; không xác nhận lại mọi thuộc tính point-in-time hay thay thế ablation. Với run mới phải tiếp tục áp dụng EXCLUDE; không tái sử dụng checkpoint cũ như thể đã train với layout mới.')
doc.add_paragraph('Bằng chứng machine-readable: LO/resuilt/static_contract_history_20261019/history_audit.json; script tools/audit_lo_split_rule_history.py. Hai cell ablation bỏ static và quyết định timeline_source/window_duration_seconds vẫn là các việc riêng chưa đóng.')
out=root/'LO/docs/REPORT_LO_FULL_E2E_SEED42_STATIC_HISTORY_2026-10-19.docx'
doc.save(out)
print(out)
