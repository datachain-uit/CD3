"""Append the final, decision-oriented CQ seed-42 evaluation to REPORT_CQ_V1."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "CQ" / "docs" / "REPORT_CQ_V1.docx"


def heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(14)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(5)


def table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    value = doc.add_table(rows=1, cols=len(headers))
    value.style = "Table Grid"
    for cell, header in zip(value.rows[0].cells, headers):
        cell.text = header
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = value.add_row().cells
        for cell, item in zip(cells, row):
            cell.text = str(item)
    doc.add_paragraph()


def main() -> None:
    doc = Document(REPORT)
    if any("ĐÁNH GIÁ CUỐI — CQ V2.2" in paragraph.text for paragraph in doc.paragraphs):
        print("Final evaluation already present; no duplicate appended.")
        return

    doc.add_page_break()
    heading(doc, "PHẦN C — ĐÁNH GIÁ CUỐI — CQ V2.2 / RNN SEED 42")
    doc.add_paragraph(
        "Kết luận này áp dụng cho release hybrid V2.2 và RNN seed 42. Nó tách rõ: "
        "(i) tính hợp lệ của dữ liệu và giao thức; (ii) kết quả screening một seed; và "
        "(iii) quyết định tái lặp tiếp theo. TEST chỉ dùng để báo cáo sau khi grid đã khóa, không dùng chọn pipeline."
    )

    heading(doc, "C.1. Điều kiện khóa release và khả năng tái lập")
    table(doc, ["Hạng mục", "Đánh giá", "Bằng chứng"], [
        ["Split / phase", "Đạt", "split_version=v2_2; wide_prefix_v2_2; hybrid; offering × timeline_source là đơn vị gán."],
        ["Tách dữ liệu", "Đạt", "W1: A/B/C; W2: A+B+C/D/E; W3: A+B+C+D+E/F/G. TEST Pj chỉ nhận P1…Pj."],
        ["Input thực nghiệm", "Đạt", "V0 input: W1 3,218,552/351,378/356,190; W2 3,926,120/360,970/360,964; W3 4,648,054/360,064/360,062 (train/validation/test mỗi prefix)."],
        ["Tính đầy đủ", "Đạt", "51/51 RNN + sanity cell hoàn tất; 204 TEST metrics và 204 sanity/AccTEMPO cell đã export ở L2."],
        ["Leakage / future", "Đạt theo contract", "Preprocess fit TRAIN-only; validation/test transform-only; TEST không refit. s_leak là QA cấp recipe, không nhân theo kiến trúc."],
    ])

    heading(doc, "C.2. Kết quả RNN seed 42 và diễn giải")
    doc.add_paragraph(
        "Metric chọn checkpoint là trung bình Validation Macro-F1 trên P1–P4. Trong các imputer-only có log validation đầy đủ, "
        "Median (V1) cao nhất; MICE (V13) và ExtraTrees (V9) rất sát. V0 là baseline bắt buộc, nhưng W1 của V0 chạy trước khi epoch telemetry được bổ sung nên Validation Macro-F1 W1 không còn trong log cục bộ; TEST artifact vẫn hợp lệ và đầy đủ."
    )
    table(doc, ["Pipeline", "Validation Macro-F1 TB", "W1", "W2", "W3", "Quyết định"], [
        ["V0 — zero-fill + masks", "0.7533*", "không còn telemetry", "0.6658", "0.8408", "Giữ baseline"],
        ["V1 — median", "0.7599", "0.7754", "0.6530", "0.8514", "Shortlist chính"],
        ["V9 — ExtraTrees", "0.7561", "0.7749", "0.6421", "0.8513", "Shortlist chính"],
        ["V13 — MICE", "0.7563", "0.7733", "0.6471", "0.8486", "Shortlist chính"],
    ])
    doc.add_paragraph(
        "* Trung bình V0 chỉ trên W2–W3 có telemetry. Không so sánh nó như một giá trị ba-window hoàn chỉnh. "
        "W2 thấp hơn W1/W3 trên phần lớn pipeline là temporal difficulty/label-drift cần báo cáo, không phải lý do trộn hoặc loại window."
    )
    doc.add_paragraph(
        "Trong TEST đã khóa, V2 (Median + CDSMOTE) có Macro-F1 trung bình cao nhất 0.6926. Số này là mô tả hậu nghiệm của grid; "
        "không đảo quyết định shortlist dựa trên validation. AccTEMPO-M3, S_perf, S_cal, S_san+ v2 và s_leak đã được lưu cho từng W×P×V ở Phần B, "
        "vì vậy không suy luận chất lượng prediction chỉ từ Accuracy."
    )

    heading(doc, "C.3. Quyết định cuối cho pha tiếp theo")
    table(doc, ["Quyết định", "Trạng thái", "Lý do"], [
        ["Khóa dữ liệu CQ V2.2", "Chấp nhận", "Release, split, input, manifest và L2 facts thống nhất; không tái dùng V1."],
        ["Chọn pipeline cuối", "Chưa khóa", "Seed 42 là screening; chưa có bất định đa seed hoặc so sánh kiến trúc đầy đủ."],
        ["Shortlist tái lặp", "V0, V1, V9, V13", "Giữ baseline; V1/V9/V13 là ba imputer-only có validation cao nhất và sát nhau."],
        ["Augmentation", "Giữ trong báo cáo so sánh", "Không nhân toàn bộ augmentation grid sang các seed/kiến trúc trước khi có bằng chứng validation đa seed."],
        ["Bước kế", "Seed 43/44 trên shortlist; sau đó LSTM/GRU/BiLSTM cho shortlist đã khóa", "Báo mean/CI paired theo window×phase, sau đó mở TEST duy nhất cho kết luận cuối."],
    ])

    heading(doc, "C.4. Kết luận")
    doc.add_paragraph(
        "CQ đủ điều kiện đi tiếp sang pha tái lặp mô hình. Kết quả hiện tại đáng tin cậy như một benchmark RNN seed 42 có provenance đầy đủ, "
        "nhưng chưa phải tuyên bố mô hình thắng cuộc. Kết luận khoa học đúng là Median, ExtraTrees và MICE là shortlist thực nghiệm; "
        "còn mọi so sánh cuối phải giữ split V2.2, đánh giá theo W1–W3 và P1–P4, và báo khoảng bất định đa seed."
    )
    doc.save(REPORT)
    print(REPORT)


if __name__ == "__main__":
    main()
