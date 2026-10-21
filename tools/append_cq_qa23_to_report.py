"""Append the accepted CQ QA23 evidence to the clean-grain RNN report."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "CQ" / "docs" / "REPORT_CQ_V2.2_CLEAN_GRAIN_RNN_SEED42_V2_2026-10-07.docx"
# QA23 is additional acceptance evidence for the same clean-grain V2 release.
# It belongs in the V2 report rather than creating a nominal new experiment.
OUTPUT = SOURCE
QA_ROOT = ROOT / "CQ" / "resuilt" / "qa23_v2_2"


def set_cell_text(cell, value: object, bold: bool = False) -> None:
    cell.text = str(value)
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(10)
            run.bold = bold


def add_section_heading(document: Document, text: str) -> None:
    """The supplied report template has no built-in Heading 1 style."""
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(14)


def main() -> None:
    if not SOURCE.is_file():
        raise FileNotFoundError(SOURCE)

    rows = []
    for phase in ("P1", "P2", "P3", "P4"):
        summary = pd.read_parquet(QA_ROOT / f"phase_id={phase}" / "qa23_summary")
        if len(summary) != 1:
            raise ValueError(f"Expected one QA23 summary row for {phase}, got {len(summary)}")
        row = summary.iloc[0]
        if not bool(row["qa23_ok"]):
            raise ValueError(f"QA23 is not accepted for {phase}")
        rows.append((
            phase,
            int(row["audited_feature_count"]),
            int(row["unreviewed_proxy_count"]),
            int(row["proxy_in_imputer_predictors"]),
            "PASS" if bool(row["qa23_ok"]) else "FAIL",
        ))

    document = Document(SOURCE)
    add_section_heading(document, "5. QA23 — kiểm soát label proxy theo phase")
    document.add_paragraph(
        "QA23 được chạy lại trên phase views V2.2 clean-grain, chỉ dùng TRAIN và "
        "đánh giá dictionary tại từng prefix P1–P4. Mục tiêu là chặn biến tạo nhãn "
        "hoặc proxy của nhãn khỏi predictor của imputer/model; đây không phải phép "
        "kiểm tra hay biện pháp cân bằng nhãn."
    )

    table = document.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    headers = ("Phase", "Biến đã audit", "Proxy chưa review", "Proxy trong imputer", "Kết quả")
    for cell, header in zip(table.rows[0].cells, headers):
        set_cell_text(cell, header, bold=True)
    for result in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, result):
            set_cell_text(cell, value)

    document.add_paragraph(
        "Kết quả: cả bốn phase đều đạt QA23. `CQ_distance_euclidean_final`, biến "
        "distance cuối kỳ dùng để tạo nhãn CQ, được phân loại label-derived và không "
        "có trong imputer predictors. Các prefix-observed behavioural measurements "
        "(problem/exercise counts, attempts, score sum, correct counts và catalog "
        "coverage đến đúng phase) được giữ sau review với rationale cố định: chúng là "
        "hành vi quan sát được trước cutoff, không phải thành phần final tạo nhãn."
    )
    document.add_paragraph(
        "Do QA23 chỉ kiểm soát leakage, mất cân bằng nhãn vẫn được đánh giá độc lập "
        "qua support từng lớp, Macro-F1, balanced accuracy, recall theo lớp và "
        "augmentation chỉ trên TRAIN."
    )
    document.add_paragraph(
        "Quyết định nghiệm thu: feature contract CQ_V2_2 được chấp nhận cho các "
        "experiment kế tiếp; không cần chạy lại S0/S1/S2 hoặc RNN chỉ vì QA23 được cập nhật."
    )
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
