"""Generate the standalone LO V2.2/final_score readiness report for review."""
from __future__ import annotations

from pathlib import Path

import polars as pl
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "LO" / "docs" / "REPORT_LO_SPLIT_V2_2_FINAL_SCORE_READINESS_2026-09-30.docx"
R2 = ROOT / "LO" / "resuilt" / "split_registry_v2_2_overlap_audit_v1_r2"
LONG = ROOT / "LO" / "resuilt" / "long_offering_assignment_audit_v1"
SCORE = ROOT / "LO" / "resuilt" / "final_score_contract_audit_v1"
NAVY, BLUE, GREY, GREEN, AMBER, RED = "17365D", "1F4E78", "F2F2F2", "E2F0D9", "FFF2CC", "FCE4D6"


def read(root: Path, name: str) -> pl.DataFrame:
    files = list((root / name).glob("*.parquet"))
    if not files:
        raise RuntimeError(f"Missing parquet: {root / name}")
    return pl.read_parquet(max(files, key=lambda path: path.stat().st_mtime))


def shade(cell, color: str) -> None:
    props = cell._tc.get_or_add_tcPr(); node = OxmlElement("w:shd")
    node.set(qn("w:fill"), color); props.append(node)


def table(doc: Document, headers: list[str], rows: list[list[str]], *, font_size: float = 9.5) -> None:
    value = doc.add_table(rows=1, cols=len(headers)); value.autofit = True
    try: value.style = "Table Grid"
    except KeyError: pass
    for i, text in enumerate(headers):
        cell = value.rows[0].cells[i]; cell.text = ""; run = cell.paragraphs[0].add_run(text)
        run.bold = True; run.font.size = Pt(font_size); run.font.color.rgb = RGBColor(255, 255, 255); shade(cell, BLUE)
    for rindex, row in enumerate(rows):
        cells = value.add_row().cells
        for i, item in enumerate(row):
            cells[i].text = ""; cells[i].paragraphs[0].add_run(str(item)).font.size = Pt(font_size)
            if rindex % 2: shade(cells[i], GREY)
    doc.add_paragraph()


def heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_paragraph(); run = paragraph.add_run(text); run.bold = True
    run.font.size = Pt(15 if level == 1 else 12.5); run.font.color.rgb = RGBColor.from_string(NAVY)


def pct(x: float) -> str: return f"{x * 100:.3f}%"


def slice_for_block(block: str) -> str:
    return "A" if block == "A" else "T1" if block in {"B", "C"} else "T2" if block in {"D", "E"} else "T3"


def main() -> None:
    contract = read(SCORE, "label_semantics_contract").row(0, named=True)
    exclusion = read(SCORE, "exclusion_reason_summary")
    profiles = read(SCORE, "score_range_by_profile")
    version = read(SCORE, "label_version_summary").row(0, named=True)
    block = read(R2, "block_label_audit")
    label = read(R2, "split_label_distribution")
    split = read(R2, "split_actual_ratios")
    registry = read(R2, "split_registry")
    qa = read(R2, "qa22_v2_2")
    duration = read(LONG, "duration_by_timeline")
    assignment = read(LONG, "enrollment_assignment_qa").row(0, named=True)
    unit = read(LONG, "unit_size_distribution")
    strict = read(LONG, "strict_by_window_split_phase")

    eligible = int(version["eligible_enrollment_count"])
    total = int(version["enrollment_count"])
    score_total = profiles.select(pl.sum("n_final_score_gt_100").alias("gt100")).item()
    doc = Document(); section = doc.sections[0]
    section.top_margin = Cm(1.5); section.bottom_margin = Cm(1.5); section.left_margin = Cm(1.55); section.right_margin = Cm(1.55)
    doc.styles["Normal"].font.name = "Aptos"; doc.styles["Normal"].font.size = Pt(10.5)
    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("LO · Split V2.2 / final_score readiness · 30/09/2026").font.size = Pt(8)
    title = doc.add_paragraph(); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("BÁO CÁO LO: SPLIT V2.2 VÀ SẴN SÀNG KHÓA NHÃN FINAL_SCORE")
    run.bold = True; run.font.size = Pt(18); run.font.color.rgb = RGBColor.from_string(NAVY)
    sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run("Scenario hybrid · offering-level temporal split · review copy").italic = True
    verdict = doc.add_paragraph(); verdict.add_run("Kết luận điều hành: ").bold = True
    verdict.add_run("Split V2.2 đạt và LO có thể materialize/S0. Tuy nhiên chưa khóa feature dictionary, imputation hoặc model cho đến khi xác nhận quy ước final_score vượt 100 do thành phần exam hiện đang cộng theo tổng ratio.")
    shade(verdict._tc if False else verdict, GREEN) if False else None

    heading(doc, "1. Phạm vi, cohort và quyết định hiện tại")
    table(doc, ["Hạng mục", "Giá trị"], [
        ["Split chính", "split_version = v2.2; scenario = hybrid; CQ calendar anchor"],
        ["Đơn vị gán", "offering_id × timeline_source; không tách offering giữa split"],
        ["LO label field", contract["label_field"]],
        ["Tên nghĩa nghiệp vụ", "final_score — điểm tổng hợp hoạt động được chấm"],
        ["Tên field đang lưu", contract["persisted_score_field"] + " (compatibility field)"],
        ["Label rule version", version["label_rule_version"]],
        ["Raw enrollment", f"{total:,}"], ["Cohort LO eligible", f"{eligible:,} ({eligible / total * 100:.3f}% raw enrollment)"],
    ])
    doc.add_paragraph("LO không được mô tả là final grade quan sát độc lập. Đây là outcome score vận hành từ activity được chấm; cohort thực nghiệm chỉ gồm course có weight hợp lệ.")

    heading(doc, "2. Công thức final_score và contract chống leakage")
    table(doc, ["Mục", "Quy ước hiện hành"], [
        ["Formula", contract["formula"]], ["Bands", contract["class_thresholds"]],
        ["Exam partition", contract["exam_partition_rule"]],
        ["Dictionary rule", "Mọi component, weight, denominator và derived component của final_score phải is_label_proxy=1; loại khỏi imputer predictor và model feature."],
        ["Suffix policy", "Áp dụng cả cột không suffix và mọi biến thể _P1…_P4."],
    ])
    doc.add_paragraph("Danh sách proxy canonical đã được xuất riêng trong final_score_contract_audit_v1/proxy_feature_contract/. Bao gồm score, ba weight, watch_percent, watched_videos/video_counts, assignment/exam attempt và ratio, cùng các derived fallback.")

    heading(doc, "3. Coverage cohort và lý do loại trừ")
    exclusion_rows = [[r["proxy_status"], f"{r['enrollment_count']:,}", f"{r['course_count']:,}", pct(r["enrollment_share"])] for r in exclusion.sort("enrollment_count", descending=True).iter_rows(named=True)]
    table(doc, ["Trạng thái", "Enrollment", "Course", "Tỷ lệ raw"], exclusion_rows)
    doc.add_paragraph("Coverage 98.7–100% theo A/T1/T2/T3 trước đây là coverage bên trong cohort/split LO. Nó không đồng nghĩa 100% raw enrollment có nhãn: 72.97% raw enrollment thiếu course_limit weights.")

    heading(doc, "4. Course assessment profiles và score range")
    profile_rows = [[r["assessment_modality_signature"], f"{r['eligible_enrollment_count']:,}",
                     f"{r['eligible_enrollment_count'] / eligible * 100:.2f}%", f"{r['final_score_p99']:.2f}",
                     f"{r['final_score_max']:.2f}", f"{r['n_final_score_gt_100']:,}",
                     f"{r['n_E']:,}", f"{r['n_G']:,}"]
                    for r in profiles.sort("eligible_enrollment_count", descending=True).iter_rows(named=True)]
    table(doc, ["Assessment signature", "Eligible", "% cohort", "p99", "Max", "Score >100", "E", "G"], profile_rows)
    alert = doc.add_paragraph(); alert.add_run("Điểm cần xác nhận trước khi khóa LO: ").bold = True
    alert.add_run(f"{score_total:,} enrollment ({score_total / eligible * 100:.3f}% cohort) có final_score >100; max = 2,810. Các trường hợp này đến từ profile có EXAM, do exam dùng total_correct_ratio_exam. Nếu final_score phải là thang 0–100, cần đổi/chuẩn hóa component exam rồi rebuild label và split LO. Nếu đây là quy ước chủ đích, phải ghi rõ score là weighted activity aggregate uncapped.")

    heading(doc, "5. Temporal split V2.2: integrity và tỷ lệ nhãn")
    qa_rows = [[r["window_id"], f"{r['n_enrollments_validation']:,}", f"{r['n_enrollments_test']:,}", "Đạt" if r["end_order_ok"] else "Không", "Đạt" if r["qa22_ok"] else "Không"] for r in qa.sort("window_id").iter_rows(named=True)]
    table(doc, ["Window", "Validation", "Test", "Label-end order", "QA22"], qa_rows)
    slice_rows = []
    for slice_id in ("A", "T1", "T2", "T3"):
        frame = block.filter(pl.col("rolling_block").map_elements(slice_for_block, return_dtype=pl.String) == slice_id)
        n = int(frame["enrollment_count"].sum()); values = {k: int(frame[k].sum()) for k in ("I/D_count", "G_count", "E_count")}
        slice_rows.append([slice_id, f"{n:,}", f"{values['I/D_count']:,} ({values['I/D_count']/n*100:.3f}%)", f"{values['G_count']:,} ({values['G_count']/n*100:.3f}%)", f"{values['E_count']:,} ({values['E_count']/n*100:.3f}%)"])
    table(doc, ["Slice", "Enrollment", "I/D", "G", "E"], slice_rows)
    doc.add_paragraph("T1 có lớp E/G cao; T2/T3 có E/G rất thấp. Đây là temporal label drift trong cohort LO, không phải lỗi phân tầng; phải được diễn giải cùng coverage log theo lát từ S0.")

    heading(doc, "6. Nhãn theo window × split và support nhỏ")
    rows = []
    for window in ("W1", "W2", "W3"):
        for arm in ("train", "validation", "test"):
            total_row = split.filter((pl.col("window_id") == window) & (pl.col("split") == arm)).row(0, named=True)
            subset = label.filter((pl.col("window_id") == window) & (pl.col("split") == arm))
            by_class = {r["label"]: r for r in subset.iter_rows(named=True)}
            rows.append([window, arm, f"{total_row['enrollment_count']:,}",
                         f"{by_class['I/D']['label_enrollment_count']:,} ({by_class['I/D']['label_ratio']*100:.3f}%)",
                         f"{by_class['G']['label_enrollment_count']:,} ({by_class['G']['label_ratio']*100:.3f}%)",
                         f"{by_class['E']['label_enrollment_count']:,} ({by_class['E']['label_ratio']*100:.3f}%)"])
    table(doc, ["W", "Split", "Enrollment", "I/D", "G", "E"], rows)
    doc.add_paragraph("Contract đề xuất: small_support_flag = 1 khi n_support <400, tính theo từng ô window × phase × split × subgroup × class. LO W2/W3 phải ghi CI bootstrap 95%, B=2,000 paired bootstrap cho so sánh pipeline, selection_margin; báo cáo s_cal v2 cạnh s_ent v1.")

    heading(doc, "7. Assignment, strictness và offering dài")
    table(doc, ["Assignment QA", "Giá trị"], [
        ["Distinct enrollment", f"{assignment['distinct_enrollment_count']:,}"],
        ["Enrollment đa rolling block", str(assignment["n_enrollments_multiple_rolling_blocks"])],
        ["Enrollment đa split cùng window", str(assignment["n_enrollments_multiple_splits_same_window"])],
        ["Context không map block", str(assignment["n_context_rows_without_block"])],
        ["Kết quả", "Đạt" if assignment["assignment_qa_pass"] == 1 else "Không đạt"],
    ])
    duration_rows = [[r["timeline_source"], f"{r['duration_days_p50']}", f"{r['duration_days_p90']}", f"{r['duration_days_max']}", pct(r["long_offering_enrollment_share"])] for r in duration.sort("timeline_source").iter_rows(named=True)]
    table(doc, ["Timeline", "p50 ngày", "p90 ngày", "Max ngày", "Enrollment >270 ngày"], duration_rows)
    test_p1 = strict.filter((pl.col("split") == "test") & (pl.col("phase_id") == "P1")).sort("window_id")
    table(doc, ["TEST", "W1", "W2", "W3"], [["Strict P1", *[pct(r["temporal_strict_share"]) for r in test_p1.iter_rows(named=True)]]])
    doc.add_paragraph("Giữ đơn vị offering; thêm timeline_source và long_offering_flag vào prediction-level context để phân tầng metric. Không tạo subgroup bắt buộc mới chỉ vì offering dài.")

    heading(doc, "8. Quyết định và cổng đi tiếp")
    table(doc, ["Bước", "Trạng thái LO", "Điều kiện"], [
        ["Khóa split V2.2", "Đạt", "Giữ v2.2 primary; run ghi scenario, split_version, parent manifest/audit và seed."],
        ["Materialize raw phase views/test prefixes", "Được", "Giữ offering dài, mang timeline_source và long_offering_flag như context."],
        ["S0", "Được", "Báo missing_rate/event_sparsity theo A/T1/T2/T3 × timeline_source."],
        ["Feature dictionary + imputation", "Chờ", "Xác nhận final_score range/exam aggregation; áp proxy_feature_contract."],
        ["Augmentation/model", "Chờ", "Sau dictionary LO; dùng cell-level small-support contract."],
    ])
    doc.add_paragraph("Câu hỏi cần xác nhận: total_correct_ratio_exam có phải quy tắc nghiệp vụ được phép tạo final_score >100 không? Nếu có, report gọi score là uncapped activity aggregate. Nếu không, chuẩn hóa exam component và rebuild toàn bộ label/split LO trước downstream.")
    OUT.parent.mkdir(parents=True, exist_ok=True); doc.save(OUT); print(OUT)


if __name__ == "__main__": main()
