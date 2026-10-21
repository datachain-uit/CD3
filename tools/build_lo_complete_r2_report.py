"""Build a complete Vietnamese LO four-family seed-42 report from audited local facts."""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
FACTS = ROOT / "LO/resuilt/lo_recurrent_allfamilies_grid_r2_20261021"
SOURCE = ROOT / "LO/docs/REPORT_LO_TONG_HOP_DAY_DU_SEED42_2026-10-19_ABLATION.docx"
TARGET = ROOT / "LO/docs/REPORT_LO_TONG_HOP_HOAN_CHINH_SEED42_2026-10-21.docx"


def add_table(document: Document, frame: pd.DataFrame, columns: list[tuple[str, str]], size: float = 7.5) -> None:
    table = document.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    for cell, (_, label) in zip(table.rows[0].cells, columns):
        cell.text = label
    for _, row in frame.iterrows():
        cells = table.add_row().cells
        for cell, (name, _) in zip(cells, columns):
            value = row[name]
            cell.text = "" if pd.isna(value) else (f"{value:.4f}" if isinstance(value, float) else str(value))
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(size)


def main() -> None:
    facts = json.loads((FACTS / "summary.json").read_text(encoding="utf-8"))
    document = Document(SOURCE)
    document.add_page_break()
    document.add_paragraph("BẢN TỔNG HỢP HOÀN CHỈNH — LO V3.1, SEED 42")
    document.add_paragraph(
        "Bản này tổng hợp evidence LO clean-grain: bốn họ RNN/LSTM/GRU/BiLSTM, V0–V16, W1–W3, P1–P4 "
        "và hai split TEST/VALIDATION. Có 1.632 cell hiệu năng và 204 resource manifest. "
        "AccTEMPO là metric hiệu năng chính; Macro-F1, balanced accuracy, s_ent và recall lớp được báo cáo để đọc "
        "đúng tình trạng mất cân bằng nhãn mạnh của LO."
    )
    document.add_paragraph("1. Tình trạng audit prediction và temporal slice")
    document.add_paragraph(
        "Full prediction-grain audit gồm 1.632 file (bốn họ × 51 run × TEST/VALIDATION × P1–P4): "
        "tất cả có ID không null/không trùng và cờ STRICT hợp lệ. Slice audit có 1.632 cell; "
        "STRICT + OVERLAP bằng ALL ở toàn bộ cell. Vì vậy bảng bên dưới dùng đúng cohort prediction đã audit."
    )
    add_table(document, pd.DataFrame([{
        "Evaluation cells": facts["evaluation_cells"], "Resource runs": facts["resource_runs"],
        "Prediction files": facts["prediction_grain_files"], "Prediction grain": "Đạt" if facts["prediction_grain_passed"] else "Không đạt",
        "Slice partition": "Đạt" if facts["all_slice_partitions_passed"] else "Không đạt",
    }]), [(x, x) for x in ["Evaluation cells", "Resource runs", "Prediction files", "Prediction grain", "Slice partition"]], size=9)
    document.add_paragraph(
        "Prefix-validation audit cũng đã xác nhận source prefix masking trên toàn mẫu và static parity TEST. "
        "15 cột static lịch sử được ghi rõ trong phụ lục ablation; chúng không được diễn giải tự động là leakage chỉ từ tên cột."
    )

    document.add_page_break()
    document.add_paragraph("2. Bảng đầy đủ hiệu năng theo family × pipeline × window × phase")
    document.add_paragraph(
        "Mỗi bảng có 204 dòng = 17 pipeline × 3 window × 4 phase. STRICT là tỷ lệ prediction STRICT/ALL "
        "theo đúng split và phase; không gộp TEST với VALIDATION."
    )
    cells = pd.read_parquet(FACTS / "lo_family_pipeline_window_phase_metrics.parquet")
    order = {f"V{i}": i for i in range(17)}
    columns = [("pipeline_id", "Pipeline"), ("window_id", "W"), ("phase_id", "P"),
               ("acctempo_m3", "AccTEMPO"), ("macro_f1", "Macro-F1"), ("balanced_acc", "Bal. acc."),
               ("s_ent", "s_ent"), ("c0_recall", "Recall I/D"), ("c1_recall", "Recall G"),
               ("c2_recall", "Recall E"), ("strict_share", "STRICT")]
    for split in ("TEST", "VALIDATION"):
        for family in ("RNN", "LSTM", "GRU", "BILSTM"):
            document.add_paragraph(f"{split} — {family}")
            frame = cells[(cells.eval_split == split) & (cells.model_name == family)].copy()
            frame["_order"] = frame.pipeline_id.map(order)
            add_table(document, frame.sort_values(["_order", "window_id", "phase_id"]).drop(columns="_order"), columns)

    document.add_page_break()
    document.add_paragraph("3. Thời gian train và telemetry")
    document.add_paragraph(
        "Bảng dùng train_seconds = tổng khoảng train theo epoch trong train_history. Các run LO lịch sử này có "
        "energy_measurement = NOT_MEASURED; vì không có telemetry NVML phù hợp, báo cáo không suy ra kWh hoặc CO₂."
    )
    time = pd.read_parquet(FACTS / "lo_train_time_family_pipeline.parquet")
    time["_order"] = time.pipeline_id.map(order)
    add_table(document, time.sort_values(["model_name", "_order"]).drop(columns="_order"), [
        ("model_name", "Họ"), ("pipeline_id", "Pipeline"), ("runs", "Số W"),
        ("train_seconds_mean", "Train TB (s)"), ("train_seconds_sum", "Train tổng (s)"),
        ("batch_size", "Batch"), ("telemetry", "Telemetry"),
    ])
    document.add_paragraph("4. Static ablation và giới hạn diễn giải")
    document.add_paragraph(
        "Ablation W1/LSTM cho V0 và V8 đã có trong report nguồn: baseline historical, no_static và gender + "
        "window_duration_seconds. Nó dùng để đánh giá trực tiếp nhánh static, không thay đổi hay ghi đè lưới 204 run lịch sử. "
        "Kết quả primary vẫn giữ layout lịch sử cho khả năng so sánh; mọi lưới mới phải khóa static contract trước khi chạy."
    )
    document.add_paragraph(
        "Chi tiết máy đọc: LO/resuilt/lo_recurrent_allfamilies_grid_r2_20261021/ chứa parquet/CSV đầy đủ; "
        "review_20261019_window_phase_r2/ giữ evidence slice/sanity và lo_prediction_grain_20261019/ giữ audit prediction-grain."
    )
    document.save(TARGET)
    print(TARGET)


if __name__ == "__main__":
    main()
