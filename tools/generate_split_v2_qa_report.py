"""Generate the formatted Word version of the Split V2 QA incident report."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "SPLIT_V2_QA_INCIDENT_REPORT_2026-09-28.docx"


def shade(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    properties.append(element)


def border(cell, color="D9E2F3"):
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), color)


def set_cell(cell, value, bold=False, color=None):
    cell.text = str(value)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.runs[0]
    run.font.size = Pt(8.5)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    border(cell)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.55)
section.bottom_margin = Inches(0.55)
section.left_margin = Inches(0.62)
section.right_margin = Inches(0.62)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
styles["Normal"].font.size = Pt(9.5)
for style_name, size, color in (("Title", 21, "17365D"), ("Heading 1", 14, "17365D"), ("Heading 2", 11, "1F4E79")):
    styles[style_name].font.name = "Aptos Display"
    styles[style_name].font.size = Pt(size)
    styles[style_name].font.color.rgb = RGBColor.from_string(color)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("BÁO CÁO QA SỰ CỐ SPLIT V2")
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.add_run("CQ và LO · hybrid scenario · đánh giá ngày 28/09/2026").italic = True

box = doc.add_table(rows=1, cols=1)
box.alignment = WD_TABLE_ALIGNMENT.CENTER
cell = box.cell(0, 0)
shade(cell, "FDE9E7")
border(cell, "C00000")
set_cell(cell, "KẾT LUẬN: Không khóa Split V2 hiện tại. Không dùng phase_views_v2 này cho S0, imputation, augmentation, model hoặc meta-release.", bold=True, color="9C0006")

doc.add_heading("1. Quy ước lý thuyết theo tài liệu v1.6", level=1)
doc.add_paragraph("Nguồn quy ước: CHANGE_SPEC v1.6, Mo_ta-meta-dataset v1.6 và tempo_fix/split_v2.py (Mục 2.2, 5.4, QA22).")
add_bullet(doc, "Chia ở cấp offering. A là khoảng 60% enrollment sớm nhất khi xếp theo offering_end_date.")
add_bullet(doc, "Future pool giữ thứ tự thời gian và được cắt thành ba lát có khối lượng enrollment xấp xỉ nhau: T1→B/C, T2→D/E, T3→F/G.")
add_bullet(doc, "Chỉ phân tầng validation/test trong cùng một lát theo nhãn hiếm và quy mô offering. Không san bằng nhãn giữa các lát vì đó là label drift theo thời gian.")
add_bullet(doc, "QA22 dùng cutoff P1 = enroll_time + 25% × (offering_end_date − enroll_time). ev_train_overlap_share phải ≤ 0.10.")
add_bullet(doc, "Nếu cần TEST_SIDE_PURGE và purge loại quá 50% evaluation rows, split không được khóa; phải đổi số lát hoặc biện pháp đã đăng ký.")

doc.add_heading("Giải thích trực quan", level=2)
doc.add_paragraph("V1 vẫn có A là phần quá khứ đầu tiên. Vấn đề là toàn bộ dữ liệu sau A bị trộn rồi chia B–G theo nhãn; vì thế B/C có thể thuộc thời gian muộn hơn D/E. Khi W2 dùng A+B+C để dự đoán D/E, model có thể học từ tương lai để dự đoán quá khứ.")
visual = doc.add_table(rows=3, cols=2)
visual.style = "Table Grid"
visual.alignment = WD_TABLE_ALIGNMENT.CENTER
for row, values in enumerate((
    ("V1", "A (quá khứ) | future pool bị trộn → B, C, D, E, F, G"),
    ("V2", "A | T1: B/C | T2: D/E | T3: F/G"),
    ("Overlap", "Train: bắt đầu tháng 1, kết thúc tháng 9; Test: bắt đầu tháng 6, kết thúc tháng 12 → test đã tồn tại khi train chưa kết thúc → purge."),
)):
    set_cell(visual.rows[row].cells[0], values[0], bold=True, color="FFFFFF")
    shade(visual.rows[row].cells[0], "1F4E79")
    set_cell(visual.rows[row].cells[1], values[1])

doc.add_heading("2. Hiện trạng sau TEST_SIDE_PURGE", level=1)
doc.add_paragraph("Overlap P1 sau purge bằng 0, nhưng điều này đạt được bằng cách loại gần toàn bộ validation/test. Tất cả window còn tập đánh giá đều vượt cờ purge_gt_half_eval.")

headers = ["Task", "Window", "Evaluation trước purge", "Còn lại", "Purge", "Đánh giá"]
rows = [
    ("CQ", "W1", "353,784", "0", "100.00%", "Không hợp lệ"),
    ("CQ", "W2", "360,967", "33,742", "90.65%", "Không hợp lệ"),
    ("CQ", "W3", "360,063", "0", "100.00%", "Không hợp lệ"),
    ("LO", "W1", "391,979", "23,277", "94.06%", "Không hợp lệ"),
    ("LO", "W2", "397,783", "0", "100.00%", "Không hợp lệ"),
    ("LO", "W3", "411,752", "0", "100.00%", "Không hợp lệ"),
]
table = doc.add_table(rows=1, cols=len(headers))
table.style = "Table Grid"
table.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, text in enumerate(headers):
    shade(table.rows[0].cells[i], "1F4E79")
    set_cell(table.rows[0].cells[i], text, bold=True, color="FFFFFF")
for row in rows:
    cells = table.add_row().cells
    for i, value in enumerate(row):
        set_cell(cells[i], value, color="9C0006" if i == 5 else None)
        if i == 5:
            shade(cells[i], "FDE9E7")

doc.add_paragraph("CQ W2 chỉ còn 21,125 validation và 12,617 test rows. LO W1 chỉ còn 115 validation rows và 23,162 test rows; validation này không đủ tin cậy cho model selection.")

doc.add_heading("2.2. Đối chiếu S0 V1 với Split V2", level=2)
doc.add_paragraph("S0 full-census V1 xác nhận mỗi window trước đây có tập evaluation đầy đủ: CQ có 358,272 rows cho validation và 358,272 rows cho test; LO có 400,504 rows cho validation và 400,504 rows cho test. Đây là baseline để đánh giá tác động của split mới, không phải dữ liệu được phép ghép vào meta-release V2.")
compare_headers = ["Task", "S0 V1: eval/window", "V2 trước purge", "V2 sau purge", "Nhận định"]
compare_rows = [
    ("CQ", "716,544", "353,784–360,967", "0–33,742", "W1/W3 rỗng; W2 mất 90.65%"),
    ("LO", "801,008", "391,979–411,752", "0–23,277", "W2/W3 rỗng; W1 mất 94.06%"),
]
compare = doc.add_table(rows=1, cols=len(compare_headers))
compare.style = "Table Grid"
compare.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, text in enumerate(compare_headers):
    shade(compare.rows[0].cells[i], "5B9BD5")
    set_cell(compare.rows[0].cells[i], text, bold=True, color="FFFFFF")
for row in compare_rows:
    cells = compare.add_row().cells
    for i, value in enumerate(row):
        set_cell(cells[i], value, color="9C0006" if i == 4 else None)
        if i == 4:
            shade(cells[i], "FDE9E7")

doc.add_heading("3. Nguyên nhân và tác động", level=1)
doc.add_paragraph("Sort theo offering_end_date không bảo đảm offering của lát sau bắt đầu sau khi train kết thúc. Offering giữa nhiều course và pseudo-offering có các khoảng thời gian chồng lấp; TEST_SIDE_PURGE vì vậy xóa gần hết B–G.")
add_bullet(doc, "split_registry chỉ còn W2 cho CQ và W1 cho LO vì các window khác không còn validation/test.")
add_bullet(doc, "label_drift_window_prev_jsd hiện không diễn giải được khi test window rỗng.")
add_bullet(doc, "phase_views_v2 đã materialize từ manifest này không được đưa vào pipeline tiếp theo.")

doc.add_heading("4. Quyết định và hành động tiếp theo", level=1)
add_bullet(doc, "Không khóa split, không chạy S0/imputation/augmentation/model trên output V2 hiện tại.")
add_bullet(doc, "Thiết kế lại ranh giới bằng temporal cohorts của interval [offering_start_date, offering_end_date] không chồng nhau; hoặc đánh giá TRAIN_EMBARGO có đăng ký trước.")
add_bullet(doc, "Builder mới chỉ cho trạng thái ACCEPTED khi mọi window còn validation/test đủ lớn, ev_train_overlap_share ≤ 0.10 và purge_gt_half_eval = false.")

doc.add_paragraph("Artifact đã kiểm tra: split_registry, split_actual_ratios, split_label_distribution, qa22_before_purge, qa22_after_purge, exclusions_test_side_purge và block_label_audit của CQ/LO.").italic = True
doc.save(OUT)
print(OUT)
