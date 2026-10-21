"""Add an honest v1.6 storage-contract compliance note to the current reports."""
from pathlib import Path

from docx import Document
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]


def append(path: Path, task: str) -> None:
    doc = Document(path)
    marker = "ĐỐI CHIẾU CẤU TRÚC VỚI META-DATASET V1.6"
    if any(marker in p.text for p in doc.paragraphs):
        return
    p = doc.add_paragraph(); r = p.add_run(marker); r.bold = True; r.font.size = Pt(13)
    doc.add_paragraph(
        "Đối chiếu theo bốn tầng của Mo_ta_meta-dataset_v1.6. ‘Đạt’ nghĩa là artifact vật lý hoặc "
        "provenance tương ứng đang được ghi; ‘chưa hoàn tất’ không phải lỗi model mà là hạng mục của "
        "giai đoạn meta-dataset sau model."
    )
    table = doc.add_table(rows=1, cols=3); table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, ["Yêu cầu v1.6", "Trạng thái hiện tại", "Ghi chú"]): cell.text = value
    rows = [
        ["L0: release/split/pipeline/metric/environment registry", "Đạt ở mức provenance", "data_release, split_registry, pipeline_registry, metric_definitions, feature_dictionary theo window, state/run/model registry và environment attempt được ghi."],
        ["L1: run artifact bất biến", "Đạt", "run_manifest, preprocessor/model input, balanced train + synthetic ledger, checkpoint, history, predictions, probability/confusion matrices và timing nằm dưới run_id/attempt_id."],
        ["L2: DQ, fidelity, model metrics, sanity, resource, drift", "Đạt phần đang chạy", "S0/S1/S2/S3 ghi các fact partitioned. Raw probability matrices hiện giữ ở L1 để tái lập; L2 đã có confusion/calibration/metrics nhưng chưa materialize probability-row fact tách riêng."],
        ["L3: meta_dataset_seed_level/seed_agg/dictionary/release manifest", "Chưa hoàn tất", "Hiện có L3 imputation_diagnostic_seed_level. Meta-dataset final chỉ được dựng sau khi khóa đủ task × model × seed × pipeline và QA pairing/decomposition."],
        ["Immutability và định danh", "Đạt", "run_id là hash cấu hình logic; attempt_id là retry đơn điệu. Retry không ghi đè artifact cũ."],
    ]
    for values in rows:
        for cell, value in zip(table.add_row().cells, values): cell.text = value
    doc.add_paragraph(
        "Hệ quả: cấu trúc hiện tại đủ để tiếp tục S3 model và audit/reproduce run. Để đạt trạng thái "
        "meta-release v1.6 hoàn chỉnh cần materialize L2 probability fact (nếu yêu cầu query trực tiếp) "
        "và thực hiện G3–G4 để dựng/khóa L3 final cùng meta_release_manifest."
    )
    doc.save(path)


append(ROOT / "CQ" / "docs" / "REPORT_CQ_V2.2.docx", "CQ")
append(ROOT / "LO" / "docs" / "REPORT_LO_V3_1.docx", "LO")
print("v1.6 compliance notes appended")
