"""Create a clearly labelled snapshot of the retired CQ/LO V1 protocol."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "REPORT_PROTOCOL_LEGACY_CQ_LO_V1.docx"


def title(doc: Document, text: str, size: int = 18) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if size >= 18 else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)


def heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(13)


def table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    tab = doc.add_table(rows=1, cols=len(headers))
    tab.style = "Table Grid"
    for cell, value in zip(tab.rows[0].cells, headers):
        cell.text = value
    for values in rows:
        for cell, value in zip(tab.add_row().cells, values):
            cell.text = value


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(10.5)

title(doc, "SNAPSHOT GIAO THỨC CŨ — CQ V1 VÀ LO V1")
doc.add_paragraph("Mục đích: lưu lại đúng cấu hình và cách chia từng được sử dụng trước khi chuyển sang split V2.2/V3.1. Tài liệu này chỉ để đối chiếu, không phải giao thức được phép dùng cho thực nghiệm/meta-release hiện hành.")

heading(doc, "1. Trạng thái và phạm vi")
table(doc, ["Hạng mục", "CQ legacy V1", "LO legacy V1"], [
    ["Trạng thái", "Retired; được thay bằng split V2.2", "Retired; được thay bằng LO V3.1"],
    ["Scenario", "hybrid", "hybrid"],
    ["Phase", "P1=25%, P2=50%, P3=75%, P4=90%", "P1=25%, P2=50%, P3=75%, P4=90%"],
    ["Đơn vị chia", "offering_id + timeline_source; pseudo-offering cho hybrid fallback", "offering_id + timeline_source; pseudo-offering cho hybrid fallback"],
    ["Mục đích dùng lại", "đối chiếu provenance và giải thích artifact cũ", "đối chiếu provenance và giải thích artifact cũ"],
])

heading(doc, "2. Nhãn và feature regime cũ")
table(doc, ["Thành phần", "CQ legacy V1", "LO legacy V1"], [
    ["Label", "CQ_label_final", "LO_performance_label_3"],
    ["Label version", "cq_vector_proximity_v1", "lo_operational_performance_v1"],
    ["Label artifact", "labels/cq_labels_v1", "labels/lo_labels_v1"],
    ["Ý nghĩa", "nhãn vector-proximity Course Quality", "behavioral performance proxy; không phải final grade, completion hay dropout quan sát trực tiếp"],
    ["Primary regime", "CQ_RAW_EARLY", "LO_FULL_EARLY"],
    ["View/phase version", "wide_prefix_v1", "wide_prefix_v1"],
    ["Imputation split version", "label_stratified_v1", "label_stratified_v1"],
])

heading(doc, "3. Cách chia legacy V1: Method-5")
doc.add_paragraph("A được khóa trước: lấy các offering sớm nhất theo offering_end_date ở ranh giới whole-offering, đạt xấp xỉ 60% enrollment. Các offering còn lại được phân bổ B–G theo allocator deterministic, count-aware để cân bằng số enrollment và nhãn hiếm. B–G không được sắp theo thời gian.")
table(doc, ["Block", "Tỷ trọng mục tiêu", "Cách tạo"], [
    ["A", "60%", "chronological theo offering_end_date"],
    ["B–G", "mỗi block xấp xỉ 6.67%", "label-count stratified; không có chronological ordering"],
])
table(doc, ["Window", "TRAIN", "VALIDATION", "TEST"], [
    ["W1", "A", "B", "C"],
    ["W2", "A + B + C", "D", "E"],
    ["W3", "A + B + C + D + E", "F", "G"],
])

heading(doc, "4. Hệ quả đã biết của V1")
table(doc, ["Nội dung", "Diễn giải"], [
    ["W1", "Vì A là historical block và B/C từ future pool, W1 thỏa thứ tự lịch mong muốn ở mức thiết kế."],
    ["W2 và W3", "B–G là các block cân bằng nhãn nhưng không là các lát thời gian liên tiếp. Vì vậy TRAIN và validation/test có thể nằm trong cùng vùng lịch; không chứng minh strict label-end order."],
    ["Temporal QA", "V1 không có registry chứa đầy đủ window_test_start_utc, window_test_end_utc, embargo/purge, ev_train_overlap_share, n_purged hoặc label_drift_window_prev_jsd."],
    ["Quyết định", "Không gộp run V1 vào kết quả/meta-release hiện tại. V2.2/V3.1 dùng calendar/QA22 đã đăng ký thay thế."],
])

heading(doc, "5. Cấu trúc artifact legacy")
table(doc, ["Layer", "CQ", "LO"], [
    ["L0 registry", "state_manifests; split_manifest_v1", "state_manifests; split_manifest_v1"],
    ["Input/views", "tasks/CQ/hybrid/phase_views_v1; test_prefix_views_v1", "tasks/LO/hybrid/phase_views_v1; test_prefix_views_v1"],
    ["L1 imputation", "split_version=label_stratified_v1; phase_version=wide_prefix_v1", "split_version=label_stratified_v1; phase_version=wide_prefix_v1"],
    ["L2 facts", "facts theo task/regime/window/pipeline/phase", "facts theo task/regime/window/pipeline/phase"],
])

heading(doc, "6. Đối chiếu nhanh với giao thức hiện hành")
table(doc, ["Trục", "Legacy", "Hiện hành"], [
    ["CQ split", "V1 label-stratified blocks", "V2.2 temporal/QA22 release"],
    ["LO label", "operational behavioural proxy V1", "catalog-normalized final-score V3.1, loại no-scored-signal courses"],
    ["LO split", "V1 label-balanced B–G", "V3.1 registry và materialized views khóa release"],
    ["Diễn giải", "chỉ historical/provenance", "được phép dùng cho reporting và model mới"],
])

doc.add_paragraph("Nguồn đối chiếu: task.yaml CQ legacy, experiment_protocol_config.yaml V1, materialize_label_stratified_views.py và các report split V1/V2 còn lưu trong workspace. Snapshot được tạo ngày 2026-10-03.")
doc.save(OUT)
print(OUT)
