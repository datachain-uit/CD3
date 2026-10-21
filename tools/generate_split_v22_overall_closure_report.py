"""Generate the consolidated V2.2 split, overlap and LO-proxy audit report."""
from __future__ import annotations

from pathlib import Path

import polars as pl
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "REPORT_SPLIT_V2_2_OVERALL_CLOSURE_2026-09-29.docx"
NAVY, BLUE, GREY, GREEN, AMBER = "17365D", "1F4E78", "F2F2F2", "E2F0D9", "FFF2CC"


def read(root: Path, name: str) -> pl.DataFrame:
    files = list((root / name).glob("*.parquet"))
    if len(files) != 1:
        raise RuntimeError(f"{root / name}: expected one parquet; found {len(files)}")
    return pl.read_parquet(files[0])


def shade(cell, colour: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), colour)
    props.append(node)


def pshade(paragraph, colour: str) -> None:
    props = paragraph._p.get_or_add_pPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), colour)
    props.append(node)


def table(doc: Document, headers: list[str], rows: list[list[object]], widths: list[float] | None = None) -> None:
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"; t.autofit = False
    for i, text in enumerate(headers):
        cell = t.rows[0].cells[i]; cell.text = ""
        run = cell.paragraphs[0].add_run(str(text))
        run.bold = True; run.font.size = Pt(11.5); run.font.color.rgb = RGBColor(255, 255, 255)
        shade(cell, BLUE)
        if widths: cell.width = Inches(widths[i])
    for j, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            cells[i].paragraphs[0].add_run(str(value)).font.size = Pt(11.5)
            if j % 2: shade(cells[i], GREY)
            if widths: cells[i].width = Inches(widths[i])
    doc.add_paragraph()


def heading(doc: Document, text: str) -> None:
    p = doc.add_heading(text, 1)
    for run in p.runs: run.font.color.rgb = RGBColor.from_string(NAVY)


def pct(value: float) -> str:
    return f"{value * 100:.2f}%"


def main() -> None:
    audit_roots = {task: ROOT / task / "resuilt" / "split_registry_v2_2_overlap_audit_v1_r2" for task in ("CQ", "LO")}
    reg = {task: read(path, "split_registry").sort("window_id") for task, path in audit_roots.items()}
    strict = {task: read(path, "temporal_strict_share") for task, path in audit_roots.items()}
    qa = {task: read(path, "qa22_v2_2") for task, path in audit_roots.items()}
    proxy_root = ROOT / "LO" / "resuilt" / "label_coverage_audit_proxy_v1"
    proxy = read(proxy_root, "slice_timeline_summary")
    concentration = read(proxy_root, "offering_rare_label_summary")
    exclusions = read(proxy_root, "slice_exclusion_reason_summary")

    if not all(qa[t].select(pl.col("qa22_ok").all()).item() for t in qa):
        raise RuntimeError("At least one r2 registry did not pass QA22")

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.7); section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.8); section.right_margin = Cm(1.8)
    doc.styles["Normal"].font.name = "Aptos"; doc.styles["Normal"].font.size = Pt(10.5)
    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("CQ–LO · Split V2.2 closure · 2026-09-29 · Registry-derived").font.size = Pt(8)
    title = doc.add_paragraph(); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("BÁO CÁO TỔNG THỂ: KHÓA SPLIT V2.2 VÀ AUDIT NHÃN PROXY LO")
    run.bold = True; run.font.size = Pt(20); run.font.color.rgb = RGBColor.from_string(NAVY)
    subtitle = doc.add_paragraph(); subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("scenario = hybrid · shared CQ calendar · offering-level assignment · proxy-label LO").italic = True
    verdict = doc.add_paragraph(); verdict.add_run("Kết luận: ").bold = True
    verdict.add_run("Split V2.2 đạt QA22 cho CQ và LO. Audit r2 tách đúng hai khái niệm overlap. Audit proxy xác nhận E/G thấp ở LO T2–T3 là drift của nhãn proxy, không phải mất coverage proxy.")
    pshade(verdict, GREEN)

    heading(doc, "1. Phạm vi và quyết định vận hành")
    doc.add_paragraph(
        "Báo cáo này hợp nhất ba nguồn immutable đã tải về: registry Split V2.2 overlap audit r2 của CQ và LO, cùng label_coverage_audit_proxy_v1 của LO. "
        "LO được diễn giải đúng scope là nhãn hiệu năng hành vi proxy; báo cáo không khẳng định điều gì về final grade quan sát được."
    )
    table(doc, ["Hạng mục", "Quyết định"], [
        ["Scenario", "hybrid"], ["Split chính", "split_version = v2.2; reference contract"],
        ["Đơn vị gán", "offering_id × timeline_source; không tách một offering qua các split"],
        ["Calendar", "CQ anchor; LO dùng chung calendar"], ["LO target", "LO operational behavioural proxy, không phải final grade"],
        ["Purge", "NONE; overlap và strict được lưu để diễn giải"],
        ["Artifact audit", "split_registry_v2_2_overlap_audit_v1_r2 + label_coverage_audit_proxy_v1"],
    ], [2.3, 4.8])

    heading(doc, "2. Cơ chế split và QA22")
    doc.add_paragraph("A là lịch sử sớm theo offering_end_date. Future pool gồm T1=B/C, T2=D/E, T3=F/G; W1=A|B|C, W2=A+B+C|D|E, W3=A–E|F|G. QA22 yêu cầu label-end order và validation/test không rỗng trong V2.2 reference.")
    qa_rows = []
    for task in ("CQ", "LO"):
        for row in qa[task].sort("window_id").iter_rows(named=True):
            qa_rows.append([task, row["window_id"], f"{row['n_enrollments_validation']:,}", f"{row['n_enrollments_test']:,}", "Đạt" if row["end_order_ok"] else "Không đạt", "Đạt" if row["qa22_ok"] else "Không đạt"])
    table(doc, ["Task", "W", "Validation", "Test", "End order", "QA22"], qa_rows, [.55, .35, 1.0, 1.0, .8, .55])

    heading(doc, "3. Overlap và strict: hai metric khác nhau")
    doc.add_paragraph("P1-cutoff overlap đo cutoff P1 trước tau_w, nên là phần bù của strict P1. Offering-start overlap đo cohort có offering bắt đầu trước tau_w; đây là metric độc lập được dùng làm ev_train_overlap_share trong r2.")
    rows = []
    for task in ("CQ", "LO"):
        for row in reg[task].iter_rows(named=True):
            rows.append([task, row["window_id"], pct(row["test_p1_cutoff_overlap_share"]), pct(row["test_offering_start_overlap_share"]), pct(row["validation_p1_cutoff_overlap_share"]), pct(row["validation_offering_start_overlap_share"])])
    table(doc, ["Task", "W", "TEST P1 cutoff", "TEST offering start", "VAL P1 cutoff", "VAL offering start"], rows, [.5, .35, 1.25, 1.45, 1.2, 1.4])
    doc.add_paragraph("Ví dụ W2 có P1-cutoff overlap thấp (CQ 15.11%, LO 4.40%) nhưng offering-start overlap rất cao (CQ 93.01%, LO 92.63%). Vì vậy không được thay thế một metric bằng metric kia.")

    heading(doc, "4. Audit coverage của nhãn proxy LO")
    relevant = proxy.filter(pl.col("temporal_slice").is_in(["A", "T1", "T2", "T3"]))
    proxy_totals = (relevant.group_by("temporal_slice").agg(
        pl.sum("n_enrollments_total").alias("total"), pl.sum("n_proxy_eligible").alias("eligible"),
        pl.sum("n_E").alias("E"), pl.sum("n_G").alias("G"), pl.sum("n_I_D").alias("I_D"))
        .with_columns((pl.col("eligible") / pl.col("total")).alias("eligible_share"),
                      (pl.col("E") / pl.col("eligible")).alias("E_share"),
                      (pl.col("G") / pl.col("eligible")).alias("G_share"))
        .sort("temporal_slice"))
    rows = []
    for row in proxy_totals.iter_rows(named=True):
        rows.append([row["temporal_slice"], f"{row['total']:,}", pct(row["eligible_share"]), f"{row['E']:,}", pct(row["E_share"]), f"{row['G']:,}", pct(row["G_share"])])
    table(doc, ["Lát", "Enrollment", "Proxy eligible", "E", "E share", "G", "G share"], rows, [.45, 1.15, 1.2, .85, .9, .85, .9])
    doc.add_paragraph("Eligibility của proxy luôn cao: A 99.992%, T1 100.000%, T2 99.924%, T3 98.709%. E/G giảm sau T1 không phải do proxy_exclusion_reason tăng mạnh; đó là temporal drift của nhãn proxy đã chấp nhận.")

    heading(doc, "5. Tập trung nhãn hiếm LO ở cấp offering")
    rare = (concentration.group_by("temporal_slice").agg(pl.len().alias("n_offerings"), pl.sum("n_E").alias("E"), pl.sum("n_G").alias("G"), pl.max("n_E").alias("top_E"), pl.max("n_G").alias("top_G"))
        .with_columns((pl.col("top_E") / pl.col("E")).alias("top_E_share"), (pl.col("top_G") / pl.col("G")).alias("top_G_share"))
        .sort("temporal_slice"))
    rows = [[r["temporal_slice"], f"{r['n_offerings']:,}", f"{r['E']:,}", pct(r["top_E_share"]), f"{r['G']:,}", pct(r["top_G_share"])] for r in rare.iter_rows(named=True)]
    table(doc, ["Lát", "Offerings", "E", "E trong offering lớn nhất", "G", "G trong offering lớn nhất"], rows, [.5, 1.0, .75, 1.6, .75, 1.6])
    doc.add_paragraph("T1, T2 và T3 có E/G tập trung cao. Riêng T3 có 2,000/2,343 E (85.36%) trong một offering. Vì anti-leakage cấm tách offering giữa validation/test, pairing không thể cân bằng hoàn toàn nhãn E/G mà vẫn giữ nguyên contract split.")

    heading(doc, "6. Policy đánh giá LO khi support nhỏ")
    table(doc, ["Quy tắc", "Áp dụng"], [
        ["small_support_flag", "1 khi support lớp × split < 100"],
        ["Metric chính", "macro-F1; thêm recall/F1 từng lớp, confusion matrix và support"],
        ["Bất định", "bootstrap confidence interval cho metric có small_support_flag"],
        ["So sánh pipeline LO W2/W3", "paired bootstrap delta, B = 2,000"],
        ["Chọn checkpoint", "validation macro-F1; lưu selection_margin"],
        ["Augmentation", "TRAIN-only; tuyệt đối không augmentation validation/test"],
    ], [2.25, 4.85])
    doc.add_paragraph("LO W2/W3 được phép chạy theo scope proxy, nhưng kết quả lớp hiếm không được dùng đơn lẻ làm kết luận khi có cờ support nhỏ.")

    heading(doc, "7. Gate tiếp theo và cấu trúc lưu trữ")
    table(doc, ["Bước", "CQ", "LO"], [
        ["Khóa registry/protocol V2.2", "Được", "Được"],
        ["Materialize phase views + test prefixes", "Được", "Được"],
        ["S0/data quality và imputation", "Được", "Được"],
        ["Augmentation TRAIN-only", "Được", "Được"],
        ["Model W1", "Được", "Được"],
        ["Model W2/W3", "Được", "Được, kèm small-support policy"],
    ], [2.7, 2.0, 2.4])
    doc.add_paragraph("Cây evidence hiện tại:\n"
        "split_registry_v2_2_overlap_audit_v1_r2/\n├── manifest/ · qa22_v2_2/ · split_registry/\n├── temporal_overlap_audit_v1/ · temporal_strict_context/ · temporal_strict_share/\n└── split_label_distribution/ · duration_summary/ · calendar_boundaries/\n\n"
        "label_coverage_audit_proxy_v1/\n├── slice_timeline_summary/ · slice_exclusion_reason_summary/\n├── offering_rare_label_summary/ · rare_label_concentration/\n└── assignment_coverage_qa/ · audit_manifest/")
    doc.add_paragraph("Để bắt đầu downstream, materializer phải nhận manifest V2.2 đã khóa. Mọi run sau đó phải ghi scenario=hybrid, split_version=v2.2, parent manifest/audit identifier và seed.")

    OUT.parent.mkdir(exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
