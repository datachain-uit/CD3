"""Append evidence from the CQ static, checkpoint and telemetry audits."""
import json
from pathlib import Path
from docx import Document
from docx.shared import Pt

root=Path(__file__).resolve().parents[1]
base=root/'CQ/resuilt/cq_batch2048_closure_audit_20261019'
s=json.loads((base/'summary.json').read_text(encoding='utf-8'))
a=json.loads((base/'static_audit.json').read_text(encoding='utf-8'))
p=json.loads((base/'static_probe.json').read_text(encoding='utf-8'))
doc=Document(root/'CQ/docs/REPORT_CQ_FULL_E2E_SEED42_UPDATED_2026-10-19.docx')
def table(headers, rows):
    t=doc.add_table(rows=1,cols=len(headers)); t.style='Table Grid'
    for c,v in zip(t.rows[0].cells,headers): c.text=str(v)
    for row in rows:
        for c,v in zip(t.add_row().cells,row): c.text=str(v)
    for row in t.rows:
        for c in row.cells:
            for para in c.paragraphs:
                for run in para.runs: run.font.size=Pt(11)
    doc.add_paragraph('')
doc.add_page_break(); doc.add_heading('Phụ lục CQ: static, checkpoint và telemetry — 19/10',0)
doc.add_paragraph('Phụ lục này cập nhật các mục còn mở của phần trước, không thay thế các giới hạn về thời điểm sẵn có đặc trưng, replay theo hash và đánh giá nhiều seed.')
doc.add_heading('1. Audit 13 cột static W1/V0',1)
doc.add_paragraph('Layout LSTM batch 2.048; TRAIN lấy mẫu ngẫu nhiên 400.000 dòng với seed 42, VALIDATION và TEST P1 đầy đủ. Cờ theo course_id_code là thống kê chẩn đoán, không chứng minh availability tại thời điểm dự đoán. Contract CQ hiện vẫn là provisional và cần phê duyệt khoa học.')
table(['Cột','Trạng thái registry','Cờ'],[[r['column'],r['allowlist_status'],', '.join(r['flags']) or '—'] for r in a['columns']])
doc.add_paragraph('primary_risk_set hiện trạng thái REVIEW, được audit biểu thị UNLISTED; split_rule_version thuộc EXCLUDE nhưng nằm trong layout lịch sử và hằng trên mẫu. Chưa suy ra có leakage hoặc tác động dự đoán từ các cờ này. Mã nguồn cho thấy window_duration_seconds được tính từ window_end_exclusive và enroll_time; chưa chứng minh mọi giá trị lịch và metadata catalog đã sẵn có ở thời điểm enrollment.')
doc.add_heading('2. Static-only probe CQ',1)
doc.add_paragraph('Logistic regression, fit trên mẫu TRAIN nêu trên; đánh giá toàn bộ VALIDATION và TEST P1. Ba lớp cố định; probe không chứng minh quan hệ nhân quả với kết quả RNN.')
rr=[]
for group,b in p['feature_sets'].items():
    for w in ('none','balanced'):
        for split,x in b[f'class_weight_{w}'].items():
            rr.append([group,w,split,x['n_rows'],f"{x['macro_f1']:.4f}",f"{x['balanced_accuracy']:.4f}",f"{x['log_loss']:.4f}"])
table(['Nhóm','Weight','Split','N','Macro-F1','Bal. acc.','Log loss'],rr)
doc.add_heading('3. Validation prefix và lựa chọn checkpoint',1)
doc.add_paragraph(f"Đối chiếu toàn bộ {s['n_runs']} manifest/train history: {s['epoch_history_match']} epoch được chọn khớp rule; sai khác lớn nhất giữa mean validation F1 tại epoch chọn và metrics export là {s['max_history_vs_validation_f1_difference']:.2g}. Replay CPU W1/V0 cho từng family trên 512 dòng validation đầu, P1–P4: prediction agreement 100%, sai khác xác suất tối đa khoảng 0,000502; sai khác tối đa logits prefix độc lập so với full sequence dưới 1e-6. Ghép hiện theo thứ tự dòng, label khớp, chưa xác minh enrollment hash do thiếu salt; mẫu và CPU replay không thay thế replay toàn bộ trên GPU cùng môi trường.")
doc.add_heading('4. Telemetry GPU',1)
doc.add_paragraph(f"Cả {s['energy_nonnull']}/{s['n_runs']} manifest recurrent batch 2.048 ghi NVML_POWER_SAMPLING_ESTIMATE, {s['energy_sample_count_range'][0]}–{s['energy_sample_count_range'][1]} mẫu/run, missing sample rate tối đa {s['energy_missing_rate_max']:.1%}. Đây là ước lượng điện năng GPU trong cửa sổ đo, không phải đo điện năng toàn hệ thống; chưa suy CO₂. Không áp kết luận này cho run khác hoặc batch lịch sử.")
doc.add_paragraph('Bằng chứng machine-readable: CQ/resuilt/cq_batch2048_closure_audit_20261019/{summary.json, manifest_history_energy_153.csv, static_audit.json, static_probe.json}.')
out=root/'CQ/docs/REPORT_CQ_FULL_E2E_SEED42_CLOSURE_AUDIT_2026-10-19.docx'; doc.save(out); print(out)
