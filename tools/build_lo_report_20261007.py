from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "LO" / "docs" / "REPORT_LO_V3_1.docx"
OUTPUT = ROOT / "LO" / "docs" / "REPORT_LO_V3.1_CLEAN_GRAIN_EXECUTION_2026-10-07.docx"


def clear(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def run(cell, text: object, bold: bool = False, size: float = 10.5) -> None:
    cell.text = ""
    item = cell.paragraphs[0].add_run(str(text))
    item.bold = bold
    item.font.name = "Times New Roman"
    item._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    item.font.size = Pt(size)


def shade(cell, color: str = "D9EAF7") -> None:
    props = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), color)
    props.append(node)


def table(document: Document, headers: list[str], rows: list[list[object]]) -> None:
    result = document.add_table(rows=1, cols=len(headers))
    result.style = "Table Grid"
    for cell, value in zip(result.rows[0].cells, headers):
        run(cell, value, True)
        shade(cell)
    for values in rows:
        for cell, value in zip(result.add_row().cells, values):
            run(cell, value)
    document.add_paragraph()


def text(document: Document, value: str) -> None:
    document.add_paragraph(value)


def bullet(document: Document, value: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.add_run("• " + value)


def heading(document: Document, value: str, level: int = 1) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(5)
    item = paragraph.add_run(value)
    item.bold = True
    item.font.name = "Times New Roman"
    item._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    item.font.size = Pt(14 if level == 1 else 12)


def main() -> None:
    document = Document(TEMPLATE)
    clear(document)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(11)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("BÁO CÁO LO — RELEASE V3.1 CLEAN-GRAIN\nKẾ HOẠCH TÁI THỰC THI S0/S1/S2/S3")
    sub = document.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run("Phạm vi: learner–course enrollment | LO_V3_1 | cập nhật 07/10/2026")
    text(document, "Đây là report mới, tách khỏi các run LO legacy và các split/reference release trước đó. Nó khóa đúng semantics nhãn final-score catalog-normalized V3.1, split V3.1 và clean-grain input; chỉ artifact khớp đầy đủ release identity mới được phép đi tiếp.")

    heading(document, "Tóm tắt điều hành")
    table(document, ["Hạng mục", "Trạng thái hiện tại", "Kết luận"], [
        ["Release", "LOCKED", "LO_V3_1; split v3_1; phase_views_v3_1_scored_signal_excluded."],
        ["Label", "LOCKED", "catalog-normalized activity-derived final_score; loại course không có scored signal."],
        ["Clean grain", "VERIFIED", "Materializer phải nhận feature V2 không duplicate; one enrollment per row là hard gate."],
        ["S0/S1", "CHƯA CÓ SUCCESS HỢP LỆ", "Volume hiện tại: 0/15 S1 manifest khớp LO_V3_1; legacy attempts bị loại khỏi resolver."],
        ["S2/S3", "PENDING", "Chỉ bắt đầu sau S0 + 15 S1 success và grain audit các model input."],
        ["Outcome chính", "LOCKED", "AccTEMPO-M3; các metric classification/calibration là hỗ trợ."],
    ])

    heading(document, "1. Scope, grain và nguyên tắc anti-leakage")
    text(document, "LO dự báo performance label ở cấp enrollment_id. user_id, course_id, offering/timeline context là metadata hoặc categorical context; chúng không được tạo duplicate analytical rows. Mọi view/model input phải thỏa exact one-row-per-enrollment trong mỗi window × split × phase.")
    bullet(document, "TRAIN-only fit: imputer/scaler chỉ fit trên TRAIN của window hiện hành.")
    bullet(document, "Validation/test không augment, không refit, và chỉ inference qua checkpoint đã chọn từ validation.")
    bullet(document, "TEST P1–P4 tuân thủ prefix availability; future slots không được impute hoặc materialize thành giá trị quan sát.")
    bullet(document, "Mọi legacy run không có release_id/split/phase/label-threshold identity đúng V3.1 bị resolver từ chối.")

    heading(document, "2. Nhãn LO V3.1: final-score catalog-normalized")
    text(document, "Nhãn primary là activity-derived final_score từ các hoạt động được chấm trong course (thành phần video/assignment/final-chapter exam theo catalog khả dụng). Đây không phải observed final grade độc lập, completion hay dropout. Catalog normalization dùng mẫu số activity có điểm có thể quan sát theo course; course không có scored signal bị loại khỏi label universe.")
    table(document, ["Thành phần", "Contract V3.1"], [
        ["Label artifact", "lo_final_score_catalog_normalized_v3_1"],
        ["Decision", "primary_catalog_normalized_activity_score_label__exclude_no_scored_signal_courses"],
        ["Classes", "I/D (default), G, E; threshold set phải đi cùng label artifact/manifest."],
        ["Exclusion", "No scored signal course không có mẫu số hợp lệ; không gán mặc định I/D."],
        ["Required audit", "course_score_ceiling_observed, catalog denominator, eligibility/exclusion theo A/T1/T2/T3."],
    ])
    text(document, "Các tỉ lệ lớp phải luôn xuất theo slice A/T1/T2/T3 và train/validation/test của W1–W3 sau join nhãn. Không báo cáo một label ratio chung thay cho temporal drift hoặc rare-class support từng window.")

    heading(document, "3. Split V3.1 và data release")
    text(document, "LO dùng shared calendar logic đã khóa, với release-specific V3.1 scored-signal exclusion. Unit split là offering/course timeline unit được ghi trong manifest; registry là nguồn duy nhất cho window boundary, purge metadata, label drift và split ratio sau join nhãn.")
    bullet(document, "Không tái dùng split_manifest V1/V2.2/reference cho V3.1 chỉ vì tên thư mục tương tự.")
    bullet(document, "Block/slice audit phải chứng minh enrollment không xuất hiện ở hai block và duration/long-offering được giữ làm context." )
    bullet(document, "Mọi run ghi release_id=LO_V3_1, split_version=v3_1 và phase version V3.1 hiện hành." )

    heading(document, "4. Clean-grain materialization và preflight")
    text(document, "Nguồn cumulative_phase_features_v1 từng có exact duplicate toàn bảng; release clean hiện phải đọc cumulative_phase_features_v2. Materializer mới dừng cứng nếu V2 vẫn có duplicate thay vì để downstream imputation/SMOTE hấp thụ duplication.")
    table(document, ["Gate", "Điều kiện đạt", "Nếu không đạt"], [
        ["Feature source", "V2 clean, before=after exact-deduplication", "Dừng materialize."],
        ["Joined base", "enrollment_id unique sau label/window/manifest join", "Dừng materialize."],
        ["Input release", "phase views và test prefixes cùng LO_V3_1", "Không chạy S0/S1."],
        ["S1 output", "train/validation/test_P1..P4 unique; preprocessor + manifest SUCCESS", "Không augment/model."],
    ])
    text(document, "Lineage cần lưu source path, source count, post-grain count và hashes. Không inject sidecar vào input namespace đang có run thành công, vì sẽ thay đổi data_release_id; sidecar phải được upload cùng lúc khi khóa release tiếp theo.")

    heading(document, "5. S0/S1 imputation contract")
    text(document, "Sau preflight, chạy S0 ALL_ROWS rồi 15 S1 cells (5 variants × W1–W3). Transform áp dụng full TRAIN/VALIDATION/TEST prefix; MICE chỉ giới hạn sample fit, không giới hạn số rows transform/profile.")
    table(document, ["Pipeline", "Variant", "Runtime contract"], [
        ["V0", "zero-fill + masks", "Baseline; masks không impute/scaled."],
        ["V1", "median", "TRAIN pooled P1–P4 fit; transform full rows."],
        ["V5", "mean", "TRAIN pooled P1–P4 fit; transform full rows."],
        ["V9", "ExtraTrees light", "Runtime-light estimator theo registry."],
        ["V13", "MICE", "3 iterations; tối đa 1,000,000 TRAIN rows fit; transform full rows."],
    ])
    text(document, "QA quyết định SUCCESS: future_imputed_cells=0; không unfittable column; train/validation/test_P4 không còn eligible numeric null; residual null ở TEST P1–P3 chỉ được chấp nhận khi là unavailable future phase theo mask contract.")

    heading(document, "6. S2 augmentation và S3 model plan")
    text(document, "Khi và chỉ khi toàn bộ S1 parent success, materialize S2 TRAIN-only với seed augmentation cố định 42. Mapping pipeline đồng nhất CQ: V1→V2/V3/V4; V5→V6/V7/V8; V9→V10/V11/V12; V13→V14/V15/V16 tương ứng CDSMOTE/SASMOTE/RadiusSMOTE.")
    bullet(document, "Synthetic ledger bắt buộc có parent rows, alpha, method, parent manifest hash và release identity.")
    bullet(document, "QA S2: row reconciliation, mask/availability preservation, phase boundary, quality/fidelity facts và Radius fallback-nearest-neighbour rate." )
    bullet(document, "S3: RNN → LSTM/GRU/BiLSTM; GPU L4 + BF16; model seeds 42/43/44, fixed augmentation seed 42." )

    heading(document, "7. Outcome và điều kiện báo cáo")
    text(document, "AccTEMPO-M3 là outcome performance chính trên TEST. Macro-F1, balanced accuracy, MCC, recall mỗi lớp, calibration NLL/Brier/ECE, probability matrix và confusion matrix là output bắt buộc hỗ trợ diễn giải class imbalance. S_san+ v1 là metric data-quality/sanity chính; s_cal v2 là sensitivity, không thay outcome AccTEMPO-M3.")
    bullet(document, "Checkpoint/pipeline shortlist chỉ chọn theo VALIDATION, không theo TEST." )
    bullet(document, "Report phải tách W1/W2/W3 và P1/P2/P3/P4; log support rare class và subgroup flags." )
    bullet(document, "s_leak/s_eff/s_drift sử dụng real rows và recipe facts, không nhân vô nghĩa theo mọi kiến trúc model." )

    heading(document, "8. Storage contract và thứ tự tái thực thi")
    tree = """tempo-data-v1/
├── input/LO_v3_1/
│   ├── release_manifest.json
│   ├── phase_views_v3_1_scored_signal_excluded/
│   └── test_prefix_views_v3_1_scored_signal_excluded/P1..P4/
└── meta_release=imputation-v1/
    ├── L0_registry/                       # S0 state + environment manifests
    ├── L1_runs/task=LO/feature_regime=LO_FULL_EARLY/window_id=W*/
    │   ├── pipeline_id=V0,V1,V5,V9,V13/model_name=IMPUTATION_ONLY/
    │   ├── pipeline_id=V2..V16/model_name=BALANCED_TRAIN_ONLY/
    │   └── pipeline_id=V0..V16/model_name=RNN|LSTM|GRU|BILSTM/
    ├── L2_facts/                          # DQ, imputation fidelity, augmentation/model/sanity/resource facts
    └── L3_meta/                           # aggregate only after paired runs complete"""
    p = document.add_paragraph()
    item = p.add_run(tree)
    item.font.name = "Consolas"
    item.font.size = Pt(8.5)
    table(document, ["Thứ tự", "Việc"], [
        ["1", "Verify LO clean-grain materialization + input release manifest."],
        ["2", "S0 ALL_ROWS cho LO_V3_1."],
        ["3", "S1 V0/V1/V5/V9/V13 × W1–W3; audit 15/15 SUCCESS."],
        ["4", "S2 V2–V16, augmentation seed 42; audit 36/36 SUCCESS."],
        ["5", "S3 RNN seed 42, sau đó paired stability/model grid theo validation shortlist."],
    ])
    text(document, "Kết luận: LO V3.1 đã có contract khoa học và software gates để tái thực thi an toàn, nhưng chưa đủ điều kiện diễn giải model result trên Volume hiện tại. Điều kiện tiếp theo là materialize/upload clean-grain release, hoàn tất S0/S1 thành công và chỉ sau đó mới bắt đầu augmentation/model.")

    document.core_properties.title = "LO V3.1 clean-grain execution report"
    document.core_properties.author = "CD3 experimental pipeline"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
