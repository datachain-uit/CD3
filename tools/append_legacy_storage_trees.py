"""Append task-specific V1 storage trees to the two legacy protocol reports."""
from pathlib import Path

from docx import Document
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]


def append_tree(path: Path, task: str, regime: str, label: str) -> None:
    doc = Document(path)
    marker = "5. Cấu trúc lưu trữ legacy V1"
    if any(p.text == marker for p in doc.paragraphs):
        return
    p = doc.add_paragraph()
    r = p.add_run(marker); r.bold = True; r.font.size = Pt(13)
    doc.add_paragraph(
        "Cây dưới đây là contract lưu trữ cho artifact V1 lịch sử. `{...}` là định danh bất biến; "
        "mỗi retry dùng attempt_id riêng và không ghi đè attempt cũ. L3 chỉ tổng hợp L2, không thay thế L1/L2."
    )
    root_name = "CQ" if task == "CQ" else "LO"
    tree = f'''tempo-data-v1/
├── input/                                      # immutable source supplied to experiments
│   └── {root_name}/
│       ├── phase_views_v1/                     # TRAIN + VALIDATION, full P1–P4
│       ├── test_prefix_views_v1/
│       │   ├── P1/                             # TEST observes P1 only
│       │   ├── P2/                             # TEST observes P1–P2
│       │   ├── P3/                             # TEST observes P1–P3
│       │   └── P4/                             # TEST observes P1–P4
│       └── release_manifest.json               # split/phase/schema/label provenance
│
└── meta_release=imputation-v1/
    ├── L0_registry/                            # immutable registry and provenance
    │   ├── state_manifests/s0_{root_name}.json # S0 quality/profile state
    │   ├── run_registry/run_id={{run_id}}.json
    │   ├── model_registry/model_name=RNN/...
    │   └── environments/attempt_id={{attempt_id}}.json
    │
    ├── L1_runs/                                # immutable executable artifacts
    │   └── task={task}/feature_regime={regime}/window_id=W1/
    │       ├── pipeline_id=V0/
    │       │   ├── model_name=IMPUTATION_ONLY/seed=20260922/
    │       │   │   └── run_id={{imputation_sha}}/attempt_id={{attempt}}-NNNN/
    │       │   │       ├── run_manifest.json
    │       │   │       ├── fitted_preprocessors/
    │       │   │       └── model_inputs/
    │       │   │           ├── train.parquet
    │       │   │           ├── validation.parquet
    │       │   │           └── test_P1.parquet ... test_P4.parquet
    │       │   └── model_name=RNN/seed=42/
    │       │       └── run_id={{model_sha}}/attempt_id={{attempt}}-NNNN/
    │       │           ├── run_manifest.json
    │       │           ├── checkpoint/model.pt
    │       │           ├── train_history.parquet + feature_layout.json
    │       │           ├── predictions/
    │       │           │   ├── eval_split=VALIDATION/phase_id=P1 ... P4/
    │       │           │   └── eval_split=TEST/phase_id=P1 ... P4/
    │       │           ├── probability_matrices/{{VALIDATION,TEST}}/phase_id=P{{1..4}}/
    │       │           ├── confusion_matrices/{{VALIDATION,TEST}}/phase_id=P{{1..4}}/
    │       │           └── timing.parquet
    │       └── phase_id=P4/pipeline_id=V2/     # augmentation only uses TRAIN
    │           └── model_name=BALANCED_TRAIN_ONLY/seed=42/
    │               └── run_id={{augmentation_sha}}/attempt_id={{attempt}}-NNNN/
    │                   ├── run_manifest.json
    │                   ├── balanced_train.parquet
    │                   ├── synthetic_ledger.parquet
    │                   └── augmentation_quality.json
    │
    ├── L2_facts/                               # queryable facts; partitioned by run/phase
    │   ├── dataset_quality_summary/             # before/after imputation quality
    │   ├── dq_feature_profile/  imputation_fidelity/
    │   ├── augmentation_quality/  resource_usage/
    │   ├── metrics_overall/  metrics_per_class/  calibration/
    │   ├── confusion_matrix/  checkpoint_selection/
    │   ├── sanity_components/  drift_measures/
    │   └── .../task={task}/feature_regime={regime}/window_id=W1/
    │       └── phase_id=P1/.../run_id={{...}}/attempt_id={{...}}/
    │
    └── L3_meta/                                # aggregated meta-dataset; never replaces L1/L2
        └── task={task}/label={label}/...'''
    p = doc.add_paragraph()
    run = p.add_run(tree)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    doc.add_paragraph(
        "Lưu ý historical: các path `input/" + root_name + "/..._v1` và `wide_prefix_v1` chỉ mô tả release cũ. "
        + ("CQ hiện hành khóa ở input/CQ_v2_2 và split V2.2." if task == "CQ" else
           "LO hiện hành khóa ở input/LO_v3_1, label catalog-normalized và split/view V3.1.")
    )
    doc.save(path)


append_tree(ROOT / "CQ" / "docs" / "REPORT_CQ_PROTOCOL_V1_LEGACY.docx", "CQ", "CQ_RAW_EARLY", "CQ_label_final")
append_tree(ROOT / "LO" / "docs" / "REPORT_LO_PROTOCOL_V1_LEGACY.docx", "LO", "LO_FULL_EARLY", "LO_performance_label_3")
print("legacy storage trees appended")
