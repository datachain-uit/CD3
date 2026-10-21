"""Replace the provisional CQ V0 validation row after its W1 rerun."""
from pathlib import Path

from docx import Document


report = Path(__file__).resolve().parents[1] / "CQ" / "docs" / "REPORT_CQ_V1.docx"
doc = Document(report)
old = "V0 — zero-fill + masks"
found = False
for table in doc.tables:
    for row in table.rows[1:]:
        if row.cells and row.cells[0].text.strip() == old:
            # Mean of immutable best validation Macro-F1: W1=.7736596454,
            # W2=.6658 and W3=.8408.
            values = [old, "0.7601", "0.7737", "0.6658", "0.8408", "Giữ baseline"]
            if len(row.cells) != len(values):
                raise RuntimeError("Unexpected CQ validation table schema")
            for cell, value in zip(row.cells, values):
                cell.text = value
            found = True

if not found:
    raise RuntimeError("CQ V0 validation row was not found")

for paragraph in doc.paragraphs:
    if "W1 của V0 chạy trước khi epoch telemetry" in paragraph.text:
        paragraph.text = (
            "V0 là baseline bắt buộc. W1–V0 đã được chạy lại thành công với epoch logging đầy đủ: "
            "Validation Macro-F1 trung bình P1–P4 = 0.7737 tại epoch 8. Vì vậy bảng validation "
            "của V0 hiện dùng đầy đủ W1–W3, trung bình 0.7601; TEST artifact được giữ theo attempt mới."
        )

doc.save(report)
print(report)
