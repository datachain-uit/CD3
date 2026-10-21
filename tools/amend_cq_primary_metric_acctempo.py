"""Amend the CQ final assessment: AccTEMPO-M3 is the primary outcome metric."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt


REPORT = Path(__file__).resolve().parents[1] / "CQ" / "docs" / "REPORT_CQ_V1.docx"


def add_heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(14)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(5)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    output = doc.add_table(rows=1, cols=len(headers))
    output.style = "Table Grid"
    for cell, value in zip(output.rows[0].cells, headers):
        cell.text = value
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in rows:
        for cell, value in zip(output.add_row().cells, row):
            cell.text = str(value)
    doc.add_paragraph()


def main() -> None:
    doc = Document(REPORT)
    if any("C.5. Thứ bậc metric chính" in paragraph.text for paragraph in doc.paragraphs):
        print("AccTEMPO amendment already present; no duplicate appended.")
        return

    # Correct the two statements in the previous final section that promoted
    # Macro-F1 beyond its legitimate role (checkpoint-selection support).
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("Trong TEST đã khóa, V2"):
            paragraph.text = (
                "Theo độ đo chính AccTEMPO-M3 trên TEST đã khóa, V14 (MICE + CDSMOTE) cao nhất 0.7465; "
                "V6 = 0.7418, V10 = 0.7409 và V2 = 0.7345. Đây là mô tả hậu nghiệm của grid. "
                "Không dùng TEST để chọn pipeline hoặc checkpoint."
            )
        if paragraph.text.startswith("CQ đủ điều kiện đi tiếp sang pha tái lặp mô hình"):
            paragraph.text = (
                "CQ đủ điều kiện đi tiếp sang pha tái lặp mô hình. AccTEMPO-M3 là độ đo chính để đánh giá kết quả; "
                "Macro-F1, balanced accuracy, MCC và các chỉ số xác suất là độ đo hỗ trợ. V14 hiện dẫn đầu AccTEMPO-M3 "
                "trên TEST seed 42, nhưng chưa được tuyên bố mô hình thắng cuộc vì quyết định cuối phải dựa trên "
                "AccTEMPO-M3 của VALIDATION đa seed, sau đó mới mở TEST để báo cáo."
            )

    add_heading(doc, "C.5. Thứ bậc metric chính: AccTEMPO-M3")
    doc.add_paragraph(
        "Từ thời điểm khóa báo cáo này, AccTEMPO-M3 là PRIMARY outcome metric. Công thức phiên bản "
        "acctempo_m3_v1_jsd_base2_six_components kết hợp S_perf (0.6) và S_san+ (0.4); vì vậy nó đồng thời "
        "phản ánh chất lượng phân loại và các điều kiện prediction sanity đã đăng ký. Accuracy không là metric chính."
    )
    add_table(doc, ["Mức sử dụng", "Metric", "Quy tắc"], [
        ["PRIMARY — kết quả", "AccTEMPO-M3", "Báo riêng từng V × W × P; tổng hợp sau đó theo protocol đã khóa."],
        ["Hỗ trợ phân loại", "Macro-F1, balanced accuracy, MCC, PR/ROC-AUC", "Giải thích hiệu năng lớp hiếm và xác suất; không thay AccTEMPO-M3."],
        ["Hỗ trợ sanity", "S_perf, S_san+, S_cal, s_leak", "Giải thích cấu phần AccTEMPO-M3 và kiểm tra contract."],
        ["Checkpoint hiện hữu", "Validation Macro-F1", "Chỉ là objective kỹ thuật đã dùng để chọn epoch trong grid seed 42; không phải metric chính của kết luận."],
    ])
    add_table(doc, ["Xếp hạng TEST seed 42 theo AccTEMPO-M3", "Giá trị", "Diễn giải"], [
        ["V14 — MICE + CDSMOTE", "0.7465", "Cao nhất toàn bộ 12 TEST cell; dẫn đầu W1 và W3."],
        ["V6 — Mean + CDSMOTE", "0.7418", "Thứ hai."],
        ["V10 — ExtraTrees + CDSMOTE", "0.7409", "Thứ ba; dẫn đầu W2."],
        ["V2 — Median + CDSMOTE", "0.7345", "Thứ tư; trước đó chỉ dẫn theo TEST Macro-F1."],
    ])
    doc.add_paragraph(
        "Quyết định anti-leakage: không chuyển V14 thành winner cuối chỉ từ TEST. Sanity job đã ghi AccTEMPO-M3 cho VALIDATION, "
        "nhưng export hiện tại chỉ tải bản TEST. Trước vòng seed 43/44, cần export các VALIDATION AccTEMPO facts và dùng "
        "mean Validation AccTEMPO-M3 P1–P4 làm metric shortlist; TEST giữ nguyên vai trò xác nhận cuối."
    )
    doc.save(REPORT)
    print(REPORT)


if __name__ == "__main__":
    main()
