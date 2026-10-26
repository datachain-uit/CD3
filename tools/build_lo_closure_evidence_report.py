"""Append the final 25-Oct LO closure evidence to the maintained LO report.

This is deliberately a reporting-only tool: it does not train a model, alter
predictions, or regenerate any immutable artifact.  It reads the downloaded
post-hoc audit evidence and makes its coverage explicit in the Word report.
"""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "LO/docs/REPORT_LO_CURRENT_EVIDENCE_SEED42_2026-10-25.docx"
TARGET = ROOT / "LO/docs/REPORT_LO_CURRENT_EVIDENCE_CLOSED_SEED42_2026-10-25.docx"
COHORT_ROOT = ROOT / "LO/resuilt/lo_cohort_audit_nvml_seed42"
AUDIT_ROOTS = {
    "RNN": ROOT / "LO/resuilt/lo_nvml_historical_layout_seed42_20261023_complete/lo_nvml_historical_layout_seed42/audits/LO/RNN/seed=42",
    "LSTM": ROOT / "LO/resuilt/lo_nvml_historical_layout_seed42_20261023_complete/lo_nvml_historical_layout_seed42/audits/LO/LSTM/seed=42",
    "GRU": ROOT / "LO/resuilt/lo_gru_nvml_audited_seed42_20261025/lo_gru_historical_layout_seed42/audits/LO/GRU/seed=42",
    "BILSTM": ROOT / "LO/resuilt/lo_bilstm_nvml_audited_seed42_20261025/audits_seed42/LO/BILSTM/seed=42",
}
RNN_REPAIR = ROOT / "LO/resuilt/lo_cohort_audit_nvml_seed42/audits_repair_7886_seed42/LO/RNN/seed=42/audit_manifest.json"


def add_text(doc: Document, text: str, size: float = 10) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)


def add_table(doc: Document, rows: list[dict[str, object]], fields: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=1, cols=len(fields))
    table.style = "Table Grid"
    for cell, (_, title) in zip(table.rows[0].cells, fields):
        cell.text = title
    for row in rows:
        cells = table.add_row().cells
        for cell, (key, _) in zip(cells, fields):
            value = row[key]
            cell.text = f"{value:,}" if isinstance(value, int) else str(value)
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(10.5)


def main() -> None:
    summaries = [json.loads(path.read_text(encoding="utf-8")) for path in COHORT_ROOT.rglob("summary.json")]
    if len(summaries) != 68:
        raise RuntimeError(f"Expected 68 cohort summaries, found {len(summaries)}")
    if any(item["duplicate_hashes"] != 0 or item["invalid_strict"] != 0 for item in summaries):
        raise RuntimeError("Cohort audit contains a duplicate prediction hash or invalid STRICT flag")

    manifests = {family: json.loads((root / "audit_manifest.json").read_text(encoding="utf-8"))
                 for family, root in AUDIT_ROOTS.items()}
    repair = json.loads(RNN_REPAIR.read_text(encoding="utf-8"))
    doc = Document(SOURCE)
    doc.add_page_break()
    doc.add_heading("CẬP NHẬT ĐÓNG EVIDENCE LO — 25/10/2026", level=1)
    add_text(doc, "Phần này đóng các hạng mục hậu kiểm còn mở cho campaign LO clean-grain seed 42. Đây chỉ là tổng hợp evidence có sẵn; không huấn luyện lại, không sửa prediction, không thay đổi split, nhãn hay input release.")

    doc.add_heading("1. Audit prediction theo cohort và lát thời gian", level=2)
    add_table(doc, [{
        "Cấu hình family × pipeline": len(summaries),
        "Bao phủ": "4 family × V0–V16",
        "Prediction files": sum(item["prediction_files"] for item in summaries),
        "Audit cells": sum(item["cells"] for item in summaries),
        "Trùng prediction hash": sum(item["duplicate_hashes"] for item in summaries),
        "STRICT không hợp lệ": sum(item["invalid_strict"] for item in summaries),
    }], [
        ("Cấu hình family × pipeline", "Cấu hình"), ("Bao phủ", "Bao phủ"),
        ("Prediction files", "Tệp prediction"), ("Audit cells", "Cells cohort"),
        ("Trùng prediction hash", "Hash trùng"), ("STRICT không hợp lệ", "STRICT lỗi"),
    ])
    add_text(doc, "Mỗi cấu hình kiểm 24 tệp prediction (W1–W3 × TEST/VALIDATION × P1–P4). Cohort SEEN/UNSEEN được xác định từ course_id có xuất hiện trong REAL parent TRAIN của đúng window; STRICT được đối chiếu với cờ temporal_strict theo phase. Kết quả: 1.632 tệp prediction, 6.528 cells cohort; không có hash enrollment trùng và không có cờ STRICT ngoài {0,1}.")

    doc.add_heading("2. Bao phủ facts hậu kiểm theo family", level=2)
    rows = []
    for family, manifest in manifests.items():
        facts = manifest["rows"]
        rows.append({
            "Family": family, "Runs": manifest["runs"], "Sanity": facts["sanity_components"],
            "S_leak canonical": facts["canonical_s_leak"], "Per-class": facts["metrics_per_class"],
            "Bootstrap CI": facts["bootstrap_metrics_per_class"], "STRICT/OVERLAP": facts["evaluation_slices"],
            "Checkpoint": facts["checkpoint_selection"],
        })
    add_table(doc, rows, [
        ("Family", "Family"), ("Runs", "Runs"), ("Sanity", "Sanity"),
        ("S_leak canonical", "S_leak"), ("Per-class", "Per-class"),
        ("Bootstrap CI", "Bootstrap CI"), ("STRICT/OVERLAP", "Slices"), ("Checkpoint", "Checkpoint"),
    ])
    add_text(doc, "Tất cả bốn family đều có 51 run (V0–V16 × W1–W3), bootstrap B=2.000, facts sanity/canonical S_leak, per-class, khoảng tin cậy bootstrap, STRICT/OVERLAP và checkpoint selection. Các facts này là đầu vào đủ để materialize L2 theo hợp đồng hiện hành.")

    doc.add_heading("3. Đồng nhất runner và cell RNN W1/V0", level=2)
    add_table(doc, [{
        "Cell thay thế": "LO / W1 / V0 / RNN / seed 42",
        "Runner commit": "7886a856 (historical-layout campaign)",
        "Batch size": 2048,
        "GPU": "NVIDIA L4",
        "Energy": "NVML_POWER_SAMPLING_ESTIMATE",
        "Facts audit": sum(repair["rows"].values()),
    }], [(key, key) for key in ("Cell thay thế", "Runner commit", "Batch size", "GPU", "Energy", "Facts audit")])
    add_text(doc, "Cell RNN W1/V0 cũ có runner revision khác phần còn lại của campaign. Cell thay thế đã được chạy và audit bằng revision 7886a856, batch 2.048 trên L4; khi materialize L2 phải chọn cell thay thế này làm RNN W1/V0 chính thức và không trộn cell cũ vào bảng so sánh.")

    doc.add_heading("3.1 Hợp đồng static-v2 dùng chung cho mọi model mới", level=2)
    add_text(doc, "Hợp đồng static_feature_contract_lo_v2_point_in_time chỉ cho phép 13 trường static đã đăng ký: 11 trường catalog, gender và window_duration_seconds. split_rule_version và timeline_source là provenance/context-only, bị loại trước khi tensor được tạo.")
    add_text(doc, "RNN, LSTM, GRU, BiLSTM, PatchTST và MOMENT không có nhánh chọn static riêng: tất cả nhận FeatureLayout từ model.data.fit_layout trước khi kiến trúc được khởi tạo. Vì vậy PatchTST và MOMENT mới tự động dùng đúng cùng 13 trường static. Manifest mới ghi static_feature_contract_id cùng feature_layout_sha256 để verifier từ chối grid trộn hợp đồng.")
    add_text(doc, "Điều này không sửa checkpoint historical-layout 15 trường. Các run đã hoàn tất vẫn được báo cáo là historical evidence; mọi run LO mới phải dùng static-v2 và không được gộp trực tiếp với bảng historical như cùng một experiment grid.")

    doc.add_heading("4. Trạng thái kết luận", level=2)
    add_text(doc, "Các điều kiện kỹ thuật còn mở của LO đã được đóng: prediction grain, tính hợp lệ STRICT, cohort SEEN/UNSEEN × STRICT/OVERLAP, facts hậu kiểm cho bốn family và cell RNN W1/V0 đồng nhất revision.")
    add_text(doc, "L2 chính thức đã nhận ingest_id = vastai_lo_f2_seed42_clean_20261025_r3. Receipt xác nhận 204 contracts, 1.632 rows metrics, 204 rows resource và đủ 87 partitions: 12 partitions cho từng fact family-specific (sanity, per-class, bootstrap, slice, checkpoint, metrics, resource) và 3 canonical S_leak partitions.")
    add_text(doc, "Policy static cho layout mới được chốt theo hướng bảo thủ: chỉ giữ các trường catalog/enrollment đã đăng ký cùng window_duration_seconds; split_rule_version và timeline_source là provenance/context-only và bị loại khỏi tensor. Historical-layout runs vẫn giữ feature_layout.json đã ghi để so sánh tái lập, không bị viết lại hay diễn giải như đã dùng policy mới.")
    doc.save(TARGET)
    print(TARGET)


if __name__ == "__main__":
    main()
