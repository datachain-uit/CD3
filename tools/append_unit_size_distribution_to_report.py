"""Append the downloaded CQ/LO unit-size audit tables to the V2.2 report."""
from __future__ import annotations

from pathlib import Path

import polars as pl
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "docs" / "REPORT_SPLIT_V2_2_OVERALL_CLOSURE_2026-09-29.docx"
BLUE, GREY = "1F4E78", "F2F2F2"


def read(task: str) -> pl.DataFrame:
    folder = ROOT / task / "resuilt" / "long_offering_assignment_audit_v1" / "unit_size_distribution"
    files = list(folder.glob("*.parquet"))
    if len(files) != 1:
        raise RuntimeError(f"{folder}: expected exactly one parquet, found {len(files)}")
    return pl.read_parquet(files[0]).sort(["temporal_slice", "timeline_source"])


def shade(cell, colour: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd"); node.set(qn("w:fill"), colour); props.append(node)


def table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    item = doc.add_table(rows=1, cols=len(headers)); item.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = item.rows[0].cells[index]; cell.text = ""
        run = cell.paragraphs[0].add_run(header); run.bold = True; run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(255, 255, 255); shade(cell, BLUE)
    for row_index, row in enumerate(rows):
        cells = item.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = ""; cells[index].paragraphs[0].add_run(value).font.size = Pt(10.5)
            if row_index % 2: shade(cells[index], GREY)
    doc.add_paragraph()


def short_timeline(value: str) -> str:
    return {"course_specific_timeline": "course-specific", "enrollment_anchored_proxy": "anchored proxy",
            "global_template_fallback": "global fallback"}[value]


def main() -> None:
    if not REPORT.exists(): raise FileNotFoundError(REPORT)
    doc = Document(REPORT)
    title = doc.add_heading("8.7 Phân bố kích thước đơn vị gán", level=2)
    for run in title.runs: run.font.color.rgb = RGBColor.from_string("17365D")
    doc.add_paragraph(
        "Một đơn vị gán là offering_id × timeline_source. Bảng tách timeline để tránh diễn giải sai: "
        "nhiều đơn vị nhỏ không đồng nghĩa nhiều enrollment nằm trong đơn vị nhỏ."
    )
    for task in ("CQ", "LO"):
        frame = read(task)
        rows = []
        for r in frame.iter_rows(named=True):
            rows.append([r["temporal_slice"], short_timeline(r["timeline_source"]),
                         f"{r['offering_timeline_count']:,}", f"{r['unit_enrollment_p50']:,}",
                         f"{r['unit_enrollment_p90']:,}", f"{r['unit_enrollment_max']:,}",
                         f"{r['share_units_le_5'] * 100:.2f}%", f"{r['share_enrollments_units_le_5'] * 100:.3f}%"])
        doc.add_heading(f"{task}: unit size theo lát và timeline", level=3)
        table(doc, ["Slice", "Timeline", "Units", "p50", "p90", "Max", "Units ≤5", "Enrollment trong units ≤5"], rows)
    doc.add_paragraph(
        "Kết luận: tồn tại nhiều unit ≤5 ở enrollment-anchored proxy (ví dụ CQ A 54.55%, CQ T3 51.86%), "
        "nhưng chúng chỉ chứa lần lượt 1.713% và 2.526% enrollment của timeline tương ứng. Với course-specific timeline, "
        "tỷ lệ enrollment trong unit ≤5 luôn rất thấp. Vì vậy validation/test gần trùng phân bố không thể được quy hoàn toàn "
        "cho các unit cực nhỏ; diễn giải seen_course vẫn phải báo cáo theo timeline_source."
    )
    doc.save(REPORT); print(REPORT)


if __name__ == "__main__": main()
