"""Append the completed locked LO V3.1 RNN seed-42 facts to the main report."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "LO" / "docs" / "REPORT_LO_V3_1.docx"
FACTS = next((ROOT / "downloads" / "analysis" / "lo_rnn_seed42_v3_1_files").iterdir())


def fmt(value: object, digits: int = 4) -> str:
    return "—" if pd.isna(value) else f"{float(value):.{digits}f}"


def pipe_order(series: pd.Series) -> list[str]:
    return sorted(series.unique(), key=lambda item: int(str(item)[1:]))


def add_title(doc: Document, text: str, size: int = 14) -> None:
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    for row in rows:
        for cell, value in zip(table.add_row().cells, row):
            cell.text = value


def family(pipeline: str) -> str:
    idx = int(pipeline[1:])
    if idx == 0:
        return "V0: zero-fill + missing/availability masks"
    parents = {1: "Median", 5: "Mean", 9: "ExtraTrees", 13: "MICE"}
    suffixes = {0: "không augment", 1: "CDSMOTE", 2: "SASMOTE", 3: "RadiusSMOTE"}
    parent = idx - ((idx - 1) % 4)
    return f"{parents[parent]} + {suffixes[idx - parent]}"


metrics = pd.read_parquet(FACTS / "metrics_test_long.parquet")
sanity = pd.read_parquet(FACTS / "sanity_test_long.parquet")
resource = pd.read_parquet(FACTS / "resource_long.parquet")
runs = pd.read_parquet(FACTS / "run_inventory.parquet")

if (len(metrics), len(sanity), len(resource), len(runs)) != (204, 204, 51, 51):
    raise RuntimeError("LO export is incomplete; refuse to update the report.")

doc = Document(REPORT)

# Correct the provisional completion note added before W3/V16 completed.
for paragraph in doc.paragraphs:
    if paragraph.text.startswith("Tổng: 50/51 model cell"):
        paragraph.text = (
            "Tổng: 51/51 model cell đã SUCCESS và có checkpoint, train history, "
            "validation/test predictions P1–P4, probability matrix, confusion matrix "
            "và TEST artifact. Không có cell bị loại khỏi tổng hợp."
        )
    elif paragraph.text.startswith("Cấu trúc và metric tuân theo CQ:"):
        paragraph.text = (
            "Cấu trúc và metric tuân theo CQ: V0 baseline masks; V1/V5/V9/V13 là "
            "median/mean/ExtraTrees/MICE; V2–V4, V6–V8, V10–V12, V14–V16 là "
            "TRAIN-only CDSMOTE/SASMOTE/RadiusSMOTE. AccTEMPO-M3 là PRIMARY outcome "
            "metric; Macro-F1, balanced accuracy, MCC, PR/ROC-AUC và calibration là "
            "metric hỗ trợ. L2 facts đã được export và kiểm tra đủ 204 TEST + 204 sanity cell."
        )
for table in doc.tables:
    if len(table.columns) == 4 and table.rows and table.cell(0, 0).text == "Window":
        for row in table.rows[1:]:
            if row.cells[0].text == "W3":
                row.cells[1].text, row.cells[2].text, row.cells[3].text = "17/17", "—", "Đủ V0–V16"

doc.add_page_break()
add_title(doc, "KẾT QUẢ RNN SEED 42 — LO V3.1 (HOÀN TẤT)")
doc.add_paragraph(
    "Nguồn facts bất biến: L2 report export RNN_seed42_v3_1, tạo ngày 2026-10-03. "
    "Khóa release: split_version=v3_1; phase_version=wide_prefix_v3_1; nhãn "
    "lo_final_score_catalog_normalized_v3_1. Mỗi cell sử dụng shared recurrent model P1–P4, "
    "NVIDIA L4, bf16 AMP, seed=42. TRAIN chỉ dùng dữ liệu train của chính window; validation "
    "chọn checkpoint; TEST chỉ inference và đánh giá."
)
add_table(doc, ["Kiểm tra", "Kết quả"], [
    ["Model cells", "51/51 SUCCESS (V0–V16 × W1–W3)"],
    ["TEST performance facts", "204/204 = 51 cells × P1–P4"],
    ["TEST sanity facts", "204/204 = 51 cells × P1–P4"],
    ["Resource facts", "51/51"],
    ["Test rows / phase", "W1: 186,328; W2: 390,978; W3: 252,432"],
])

add_title(doc, "1. AccTEMPO-M3 — độ đo PRIMARY", 12)
doc.add_paragraph(
    "AccTEMPO-M3 là outcome metric chính để báo cáo chất lượng cuối. Bảng dưới là TEST, "
    "mỗi giá trị là một run cụ thể; chúng dùng để mô tả kết quả khóa. Không dùng TEST để chọn "
    "cấu hình/seed ở các vòng tiếp theo: shortlist phải dùng Validation AccTEMPO-M3, rồi TEST "
    "chỉ xác nhận một lần."
)
for phase in ["P1", "P2", "P3", "P4"]:
    subset = sanity[sanity.phase_id == phase].pivot(index="pipeline_id", columns="window_id", values="acctempo_m3")
    subset = subset.reindex(pipe_order(sanity.pipeline_id))
    rows = [[pipe, family(pipe), *(fmt(subset.loc[pipe, window]) for window in ["W1", "W2", "W3"])] for pipe in subset.index]
    add_title(doc, f"AccTEMPO-M3 TEST — {phase}", 11)
    add_table(doc, ["Pipeline", "Tiền xử lý / augment", "W1", "W2", "W3"], rows)

overall = sanity.groupby("pipeline_id", as_index=False).acctempo_m3.mean().sort_values("acctempo_m3", ascending=False)
add_title(doc, "2. Tổng hợp AccTEMPO-M3 trên 12 TEST cells", 12)
add_table(doc, ["Hạng", "Pipeline", "Cấu hình", "AccTEMPO-M3 mean"], [
    [str(rank), row.pipeline_id, family(row.pipeline_id), fmt(row.acctempo_m3)]
    for rank, (_, row) in enumerate(overall.iterrows(), start=1)
])

winner = sanity.loc[sanity.groupby(["window_id", "phase_id"]).acctempo_m3.idxmax()].sort_values(["window_id", "phase_id"])
add_title(doc, "3. Pipeline cao nhất theo window và phase", 12)
add_table(doc, ["Window", "Phase", "Pipeline", "Cấu hình", "AccTEMPO-M3"], [
    [row.window_id, row.phase_id, row.pipeline_id, family(row.pipeline_id), fmt(row.acctempo_m3)]
    for _, row in winner.iterrows()
])

core = metrics.groupby("pipeline_id", as_index=False)[
    ["f1_macro", "balanced_accuracy", "mcc", "pr_auc_macro", "roc_auc_macro_ovr", "multiclass_nll", "multiclass_brier", "top_label_ece_15"]
].mean()
core = core.set_index("pipeline_id").loc[pipe_order(metrics.pipeline_id)].reset_index()
add_title(doc, "4. Performance TEST hỗ trợ (mean của W1–W3 và P1–P4)", 12)
add_table(doc, ["Pipeline", "Macro-F1", "Balanced Acc.", "MCC", "PR-AUC", "ROC-AUC", "NLL", "Brier", "ECE-15"], [
    [row.pipeline_id, *(fmt(row[field]) for field in ["f1_macro", "balanced_accuracy", "mcc", "pr_auc_macro", "roc_auc_macro_ovr", "multiclass_nll", "multiclass_brier", "top_label_ece_15"])]
    for _, row in core.iterrows()
])

sanity_mean = sanity.groupby("pipeline_id", as_index=False)[
    ["s_nan", "s_maj_jsd", "s_ent", "s_drift", "s_eff", "s_leak", "s_san_plus", "s_cal", "s_perf", "acctempo_m3"]
].mean().set_index("pipeline_id").loc[pipe_order(sanity.pipeline_id)].reset_index()
add_title(doc, "5. Sanity TEST (mean của W1–W3 và P1–P4)", 12)
add_table(doc, ["Pipeline", "S_nan", "S_maj/JSD", "S_ent", "S_drift", "S_eff", "S_leak", "S_SAN+", "S_cal", "S_perf", "AccTEMPO"], [
    [row.pipeline_id, *(fmt(row[field]) for field in ["s_nan", "s_maj_jsd", "s_ent", "s_drift", "s_eff", "s_leak", "s_san_plus", "s_cal", "s_perf", "acctempo_m3"])]
    for _, row in sanity_mean.iterrows()
])

summary = resource.agg({"time_run_total_s": ["mean", "median", "min", "max"], "peak_ram_mb": ["mean", "median", "max"]})
epoch = runs.best_epoch
add_title(doc, "6. Tài nguyên và tái lập", 12)
add_table(doc, ["Hạng mục", "Giá trị"], [
    ["GPU", "NVIDIA L4; 51/51 run dùng GPU; bf16 AMP"],
    ["Thời gian run / cell", f"mean {summary.loc['mean', 'time_run_total_s']/60:.2f} phút; median {summary.loc['median', 'time_run_total_s']/60:.2f} phút; min–max {summary.loc['min', 'time_run_total_s']/60:.2f}–{summary.loc['max', 'time_run_total_s']/60:.2f} phút"],
    ["Peak RAM / cell", f"mean {summary.loc['mean', 'peak_ram_mb']/1024:.1f} GiB; median {summary.loc['median', 'peak_ram_mb']/1024:.1f} GiB; max {summary.loc['max', 'peak_ram_mb']/1024:.1f} GiB"],
    ["Best epoch", f"mean {epoch.mean():.1f}; median {epoch.median():.0f}; min–max {epoch.min():.0f}–{epoch.max():.0f} (max_epochs=30, patience=5)"],
    ["Năng lượng/CO2/chi phí", "Telemetry/billing Modal không cung cấp số đo trực tiếp; các trường được giữ NULL có lý do, không suy diễn chi phí."],
])
doc.add_paragraph(
    "Kết luận thực nghiệm seed 42: toàn bộ ma trận RNN LO đã hoàn tất và artifact tái lập đủ. "
    "Kết quả mô tả mạnh nhất theo AccTEMPO-M3 trung bình là V2 (Median + CDSMOTE, 0.6020), "
    "kế đến V10 (ExtraTrees + CDSMOTE, 0.6010). Đây chưa phải lựa chọn cuối cùng vì một seed "
    "và bảng TEST không được dùng để chọn mô hình; bước tiếp theo là chạy seed đã đăng ký, "
    "shortlist theo Validation AccTEMPO-M3, sau đó chỉ báo cáo TEST cho cấu hình đã khóa."
)

doc.save(REPORT)
print(REPORT)
