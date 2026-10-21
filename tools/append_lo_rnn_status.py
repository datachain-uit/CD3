from pathlib import Path
from docx import Document
from docx.shared import Pt

report = Path(__file__).resolve().parents[1] / 'LO' / 'docs' / 'REPORT_LO_V3_1.docx'
doc = Document(report)
if not any('CẬP NHẬT RNN SEED 42' in p.text for p in doc.paragraphs):
    doc.add_page_break()
    p=doc.add_paragraph(); r=p.add_run('CẬP NHẬT RNN SEED 42 — LO V3.1'); r.bold=True; r.font.size=Pt(14)
    doc.add_paragraph('Release model: split_version=v3_1, phase_version=wide_prefix_v3_1, nhãn lo_final_score_catalog_normalized_v3_1. Kiến trúc shared RNN P1–P4, L4, bf16 AMP, seed 42; TEST chỉ inference, không refit.')
    t=doc.add_table(rows=1, cols=4); t.style='Table Grid'
    for c,x in zip(t.rows[0].cells,['Window','Hoàn tất','Còn lại','Trạng thái']): c.text=x
    for row in [('W1','17/17','—','Đủ V0–V16'),('W2','17/17','—','Đủ V0–V16'),('W3','16/17','V16 = MICE + RadiusSMOTE','Đang chạy; chưa có SUCCESS manifest')]:
        for c,x in zip(t.add_row().cells,row): c.text=x
    doc.add_paragraph('Tổng: 50/51 model cell đã SUCCESS và có checkpoint, train history, validation/test predictions P1–P4, probability matrix, confusion matrix và TEST artifact. Không dùng cell W3–V16 trong bất kỳ trung bình/xếp hạng nào trước khi nó hoàn tất.')
    doc.add_paragraph('Cấu trúc và metric tuân theo CQ: V0 baseline masks; V1/V5/V9/V13 là median/mean/ExtraTrees/MICE; V2–V4, V6–V8, V10–V12, V14–V16 là TRAIN-only CDSMOTE/SASMOTE/RadiusSMOTE. AccTEMPO-M3 là PRIMARY outcome metric; Macro-F1, balanced accuracy, MCC, PR/ROC-AUC và calibration là metric hỗ trợ. Sau khi V16 xong sẽ export L2 facts, kiểm tra 204 TEST + 204 sanity cell, rồi chỉ dùng Validation AccTEMPO-M3 để shortlist đa seed; TEST là xác nhận cuối.')
    doc.save(report)
print(report)
