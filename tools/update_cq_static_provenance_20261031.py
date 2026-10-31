"""Append the CQ static availability and historical runner investigation."""
from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'CQ/docs/REPORT_CQ_PATCHTST_V0_CURRENT_2026-10-31.docx'
DEST = ROOT / 'CQ/docs/REPORT_CQ_PATCHTST_V0_CURRENT_2026-10-31_STATIC_PROVENANCE.docx'
NOTE = ROOT / 'CQ/resuilt/cq_patchtst_v0_current_seed42_r21_report/static_provenance_review_20261031.md'

paragraphs = [
    ('Kết luận', 'CQ V0 hiện có kết quả và audit cho W1–W3; chưa thể xác nhận toàn bộ 11 cột static sẵn có tại từng thời điểm dự đoán chỉ từ source. Kiểm tra epoch 1 cùng run hiện hành khớp; không suy rộng thành replay toàn bộ run hoặc tương đương mọi revision.'),
    ('Nguồn các đặc trưng static', 'ex_counts và resource_count được đếm từ course.csv theo course_id. teacher_count, teacher_org_count và teacher_avg_bio_length_chars được tổng hợp từ quan hệ course–teacher và hồ sơ giảng viên. school_count, school_avg_bio_length_chars và school_avg_motto_length_chars được tổng hợp tương tự từ quan hệ course–school và hồ sơ trường. Các bảng tổng hợp được join theo course_id; source hiện không chứng minh snapshot từng catalog/hồ sơ đã tồn tại trước cutoff của mỗi phase.'),
    ('Thông tin người học và lịch', 'gender lấy từ enrollments; cần xác nhận thời điểm thu thập. window_duration_seconds bằng window_end_exclusive trừ enroll_time; window_end_exclusive được dựng từ offering_end_date, kể cả nhánh proxy. Cần xác nhận ngày kết thúc hoặc quy tắc proxy được biết tại thời điểm dự đoán và không được cập nhật bằng thông tin tương lai. Không kết luận cột này tự động bị leakage chỉ từ phép tính.'),
    ('primary_risk_set', 'Materializer CQ gán hằng 1. Full scan release và 18 input V0 đã kiểm tra; sau preprocessing cột và mask đều bằng 0. Kết luận này áp dụng cho input CQ V0 đã kiểm tra, không tự chứng minh tiêu chí chọn cohort không dùng thông tin tương lai.'),
    ('Runner RNN lịch sử', 'Hash runner ghi trong metadata_contract là 4077c6c91de8f16053653ae315317fb11746a5fa8924873b66be6cbaad14c78e. Hash này khớp byte file tools/vastai_train_shared.py của commit eb8051626fa82b2b8012a48716e95d7f3dd9575a sau chuyển LF sang CRLF. Commit đó có thời điểm 21/10, sau thời điểm manifest run 20/10; do đó không thể gán commit đó cho run hoặc suy ra toàn bộ training loop lịch sử. Receipt trước đây ghi unmatched vì chỉ so byte blob LF; phép so CRLF bổ sung bằng chứng về nội dung runner, không đóng kiểm định trước/sau loop.'),
    ('Việc cần chốt', 'Duyệt bằng chứng snapshot hoặc hợp đồng nguồn cho catalog, hồ sơ, gender và lịch; kiểm tra độc lập tiêu chí chọn cohort. Trước khi mở 48 ô F2 còn lại, thống nhất contract static và phạm vi claim. Không thay đổi layout của ba run đã hoàn thành hoặc tự gộp pilot cũ khác layout.'),
]

def main():
    doc = Document(SOURCE)
    doc.add_heading('11. Kiểm tra bổ sung static và provenance', 1)
    for heading, content in paragraphs:
        doc.add_heading(heading, 2)
        doc.add_paragraph(content)
    doc.save(DEST)
    NOTE.parent.mkdir(parents=True, exist_ok=True)
    NOTE.write_text('# Kiểm tra CQ static và provenance — 31/10/2026\n\n' + ''.join(f'## {h}\n\n{p}\n\n' for h,p in paragraphs) + 'Nguồn: CQ/Feature_extraction/features/build_scenario_phase_features.py; CQ/Feature_extraction/processing/process_course.py, process_teacher.py, process_school.py; CQ/Feature_extraction/views/materialize_label_stratified_views_v2_2.py; CQ/resuilt/cq_rnn_matched_historical_20261029/paired_replay_equivalence.json.\n', encoding='utf-8')
    print(DEST)
    print(NOTE)

if __name__ == '__main__':
    main()
