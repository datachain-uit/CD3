"""Append the CQ all-family r3 evidence closure to the current Vietnamese report."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
FACTS = ROOT / "CQ/resuilt/cq_recurrent_allfamilies_grid_r3_20261021"
SOURCE = ROOT / "CQ/docs/REPORT_CQ_TONG_HOP_DAY_DU_SEED42_2026-10-19_RESOURCE.docx"
TARGET = ROOT / "CQ/docs/REPORT_CQ_TONG_HOP_DAY_DU_SEED42_GRID_R3_2026-10-21.docx"


def _add_table(document: Document, frame: pd.DataFrame, headers: list[tuple[str, str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, (column, label) in zip(table.rows[0].cells, headers):
        cell.text = label
    for _, row in frame.iterrows():
        cells = table.add_row().cells
        for cell, (column, _) in zip(cells, headers):
            value = row[column]
            if isinstance(value, float):
                cell.text = "" if pd.isna(value) else f"{value:.4f}"
            else:
                cell.text = str(value)
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)


def main() -> None:
    if not FACTS.is_dir():
        raise FileNotFoundError(FACTS)
    document = Document(SOURCE)
    document.add_page_break()
    # The existing Vietnamese template uses localized style names, so avoid
    # assuming that python-docx can resolve the English ``Heading 1`` style.
    document.add_paragraph("PHỤ LỤC R3 — HOÀN TẤT EXPORT BỐN HỌ RECURRENT")
    document.add_paragraph(
        "Phụ lục này thay thế các diễn giải cũ về coverage tài nguyên. Phạm vi là CQ, seed 42, "
        "V0–V16, W1–W3, P1–P4 và hai split TEST/VALIDATION. Bảng chi tiết ở đúng hạt "
        "họ × pipeline × window × phase được lưu dưới dạng CSV và Parquet trong "
        "CQ/resuilt/cq_recurrent_allfamilies_grid_r3_20261021/."
    )
    document.add_paragraph("1. Hạt kết quả TEST và VALIDATION")
    document.add_paragraph(
        "Export có 1.632 cell = 4 họ × 17 pipeline × 3 window × 4 phase × 2 split. "
        "Mỗi cell gồm AccTEMPO, Macro-F1, balanced accuracy, s_ent và recall ba lớp. "
        "Tỷ lệ STRICT W2 được tính STRICT/ALL riêng cho từng pipeline, phase và split; W1/W3 "
        "để trống vì không phải đại lượng W2."
    )
    mean = pd.read_parquet(FACTS / "cq_family_window_phase_mean.parquet")
    mean = mean.sort_values(["eval_split", "model_name", "window_id", "phase_id"])
    _add_table(document, mean, [
        ("eval_split", "Split"), ("model_name", "Họ"), ("window_id", "Window"), ("phase_id", "Phase"),
        ("acctempo_m3", "AccTEMPO"), ("f1_macro", "Macro-F1"), ("balanced_accuracy", "Bal. acc."),
        ("s_ent", "s_ent"), ("c0_recall", "Recall c0"), ("c1_recall", "Recall c1"),
        ("c2_recall", "Recall c2"), ("w2_strict_share", "STRICT W2"),
    ])

    document.add_paragraph("2. Invariant STRICT/OVERLAP (audit r3)")
    document.add_paragraph(
        "Audit prediction-level r3 chỉ kiểm tra invariant, không chạy lại train hay bootstrap CI. "
        "Ở từng cell, audit kiểm tra enrollment_id_hash không null/không trùng, cờ temporal_strict_Pk thuộc {0,1}, "
        "và ma trận confusion STRICT + OVERLAP bằng chính xác ALL."
    )
    audit_rows = []
    root = ROOT / "CQ/resuilt/cq_recurrent_batch2048_slice_r3_partition_20261021"
    for family in ("LSTM", "GRU", "BILSTM"):
        summary = pd.read_json(root / family / "summary.json", typ="series")
        audit_rows.append({"Họ": family, "Run": summary["runs"], "File prediction": summary["files"],
                           "Dòng prediction": summary["rows"], "Cell lỗi": summary["failed_cells"],
                           "Kết quả": "Đạt" if summary["all_passed"] else "Không đạt"})
    _add_table(document, pd.DataFrame(audit_rows), [(key, key) for key in audit_rows[0]])
    document.add_paragraph(
        "RNN đã đạt audit cùng invariant trong bundle historical-layout seed 42 trước đó; vì vậy coverage hiện là "
        "204/204 run recurrent."
    )

    document.add_paragraph("3. Thời gian train và năng lượng GPU")
    document.add_paragraph(
        "Bảng sau dùng train_seconds = tổng các khoảng train theo epoch trong train_history, không dùng elapsed_seconds "
        "toàn chương trình. energy_gpu_estimated_kwh là ước lượng GPU-only từ lấy mẫu NVML, không phải năng lượng toàn hệ thống "
        "và không suy ra CO₂. Coverage NVML là 204/204 manifest, không phải chỉ một manifest."
    )
    energy = pd.read_parquet(FACTS / "cq_train_time_energy_family_pipeline.parquet")
    energy = energy.sort_values(["model_name", "pipeline_id"], key=lambda s: s.map(lambda x: int(str(x)[1:])) if s.name == "pipeline_id" else s)
    _add_table(document, energy, [
        ("model_name", "Họ"), ("pipeline_id", "Pipeline"), ("runs", "Số W"),
        ("train_seconds_mean", "Train TB (s)"), ("train_seconds_sum", "Train tổng (s)"),
        ("energy_gpu_estimated_kwh_mean", "GPU TB (kWh)"), ("energy_gpu_estimated_kwh_sum", "GPU tổng (kWh)"),
        ("nvml_samples_min", "NVML mẫu min"), ("nvml_missing_rate_max", "Tỷ lệ thiếu max"),
    ])
    document.add_paragraph(
        "Các tệp nguồn: cq_family_pipeline_window_phase_metrics.{csv,parquet}; "
        "cq_family_window_phase_mean.{csv,parquet}; cq_train_time_energy_family_pipeline.{csv,parquet}; "
        "và strict_overlap_partition_cells.parquet cho từng họ cũ."
    )
    document.add_paragraph("4. Receipt L2 và tách lineage runner")
    document.add_paragraph(
        "Không gộp cưỡng bức các runner revision khác nhau. RNN historical-layout đã có receipt L2 riêng "
        "vastai_cq_rnn_historical_layout_seed42_20261020; 153 run LSTM/GRU/BiLSTM được ingest bằng receipt "
        "vastai_cq_f1_batch2048_oldfamilies_r2_20261021 (1.224 metrics rows, 153 resource rows, 66 partition). "
        "Cả hai receipt đều khóa CQ_V2_2, NVIDIA L4 và seed 42. Bảng trong phụ lục gộp theo family chỉ để so sánh "
        "kết quả; lineage gốc vẫn được giữ theo receipt và runner revision."
    )
    TARGET.parent.mkdir(parents=True, exist_ok=True)
    document.save(TARGET)
    print(TARGET)


if __name__ == "__main__":
    main()
