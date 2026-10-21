"""Build the CQ V2.2 clean-grain RNN seed-42 execution report.

Inputs are the immutable L2 facts downloaded from Modal.  The report never
reads legacy V1/V2.2-pre-clean-grain model artifacts.
"""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "CQ" / "docs" / "REPORT_CQ_V2.2.docx"
FACTS = ROOT / "_tmp_cq_rnn_facts" / "report_id=RNN_seed42_v2_2"
OUTPUT = ROOT / "CQ" / "docs" / "REPORT_CQ_V2.2_CLEAN_GRAIN_RNN_SEED42_2026-10-07.docx"


def clear_body(doc: Document) -> None:
    body = doc._element.body
    for node in list(body):
        if node.tag != qn("w:sectPr"):
            body.remove(node)


def run_font(run, size: float, bold: bool = False) -> None:
    run.bold = bold
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12 if level == 1 else 7)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    run_font(run, 14 if level == 1 else 12, bold=True)


def add_paragraph(doc: Document, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run_font(run, 11, bold=bold)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    run = p.add_run("• " + text)
    run_font(run, 11)


def shade(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    props.append(element)


def put(cell, value: object, bold: bool = False) -> None:
    cell.text = ""
    run = cell.paragraphs[0].add_run(str(value))
    run_font(run, 10.5, bold=bold)


def add_table(doc: Document, headers: list[str], rows: list[list[object]]) -> None:
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.style = "Table Grid"
    for cell, value in zip(tbl.rows[0].cells, headers):
        put(cell, value, bold=True)
        shade(cell, "D9EAF7")
    for values in rows:
        cells = tbl.add_row().cells
        for cell, value in zip(cells, values):
            put(cell, value)
    doc.add_paragraph()


def f(value: float, digits: int = 4) -> str:
    return f"{float(value):.{digits}f}"


def main() -> None:
    required = [
        FACTS / "metrics_test_long.parquet",
        FACTS / "sanity_test_long.parquet",
        FACTS / "sanity_summary_p1_p4_mean.parquet",
        FACTS / "resource_long.parquet",
        FACTS / "run_inventory.parquet",
    ]
    missing = [str(path) for path in required if not path.exists()]
    if missing:
        raise FileNotFoundError("Missing exported facts: " + ", ".join(missing))

    metrics = pd.read_parquet(required[0])
    sanity = pd.read_parquet(required[2])
    resources = pd.read_parquet(required[3])
    inventory = pd.read_parquet(required[4])
    if len(metrics) != 204 or len(sanity) != 51 or len(resources) != 51 or len(inventory) != 51:
        raise ValueError("Unexpected fact cardinality; expected 204 TEST metrics and 51 model/sanity/resource rows")

    rank = (sanity.groupby("pipeline_id", as_index=False)
            [["acctempo_m3", "s_perf", "s_san_plus", "s_cal"]].mean()
            .sort_values("acctempo_m3", ascending=False))
    by_window = sanity.pivot(index="pipeline_id", columns="window_id", values="acctempo_m3")
    phase = (metrics.groupby(["pipeline_id", "phase_id"], as_index=False)
             [["f1_macro", "balanced_accuracy", "mcc", "top_label_ece_15"]].mean())

    doc = Document(TEMPLATE)
    clear_body(doc)
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(11)
    doc.styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("BÁO CÁO CQ — V2.2 CLEAN-GRAIN\nRNN SEED 42: S0–S3 HOÀN TẤT")
    run_font(run, 18, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Đơn vị: learner–course enrollment | scenario=hybrid | cập nhật 07/10/2026")
    run_font(run, 11)
    add_paragraph(doc, "Phạm vi báo cáo là release CQ_V2_2 clean-grain, split_version=v2_2 và phase_version=wide_prefix_v2_2. Báo cáo chỉ dùng facts của 51 run RNN SUCCESS seed 42 vừa export; không trộn kết quả legacy.")

    add_heading(doc, "Tóm tắt điều hành")
    add_table(doc, ["Hạng mục", "Trạng thái", "Kết luận"], [
        ["Input và grain", "LOCKED / GRAIN_OK", "CQ_V2_2; một dòng mỗi enrollment ở mọi TRAIN/VALIDATION/TEST prefix."],
        ["S0/S1 imputation", "SUCCESS 15/15", "V0, V1, V5, V9, V13 × W1–W3; future_imputed_cells=0."],
        ["S2 augmentation", "SUCCESS 36/36", "TRAIN-only; 3 phương pháp × 4 parent imputer × 3 window; seed augmentation cố định 42."],
        ["S3 RNN", "SUCCESS 51/51", "V0–V16 × W1–W3, RNN seed 42; đủ TEST metrics, sanity và resource facts."],
        ["Outcome chính", "AccTEMPO-M3", "Đọc trên TEST đã khóa; Macro-F1, balanced accuracy, MCC và calibration là chỉ số hỗ trợ."],
        ["Kết quả sơ bộ", "V14 cao nhất", f"Mean TEST AccTEMPO-M3 P1–P4/W1–W3 = {f(rank.iloc[0].acctempo_m3)}; chưa dùng TEST để khóa lựa chọn recipe."],
    ])

    add_heading(doc, "1. Release, split và contract đánh giá")
    add_paragraph(doc, "CQ dự báo nhãn vận hành ở grain learner–course enrollment. CQ là calendar anchor của hybrid temporal split V2.2. Split unit là offering_id × timeline_source; A là lịch sử, các lát T1/T2/T3 tạo validation/test cho W1/W2/W3. User_id và course_id chỉ là context/categorical codes TRAIN-only, không phải khóa nhân bản dòng.")
    add_bullet(doc, "TEST P1/P2/P3/P4 lần lượt chỉ quan sát dữ liệu đến P1/P2/P3/P4; future-phase values giữ unavailable/NULL.")
    add_bullet(doc, "Validation và TEST luôn là dữ liệu thật. Augmentation chỉ tạo thêm các dòng TRAIN.")
    add_bullet(doc, "Checkpoint chọn bằng validation; các bảng TEST dưới đây là reporting facts, không phải test-selection.")
    add_table(doc, ["Window", "TRAIN rows", "Vai trò"], [
        ["W1", "1,609,276", "Lịch sử A làm TRAIN; đánh giá ở lát W1."],
        ["W2", "1,963,060", "TRAIN tích lũy theo registry trước W2."],
        ["W3", "2,324,027", "TRAIN tích lũy theo registry trước W3."],
    ])

    add_heading(doc, "2. S0/S1/S2 trước mô hình")
    add_paragraph(doc, "S0 profile ALL_ROWS phân biệt ô thiếu đủ điều kiện với ô vắng mặt cấu trúc. Chỉ numeric missing khi modality đã quan sát mới được impute; structural absence được giữ bằng mask và không trở thành tỷ lệ missing giả.")
    add_table(doc, ["Pipeline", "Phương pháp", "Trạng thái W1–W3"], [
        ["V0", "zero-fill + masks", "SUCCESS"],
        ["V1", "median", "SUCCESS"],
        ["V5", "mean", "SUCCESS"],
        ["V9", "ExtraTrees light", "SUCCESS"],
        ["V13", "MICE: 3 vòng, fit tối đa 1,000,000 TRAIN rows", "SUCCESS"],
    ])
    add_paragraph(doc, "S1 đạt 15/15 cell với model input đầy đủ, fitted preprocessor đầy đủ, không có unfittable column và future_imputed_cells=0. Residual NULL ở TEST P1–P3 là vị trí future unavailable được bảo toàn, không phải imputation failure.")
    add_table(doc, ["Parent imputer", "CDSMOTE", "SASMOTE", "RadiusSMOTE"], [
        ["V1 median", "V2", "V3", "V4"],
        ["V5 mean", "V6", "V7", "V8"],
        ["V9 ExtraTrees", "V10", "V11", "V12"],
        ["V13 MICE", "V14", "V15", "V16"],
    ])
    add_table(doc, ["Window", "Original TRAIN", "Balanced TRAIN", "Synthetic rows"], [
        ["W1", "1,609,276", "1,866,048", "256,772"],
        ["W2", "1,963,060", "2,269,260", "306,200"],
        ["W3", "2,324,027", "2,688,368", "364,341"],
    ])

    add_heading(doc, "3. RNN seed 42 — kết quả TEST chính")
    add_paragraph(doc, "Mỗi giá trị AccTEMPO-M3 dưới đây là trung bình TEST qua P1–P4 trong từng window; cột Mean là trung bình ba window. Đây là primary outcome. V14 là MICE + CDSMOTE; V6 là mean + CDSMOTE; V10 là ExtraTrees + CDSMOTE.")
    result_rows = []
    for item in rank.itertuples(index=False):
        pipeline = item.pipeline_id
        result_rows.append([
            pipeline, f(by_window.loc[pipeline, "W1"]), f(by_window.loc[pipeline, "W2"]),
            f(by_window.loc[pipeline, "W3"]), f(item.acctempo_m3), f(item.s_perf),
        ])
    add_table(doc, ["Pipeline", "W1", "W2", "W3", "Mean AccTEMPO-M3", "Mean S_perf"], result_rows)
    add_paragraph(doc, "Diễn giải hạn chế: W2 thấp hơn W1/W3 ở hầu hết pipeline, nhất quán với temporal/label drift theo window. V14 dẫn đầu primary outcome trong RNN seed 42, nhưng kết luận ổn định và lựa chọn cuối chỉ thực hiện sau validation-governed shortlist và replication seed 43/44.")

    add_heading(doc, "4. Chỉ số hỗ trợ theo phase")
    add_paragraph(doc, "Bảng minh họa hai pipeline: V0 là baseline không augmentation; V14 là primary-outcome leader trong lần chạy RNN seed 42. Mỗi dòng trung bình qua W1–W3. ECE nhỏ hơn là tốt hơn.")
    support_rows = []
    for pipeline in ("V0", "V14"):
        for phase_id in ("P1", "P2", "P3", "P4"):
            row = phase[(phase.pipeline_id == pipeline) & (phase.phase_id == phase_id)].iloc[0]
            support_rows.append([pipeline, phase_id, f(row.f1_macro), f(row.balanced_accuracy), f(row.mcc), f(row.top_label_ece_15)])
    add_table(doc, ["Pipeline", "Phase", "Macro-F1", "Balanced accuracy", "MCC", "ECE-15"], support_rows)

    add_heading(doc, "5. Sanity, calibration và tài nguyên")
    add_paragraph(doc, "S_san+ v1 là sanity/data-quality metric chính; s_cal (v2) chỉ sensitivity report. Không metric sanity nào thay AccTEMPO-M3 làm outcome hiệu năng. s_leak, s_eff và s_drift được tính trên real rows theo contract, không tính lại riêng theo từng kiến trúc mô hình.")
    sanity_rows = []
    for pipeline in ("V0", "V6", "V10", "V14"):
        row = rank[rank.pipeline_id == pipeline].iloc[0]
        full = sanity[sanity.pipeline_id == pipeline].mean(numeric_only=True)
        sanity_rows.append([pipeline, f(row.acctempo_m3), f(full.s_san_plus), f(full.s_ent), f(full.s_leak), f(full.s_eff), f(full.s_drift), f(row.s_cal)])
    add_table(doc, ["Pipeline", "AccTEMPO", "S_san+", "s_ent", "s_leak", "s_eff", "s_drift", "s_cal"], sanity_rows)
    total_hours = resources.time_run_total_s.sum() / 3600
    add_table(doc, ["Resource fact", "Giá trị"], [
        ["Run thành công", "51/51 RNN cells trên NVIDIA L4 với AMP bfloat16."],
        ["Tổng wall-clock", f"{total_hours:.2f} giờ (52,302.5 giây) cho toàn grid."],
        ["Median / max mỗi run", f"{resources.time_run_total_s.median():.1f} / {resources.time_run_total_s.max():.1f} giây."],
        ["CPU process time", f"{resources.cpu_process_s.sum()/3600:.2f} CPU-giờ; peak RAM telemetry không được worker cung cấp."],
        ["Energy / cost", "Không có telemetry Modal billing/energy; không suy diễn CO2e hoặc USD."],
    ])

    add_heading(doc, "6. Lưu trữ, truy vết và bước kế tiếp")
    add_paragraph(doc, "Mỗi retry là attempt immutable mới. Resolver chỉ chọn attempt SUCCESS khớp release identity, split/phase version và parent manifest. Facts của báo cáo này được export tại L2, vì vậy bảng có thể tái tạo mà không đọc trực tiếp prediction rows.")
    tree = (
        "tempo-data-v1/\n"
        "├── input/CQ_v2_2/{release_manifest,phase_views,test_prefix_views/P1..P4}\n"
        "└── meta_release=imputation-v1/\n"
        "    ├── L1_runs/task=CQ/.../pipeline_id=V*/model_name={IMPUTATION_ONLY,BALANCED_TRAIN_ONLY,RNN}/\n"
        "    └── L2_facts/report_exports/task=CQ/report_id=RNN_seed42_v2_2/\n"
        "        ├── metrics_test_long.parquet (204 rows)\n"
        "        ├── sanity_test_long.parquet (204 rows)\n"
        "        ├── resource_long.parquet (51 rows)\n"
        "        ├── run_inventory.parquet (51 rows)\n"
        "        └── test_summary_p1_p4_mean.parquet, sanity_summary_p1_p4_mean.parquet"
    )
    p = doc.add_paragraph()
    run = p.add_run(tree)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    add_table(doc, ["Ưu tiên", "Bước tiếp theo", "Điều kiện"], [
        ["1", "Khóa shortlist bằng VALIDATION, không bằng TEST ranking ở báo cáo này.", "Recipe được định nghĩa trước replication."],
        ["2", "Chạy seed model 43/44 cho shortlist RNN; sau đó LSTM/GRU/BiLSTM theo cùng contract.", "Báo cáo CI/stability paired theo window/phase."],
        ["3", "Báo cáo RadiusSMOTE fallback-nearest-neighbour rate cùng hiệu năng.", "Cảnh báo nếu rate > 0.50 theo window."],
        ["4", "Tổng hợp probability/confusion, per-class và calibration từ L2 facts.", "Không dùng accuracy đơn lẻ cho dữ liệu mất cân bằng."],
    ])
    add_paragraph(doc, "Kết luận: toàn bộ chuỗi CQ V2.2 clean-grain từ S0 đến RNN seed 42 đã hoàn tất và có facts truy vết đầy đủ. V14 dẫn đầu AccTEMPO-M3 sơ bộ; kết luận khoa học cuối cùng vẫn cần validation-governed shortlist và replication, không test-select.")

    doc.core_properties.title = "CQ V2.2 clean-grain — RNN seed 42 execution report"
    doc.core_properties.subject = "Locked release, S0-S3 execution, test metrics and sanity facts"
    doc.core_properties.author = "CD3 experimental pipeline"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
