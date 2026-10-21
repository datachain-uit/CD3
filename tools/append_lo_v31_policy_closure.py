"""Append the V3.1 proxy-name and evaluation-policy closure to the LO report."""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


REPORT = Path(__file__).resolve().parents[1] / "LO" / "docs" / "REPORT_LO_V3_1.docx"
HEADING = "15. Khóa proxy contract và policy đánh giá LO"


def shade(cell, color: str) -> None:
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(element)


def heading(doc: Document, text: str, level: int) -> None:
    paragraph = doc.add_paragraph()
    paragraph.style = doc.styles["693" if level == 1 else "694"]
    paragraph.add_run(text)


def table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    result = doc.add_table(rows=1, cols=len(headers))
    result.style = "Table Grid"
    result.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, value in zip(result.rows[0].cells, headers):
        cell.text = value
        shade(cell, "1F4E78")
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(11)
    for values in rows:
        for cell, value in zip(result.add_row().cells, values):
            cell.text = value
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(11)
    doc.add_paragraph()


def main() -> None:
    document = Document(REPORT)
    if any(p.text == HEADING for p in document.paragraphs):
        raise RuntimeError("Policy-closure section already exists; refusing duplicate append.")
    heading(document, HEADING, 1)
    document.add_paragraph(
        "Mục này đóng hai ghi chú còn lại trong Danh_gia_LO_V3.1: tên vật lý của bốn proxy problem "
        "và policy metric/xấp xỉ bất định cho các cell hiếm. Các quy tắc dưới đây được đồng bộ trong "
        "registries.py, METRIC_DEF v1.2."
    )
    heading(document, "15.1. Bốn trường proxy problem đã xác nhận", 2)
    table(document, ["Trường vật lý", "Vai trò trong final_score", "Xử lý"], [
        ["assignment_problem_catalog_count", "Mẫu số assignment", "Label proxy; cấm imputer/model"],
        ["exam_problem_catalog_count", "Mẫu số exam", "Label proxy; cấm imputer/model"],
        ["assignment_best_correct_sum", "Tử số assignment", "Label proxy; cấm imputer/model"],
        ["exam_best_correct_sum", "Tử số exam", "Label proxy; cấm imputer/model"],
    ])
    document.add_paragraph(
        "Bốn trường và mọi biến thể phase _P1–_P4 được đánh is_label_proxy = 1. Chúng không thuộc "
        "LO_FULL_EARLY, LO_NO_CURRENT_GRADE hoặc predictor của imputer; vì vậy không tạo leakage từ nhãn."
    )
    heading(document, "15.2. Policy xác suất và uncertainty", 2)
    table(document, ["Bối cảnh", "Quy tắc khóa", "Mục đích"], [
        ["LO W2", "Báo log_loss và PR-AUC macro bên cạnh macro-F1", "Đánh giá xác suất trong window exploratory"],
        ["Mọi cell n_offerings_with_class ≤ 2", "Thêm offering-cluster bootstrap: resample whole offering, B=2,000, CI 95%", "Không đánh giá thấp bất định do positive cùng một offering"],
        ["LO W1 TEST — E", "Áp dụng cluster CI", "1,000/1,000 E từ một offering"],
        ["LO W3 TEST — G", "Áp dụng cluster CI", "1,074/1,079 G từ một offering"],
    ])
    document.add_paragraph(
        "Enrollment bootstrap vẫn được báo cáo song song, không bị thay thế. Với PR-AUC macro, nếu một lớp "
        "vắng mặt trong cell thì metric là NULL cùng null_reason = CLASS_ABSENT, không nội suy thành 0."
    )
    heading(document, "15.3. Trạng thái đối chiếu", 2)
    document.add_paragraph(
        "Như vậy mọi yêu cầu có thể trả lời trước model trong Danh_gia_LO_V3.1 đã được đóng: catalog denominator, "
        "coverage theo lát, tên proxy và policy đánh giá. Các audit strict P1–P4, duration/long-offering và S0/P0 là "
        "artifact hậu materialize; chúng không thay đổi label, split hay quyết định V3.1."
    )
    document.save(REPORT)
    print(REPORT)


if __name__ == "__main__":
    main()
