"""Append window/phase, label and per-class CI evidence to the current CQ/LO reports."""
from pathlib import Path
import pandas as pd
from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
LO = ROOT / 'LO/resuilt/lo_static_v2_seed42_20261026'
CQ = ROOT / 'CQ/resuilt/cq_recurrent_allfamilies_grid_r4_20261023'
FAMILIES = ('RNN', 'LSTM', 'GRU', 'BILSTM')
KEY = ['model_name', 'pipeline_id', 'window_id', 'eval_split', 'phase_id']


def table(doc, frame, columns, chunk=80):
    for start in range(0, len(frame), chunk):
        part = frame.iloc[start:start + chunk]
        tab = doc.add_table(rows=len(part) + 1, cols=len(columns))
        tab.style = 'Table Grid'
        for j, (_, title) in enumerate(columns):
            tab.cell(0, j).text = title
        for row, values in zip(tab.rows[1:], part.itertuples(index=False, name=None)):
            for cell, value in zip(row.cells, values):
                cell.text = '' if pd.isna(value) else (f'{value:.4f}' if isinstance(value, float) else str(value))
        for row in tab.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(11)


def lo_roots():
    for family in FAMILIES:
        yield (LO / '_closure_audit_temp/audits_seed42/LO/BILSTM/seed=42' if family == 'BILSTM'
               else LO / family / 'audits/LO' / family / 'seed=42')


def labels(task):
    root = ROOT / f'downloads/vastai/{"CQ_V2_2_clean_grain_seed42" if task == "CQ" else "LO_V3_1_clean_grain_seed42"}'
    label = 'CQ_label_final' if task == 'CQ' else 'LO_performance_label_3'
    records = []
    for w in ('W1', 'W2', 'W3'):
        for split, path in [('TRAIN', root / 'train' / w / 'V0/train.parquet'),
                            ('VALIDATION', root / 'evaluation' / w / 'V0/validation.parquet'),
                            ('TEST', root / 'evaluation' / w / 'V0/test_P1.parquet')]:
            s = pd.read_parquet(path, columns=[label])[label]
            for name, count in s.value_counts(dropna=False).items():
                records.append({'window': w, 'split': split, 'label': str(name), 'n': int(count), 'total': len(s), 'share_pct': 100 * count / len(s)})
    return pd.DataFrame(records)


def main(task):
    report = ROOT / ('CQ/docs/REPORT_CQ_CONSOLIDATED_EVIDENCE_2026-10-28_R20.docx' if task == 'CQ' else 'LO/docs/REPORT_LO_STATIC_V2_CONSOLIDATED_EVIDENCE_2026-10-28_R20.docx')
    doc = Document(report)
    for paragraph in doc.paragraphs:
        if 'full grid F2' in paragraph.text or 'full-grid F2' in paragraph.text:
            paragraph.text = paragraph.text.replace('full grid F2 chưa khởi chạy vì là pha thực nghiệm kế tiếp, không phải vì một lỗi replay.', 'paired replay là kiểm thử của cell CQ/W1/V0/RNN.')
            paragraph.text = paragraph.text.replace('nhưng full grid F2 vẫn chưa chạy.', 'và D23 vẫn là sensitivity pilot.')
            paragraph.text = paragraph.text.replace('thực hiện full-grid F2 v2 theo protocol đã khóa;', 'giữ protocol đã khóa và không chọn theo TEST;')
    if any(p.text.startswith('10. Phụ lục chi tiết seed 42') for p in doc.paragraphs):
        raise ValueError(f'Appendix already exists in {report}; regenerate base first')
    if task == 'CQ':
        metrics = pd.read_parquet(CQ / 'cq_family_pipeline_window_phase_metrics.parquet').rename(columns={'f1_macro':'macro_f1', 'balanced_accuracy':'balanced_acc'})
        per = pd.read_parquet(CQ / 'cq_metrics_per_class_long.parquet')
        ci = pd.read_parquet(CQ / 'cq_bootstrap_metrics_per_class_long.parquet')
    else:
        roots = list(lo_roots())
        metrics = pd.concat([pd.read_parquet(p / 'sanity_components.parquet') for p in roots], ignore_index=True)
        per = pd.concat([pd.read_parquet(p / 'metrics_per_class.parquet') for p in roots], ignore_index=True)
        ci = pd.concat([pd.read_parquet(p / 'bootstrap_metrics_per_class.parquet') for p in roots], ignore_index=True)
    for frame in (metrics, per, ci):
        frame['eval_split'] = frame.eval_split.str.upper()
    assert len(metrics) == 1632 and len(per) == 4896 and len(ci) == 14688
    doc.add_heading('10. Phụ lục chi tiết seed 42 — window × phase', level=1)
    doc.add_paragraph('Các bảng dưới đây không trung bình qua window hoặc phase. Mỗi split có 816 cell: 4 họ × 17 pipeline × 3 window × 4 phase; TEST dùng để báo cáo, VALIDATION dùng chọn checkpoint.')
    fields = KEY + ['acctempo_m3', 'macro_f1', 'balanced_acc', 's_ent', 's_san_plus', 'c0_recall', 'c1_recall', 'c2_recall']
    for split in ('TEST', 'VALIDATION'):
        doc.add_heading(f'10.1. {split} — tất cả 816 cell', level=2)
        frame = metrics[metrics.eval_split.eq(split)][fields].sort_values(KEY).copy()
        table(doc, frame, list(zip(fields, ['Model','V','W','Split','P','AccTEMPO','F1 macro','Bal.acc','s_ent','S_san+','Recall 0','Recall 1','Recall 2'])))
    doc.add_heading('11. Phân phối nhãn theo từng window và split', level=1)
    doc.add_paragraph('Đếm trực tiếp enrollment trong V0 model input: TRAIN, VALIDATION và TEST P1; tỷ lệ tính trong từng split. Các phase TEST giữ cùng cohort.')
    lab = labels(task)
    table(doc, lab, [('window','Window'),('split','Split'),('label','Nhãn'),('n','Số mẫu'),('total','Tổng split'),('share_pct','Tỷ lệ %')])
    doc.add_heading('12. Support và khoảng tin cậy recall theo lớp', level=1)
    doc.add_paragraph('Bootstrap phân tầng B=2.000, CI 95%; mỗi dòng là một cell model × pipeline × window × phase × split. Giá trị theo thứ tự lớp 0 / 1 / 2; support dưới 400 cần đọc thận trọng. CI của precision và F1 cũng được giữ trong facts machine-readable.')
    recall = ci[ci.metric_name.eq('recall')].copy()
    assert len(recall) == 4896 and recall.bootstrap_repetitions.eq(2000).all()
    recall['entry'] = recall.apply(lambda r: f"{r['class_code']}: n={int(r['n_support'])}, {r['estimate']:.3f} [{r['ci_lower']:.3f}, {r['ci_upper']:.3f}]" if pd.notna(r['ci_lower']) else f"{r['class_code']}: n={int(r['n_support'])}, CI không xác định", axis=1)
    pivot = recall.pivot(index=KEY, columns='class_index', values='entry').reset_index().sort_values(KEY)
    assert len(pivot) == 1632
    for split in ('TEST','VALIDATION'):
        doc.add_heading(f'12.1. {split} — support / recall [CI 95%]', level=2)
        frame = pivot[pivot.eval_split.eq(split)][KEY + [0,1,2]]
        table(doc, frame, list(zip(KEY+[0,1,2], ['Model','V','W','Split','P','Lớp 0','Lớp 1','Lớp 2'])))
    doc.add_heading('13. Lát đánh giá và cấu trúc lưu trữ', level=1)
    doc.add_paragraph('STRICT/OVERLAP là các lát mô tả theo thời điểm; SEEN/UNSEEN theo course_id có trong TRAIN cùng window. Các lát không thay đổi cohort TEST hoặc quy tắc chọn checkpoint. Support nhỏ/lớp vắng phải được đọc kèm cờ small_support; không suy diễn độ ổn định từ một seed.')
    if task == 'LO':
        slice_files = sorted((ROOT / 'LO/resuilt/lo_static_v2_cohort_audit_20261028').glob('*/*/seen_unseen_strict_overlap_metrics.parquet'))
        assert len(slice_files) == 68
        slices = pd.concat([pd.read_parquet(p) for p in slice_files], ignore_index=True)
        slices['classes_present'] = slices[['support_c0', 'support_c1', 'support_c2']].gt(0).sum(axis=1)
        slices.loc[slices.classes_present < 2, ['macro_f1', 'balanced_accuracy']] = float('nan')
        group = ['window_id', 'eval_split', 'phase_id', 'course_cohort', 'temporal_slice']
        summary = slices.groupby(group, as_index=False).agg(n_min=('n_eval','min'), n_max=('n_eval','max'), classes_min=('classes_present','min'), classes_max=('classes_present','max'), f1=('macro_f1','mean'), bal=('balanced_accuracy','mean'), cells=('n_eval','size'))
        summary['status'] = summary.apply(lambda r: 'SLICE_EMPTY' if r.n_max == 0 else ('ONE_CLASS' if r.classes_max < 2 else ('MIXED_CLASS_COUNTS' if r.classes_min != r.classes_max else 'OK')), axis=1)
        doc.add_paragraph('Trung bình mô tả qua 68 cấu hình; lát rỗng được ghi SLICE_EMPTY, lát dưới hai lớp không báo cáo F1/balanced accuracy. classes min–max là số lớp thực sự có mặt; không xem trung bình này là kết quả của một run riêng.')
        table(doc, summary, list(zip(group + ['n_min','n_max','classes_min','classes_max','f1','bal','cells','status'], ['W','Split','P','Course','Thời gian','n min','n max','Lớp min','Lớp max','F1 TB','Bal.acc TB','Cells','Trạng thái'])))
        unseen = slices.assign(unseen=lambda d: d.n_eval.where(d.course_cohort.eq('UNSEEN_COURSE'), 0)).groupby(['window_id','eval_split','phase_id','model_name','pipeline_id'], as_index=False).agg(unseen=('unseen','sum'), total=('n_eval','sum'))
        unseen['unseen_pct'] = 100 * unseen.unseen / unseen.total.where(unseen.total.ne(0))
        unseen = unseen.groupby(['window_id','eval_split','phase_id'], as_index=False).agg(unseen_min_pct=('unseen_pct','min'), unseen_max_pct=('unseen_pct','max'))
        doc.add_paragraph('Tỷ lệ UNSEEN theo window × split × phase; min–max qua các cấu hình.')
        table(doc, unseen, [('window_id','W'),('eval_split','Split'),('phase_id','P'),('unseen_min_pct','UNSEEN min %'),('unseen_max_pct','UNSEEN max %')])
        doc.add_paragraph('LO: cohort SEEN/UNSEEN × STRICT/OVERLAP đã audit 204 run; chi tiết 6.528 lát và prediction invariants lưu tại LO/resuilt/lo_static_v2_cohort_audit_20261028. Audit theo hash khớp với input evaluation, không chỉ so số dòng.')
    else:
        roots = [ROOT / 'CQ/resuilt/cq_f1_allfamilies_l2_ingest_r3_20261021/audits/CQ' / family / 'seed=42/evaluation_slices.parquet' for family in FAMILIES]
        slices = pd.concat([pd.read_parquet(p) for p in roots], ignore_index=True)
        group = ['window_id','eval_split','phase_id','slice']
        summary = slices.groupby(group, as_index=False).agg(n_min=('n_eval_samples','min'), n_max=('n_eval_samples','max'), f1=('macro_f1','mean'), bal=('balanced_acc','mean'), cells=('n_eval_samples','size'))
        doc.add_paragraph('STRICT/OVERLAP: trung bình mô tả qua cấu hình hiện có, không dùng để chọn model; support và số cell được nêu kèm.')
        table(doc, summary, list(zip(group + ['n_min','n_max','f1','bal','cells'], ['W','Split','P','Lát','n min','n max','F1 TB','Bal.acc TB','Cells'])))
        doc.add_paragraph('CQ: facts STRICT/OVERLAP và kiểm tra partition được lưu tại CQ/resuilt/cq_recurrent_batch2048_slice_r3_partition_20261021 và các bundle RNN tương ứng; audit cohort chi tiết chỉ được diễn giải theo phạm vi file thực tế, không mặc định đủ toàn lưới.')
    doc.add_paragraph('Cấu trúc lưu trữ: input release gồm phase views TRAIN/VALIDATION và TEST prefix P1–P4; L0 registry lưu contract/identity; L1 lưu manifest, preprocessor, checkpoint và prediction; L2 lưu metrics, sanity, per-class, CI, slices và resource; L3 dành cho meta-dataset khi hoàn thành tổng hợp. Các đường dẫn local dưới đây là artefact đã tải về, không khẳng định toàn bộ facts đã ingest vào L2 chính thức.')
    doc.add_paragraph(f'Facts chi tiết: {CQ.relative_to(ROOT) if task == "CQ" else LO.relative_to(ROOT)}. Bảng bootstrap lưu theo metric_name = precision/recall/f1, B=2.000. Dữ liệu source và kết quả được định danh bằng release, run, attempt, window, pipeline, model, seed, split và phase.')
    target = report.with_name(report.stem + '_DAY_DU.docx')
    doc.save(target)
    print(task, 'saved', target, 'metrics', len(metrics), 'class_facts', len(per), 'bootstrap_facts', len(ci))


if __name__ == '__main__':
    for task in ('CQ', 'LO'):
        main(task)
