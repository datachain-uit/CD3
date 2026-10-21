"""Append downloaded CQ/LO long-offering audit evidence to the V2.2 report."""
from __future__ import annotations

from pathlib import Path

import polars as pl
from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "docs" / "REPORT_SPLIT_V2_2_OVERALL_CLOSURE_2026-09-29.docx"
BLUE, GREY, GREEN = "1F4E78", "F2F2F2", "E2F0D9"


def read(task: str, artifact: str) -> pl.DataFrame:
    folder = ROOT / task / "resuilt" / "long_offering_assignment_audit_v1" / artifact
    files = list(folder.glob("*.parquet"))
    if len(files) != 1:
        raise RuntimeError(f"{folder}: expected exactly one parquet, found {len(files)}")
    return pl.read_parquet(files[0])


def shade(cell, colour: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), colour)
    props.append(node)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.size = Pt(11.5)
        run.font.color.rgb = RGBColor(255, 255, 255)
        shade(cell, BLUE)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            cells[i].paragraphs[0].add_run(value).font.size = Pt(11.5)
            if ridx % 2:
                shade(cells[i], GREY)
    doc.add_paragraph()


def pct(value: float) -> str:
    return f"{float(value) * 100:.2f}%"


def main() -> None:
    if not REPORT.exists():
        raise FileNotFoundError(REPORT)
    data = {
        task: {
            "duration": read(task, "duration_by_timeline"),
            "slice": read(task, "slice_assignment_summary"),
            "qa": read(task, "enrollment_assignment_qa"),
            "strict": read(task, "strict_by_window_split_phase"),
        }
        for task in ("CQ", "LO")
    }
    for task, frames in data.items():
        if frames["qa"].item(0, "assignment_qa_pass") != 1:
            raise RuntimeError(f"{task} assignment QA did not pass")

    doc = Document(REPORT)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    title = doc.add_heading("8. Long-offering và assignment audit (kết quả R2)", level=1)
    for run in title.runs:
        run.font.color.rgb = RGBColor.from_string("17365D")
    p = doc.add_paragraph()
    p.add_run("Kết luận: ").bold = True
    p.add_run("CQ và LO đều đạt assignment QA. Mỗi enrollment_id chỉ ánh xạ tới một rolling block; trong mỗi window, enrollment chỉ có một split. Các bảng dưới dùng ngưỡng long offering >270 ngày và strict = cutoff_time_phase ≥ tau_w.")
    shade_cell = p._p.get_or_add_pPr()
    shade_node = OxmlElement("w:shd"); shade_node.set(qn("w:fill"), GREEN); shade_cell.append(shade_node)

    duration_rows: list[list[str]] = []
    for task in ("CQ", "LO"):
        for r in data[task]["duration"].sort("timeline_source").iter_rows(named=True):
            duration_rows.append([
                task, r["timeline_source"], f"{r['offering_count']:,}",
                f"{r['duration_days_p50']:,}", f"{r['duration_days_p90']:,}",
                f"{r['duration_days_max']:,}", f"{r['n_long_offerings']:,}",
                pct(r["long_offering_enrollment_share"]),
            ])
    doc.add_heading("8.1 Duration theo timeline_source", level=2)
    add_table(doc, ["Task", "Timeline", "Offerings", "p50 ngày", "p90 ngày", "Max ngày", ">270d offerings", "Enrollment >270d"], duration_rows)
    doc.add_paragraph("Lưu ý: global_template_fallback có ít offering nhưng tỷ lệ enrollment long cao; cần được diễn giải theo timeline_source, không gộp mù với course-specific timeline.")

    slice_rows: list[list[str]] = []
    for task in ("CQ", "LO"):
        grouped = (data[task]["slice"].group_by("temporal_slice").agg(
            pl.sum("offering_timeline_count").alias("units"),
            pl.sum("enrollment_count").alias("enrollments"),
            pl.sum("n_long_offerings").alias("long_units"),
            pl.sum("long_offering_enrollment_count").alias("long_enrollments"))
            .with_columns((pl.col("long_enrollments") / pl.col("enrollments")).alias("long_share"))
            .sort("temporal_slice"))
        for r in grouped.iter_rows(named=True):
            slice_rows.append([task, r["temporal_slice"], f"{r['units']:,}", f"{r['enrollments']:,}", f"{r['long_units']:,}", pct(r["long_share"])])
    doc.add_heading("8.2 Quy mô A/T1/T2/T3", level=2)
    add_table(doc, ["Task", "Slice", "offering × timeline", "Enrollments", ">270d offerings", "Enrollment >270d"], slice_rows)
    doc.add_paragraph("Quy ước lát: A; T1 = B/C; T2 = D/E; T3 = F/G. Đây là bảng gộp theo lát, không nhân đôi enrollment qua các window mở rộng.")

    qa_rows: list[list[str]] = []
    for task in ("CQ", "LO"):
        r = data[task]["qa"].row(0, named=True)
        qa_rows.append([
            task, f"{r['distinct_enrollment_count']:,}", f"{r['context_row_count']:,}",
            str(r["n_enrollments_multiple_rolling_blocks"]),
            str(r["n_enrollments_multiple_splits_same_window"]),
            str(r["n_offering_timeline_multiple_rolling_blocks"]),
            str(r["n_context_rows_without_block"]), "Đạt" if r["assignment_qa_pass"] == 1 else "Không đạt",
        ])
    doc.add_heading("8.3 QA assignment", level=2)
    add_table(doc, ["Task", "Enrollment riêng", "Context rows", "Multi-block enrollment", "Multi-split / window", "Multi-block unit", "Không map block", "Kết quả"], qa_rows)

    for split, label in (("validation", "VALIDATION"), ("test", "TEST")):
        strict_rows: list[list[str]] = []
        for task in ("CQ", "LO"):
            pivot = (data[task]["strict"].filter(pl.col("split") == split)
                .pivot(on="phase_id", index="window_id", values="temporal_strict_share", aggregate_function="first")
                .sort("window_id"))
            for r in pivot.iter_rows(named=True):
                strict_rows.append([task, r["window_id"], *[pct(r[f"P{i}"]) for i in range(1, 5)]])
        doc.add_heading(f"8.{4 if split == 'validation' else 5} Strict P1–P4: {label}", level=2)
        add_table(doc, ["Task", "Window", "P1", "P2", "P3", "P4"], strict_rows)
    doc.add_paragraph("Strict giảm hoặc tăng theo window/phase là đặc tính của cutoff so với tau_w, không phải lỗi gán split. P1-cutoff overlap = 1 − strict share; không thay thế bằng offering-start overlap.")

    doc.add_heading("8.6 Artifact bổ sung", level=2)
    doc.add_paragraph(
        "long_offering_assignment_audit_v1/\n"
        "├── duration_by_timeline/\n"
        "├── slice_assignment_summary/\n"
        "├── enrollment_assignment_qa/\n"
        "├── strict_by_window_split_phase/\n"
        "└── audit_manifest/"
    )
    doc.save(REPORT)
    print(REPORT)


if __name__ == "__main__":
    main()
