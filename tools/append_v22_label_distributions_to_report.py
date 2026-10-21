"""Append label rates by temporal slice and window/split to the V2.2 report."""
from __future__ import annotations

import os
from pathlib import Path

import polars as pl
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
REPORT = Path(os.environ.get(
    "SPLIT_LABEL_REPORT_PATH",
    str(ROOT / "docs" / "Danh_gia_Split_V2.2_cap_nhat_2026-09-30.docx"),
))
BLUE, GREY = "1F4E78", "F2F2F2"
CLASSES = {"CQ": ["warning", "average", "good"], "LO": ["I/D", "G", "E"]}


def read(task: str, name: str) -> pl.DataFrame:
    folder = ROOT / task / "resuilt" / "split_registry_v2_2_overlap_audit_v1_r2" / name
    files = list(folder.glob("*.parquet"))
    if not files: raise RuntimeError(f"No parquet in {folder}")
    return pl.read_parquet(max(files, key=lambda p: p.stat().st_mtime))


def shade(cell, color: str) -> None:
    props = cell._tc.get_or_add_tcPr(); node = OxmlElement("w:shd")
    node.set(qn("w:fill"), color); props.append(node)


def heading(doc: Document, text: str, size: int = 13) -> None:
    paragraph = doc.add_paragraph(); run = paragraph.add_run(text)
    run.bold = True; run.font.size = Pt(size); run.font.color.rgb = RGBColor.from_string("17365D")


def table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    result = doc.add_table(rows=1, cols=len(headers))
    try: result.style = "Table Grid"
    except KeyError: pass
    for i, header in enumerate(headers):
        cell = result.rows[0].cells[i]; cell.text = ""
        run = cell.paragraphs[0].add_run(header); run.bold = True; run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(255, 255, 255); shade(cell, BLUE)
    for ridx, row in enumerate(rows):
        cells = result.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""; cells[i].paragraphs[0].add_run(value).font.size = Pt(10.5)
            if ridx % 2: shade(cells[i], GREY)
    doc.add_paragraph()


def slice_for_block(block: str) -> str:
    return "A" if block == "A" else "T1" if block in {"B", "C"} else "T2" if block in {"D", "E"} else "T3"


def main() -> None:
    if not REPORT.exists(): raise FileNotFoundError(REPORT)
    doc = Document(REPORT)
    heading(doc, "7.5 Phân bố nhãn và support theo split", 13)
    doc.add_paragraph("Các tỷ lệ dưới được lấy trực tiếp từ registry V2.2 overlap-audit R2. Tỷ lệ nhãn có mẫu số là enrollment trong lát hoặc trong split tương ứng; không dùng một mẫu số chung giữa CQ và LO.")

    for task in ("CQ", "LO"):
        blocks = read(task, "block_label_audit")
        classes = CLASSES[task]
        rows = []
        for slice_id in ("A", "T1", "T2", "T3"):
            selected = blocks.filter(pl.col("rolling_block").map_elements(slice_for_block, return_dtype=pl.String) == slice_id)
            n = int(selected["enrollment_count"].sum())
            values = [int(selected[f"{label}_count"].sum()) for label in classes]
            rows.append([slice_id, f"{n:,}", *[f"{value:,} ({value / n * 100:.3f}%)" for value in values]])
        heading(doc, f"{task}: tỷ lệ nhãn theo lát thời gian", 12)
        table(doc, ["Lát", "Enrollment", *classes], rows)

        labels = read(task, "split_label_distribution")
        totals = read(task, "split_actual_ratios")
        rows = []
        for window in ("W1", "W2", "W3"):
            for split in ("train", "validation", "test"):
                total = totals.filter((pl.col("window_id") == window) & (pl.col("split") == split)).row(0, named=True)
                subset = labels.filter((pl.col("window_id") == window) & (pl.col("split") == split))
                value_by_label = {r["label"]: r for r in subset.iter_rows(named=True)}
                n = int(total["enrollment_count"])
                class_cells = []
                for label in classes:
                    item = value_by_label[label]
                    class_cells.append(f"{item['label_enrollment_count']:,} ({item['label_ratio'] * 100:.3f}%)")
                rows.append([window, split, f"{n:,} ({total['split_enrollment_share'] * 100:.2f}%)", *class_cells])
        heading(doc, f"{task}: nhãn theo window × split", 12)
        table(doc, ["Window", "Split", "Enrollment (% window)", *classes], rows)

    doc.add_paragraph("Diễn giải: LO W2/W3 có tỷ lệ G/E rất nhỏ trong validation/test nên cần áp small_support_flag và CI bootstrap khi báo cáo performance theo lớp. CQ có support các lớp cao hơn ở mọi split, nhưng vẫn thể hiện label drift thật giữa các lát T1, T2 và T3.")
    doc.save(REPORT); print(REPORT)


if __name__ == "__main__": main()
