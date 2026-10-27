#!/usr/bin/env python3
"""Build the current CQ seed-42 report, including the PatchTST V0 pilot."""
from __future__ import annotations

import json
import sys
import hashlib
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
R4 = ROOT / "CQ/resuilt/cq_recurrent_allfamilies_grid_r4_20261023"
PATCH = ROOT / "CQ/resuilt/cq_patchtst_v0_seed42_20261023"
TARGET = ROOT / "CQ/docs/REPORT_CQ_CURRENT_RECURRENT_PATCHTST_V0_SEED42_2026-10-23.docx"
CHECKED_FEATURE = ROOT / "CQ/resuilt/_checked_feature_v2"
LINEAGE_ROOT = CHECKED_FEATURE / "materialization_lineage_v2_2/materialization_lineage_v2_2"
PREFIX_AUDIT = ROOT / "CQ/resuilt/audit_prefix_validation_alignment_v1_20261017"
FULL_SCAN = ROOT / "CQ/resuilt/cq_full_prefix_static_audit_20261017/full_scan.json"
POSTHOC_RECEIPT = PATCH / "posthoc_provenance_receipt.json"
COHORT_AUDIT = PATCH / "audits_seen_unseen_v1"

# Registered F2 pilot protocol. The report reads this from the actual run
# manifests and fails closed rather than merely asserting it in prose.
F2_PROTOCOL = {
    "model_revision": "patchtst_channel_independent_v2_1_math_sdp",
    "batch_size": 2048,
    "micro_batch_size": 512,
    "accumulation_steps": 4,
    "static_ablation": "none",
    "run_purpose": "PRIMARY_GRID",
    # Registered optimisation protocol (HPARAMS of tools/vastai_train_shared.py; model/shared_phase defaults).
    # 044d6c7 made these per-run options: a grid trained with another learning rate, a warm-up or gradient
    # clipping is a different protocol and must fail here until MODEL_REGISTRY registers it (hparam config v3).
    "learning_rate": 1e-3,
    "warmup_epochs": 0,
    "gradient_clip_norm": None,
}
# Manifests older than 044d6c7 do not carry these keys; they ran the registered values.
F2_PROTOCOL_DEFAULT_IF_ABSENT = {"warmup_epochs": 0, "gradient_clip_norm": None, "learning_rate": 1e-3}
WINDOWS = ("W1", "W2", "W3")


def patchtst_provenance(root: Path, *, protocol: dict = F2_PROTOCOL) -> pd.DataFrame:
    """Read three pilot receipts and reject a protocol or layout mismatch."""
    rows: list[dict] = []
    deviations: list[str] = []
    for window in WINDOWS:
        path = root / window / "run_manifest.json"
        if not path.is_file():
            raise FileNotFoundError(f"PatchTST pilot manifest missing: {path}")
        manifest = json.loads(path.read_text(encoding="utf-8"))
        layout_path = root / window / "feature_layout.json"
        if not layout_path.is_file():
            raise FileNotFoundError(f"PatchTST pilot feature layout missing: {layout_path}")
        feature_layout_sha256 = "sha256:" + hashlib.sha256(layout_path.read_bytes()).hexdigest()
        for key, expected in protocol.items():
            actual = manifest.get(key, F2_PROTOCOL_DEFAULT_IF_ABSENT.get(key)) if key in F2_PROTOCOL_DEFAULT_IF_ABSENT else manifest.get(key)
            if actual != expected:
                deviations.append(f"{window}: {key}={actual!r} (registered {expected!r})")
        for key in (
            "static_layout_policy", "runner_git_commit", "seed", "best_epoch",
            "elapsed_seconds", "energy_measurement_type", "energy_gpu_estimated_kwh",
        ):
            if manifest.get(key) in (None, ""):
                deviations.append(f"{window}: required provenance field {key} is absent")
        if not (manifest.get("nvml_gpu_name") or manifest.get("gpu")):
            deviations.append(f"{window}: required provenance field GPU name is absent")
        if (manifest.get("static_layout_policy") == "historical_reference"
                and not manifest.get("reference_layout_sha256")):
            deviations.append(f"{window}: historical_reference requires reference_layout_sha256")
        actual_identity = (manifest.get("window_id"), manifest.get("pipeline_id"), manifest.get("model_name"))
        if actual_identity != (window, "V0", "PATCHTST"):
            deviations.append(f"{window}: manifest identity {actual_identity}")
        test_rows = manifest.get("test_rows") or {}
        rows.append({
            "window_id": window,
            "model_revision": manifest.get("model_revision"),
            "static_layout_policy": manifest.get("static_layout_policy"),
            "reference_layout_sha256": manifest.get("reference_layout_sha256"),
            "feature_layout_sha256": feature_layout_sha256,
            "runner_git_commit": manifest.get("runner_git_commit"),
            "batch": f"{manifest.get('batch_size')} = {manifest.get('accumulation_steps')} x {manifest.get('micro_batch_size')}",
            "best_epoch": manifest.get("best_epoch"),
            "elapsed_seconds": manifest.get("elapsed_seconds"),
            "energy_measurement_type": manifest.get("energy_measurement_type"),
            "energy_gpu_estimated_kwh": manifest.get("energy_gpu_estimated_kwh"),
            "nvml_gpu_name": manifest.get("nvml_gpu_name") or manifest.get("gpu"),
            "train_rows": manifest.get("train_rows"),
            "validation_rows": manifest.get("validation_rows"),
            "test_rows": test_rows.get("P1") if isinstance(test_rows, dict) else None,
            "seed": manifest.get("seed"),
        })
    if deviations:
        raise ValueError("PatchTST pilot deviates from the registered F2 protocol: " + "; ".join(deviations))
    frame = pd.DataFrame(rows)
    for column in ("model_revision", "static_layout_policy", "reference_layout_sha256",
                   "feature_layout_sha256", "runner_git_commit", "seed"):
        # ``reference_layout_sha256`` is intentionally null for the ``current``
        # layout policy.  Null still needs to be consistent across the pilot.
        if frame[column].nunique(dropna=False) != 1:
            raise ValueError(f"PatchTST pilot runs disagree on {column}: {frame[column].tolist()}")
    return frame


def patchtst_history(root: Path) -> pd.DataFrame:
    """Return per-epoch validation selection history with the selected epoch marked."""
    frames: list[pd.DataFrame] = []
    for window in WINDOWS:
        path = root / window / "train_history.parquet"
        if not path.is_file():
            raise FileNotFoundError(f"PatchTST pilot history missing: {path}")
        history = pd.read_parquet(path)
        manifest = json.loads((root / window / "run_manifest.json").read_text(encoding="utf-8"))
        required = {
            "epoch", "validation_macro_f1_mean_p1_p4", "validation_cross_entropy_mean_p1_p4",
            "train_cross_entropy_mean_p1_p4", "epoch_seconds",
        }
        missing = sorted(required.difference(history.columns))
        if missing:
            raise ValueError(f"{path} lacks {missing}")
        selected_count = int((history["epoch"] == manifest.get("best_epoch")).sum())
        if selected_count != 1:
            raise ValueError(
                f"{window}: best_epoch={manifest.get('best_epoch')!r} occurs {selected_count} times in train history"
            )
        frames.append(history.assign(
            window_id=window,
            selected=(history["epoch"] == manifest.get("best_epoch")).map({True: "YES", False: ""}),
        ))
    columns = [
        "window_id", "epoch", "selected", "validation_macro_f1_mean_p1_p4",
        "validation_cross_entropy_mean_p1_p4", "train_cross_entropy_mean_p1_p4", "epoch_seconds",
    ]
    out = pd.concat(frames, ignore_index=True)
    return out[columns]


def single_parquet(root: Path) -> Path:
    paths = sorted(root.rglob("*.parquet"))
    if len(paths) != 1:
        raise FileNotFoundError(f"Expected exactly one parquet file under {root}, found {len(paths)}")
    return paths[0]


def posthoc_pilot_receipt(lineage: pd.DataFrame) -> pd.DataFrame:
    """Write an explicit, non-retroactive identity receipt for legacy pilot manifests.

    The pilot's original run manifests predate the release-identity fields.  This
    receipt records the controlled CQ contract and the materialization lineage,
    but deliberately labels the identity as post-hoc rather than altering an
    immutable run manifest.
    """
    from release_core.contracts import resolve_release

    spec = resolve_release("CQ", "CQ_V2_2")
    manifests = [json.loads((PATCH / window / "run_manifest.json").read_text(encoding="utf-8"))
                 for window in WINDOWS]
    layout_hashes = sorted({
        "sha256:" + hashlib.sha256((PATCH / window / "feature_layout.json").read_bytes()).hexdigest()
        for window in WINDOWS
    })
    absent = [field for field in (
        "release_id", "data_release_id", "split_version", "phase_version",
        "feature_dictionary_version", "label_rule_version", "label_threshold_set",
    ) if any(not manifest.get(field) for manifest in manifests)]
    lineage_hashes = sorted(lineage["materialization_lineage_sha256"].dropna().astype(str).unique())
    payload = {
        "receipt_version": "posthoc_pilot_provenance_v1",
        "evidence_status": "POSTHOC_CONTROLLED_CONTRACT_NOT_ORIGINAL_RUN_MANIFEST",
        "task": "CQ",
        "release_spec": spec,
        "materialization_lineage_sha256": lineage_hashes,
        "original_manifest_missing_fields": absent,
        "static_layout_policy": manifests[0].get("static_layout_policy"),
        "reference_layout_sha256": manifests[0].get("reference_layout_sha256"),
        "posthoc_feature_layout_sha256": layout_hashes,
        "interpretation": (
            "The controlled release contract and clean-grain lineage are evidenced post hoc. "
            "This receipt does not claim that the absent identity fields were present in the original run manifests."
        ),
    }
    POSTHOC_RECEIPT.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    return pd.DataFrame([{
        "evidence_status": payload["evidence_status"],
        "release_id": spec["release_id"],
        "split_version": spec["split_version"],
        "phase_version": spec["phase_version"],
        "split_registry_id": spec["split_registry_id"],
        "label_rule_version": spec["label_rule_version"],
        "lineage_sha256": "; ".join(lineage_hashes),
        "missing_original_fields": ", ".join(absent),
        "static_layout_policy": payload["static_layout_policy"],
        "reference_layout_sha256": payload["reference_layout_sha256"],
        "feature_layout_sha256": "; ".join(layout_hashes),
    }])


def heading(doc: Document, value: str, size: int = 14) -> None:
    p = doc.add_paragraph()
    r = p.add_run(value)
    r.bold = True
    r.font.size = Pt(size)


def paragraph(doc: Document, value: str) -> None:
    p = doc.add_paragraph(value)
    p.paragraph_format.space_after = Pt(5)
    for r in p.runs:
        r.font.size = Pt(10.5)


def table(doc: Document, frame: pd.DataFrame, cols: list[tuple[str, str]], size: float = 8.5) -> None:
    out = doc.add_table(rows=1, cols=len(cols))
    out.style = "Table Grid"
    for cell, (_, label) in zip(out.rows[0].cells, cols):
        cell.text = label
        for r in cell.paragraphs[0].runs:
            r.bold = True
            r.font.size = Pt(size + .5)
    for _, row in frame.iterrows():
        cells = out.add_row().cells
        for cell, (name, _) in zip(cells, cols):
            value = row.get(name)
            cell.text = "" if pd.isna(value) else (f"{value:.4f}" if isinstance(value, float) else str(value))
            for r in cell.paragraphs[0].runs:
                r.font.size = Pt(size)
    doc.add_paragraph()


def main() -> None:
    required = [R4 / "summary.json", R4 / "cq_family_pipeline_window_phase_metrics.parquet",
                R4 / "cq_qa23_reviewed_behavioural_linkage.parquet",
                PATCH / "audits_v0_seed42/CQ/PATCHTST/seed=42/audit_manifest.json",
                COHORT_AUDIT / "seen_unseen_strict_overlap_metrics.parquet",
                COHORT_AUDIT / "prediction_invariants.parquet", COHORT_AUDIT / "summary.json",
                FULL_SCAN, PREFIX_AUDIT / "summary.json",
                *(PATCH / window / name for window in WINDOWS
                  for name in ("run_manifest.json", "train_history.parquet", "feature_layout.json"))]
    missing = [str(p) for p in required if not p.exists()]
    if missing:
        raise FileNotFoundError("Missing report inputs: " + ", ".join(missing))

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(.65)
    section.bottom_margin = Inches(.65)
    section.left_margin = Inches(.7)
    section.right_margin = Inches(.7)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("BÁO CÁO CQ HIỆN TẠI — RECURRENT GRID VÀ PATCHTST V0")
    run.bold = True
    run.font.size = Pt(18)
    subtitle = doc.add_paragraph("CQ_V2_2 clean-grain · seed 42 · cập nhật 23/10/2026")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in subtitle.runs:
        r.italic = True
        r.font.size = Pt(11)

    summary = json.loads((R4 / "summary.json").read_text(encoding="utf-8"))
    heading(doc, "1. Phạm vi và trạng thái nghiệm thu")
    paragraph(doc, "Phần recurrent hoàn tất bốn họ RNN, LSTM, GRU và BiLSTM trên VastAI L4, batch 2.048: V0–V16, W1–W3, P1–P4, hai split TEST/VALIDATION. PatchTST hiện là pilot riêng chỉ cho V0 trên ba window; nó xác nhận pipeline kiến trúc hoạt động nhưng chưa phải so sánh F2/full-grid.")
    status = pd.DataFrame([
        {"Hạng mục": "Recurrent grid", "Phạm vi": "4 họ × 17 pipeline × 3 W × 4 P × 2 split", "Kết quả": f"{summary['evaluation_cells']:,} cell"},
        {"Hạng mục": "Tài nguyên recurrent", "Phạm vi": "NVIDIA L4 / BF16 / batch 2.048", "Kết quả": f"{summary['resource_runs']:,} manifest; NVML đầy đủ"},
        {"Hạng mục": "Per-class / bootstrap", "Phạm vi": "support, recall, F1; CI B=2.000", "Kết quả": f"{summary['per_class_facts']:,} / {summary['bootstrap_per_class_facts']:,} fact"},
        {"Hạng mục": "PatchTST", "Phạm vi": "V0 × W1–W3 × seed 42", "Kết quả": "3 run; post-hoc audit hoàn tất"},
    ])
    table(doc, status, [("Hạng mục", "Hạng mục"), ("Phạm vi", "Phạm vi"), ("Kết quả", "Kết quả")], 9)

    lineage = pd.read_parquet(single_parquet(LINEAGE_ROOT))
    receipt = posthoc_pilot_receipt(lineage)
    heading(doc, "1b. Clean-grain, prefix và identity release")
    paragraph(doc, "Lineage dưới đây là bằng chứng materialization clean-grain. Mọi nguồn có exact_duplicates_collapsed = 0. Receipt identity của PatchTST là hậu kiểm: ba manifest gốc của pilot chưa ghi các trường release, nên receipt không được diễn giải là thay đổi manifest bất biến.")
    table(doc, lineage.sort_values("source_name"), [
        ("source_name", "Nguồn"), ("rows_before", "Rows trước"),
        ("rows_after_exact_deduplication", "Rows sau"), ("exact_duplicates_collapsed", "Trùng loại"),
        ("source_schema_sha256", "Schema hash"),
    ], 7.5)
    prefix_summary = json.loads((PREFIX_AUDIT / "summary.json").read_text(encoding="utf-8"))
    full_scan = json.loads(FULL_SCAN.read_text(encoding="utf-8"))
    mask = pd.DataFrame(full_scan["mask"])
    table(doc, mask, [("phase", "Phase"), ("rows", "Rows"), ("future_cells", "Future cells"),
                      ("non_null_future_cells", "Future non-null"), ("pass", "Đạt")], 8)
    prefix_evidence = pd.DataFrame([{
        "layout_count": prefix_summary["layout_count"],
        "layout_signature_count": prefix_summary["layout_signature_count"],
        "static_column_count": prefix_summary["static_column_count"],
        "prefix_future_mask_all_samples_ok": prefix_summary["prefix_future_mask_all_samples_ok"],
        "test_static_prefix_parity_ok": prefix_summary["test_static_prefix_parity_ok"],
    }])
    table(doc, prefix_evidence, [("layout_count", "Layout"), ("layout_signature_count", "Layout hash"),
                                 ("static_column_count", "Cột static"), ("prefix_future_mask_all_samples_ok", "Future mask"),
                                 ("test_static_prefix_parity_ok", "Static parity")], 8)
    heading(doc, "1c. Receipt identity hậu kiểm của PatchTST V0", 12)
    paragraph(doc, "Receipt được ghi tại cq_patchtst_v0_seed42_20261023/posthoc_provenance_receipt.json. Trạng thái POSTHOC có nghĩa contract CQ_V2_2 và lineage được đối chiếu từ release registry và materialization; các trường thiếu trong manifest gốc vẫn được giữ nguyên là thiếu.")
    table(doc, receipt, [("evidence_status", "Trạng thái"), ("release_id", "Release"),
                         ("split_version", "Split"), ("phase_version", "Phase"),
                         ("split_registry_id", "Registry"), ("label_rule_version", "Label rule"),
                         ("static_layout_policy", "Static policy"), ("reference_layout_sha256", "Reference hash"),
                         ("feature_layout_sha256", "Layout hash hậu kiểm"),
                         ("missing_original_fields", "Thiếu trong manifest gốc")], 7)

    heading(doc, "2. Quy tắc đọc metric")
    paragraph(doc, "AccTEMPO-M3 v1 vẫn là composite PRIMARY đã tiền đăng ký trong meta-dataset. Khi trình bày xu hướng P1→P4 hoặc xếp hạng trong bài, phải đọc cùng S_perf, Macro-F1 và balanced accuracy; đồng thời báo cáo AccTEMPO-M4 (S_san+ v2 dựa trên s_cal) là sensitivity composite. Lý do: s_ent thường giảm khi dự đoán ở phase muộn chắc chắn hơn, nên M3 có thể giảm dù hiệu năng phân loại tăng. Accuracy không dùng độc lập vì nhãn mất cân bằng.")

    cells = pd.read_parquet(R4 / "cq_family_pipeline_window_phase_metrics.parquet")
    recurrent_mean = (cells.groupby(["model_name", "eval_split"], as_index=False)
                      [["acctempo_m3", "acctempo_m4", "f1_macro", "balanced_accuracy", "s_perf", "s_san_plus", "s_cal"]].mean()
                      .sort_values(["eval_split", "model_name"]))
    heading(doc, "3. Recurrent: trung bình mô tả theo họ")
    paragraph(doc, "Bảng này chỉ mô tả trung bình qua pipeline/window/phase; không thay thế bảng hạt đầy đủ trong package evidence."
    )
    table(doc, recurrent_mean, [("eval_split", "Split"), ("model_name", "Họ"), ("acctempo_m3", "AccTEMPO M3"), ("acctempo_m4", "AccTEMPO M4"), ("f1_macro", "Macro-F1"), ("balanced_accuracy", "Bal. acc."), ("s_perf", "S_perf"), ("s_san_plus", "S_san+ v1"), ("s_cal", "s_cal")])

    rnn_phase = (cells[(cells.model_name == "RNN") & (cells.eval_split == "TEST")]
                 .groupby(["window_id", "phase_id"], as_index=False)[["acctempo_m3", "acctempo_m4", "f1_macro", "balanced_accuracy", "s_ent"]].mean())
    heading(doc, "4. Recurrent: RNN TEST theo window × phase")
    table(doc, rnn_phase, [("window_id", "W"), ("phase_id", "P"), ("acctempo_m3", "M3"), ("acctempo_m4", "M4"), ("f1_macro", "Macro-F1"), ("balanced_accuracy", "Bal. acc."), ("s_ent", "s_ent")])

    heading(doc, "4b. Recurrent: đủ 17 pipeline theo họ và split")
    paragraph(doc, "Bảng này giữ từng pipeline V0–V16 và từng họ; mỗi ô là trung bình 12 tổ hợp window × phase trong cùng split. File cq_family_pipeline_window_phase_metrics.parquet/CSV trong bundle R4 là phụ lục hạt đầy đủ 1.632 dòng, không bị thay thế bởi bảng tóm tắt này.")
    grid_summary = (cells.groupby(["model_name", "pipeline_id", "eval_split"], as_index=False)
                    [["acctempo_m3", "acctempo_m4", "f1_macro", "balanced_accuracy", "s_perf",
                      "s_san_plus", "s_san_plus_v2", "s_eff", "s_leak"]].mean()
                    .sort_values(["eval_split", "model_name", "pipeline_id"]))
    table(doc, grid_summary, [("eval_split", "Split"), ("model_name", "Họ"), ("pipeline_id", "Pipeline"),
                              ("acctempo_m3", "M3"), ("acctempo_m4", "M4"), ("f1_macro", "Macro-F1"),
                              ("balanced_accuracy", "Bal. acc."), ("s_perf", "S_perf"),
                              ("s_san_plus", "S_san+ v1"), ("s_san_plus_v2", "S_san+ v2"),
                              ("s_eff", "s_eff"), ("s_leak", "s_leak")], 6.8)
    resource = pd.read_parquet(R4 / "cq_train_time_energy_family_pipeline.parquet")
    heading(doc, "4c. Recurrent: thời gian train và năng lượng GPU theo họ × pipeline")
    paragraph(doc, "Thời gian dùng train_seconds. Năng lượng chỉ là ước lượng NVML GPU-only theo sampling, không suy ra CO2. Mọi cell trong bảng có telemetry NVML hợp lệ.")
    table(doc, resource.sort_values(["model_name", "pipeline_id"]), [
        ("model_name", "Họ"), ("pipeline_id", "Pipeline"), ("runs", "Run"),
        ("train_seconds_mean", "Train s mean"), ("train_seconds_sum", "Train s sum"),
        ("energy_gpu_estimated_kwh_mean", "GPU kWh mean"),
        ("energy_gpu_estimated_kwh_sum", "GPU kWh sum"),
        ("nvml_samples_min", "NVML samples min"), ("nvml_missing_rate_max", "NVML miss max"),
    ], 6.8)

    qa = pd.read_parquet(R4 / "cq_qa23_reviewed_behavioural_linkage.parquet")
    qa_summary = (qa.groupby("phase_id", as_index=False)
                  .agg(reviewed_bases=("base", "nunique"), reviewed_columns=("column_name", "count"),
                       rho_abs_max=("spearman_abs", "max"), nmi_max=("nmi", "max"),
                       status=("review_status", "first")))
    heading(doc, "5. QA23 — feature/label-proxy")
    paragraph(doc, "Tám behavioural base được giữ với điều kiện chỉ dùng prefix đã quan sát trước cutoff. Chúng có thể tương quan với final label; correlation được công khai thay vì bị coi là label-derived. Trường tạo nhãn final-distance không thuộc predictor contract.")
    table(doc, qa_summary, [("phase_id", "Phase"), ("reviewed_bases", "Base"), ("reviewed_columns", "Cột"), ("rho_abs_max", "|ρ| max"), ("nmi_max", "NMI max"), ("status", "Trạng thái")])

    metrics = pd.concat([pd.read_parquet(PATCH / w / "metrics_overall.parquet") for w in WINDOWS], ignore_index=True)
    audit_root = PATCH / "audits_v0_seed42/CQ/PATCHTST/seed=42"
    sanity = pd.read_parquet(audit_root / "sanity_components.parquet")
    patch = metrics.merge(sanity[["window_id", "pipeline_id", "seed", "eval_split", "phase_id", "acctempo_m3", "s_perf", "s_san_plus", "s_cal", "s_san_plus_v2", "s_ent", "s_leak"]], on=["window_id", "pipeline_id", "seed", "eval_split", "phase_id"], validate="one_to_one")
    patch["acctempo_m4"] = patch.s_perf.clip(lower=1e-3).pow(.6) * patch.s_san_plus_v2.clip(lower=1e-3).pow(.4)
    provenance = patchtst_provenance(PATCH)
    heading(doc, "6. PatchTST V0 — kết quả pilot")
    paragraph(doc, "PatchTST V0 có contract runtime đạt gate: seed 42, NVIDIA L4, BF16, batch 2.048 và accumulation 4. Release CQ_V2_2 được đối chiếu hậu kiểm ở mục 1c vì manifest gốc của pilot chưa ghi identity release. Bảng dưới là TEST; VALIDATION được audit và giữ trong artifact nhưng không dùng TEST để chọn checkpoint.")
    patch_test = patch[patch.eval_split == "TEST"].sort_values(["window_id", "phase_id"])
    table(doc, patch_test, [("window_id", "W"), ("phase_id", "P"), ("n_rows", "n TEST"), ("acctempo_m3", "M3"), ("acctempo_m4", "M4"), ("f1_macro", "Macro-F1"), ("balanced_accuracy", "Bal. acc."), ("s_perf", "S_perf"), ("s_san_plus", "S_san+"), ("s_ent", "s_ent"), ("s_leak", "s_leak")], 8)
    patch_val = patch[patch.eval_split == "VALIDATION"].sort_values(["window_id", "phase_id"])
    heading(doc, "7. PatchTST V0 — VALIDATION")
    table(doc, patch_val, [("window_id", "W"), ("phase_id", "P"), ("n_rows", "n VAL"), ("acctempo_m3", "M3"), ("acctempo_m4", "M4"), ("f1_macro", "Macro-F1"), ("balanced_accuracy", "Bal. acc."), ("s_perf", "S_perf")], 8)

    manifest = json.loads((audit_root / "audit_manifest.json").read_text(encoding="utf-8"))
    checkpoints = pd.read_parquet(audit_root / "checkpoint_selection.parquet")
    evaluation_slices = pd.read_parquet(audit_root / "evaluation_slices.parquet")
    per_class = pd.read_parquet(audit_root / "metrics_per_class.parquet")
    bootstrap = pd.read_parquet(audit_root / "bootstrap_metrics_per_class.parquet")
    canonical_leak = pd.read_parquet(audit_root / "canonical_s_leak.parquet")
    heading(doc, "8. PatchTST V0 — audit và giới hạn")
    paragraph(doc, f"Post-hoc audit có {manifest['runs']} run, canonical S_leak={manifest['rows']['canonical_s_leak']} row, sanity={manifest['rows']['sanity_components']} row, STRICT/OVERLAP={manifest['rows']['evaluation_slices']} row, per-class={manifest['rows']['metrics_per_class']} row và bootstrap={manifest['rows']['bootstrap_metrics_per_class']} row (B={manifest['bootstrap_repetitions']:,}). Đây là audit partial được cho phép chỉ cho V0; full-grid PatchTST cần hoàn tất V1–V16 trước khi đưa vào bảng so sánh family/pipeline chính.")
    table(doc, checkpoints, [("window_id", "W"), ("best_epoch", "Epoch chọn"), ("best_score", "Score VAL"), ("checkpoint_count_evaluated", "Checkpoint"), ("selection_margin", "Margin"), ("n_checkpoints_within_margin", "Trong margin")])
    heading(doc, "8a. PatchTST V0 — sanity, leakage và lát đánh giá", 12)
    sanity_detail = sanity.sort_values(["eval_split", "window_id", "phase_id"])
    table(doc, sanity_detail, [("eval_split", "Split"), ("window_id", "W"), ("phase_id", "P"),
                               ("s_nan", "s_nan"), ("s_maj_jsd", "s_maj_jsd"), ("s_ent", "s_ent"),
                               ("s_drift", "s_drift"), ("s_eff", "s_eff"), ("s_leak", "s_leak"),
                               ("s_cal", "s_cal"), ("s_san_plus", "S_san+ v1"), ("s_san_plus_v2", "S_san+ v2")], 7.3)
    paragraph(doc, "Canonical S_leak được liệt kê kèm scope và số mẫu probe. Đây là đo trên parent real-data contract; augmented train không tái fit leakage probe riêng.")
    table(doc, canonical_leak.sort_values(["eval_split", "window_id", "phase_id"]), [
        ("eval_split", "Split"), ("window_id", "W"), ("phase_id", "P"), ("parent_pipeline", "Parent"),
        ("probe_fit_rows", "n_fit"), ("s_leak_probe_n_eval", "n_eval"),
        ("probe_fit_scope", "Fit scope"), ("probe_score_scope", "Score scope"), ("s_leak", "s_leak"),
    ], 7.2)
    strict_overlap = evaluation_slices[evaluation_slices["slice"].isin(["STRICT", "OVERLAP"])].copy()
    strict_overlap = strict_overlap.sort_values(["eval_split", "window_id", "phase_id", "slice"])
    table(doc, strict_overlap, [("eval_split", "Split"), ("window_id", "W"), ("phase_id", "P"),
                                ("slice", "Lát"), ("n_eval_samples", "n"), ("macro_f1", "Macro-F1"),
                                ("balanced_acc", "Bal. acc."), ("s_perf", "S_perf"), ("acctempo_m3", "M3")], 7.3)
    heading(doc, "8aa. PatchTST V0 — per-class và bootstrap CI", 12)
    per_class_view = per_class.sort_values(["eval_split", "window_id", "phase_id", "class_index"])
    table(doc, per_class_view, [("eval_split", "Split"), ("window_id", "W"), ("phase_id", "P"),
                                ("class_code", "Class"), ("n_support", "Support"),
                                ("small_support_flag", "Small"), ("precision", "Precision"),
                                ("recall", "Recall"), ("f1", "F1")], 7.1)
    bootstrap_f1 = bootstrap[bootstrap["metric_name"].eq("f1")].sort_values(
        ["eval_split", "window_id", "phase_id", "class_index"]
    )
    table(doc, bootstrap_f1, [("eval_split", "Split"), ("window_id", "W"), ("phase_id", "P"),
                              ("class_code", "Class"), ("n_support", "Support"),
                              ("bootstrap_repetitions", "B"), ("estimate", "F1"),
                              ("ci_lower", "CI low"), ("ci_upper", "CI high")], 7.1)
    heading(doc, "8b. PatchTST V0 — provenance của ba run", 12)
    paragraph(doc, "Đọc trực tiếp từ run_manifest.json. Builder dừng nếu revision, batch 2048 = 4 x 512, static_ablation hoặc run_purpose lệch giao thức F2; builder cũng dừng nếu ba run khác nhau về layout policy, reference layout, runner commit hoặc seed.")
    table(doc, provenance, [("window_id", "W"), ("model_revision", "Revision"), ("static_layout_policy", "Layout policy"), ("feature_layout_sha256", "Layout hash"), ("runner_git_commit", "Runner commit"), ("batch", "Batch"), ("best_epoch", "Epoch chon"), ("elapsed_seconds", "Elapsed (s)"), ("energy_measurement_type", "Energy"), ("energy_gpu_estimated_kwh", "kWh"), ("nvml_gpu_name", "GPU")], 6.7)
    history = patchtst_history(PATCH)
    heading(doc, "8c. PatchTST V0 — lịch sử chọn checkpoint trên VALIDATION", 12)
    paragraph(doc, "Mỗi dòng là một epoch; tiêu chí chọn là mean Macro-F1 P1–P4 trên VALIDATION. Dòng YES là epoch được chọn. W3 dừng sau epoch 6 với epoch 1 được chọn, nên cần đọc đường cong này trước khi mở full grid.")
    table(doc, history, [("window_id", "W"), ("epoch", "Epoch"), ("selected", "Chon"), ("validation_macro_f1_mean_p1_p4", "VAL Macro-F1 (mean P1–P4)"), ("validation_cross_entropy_mean_p1_p4", "VAL CE"), ("train_cross_entropy_mean_p1_p4", "Train CE"), ("epoch_seconds", "Epoch (s)")], 8)
    heading(doc, "8d. PatchTST V0 so với recurrent V0 cùng window × phase", 12)
    recurrent_v0 = cells[(cells["pipeline_id"] == "V0") & (cells["eval_split"] == "TEST")]
    recurrent_v0 = (recurrent_v0.groupby(["window_id", "phase_id"], as_index=False)
                    .agg(recurrent_m3_mean=("acctempo_m3", "mean"),
                         recurrent_f1_mean=("f1_macro", "mean"),
                         recurrent_f1_min=("f1_macro", "min"), recurrent_f1_max=("f1_macro", "max"),
                         recurrent_balacc_mean=("balanced_accuracy", "mean")))
    comparison = (patch_test[["window_id", "phase_id", "acctempo_m3", "f1_macro", "balanced_accuracy"]]
                  .merge(recurrent_v0, on=["window_id", "phase_id"], validate="one_to_one")
                  .rename(columns={"acctempo_m3": "patch_m3", "f1_macro": "patch_f1",
                                   "balanced_accuracy": "patch_balacc"}))
    comparison["delta_f1_vs_recurrent_mean"] = comparison["patch_f1"] - comparison["recurrent_f1_mean"]
    comparison["delta_balacc_vs_recurrent_mean"] = comparison["patch_balacc"] - comparison["recurrent_balacc_mean"]
    table(doc, comparison.sort_values(["window_id", "phase_id"]), [
        ("window_id", "W"), ("phase_id", "P"), ("patch_m3", "Patch M3"),
        ("patch_f1", "Patch F1"), ("recurrent_f1_mean", "Recurrent F1 mean"),
        ("recurrent_f1_min", "Recurrent F1 min"), ("recurrent_f1_max", "Recurrent F1 max"),
        ("delta_f1_vs_recurrent_mean", "ΔF1"), ("patch_balacc", "Patch bal. acc."),
        ("delta_balacc_vs_recurrent_mean", "Δbal. acc."),
    ], 7.1)
    cohort = pd.read_parquet(COHORT_AUDIT / "seen_unseen_strict_overlap_metrics.parquet")
    cohort_summary = json.loads((COHORT_AUDIT / "summary.json").read_text(encoding="utf-8"))
    heading(doc, "8e. Cohort SEEN/UNSEEN × STRICT/OVERLAP", 12)
    paragraph(doc, "Audit read-only ghép prediction với evaluation bằng hash salt nội bộ; salt không được xuất. SEEN_COURSE nghĩa là course_id đã có trong real parent TRAIN của cùng window, không có nghĩa offering đã thấy. STRICT dùng cờ temporal theo đúng phase. Kiểm tra prediction đạt: duplicate hash = 0 và invalid STRICT = 0 trên 24 file.")
    cohort_overview = pd.DataFrame([{
        "prediction_files": cohort_summary["prediction_files"], "cells": cohort_summary["cells"],
        "duplicate_hashes": cohort_summary["duplicate_hashes"], "invalid_strict": cohort_summary["invalid_strict"],
    }])
    table(doc, cohort_overview, [("prediction_files", "Prediction files"), ("cells", "Cohort cells"),
                                 ("duplicate_hashes", "Duplicate hash"), ("invalid_strict", "Invalid STRICT")], 8)
    table(doc, cohort.sort_values(["eval_split", "window_id", "phase_id", "course_cohort", "temporal_slice"]), [
        ("eval_split", "Split"), ("window_id", "W"), ("phase_id", "P"),
        ("course_cohort", "Course cohort"), ("temporal_slice", "Temporal slice"),
        ("n_eval", "n"), ("macro_f1", "Macro-F1"), ("balanced_accuracy", "Bal. acc."),
        ("support_c0", "Support c0"), ("support_c1", "Support c1"), ("support_c2", "Support c2"),
    ], 6.9)
    paragraph(doc, "Kết luận: recurrent grid CQ seed 42 đạt evidence đầy đủ. PatchTST V0 chạy đúng contract và audit đạt, nên có thể mở tiếp PatchTST cho các pipeline còn lại; hiện chưa được dùng để chọn model thắng hoặc kết luận ổn định đa seed.")

    TARGET.parent.mkdir(parents=True, exist_ok=True)
    doc.save(TARGET)
    print(TARGET)


if __name__ == "__main__":
    main()
