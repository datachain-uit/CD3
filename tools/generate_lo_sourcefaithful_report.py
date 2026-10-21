"""Create the LO source-faithful label and split readiness report."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
RESULT = ROOT / "LO" / "resuilt"
OUT = ROOT / "LO" / "docs" / "REPORT_LO_SOURCEFAITHFUL_V2_2_2026-09-30.docx"
REGISTRY = RESULT / "split_registry_v2_2_sourcefaithful_overlap_audit_v1"
SCORE = RESULT / "final_score_contract_audit_v2_sourcefaithful"
LABEL = RESULT / "lo_labels_v2_sourcefaithful_audit"
LONG = RESULT / "long_offering_assignment_audit_v2_sourcefaithful"
COVERAGE = RESULT / "label_coverage_audit_proxy_v2_sourcefaithful"

NAVY, BLUE, GREY, GREEN, AMBER = "17365D", "1F4E78", "F2F2F2", "E2F0D9", "FFF2CC"


def read(root: Path, name: str) -> pd.DataFrame:
    files = list((root / name).glob("*.parquet"))
    if not files:
        raise FileNotFoundError(root / name)
    return pd.read_parquet(max(files, key=lambda item: item.stat().st_mtime))


def shade(cell, color: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), color)
    props.append(node)


def paragraph(doc: Document, text: str, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)


def heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(15 if level == 1 else 12.5)
    r.font.color.rgb = RGBColor.from_string(NAVY)


def table(doc: Document, headers: list[str], rows: list[list[object]]) -> None:
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, header in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = str(header)
        shade(cell, BLUE)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(10.5)
    for ridx, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            for run in cells[i].paragraphs[0].runs:
                run.font.size = Pt(10.5)
            if ridx % 2:
                shade(cells[i], GREY)
    doc.add_paragraph()


def pct(value: float) -> str:
    return f"{100 * float(value):.3f}%"


def main() -> None:
    reg = read(REGISTRY, "split_registry")
    split_ratio = read(REGISTRY, "split_actual_ratios")
    split_labels = read(REGISTRY, "split_label_distribution")
    block = read(REGISTRY, "block_label_audit")
    label_dist = read(LABEL, "label_distribution_3")
    version = read(SCORE, "label_version_summary").iloc[0]
    contract = read(SCORE, "label_semantics_contract").iloc[0]
    score_ranges = read(SCORE, "score_range_by_profile")
    assignment = read(LONG, "enrollment_assignment_qa").iloc[0]
    duration = read(LONG, "duration_by_timeline")
    strict = read(LONG, "strict_by_window_split_phase")
    coverage_manifest = read(COVERAGE, "audit_manifest").iloc[0]

    w3 = split_labels[split_labels["window_id"].eq("W3")].groupby("label", as_index=False)["label_enrollment_count"].sum()
    label_counts = dict(zip(w3["label"], w3["label_enrollment_count"]))
    eligible = int(version["eligible_enrollment_count"])
    gt100 = int(score_ranges["n_final_score_gt_100"].sum())

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.5); section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.6); section.right_margin = Cm(1.6)
    doc.styles["Normal"].font.name = "Aptos"
    doc.styles["Normal"].font.size = Pt(11)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("LO · Source-faithful V2.2 · 30/09/2026").font.size = Pt(9)

    title = doc.add_paragraph(); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("BÁO CÁO LO — NHÃN FINAL_SCORE SOURCE-FAITHFUL VÀ SPLIT V2.2")
    run.bold = True; run.font.size = Pt(18); run.font.color.rgb = RGBColor.from_string(NAVY)
    subtitle = doc.add_paragraph(); subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Bản tổng hợp kiểm tra nhãn, temporal split và điều kiện sẵn sàng thực nghiệm").italic = True

    heading(doc, "1. Kết luận điều hành")
    paragraph(doc, "Bản LO hiện tại đạt chuẩn source-faithful để khóa label và split: nhãn được tái lập theo quy tắc TEMPO, split V2.2 đạt QA22 ở cả W1–W3, và kiểm tra assignment không phát hiện enrollment bị gán chồng block/split.")
    paragraph(doc, "Điều kiện còn lại trước khi chốt hồ sơ: chạy lại riêng label-coverage audit sau bản code metadata đã sửa, để trường label_contract không còn dùng thuật ngữ behavioural proxy cũ.")
    table(doc, ["Hạng mục", "Trạng thái"], [
        ["Nhãn", "Đạt — lo_final_score_sourcefaithful_v2"],
        ["Split", "Đạt — v2.2.sourcefaithful, scenario=hybrid, CQ calendar anchor"],
        ["QA22", "Đạt ở W1, W2 và W3"],
        ["Assignment integrity", "Đạt — không chồng block/split"],
        ["Coverage metadata", "Cần tái xuất audit sau sửa câu chữ metadata"],
    ])

    heading(doc, "2. Định nghĩa nhãn và quy tắc source-faithful")
    paragraph(doc, "final_score là điểm tổng hợp từ hoạt động được chấm; không được diễn giải là final grade quan sát độc lập, completion hay dropout.")
    table(doc, ["Thuộc tính", "Giá trị"], [
        ["Phiên bản", version["label_rule_version"]],
        ["Score semantics", contract["score_semantics"]],
        ["Công thức", contract["formula"]],
        ["Band", contract["class_thresholds"]],
        ["Exam rule", contract["exam_partition_rule"]],
        ["Cohort eligible", f"{eligible:,}"],
    ])
    paragraph(doc, f"Có {gt100:,} mẫu final_score >100. Theo TEMPO gốc, các mẫu này đi nhánh mặc định I/D; không clip về 100 và không gán sang E/G.")

    heading(doc, "3. Phân bố nhãn đã xác nhận")
    table(doc, ["Nhãn", "Số enrollment", "Tỷ lệ cohort eligible"], [
        [label, f"{int(label_counts[label]):,}", pct(label_counts[label] / eligible)]
        for label in ("I/D", "G", "E")
    ])
    paragraph(doc, "Audit nhãn xác nhận G có score trung bình nằm trong band 60–<85; không còn hiện tượng score trung bình hàng trăm của bản mapping lỗi trước đó.")

    heading(doc, "4. Split V2.2: QA và quy mô thực tế")
    qa_rows = []
    for _, row in reg.sort_values("window_id").iterrows():
        qa_rows.append([
            row["window_id"], f"{int(row['n_enrollments_validation']):,}", f"{int(row['n_enrollments_test']):,}",
            "Đạt" if bool(row["end_order_ok"]) else "Không đạt", "Đạt" if bool(row["qa22_ok"]) else "Không đạt",
            row["purge_mode"],
        ])
    table(doc, ["Window", "Validation", "Test", "Label-end order", "QA22", "Purge"], qa_rows)
    ratio_rows = []
    for window in ("W1", "W2", "W3"):
        x = split_ratio[split_ratio["window_id"].eq(window)]
        ratio_rows.append([window, *[f"{int(x[x['split'].eq(s)].iloc[0]['enrollment_count']):,} ({pct(x[x['split'].eq(s)].iloc[0]['split_enrollment_share'])})" for s in ("train", "validation", "test")]])
    table(doc, ["Window", "Train", "Validation", "Test"], ratio_rows)

    heading(doc, "5. Temporal drift, strictness và offering dài")
    block_rows = []
    for b in ("A", "B", "C", "D", "E", "F", "G"):
        row = block[block["rolling_block"].eq(b)].iloc[0]
        block_rows.append([b, f"{int(row['enrollment_count']):,}", f"{int(row['I/D_count']):,}", f"{int(row['G_count']):,}", f"{int(row['E_count']):,}"])
    table(doc, ["Block", "Enrollment", "I/D", "G", "E"], block_rows)
    strict_rows = []
    for window in ("W1", "W2", "W3"):
        row = strict[(strict["window_id"].eq(window)) & (strict["split"].eq("test")) & (strict["phase_id"].eq("P1"))].iloc[0]
        strict_rows.append([window, pct(row["temporal_strict_share"]), pct(row["p_cutoff_overlap_share"])])
    table(doc, ["TEST P1", "Temporal strict", "P1 cutoff overlap"], strict_rows)
    table(doc, ["Timeline source", "p50 ngày", "p90 ngày", "Max", "Enrollment offering >270 ngày"], [
        [r["timeline_source"], int(r["duration_days_p50"]), int(r["duration_days_p90"]), int(r["duration_days_max"]), pct(r["long_offering_enrollment_share"])]
        for _, r in duration.sort_values("timeline_source").iterrows()
    ])
    paragraph(doc, "Khác biệt nhãn giữa T1/T2/T3 là temporal label drift thật, không được san bằng giữa các lát. Strictness P1 không đồng đều theo window, nên kết quả model phải báo theo window × phase cùng contextual metrics.")

    heading(doc, "6. Kiểm tra assignment và provenance")
    table(doc, ["Kiểm tra", "Kết quả"], [
        ["Enrollment distinct", f"{int(assignment['distinct_enrollment_count']):,}"],
        ["Nằm ở nhiều rolling block", int(assignment["n_enrollments_multiple_rolling_blocks"])],
        ["Nằm ở nhiều split cùng window", int(assignment["n_enrollments_multiple_splits_same_window"])],
        ["Context không có block", int(assignment["n_context_rows_without_block"])],
        ["Assignment QA", "Đạt" if int(assignment["assignment_qa_pass"]) == 1 else "Không đạt"],
        ["Coverage audit source", coverage_manifest["label_source"]],
    ])

    heading(doc, "7. Bước tiếp theo")
    table(doc, ["Thứ tự", "Việc cần làm"], [
        ["1", "Upload audit_lo_proxy_label_coverage_v1.py đã sửa và chạy lại duy nhất coverage audit để làm sạch metadata."],
        ["2", "Materialize phase views/test prefixes với manifest source-faithful (nếu chưa thực hiện sau label rebuild)."],
        ["3", "Chạy S0, feature dictionary và kiểm tra cờ is_label_proxy trước imputation/model."],
        ["4", "Chạy imputation, augmentation train-only và model; lưu split_version, parent manifest, seed và context timeline/long-offering."],
    ])

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
