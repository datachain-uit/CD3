"""Vietnamese localisation for the appended CQ clean-grain r2 report sections."""
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "CQ" / "docs" / "REPORT_CQ_V2.2_CLEAN_GRAIN_RNN_SEED42_V2_2026-10-07.docx"


PARAGRAPHS = {
    "8. Evidence acceptance r2": "8. Bằng chứng nghiệm thu r2 — CQ V2.2 clean-grain",
    "This section supersedes the preliminary": (
        "Mục này thay thế ghi chú nghiệm thu seed 42 sơ bộ. Nội dung được tạo từ export r2 "
        "kiểm soát provenance nghiêm ngặt sau khi phủ lại S_leak canonical trên dữ liệu thật của imputer cha. "
        "Không chạy lại mô hình, điền khuyết hoặc augmentation khi lập mục này."
    ),
    "8.1 Grain and lineage": "8.1 Grain và lineage",
    "The release lineage sidecar records": (
        "Sidecar lineage của release ghi nhận không có dòng trùng tuyệt đối trong labels, features, "
        "enrollment_windows và joined_base. Audit grain chỉ đọc cũng kiểm tra trực tiếp model input V0 và "
        "balanced bên dưới. Ở mọi bảng được audit, số dòng bằng số enrollment_id phân biệt và histogram chỉ có một dòng mỗi enrollment."
    ),
    "8.2 Canonical sanity and validation": "8.2 Sanity canonical và AccTEMPO trên VALIDATION",
    "Canonical S_leak coverage is": (
        "Độ phủ S_leak canonical là 120/120 facts theo imputer cha/split/phase. Mỗi dòng ghi phạm vi fit TRAIN "
        "gộp trên dữ liệu thật của imputer cha, phạm vi đánh giá cố định, số dòng fit probe và n_eval. Số giá trị "
        "canonical trùng tối đa theo khóa imputer cha/split/phase là 1; vì vậy các biến thể augmentation không thể fit lại leakage probe riêng."
    ),
    "8.3 Checkpoint selection": "8.3 Chọn checkpoint, support theo lớp, CI và lát đánh giá",
    "Per-class facts contain": (
        "Facts theo lớp có 1.224 dòng và facts bootstrap có 3.672 dòng; toàn bộ bootstrap dùng B=2.000. "
        "Chúng lưu support, cờ small-support, precision, recall, F1 và cận khoảng tin cậy theo lớp."
    ),
    "STRICT/OVERLAP are descriptive": (
        "STRICT/OVERLAP là các lát đánh giá mô tả. Chúng không thay đổi cohort TEST đã khóa, quy tắc chọn checkpoint "
        "hay kết quả primary trên toàn bộ dòng."
    ),
    "9. RadiusSMOTE fallback audit": "9. Audit fallback RadiusSMOTE — CQ V2.2",
    "RadiusSMOTE balanced TRAIN artifacts": (
        "Artifact TRAIN đã cân bằng bằng RadiusSMOTE được tạo một lần cho từng imputer cha và window tại augmentation seed 42. "
        "Audit hậu kiểm dưới đây chỉ phát lại một cách xác định quyết định chọn láng giềng; không thay đổi dòng tổng hợp, "
        "file TRAIN đã cân bằng hoặc bất kỳ run RNN nào."
    ),
    "All 24/24 class-level cells": (
        "Cả 24/24 ô window × lớp vượt ngưỡng cảnh báo 50% đã đăng ký. Fallback luôn cao hơn ở lớp good so với average. "
        "Do đó RadiusSMOTE với bán kính 0.092 cần được đọc như biến thể sensitivity có fallback láng giềng gần nhất thường xuyên, "
        "không phải bằng chứng về nội suy local-radius ổn định. So sánh hiệu năng vẫn hợp lệ vì audit chỉ có tính mô tả và artifact đã cân bằng là bất biến."
    ),
    "9.1 S2 sampling-plan coverage": "9.1 Độ phủ sampling plan S2",
    "The S2 ledger contains": (
        "Ledger S2 có 36 plan bất biến tại seed 42: bốn imputer cha × ba window × ba phương pháp. "
        "Mọi plan dùng chính sách mục tiêu IR10_K10 đã đăng ký; VALIDATION và TEST không bị tác động."
    ),
    "10. Ma tran nghiem thu": "10. Ma trận nghiệm thu GĐ1 clean-grain (r2)",
    "Bang nay doi chieu": (
        "Bảng này đối chiếu trực tiếp với checklist trong Danh_gia_GD1_clean-grain_2026-10-07. "
        "Tất cả số liệu là facts r2 của 51 run RNN clean-grain seed 42; không có bước train, impute hay augmentation nào được chạy lại khi lập bảng."
    ),
    "Luu y provenance": (
        "Lưu ý provenance: 51 run seed 42 hoàn thành trước khi manifest có trường created_at_utc; inventory vì vậy lưu "
        "manifest_mtime_utc như timestamp provenance minh bạch. Model runner hiện ghi created_at_utc cho mọi run mới và exporter "
        "ghi created_at_source để phân biệt timestamp gốc với fallback mtime."
    ),
}

CELLS = {
    "Window": "Window", "Input": "Đầu vào", "Rows": "Số dòng", "Distinct enrollment": "Enrollment phân biệt",
    "Max rows": "Số dòng tối đa", "Histogram": "Histogram", "Pipeline": "Pipeline",
    "Validation AccTEMPO": "AccTEMPO VALIDATION", "Runs": "Số run", "Checkpoints min": "Checkpoint nhỏ nhất",
    "Checkpoints max": "Checkpoint lớn nhất", "Within-margin min": "Trong margin nhỏ nhất",
    "Within-margin max": "Trong margin lớn nhất", "Split": "Tập", "Slice": "Lát",
    "Mean n": "n trung bình", "Macro-F1": "Macro-F1", "Balanced accuracy": "Balanced accuracy",
    "Class": "Lớp", "Synthetic rows": "Dòng tổng hợp", "Parent imputers": "Imputer cha",
    "Fallback mean": "Fallback trung bình", "Range": "Khoảng", "Method": "Phương pháp",
    "Pipelines": "Pipeline", "Real TRAIN": "TRAIN thật", "Synthetic": "Tổng hợp", "Balanced TRAIN": "TRAIN đã cân bằng",
    "Muc nghiem thu": "Mục nghiệm thu", "Bang chung hien co": "Bằng chứng hiện có", "Trang thai": "Trạng thái",
    "DAT": "ĐẠT", "DAT (co canh bao)": "ĐẠT (có cảnh báo)", "KE HOACH": "KẾ HOẠCH",
}


def _replace_paragraph(paragraph, value: str) -> None:
    if not paragraph.runs:
        paragraph.text = value
        return
    paragraph.runs[0].text = value
    for run in paragraph.runs[1:]:
        run.text = ""


def main() -> None:
    document = Document(REPORT)
    changed = 0
    for paragraph in document.paragraphs:
        for prefix, replacement in PARAGRAPHS.items():
            if paragraph.text.startswith(prefix):
                _replace_paragraph(paragraph, replacement)
                changed += 1
                break
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text in CELLS:
                    cell.text = CELLS[cell.text]
                    changed += 1
                elif cell.text == "V2 balanced TRAIN":
                    cell.text = "V2 TRAIN đã cân bằng"
                    changed += 1
    if changed < 20:
        raise RuntimeError(f"Only {changed} replacements; report layout may have changed.")
    document.save(REPORT)
    print(f"Localized {changed} items: {REPORT}")


if __name__ == "__main__":
    main()
