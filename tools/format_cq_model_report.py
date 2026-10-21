from docx import Document
from docx.shared import Pt


REPORT = "CQ/docs/REPORT_CQ_V1.docx"


def remove_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)


def heading(document, value):
    paragraph = document.add_paragraph(value)
    style = next(style for style in document.styles if style.name == "Heading 1")
    paragraph._p.get_or_add_pPr().get_or_add_pStyle().val = style.style_id


def add_table(document, headers, rows):
    result = document.add_table(rows=1, cols=len(headers))
    result.style = "Table Grid"
    for cell, value in zip(result.rows[0].cells, headers):
        cell.text = value
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in rows:
        for cell, value in zip(result.add_row().cells, row):
            cell.text = str(value)
    # A dedicated empty paragraph prevents adjacent Word tables from merging
    # visually into one dense block.
    document.add_paragraph()


def add_tree(document, lines):
    for line in lines:
        paragraph = document.add_paragraph(line)
        for run in paragraph.runs:
            run.font.name = "Courier New"
            run.font.size = Pt(8)


def remove_generated_tables(document):
    """Remove all tables owned by sections 14–15 before rebuilding them.

    python-docx appends tables near the document tail, independently of the
    paragraph list. Removing only the old paragraphs therefore left duplicate
    tables after repeated report updates.
    """
    generated_headers = {
        ("Hạng mục", "Quy ước đã khóa"),
        ("Window", "Shared RNN V0", "Sanity facts", "Ghi chú"),
        ("Nhóm", "Độ đo lưu trong L2", "Mục đích"),
        ("Nhóm", "Chỉ số", "Giá trị"),
        ("Window", "Test prefix", "S_san_plus", "S_perf", "AccTEMPO_M3"),
        ("Đoạn chạy", "Thời gian (giây)", "Diễn giải"),
        ("Window", "Read s", "Train/select s", "Inference s", "Total s", "CPU process s"),
        ("Window", "GPU", "Peak GPU memory MiB", "Cost USD", "Energy kWh", "CO2e g"),
        ("Tầng", "Nội dung", "Vai trò tái lập"),
    }
    for existing in list(document.tables):
        header = tuple(cell.text for cell in existing.rows[0].cells)
        if header in generated_headers:
            existing._element.getparent().remove(existing._element)


document = Document(REPORT)
# `Document.paragraphs` does not include tables.  Remove every XML body child
# from the generated section onward so rerunning this helper replaces, rather
# than appends, both prose and tables.  Preserve the final section properties.
remove_generated_tables(document)
start_paragraph = next(paragraph for paragraph in document.paragraphs
                       if paragraph.text.startswith("14. RNN V0 và sanity components"))
body_children = list(document._element.body)
start = body_children.index(start_paragraph._element)
for child in body_children[start:]:
    if not child.tag.endswith("}sectPr"):
        document._element.body.remove(child)
for paragraph in document.paragraphs:
    if paragraph.text.startswith("Cập nhật:"):
        paragraph.text = "Cập nhật: 26/09/2026"
        break

heading(document, "14. RNN V0 và sanity components (cập nhật 26/09/2026)")
document.add_paragraph("Phạm vi: CQ pipeline V0 (RAW_CONSTANT_FILL_WITH_MASKS), RNN, seed 42, GPU L4 và AMP bfloat16. Một shared checkpoint được huấn luyện trên sequence TRAIN P1–P4 đầy đủ.")
add_table(document, ["Hạng mục", "Quy ước đã khóa"], [
    ("Chọn checkpoint", "Mean validation Macro-F1 P1–P4; nếu hòa, chọn mean cross-entropy thấp hơn."),
    ("Suy luận test", "Checkpoint đóng băng chạy TEST/P1–P4; không retrain, fine-tune, refit scaler hay imputer."),
    ("Input V0", "Numeric train-only transform + missingness mask + availability/modality mask."),
    ("Missing / structural", "Neutral-fill chỉ observed missing; giữ missing__*, phase_available_* và modality masks."),
    ("ID code", "Raw/high-cardinality user_id, course_id và ID code không đưa vào tensor RNN."),
])
document.add_paragraph("Trạng thái thực thi CQ:")
add_table(document, ["Window", "Shared RNN V0", "Sanity facts", "Ghi chú"], [
    ("W1", "Hoàn thành", "8/8", "VALIDATION và TEST × P1–P4."),
    ("W2", "Hoàn thành", "8/8", "VALIDATION và TEST × P1–P4."),
    ("W3", "Hoàn thành", "8/8", "VALIDATION và TEST × P1–P4."),
])
document.add_paragraph("Mỗi phase/split lưu prediction table, probability matrix c0/c1/c2 và confusion matrix. L2 lưu Macro-F1, balanced accuracy, MCC, Cohen kappa, G-mean, per-class precision/recall/F1/support, NLL, multiclass Brier, ECE-15, cross-entropy cùng timing/resource facts.")
document.add_paragraph("Danh mục thước đo được tổng hợp theo từng window × split × phase:")
add_table(document, ["Nhóm", "Độ đo lưu trong L2", "Mục đích"], [
    ("Phân loại tổng thể", "accuracy, balanced_accuracy, f1_macro, f1_weighted, gmean_recall, MCC, Cohen kappa", "Đánh giá đúng lớp và cân bằng giữa các lớp."),
    ("Theo lớp", "precision, recall, F1, support cho c0/c1/c2; confusion matrix 3×3", "Nhìn được lỗi nhầm lớp và hiệu suất lớp thiểu số."),
    ("Calibration", "multiclass NLL, multiclass Brier, top-label ECE-15, cross-entropy", "Đánh giá xác suất dự báo, không chỉ nhãn argmax."),
    ("Probability integrity", "probability matrix c0/c1/c2; max row-sum error; non-finite check", "Bảo đảm xác suất hợp lệ và dùng lại được để tính metric."),
    ("Sanity / leakage", "s_nan, s_maj_jsd, s_ent, s_drift, s_eff, s_leak, S_san_plus; mask-probe AUC", "Kiểm tra ổn định phân bố, thông tin và nguy cơ leakage."),
    ("Tài nguyên", "time_read_s, time_train_selection_s, time_test_inference_s, time_run_total_s, cpu_process_s, peak_gpu_memory_mb, GPU name", "So sánh thời gian/CPU/GPU công bằng giữa run."),
    ("Chi phí / năng lượng", "cost_usd_run, energy_it_run_kwh, co2e_it_run_g", "Trường chuẩn hóa, để NULL có lý do nếu Modal không cấp telemetry."),
])
document.add_paragraph("Kết quả minh họa CQ/W1 TEST/P1:")
add_table(document, ["Nhóm", "Chỉ số", "Giá trị"], [
    ("Sanity", "s_nan", "1.000000"), ("Sanity", "s_maj_jsd", "0.992046"),
    ("Sanity", "s_ent", "0.076604"), ("Sanity", "s_drift", "0.999927"),
    ("Sanity", "s_eff", "0.697147"), ("Sanity", "s_leak", "0.452014"),
    ("Composite", "S_san_plus", "0.536872"),
    ("Composite", "S_perf / AccTEMPO_M3", "0.586045 / 0.565857"),
    ("Probability integrity", "max |sum(probability) − 1|", "1.45e-07"),
    ("Leakage diagnostic", "mask-probe AUC / fit phase rows", "0.773993 / 12,874,208"),
])
document.add_paragraph("Tổng hợp thước đo TEMPO trên TEST, đọc trực tiếp từ L2 sanity facts. S_san_plus phản ánh integrity/sanity; S_perf phản ánh hiệu năng dự đoán đã chuẩn hóa; AccTEMPO_M3 là composite của hai thành phần theo contract hiện hành.")
add_table(document, ["Window", "Test prefix", "S_san_plus", "S_perf", "AccTEMPO_M3"], [
    ("W1", "P1", "0.536872", "0.586045", "0.565857"),
    ("W1", "P2", "0.507194", "0.676861", "0.603072"),
    ("W1", "P3", "0.460321", "0.758121", "0.620966"),
    ("W1", "P4", "0.432689", "0.795612", "0.623578"),
    ("W2", "P1", "0.579973", "0.695191", "0.646585"),
    ("W2", "P2", "0.540640", "0.775347", "0.671214"),
    ("W2", "P3", "0.506236", "0.835429", "0.683734"),
    ("W2", "P4", "0.477544", "0.853966", "0.676815"),
    ("W3", "P1", "0.596628", "0.558751", "0.573604"),
    ("W3", "P2", "0.568729", "0.687037", "0.637016"),
    ("W3", "P3", "0.477223", "0.800665", "0.650968"),
    ("W3", "P4", "0.459266", "0.851591", "0.665219"),
])
document.add_paragraph("Sanity chạy hậu kiểm, không train lại RNN. acctempo_m3_v1_jsd_base2_six_components gồm s_nan, s_maj_jsd, s_ent, s_drift, s_eff, s_leak; S_san_plus là geometric mean của sáu component. s_maj_jsd = 1 − JSD_base2(phân bố lớp dự đoán, phân bố lớp thật); s_drift dùng TRAIN thực trước balancing làm tham chiếu.")
document.add_paragraph("Hiệu suất train/inference đã ghi nhận ở mức run (mỗi run dùng chung cho P1–P4, nên không nhân bốn chi phí train):")
add_table(document, ["Window", "Read s", "Train/select s", "Inference s", "Total s", "CPU process s"], [
    ("W1", "7.248", "574.128", "32.752", "622.797", "555.400"),
    ("W2", "7.574", "652.299", "32.371", "700.788", "635.050"),
    ("W3", "7.110", "357.082", "32.802", "405.999", "446.160"),
])
document.add_paragraph("Accelerator, cost và carbon telemetry:")
add_table(document, ["Window", "GPU", "Peak GPU memory MiB", "Cost USD", "Energy kWh", "CO2e g"], [
    ("W1", "NVIDIA L4", "80.387", "NULL", "NULL", "NULL"),
    ("W2", "NVIDIA L4", "80.387", "NULL", "NULL", "NULL"),
    ("W3", "NVIDIA L4", "80.387", "NULL", "NULL", "NULL"),
])
document.add_paragraph("NULL ở cost_usd_run, energy_it_run_kwh và co2e_it_run_g là giá trị có chủ đích: Modal worker không cung cấp billing/energy API cho run. Báo cáo không thay bằng ước lượng carbon để tránh trộn số đo thực tế với giả định. Khi có billing export hoặc hệ số điện năng/khu vực được phê duyệt, ba trường này sẽ được materialize lại ở L2 resource_usage với provenance nguồn.")

heading(document, "15. Cấu trúc lưu trữ và truy vết tái lập")
document.add_paragraph("Output bất biến được gắn theo task, feature regime, window, pipeline, model, seed, run_id và attempt_id. run_id là SHA-256 cấu hình chuẩn hóa; attempt_id đánh số lần thực thi. Không dùng tên file tự đặt làm định danh thực nghiệm.")
add_tree(document, [
    "tempo-data-v1/",
    "├── input/                                  # immutable source supplied to experiments",
    "│   └── CQ/",
    "│       ├── phase_views_v1/                   # TRAIN + VALIDATION, full P1–P4",
    "│       └── test_prefix_views_v1/",
    "│           ├── P1/                           # TEST observes P1 only",
    "│           ├── P2/                           # TEST observes P1–P2",
    "│           ├── P3/                           # TEST observes P1–P3",
    "│           └── P4/                           # TEST observes P1–P4",
    "│",
    "└── meta_release=imputation-v1/",
    "    ├── L0_registry/                         # state, schema/input hash, environment",
    "    │   ├── state_manifests/s0_CQ.json",
    "    │   ├── run_registry/run_id={run_id}.json",
    "    │   └── environments/attempt_id={attempt_id}.json",
    "    │",
    "    ├── L1_runs/                              # immutable executable artifacts",
    "    │   └── task=CQ/feature_regime=CQ_RAW_EARLY/window_id=W1/",
    "    │       └── pipeline_id=V0/",
    "    │           ├── model_name=IMPUTATION_ONLY/seed=20260922/",
    "    │           │   └── run_id={imputation_sha}/attempt_id={attempt}-NNNN/",
    "    │           │       ├── run_manifest.json",
    "    │           │       ├── fitted_preprocessors/",
    "    │           │       └── model_inputs/{train,validation,test_P1..test_P4}.parquet",
    "    │           └── model_name=RNN/seed=42/",
    "    │               └── run_id={model_sha}/attempt_id={attempt}-NNNN/",
    "    │                   ├── run_manifest.json",
    "    │                   ├── checkpoint/model.pt + train_history.parquet",
    "    │                   ├── predictions/{VALIDATION,TEST}/phase_id=P{1..4}.parquet",
    "    │                   ├── probability_matrices/{VALIDATION,TEST}/phase_id=P{1..4}.parquet",
    "    │                   ├── confusion_matrices/{VALIDATION,TEST}/phase_id=P{1..4}.parquet",
    "    │                   └── timing.parquet",
    "    │",
    "    ├── L2_facts/                             # queryable facts, partitioned by run/phase",
    "    │   ├── dataset_quality_summary/  dq_feature_profile/  imputation_fidelity/",
    "    │   ├── metrics_overall/  calibration/  confusion_matrix/",
    "    │   ├── sanity_components/  resource_usage/  drift_measures/",
    "    │   └── .../task=CQ/.../window_id=W1/phase_id=P1/.../run_id={...}/",
    "    │",
    "    └── L3_meta/                              # G3 output; build only after paired run grid is available",
    "        ├── meta_dataset_seed_level/           # one wide row per meta-observation / seed",
    "        ├── meta_dataset_seed_agg/             # seed-aggregated primary table",
    "        ├── meta_column_dictionary.parquet     # mandatory column-role/type dictionary",
    "        └── meta_release_manifest.json         # QA01–QA21, checksum, row-count summary",
])
add_table(document, ["Tầng", "Nội dung", "Vai trò tái lập"], [
    ("L0_registry", "state manifest, input/schema hash, environment, code/config hash", "Xác định đúng release và runtime."),
    ("L1_runs", "artifact, checkpoint, manifest, probability/confusion", "Khôi phục/chứng thực một run."),
    ("L2_facts", "performance, calibration, sanity_components, resource usage", "So sánh model/phase không đọc tensor lớn."),
    ("L3_meta", "meta_dataset_seed_level, meta_dataset_seed_agg, meta_column_dictionary, meta_release_manifest", "MDS wide cho meta-model, khuyến nghị và báo cáo; dựng lại được từ L2."),
])
document.add_paragraph("Trạng thái hiện tại: L3 chưa materialize, không phải vì L3 không có nội dung mà vì mới có một phần run V0/RNN. Builder G3 chỉ được chạy khi có các run paired trước/sau can thiệp cần thiết (V0, imputer-only và balanced V0–V16), model/seed đã đăng ký và đầy đủ L2 facts. Khi đó mỗi meta-row ghép Perf_before, Perf_after, DeltaPerf, chất lượng trước/sau, chi phí và nhãn khuyến nghị; sau đó chạy QA01–QA21 rồi khóa meta-release. Ví dụ partition L2 nguồn: L2_facts/sanity_components/task=CQ/feature_regime=CQ_RAW_EARLY/window_id=W1/phase_id=P1/stage=S3_MODEL/pipeline_id=V0/model_name=RNN/run_id={run_id}/attempt_id={attempt_id}/eval_split=TEST.parquet. cost_shared_across_phases=true để không nhân bốn chi phí train cho bốn test prefix.")
document.save(REPORT)
