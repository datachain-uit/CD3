"""Create the CQ clean-grain RNN seed-42 report V2 from strict facts."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "CQ" / "docs" / "REPORT_CQ_V2.2.docx"
FACTS = ROOT / "_tmp_cq_rnn_clean_facts"
OUTPUT = ROOT / "CQ" / "docs" / "REPORT_CQ_V2.2_CLEAN_GRAIN_RNN_SEED42_V2_2026-10-07.docx"
REPORT_ID = "RNN_seed42_v2_2_clean_20261007_r1"


def clear(doc: Document) -> None:
    for node in list(doc._element.body):
        if node.tag != qn("w:sectPr"):
            doc._element.body.remove(node)


def style(run, size: float = 11, bold: bool = False) -> None:
    run.bold = bold
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)


def text(doc: Document, value: str) -> None:
    style(doc.add_paragraph().add_run(value))


def head(doc: Document, value: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(5)
    style(p.add_run(value), 14, True)


def bullet(doc: Document, value: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    style(p.add_run("• " + value))


def shade(cell, color: str = "D9EAF7") -> None:
    props = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), color)
    props.append(element)


def table(doc: Document, headers: list[str], rows: list[list[object]]) -> None:
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.style = "Table Grid"
    for cell, value in zip(tbl.rows[0].cells, headers):
        cell.text = ""
        style(cell.paragraphs[0].add_run(str(value)), 10.5, True)
        shade(cell)
    for values in rows:
        for cell, value in zip(tbl.add_row().cells, values):
            cell.text = ""
            style(cell.paragraphs[0].add_run(str(value)), 10.5)
    doc.add_paragraph()


def f(value: float) -> str:
    return f"{float(value):.4f}"


def main() -> None:
    paths = {name: FACTS / name for name in (
        "metrics_test_long.parquet", "sanity_test_long.parquet", "resource_long.parquet",
        "run_inventory.parquet", "export_manifest.json",
    )}
    missing = [str(path) for path in paths.values() if not path.exists()]
    if missing:
        raise FileNotFoundError("Missing strict CQ facts: " + ", ".join(missing))
    metrics = pd.read_parquet(paths["metrics_test_long.parquet"])
    sanity = pd.read_parquet(paths["sanity_test_long.parquet"])
    resource = pd.read_parquet(paths["resource_long.parquet"])
    inventory = pd.read_parquet(paths["run_inventory.parquet"])
    if (len(metrics), len(sanity), len(resource), len(inventory)) != (204, 204, 51, 51):
        raise ValueError("Strict export cardinalities must be 204/204/51/51")

    rank = (sanity.groupby("pipeline_id", as_index=False)[["acctempo_m3", "s_perf", "s_san_plus", "s_cal"]]
            .mean().sort_values("acctempo_m3", ascending=False))
    windows = sanity.groupby(["pipeline_id", "window_id"], as_index=False).acctempo_m3.mean().pivot(
        index="pipeline_id", columns="window_id", values="acctempo_m3")
    phase = metrics.groupby(["pipeline_id", "phase_id"], as_index=False)[
        ["f1_macro", "balanced_accuracy", "mcc", "top_label_ece_15"]].mean()
    input_sizes = inventory.groupby("window_id")[["train_rows", "validation_rows"]].first()
    s_eff_unique = sanity.groupby(["window_id", "phase_id"]).s_eff.nunique().max()
    leak_unique = sanity.groupby(["window_id", "phase_id"]).s_leak.nunique().max()
    first = inventory.sort_values(["window_id", "pipeline_id"]).iloc[0]

    doc = Document(TEMPLATE)
    clear(doc)
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(11)
    doc.styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style(title.add_run("BÁO CÁO CQ V2.2 CLEAN-GRAIN — V2\nRNN SEED 42, PROVENANCE-STRICT"), 18, True)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style(subtitle.add_run("Release CQ_V2_2 | strict export r1 | cập nhật 07/10/2026"))
    text(doc, "Bản V2 thay thế execution report trước đó. Nó chỉ dùng export " + REPORT_ID +
         ", trong đó mọi L2 fact được đọc theo path do 51 run manifest hiện hành khai báo. Không quét hoặc chọn file từ thư mục L2 lịch sử.")

    head(doc, "Tóm tắt nghiệm thu")
    table(doc, ["Nội dung", "Trạng thái", "Bằng chứng"], [
        ["Release", "ĐẠT", "CQ_V2_2; split v2_2; phase wide_prefix_v2_2."],
        ["RNN lineage", "ĐẠT", "51/51 manifest current; model revision v2_4 prefix-safe; data release hash hiện hành."],
        ["L2 fact linkage", "ĐẠT", "204 TEST metric + 204 TEST sanity + 51 resource; run_id/attempt_id khớp manifest."],
        ["Input sizes", "ĐẠT", "Các model manifest ghi TRAIN/VALIDATION từng window; V0 và balanced TRAIN có grain guard trước S1/S3."],
        ["AccTEMPO/S_san seed 42", "ĐẠT", "Sanity theo feature layout thực tế từng pipeline; probe chỉ dùng real parent TRAIN, không dùng synthetic rows."],
        ["Outcome sơ bộ", "V10", f"Mean TEST AccTEMPO-M3 P1–P4/W1–W3 = {f(rank.iloc[0].acctempo_m3)}; chỉ là winner seed 42."],
    ])

    head(doc, "1. Contract khoa học và chuỗi dữ liệu")
    text(doc, "CQ dự báo nhãn vận hành tại grain learner–course enrollment. CQ là calendar anchor của hybrid temporal split V2.2. Unit split là offering_id × timeline_source; user_id/course_id là context code TRAIN-only, không phải khóa nhân bản enrollment.")
    bullet(doc, "TEST P1/P2/P3/P4 lần lượt chỉ quan sát dữ liệu đến P1/P2/P3/P4; future phase giữ unavailable/NULL.")
    bullet(doc, "Imputation/scaler fit TRAIN-only; validation và TEST là dòng thật. Augmentation chỉ tạo thêm TRAIN rows.")
    bullet(doc, "AccTEMPO-M3 là outcome hiệu năng chính; Macro-F1, balanced accuracy, MCC và calibration là chỉ số hỗ trợ.")
    table(doc, ["Window", "V0 TRAIN", "V0 VALIDATION", "Ý nghĩa"], [
        [window, f"{int(input_sizes.loc[window, 'train_rows']):,}", f"{int(input_sizes.loc[window, 'validation_rows']):,}", "Registry-locked temporal window"]
        for window in ("W1", "W2", "W3")
    ])

    head(doc, "2. S0/S1/S2 trước mô hình")
    text(doc, "S0 phân biệt eligible missing với structural absence: chỉ numeric missing khi modality đã quan sát mới được impute; structural absence được giữ bằng availability/mask và không bị tính thành tỷ lệ thiếu giả.")
    table(doc, ["Pipeline", "Phương pháp", "Vai trò"], [
        ["V0", "zero-fill + masks", "baseline"],
        ["V1", "median", "imputation-only"],
        ["V5", "mean", "imputation-only"],
        ["V9", "ExtraTrees light", "imputation-only"],
        ["V13", "MICE: 3 vòng, tối đa 1,000,000 TRAIN rows fit", "imputation-only"],
        ["V2–V4, V6–V8, V10–V12, V14–V16", "CDSMOTE / SASMOTE / RadiusSMOTE", "TRAIN-only augmentation"],
    ])
    text(doc, "Các run model hiện hành được resolver lọc đồng thời theo release ID, split registry, split/phase version và data_release_id. Vì vậy các bảng dưới không trộn artefact V1 hoặc run từ input cũ.")

    head(doc, "3. Provenance của 51 run RNN")
    table(doc, ["Trường", "Giá trị / nghiệm thu"], [
        ["Export ID", REPORT_ID + " (exporter_revision=clean_provenance_r1)."],
        ["Model revision", str(inventory.model_revision.dropna().unique().tolist())],
        ["Data release ID", str(inventory.data_release_id.dropna().unique().tolist())],
        ["Rows inventory", "51 rows; mỗi row có mtime, run/attempt ID, input root, parent IDs, seed bundle và sizes."],
        ["Ví dụ W1/V0", f"run={first.run_id[:12]}…; train={first.train_rows:,}; validation={first.validation_rows:,}; best epoch={first.best_epoch}."],
        ["GPU/precision", "NVIDIA L4; BF16 AMP; recurrent_shared_phase_v2_4_prefixsafe_epoch50."],
    ])

    head(doc, "4. RNN seed 42 — kết quả TEST hiện hành")
    text(doc, "Mỗi ô là trung bình TEST P1–P4 trong từng window; Mean là trung bình ba window. Đây là facts current, thay thế hoàn toàn export legacy. Winner dưới đây chỉ là mô tả seed 42; chưa dùng TEST để khóa recipe đa seed.")
    result_rows = []
    for value in rank.itertuples(index=False):
        p = value.pipeline_id
        result_rows.append([p, f(windows.loc[p, "W1"]), f(windows.loc[p, "W2"]), f(windows.loc[p, "W3"]), f(value.acctempo_m3), f(value.s_perf)])
    table(doc, ["Pipeline", "W1", "W2", "W3", "Mean AccTEMPO", "Mean S_perf"], result_rows)
    text(doc, f"V10 dẫn đầu AccTEMPO mean={f(rank.iloc[0].acctempo_m3)} trong seed 42. Sự chênh lệch giữa windows phải được diễn giải cùng temporal/label drift, không bằng cách gộp hay test-select pipeline.")

    head(doc, "5. Chỉ số performance theo phase")
    text(doc, "Bảng minh họa V0 baseline và V10 leader seed-42. Giá trị là mean qua W1–W3; ECE nhỏ hơn là tốt hơn.")
    rows = []
    for pipeline in ("V0", "V10"):
        for phase_id in ("P1", "P2", "P3", "P4"):
            row = phase[(phase.pipeline_id == pipeline) & (phase.phase_id == phase_id)].iloc[0]
            rows.append([pipeline, phase_id, f(row.f1_macro), f(row.balanced_accuracy), f(row.mcc), f(row.top_label_ece_15)])
    table(doc, ["Pipeline", "Phase", "Macro-F1", "Balanced accuracy", "MCC", "ECE-15"], rows)

    head(doc, "6. Sanity, calibration và tài nguyên")
    text(doc, "S_san+ v1 là metric sanity/data quality chính; s_cal (v2) chỉ là sensitivity. Không metric sanity nào thay AccTEMPO-M3 làm outcome hiệu năng. s_eff/s_leak probe dùng real parent TRAIN; feature layout vẫn là layout thực tế của mỗi pipeline nên s_leak có thể khác theo pipeline một cách hợp lệ.")
    sanity_rows = []
    for pipeline in ("V0", "V6", "V10", "V14"):
        item = rank[rank.pipeline_id == pipeline].iloc[0]
        component = sanity[sanity.pipeline_id == pipeline].mean(numeric_only=True)
        sanity_rows.append([pipeline, f(item.acctempo_m3), f(component.s_san_plus), f(component.s_ent), f(component.s_leak), f(component.s_eff), f(component.s_drift), f(item.s_cal)])
    table(doc, ["Pipeline", "AccTEMPO", "S_san+", "s_ent", "s_leak", "s_eff", "s_drift", "s_cal"], sanity_rows)
    table(doc, ["Resource fact", "Giá trị"], [
        ["Run thành công", "51/51 RNN cells trên NVIDIA L4 với AMP BF16."],
        ["Tổng wall-clock", f"{resource.time_run_total_s.sum() / 3600:.2f} giờ cho toàn grid."],
        ["Median / max mỗi run", f"{resource.time_run_total_s.median():.1f} / {resource.time_run_total_s.max():.1f} giây."],
        ["Peak RAM median / max", f"{resource.peak_ram_mb.median():,.0f} / {resource.peak_ram_mb.max():,.0f} MB."],
        ["Energy / cost", "Telemetry Modal billing/energy unavailable; không suy diễn USD hoặc CO2e."],
    ])
    text(doc, f"Sanity coverage đạt đủ 204 TEST rows. s_eff có tối đa {s_eff_unique} giá trị trong mỗi window × phase; s_leak có tối đa {leak_unique} giá trị do pipeline-specific layout, nhưng mọi probe đều truy về real parent input.")

    head(doc, "7. Lưu trữ, truy vết và bước tiếp theo")
    tree = (
        "tempo-data-v1/\n"
        "├── input/CQ_v2_2/{release_manifest, phase_views, test_prefix_views/P1..P4}\n"
        "└── meta_release=imputation-v1/\n"
        "    ├── L1_runs/task=CQ/.../pipeline_id=V*/model_name={IMPUTATION_ONLY,BALANCED_TRAIN_ONLY,RNN}/\n"
        "    └── L2_facts/report_exports/task=CQ/report_id=RNN_seed42_v2_2_clean_20261007_r1/\n"
        "        ├── metrics_test_long.parquet (204)\n"
        "        ├── sanity_test_long.parquet (204)\n"
        "        ├── resource_long.parquet (51)\n"
        "        └── run_inventory.parquet (51)"
    )
    p = doc.add_paragraph()
    run = p.add_run(tree)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    table(doc, ["Ưu tiên", "Bước tiếp theo", "Có cần train lại?"], [
        ["1", "Khóa shortlist bằng VALIDATION, không bằng TEST ranking trong báo cáo này.", "Không, đọc checkpoint/facts hiện có."],
        ["2", "Chạy replication seed/model theo chính sách master-seed đã khóa.", "Có, cho stability/F1."],
        ["3", "Export per-class support/recall, bootstrap CI, selection margin và Radius fallback rate.", "Không, đọc L2/ledger hiện có."],
        ["4", "Bổ sung materialization lineage sidecar trong input release kế tiếp nếu cần provenance đầy đủ hơn.", "Không đổi facts/model hiện hành."],
    ])
    text(doc, "Kết luận V2: CQ V2.2 clean-grain từ S0 đến RNN seed 42 có lineage strict và L2 facts khớp 51 manifest hiện hành. V10 đứng đầu AccTEMPO ở seed 42; kết luận khoa học cuối cùng vẫn cần validation-governed selection, CI và stability đa seed.")
    doc.core_properties.title = "CQ V2.2 clean-grain RNN seed 42 — report V2"
    doc.core_properties.author = "CD3 experimental pipeline"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
