"""Integrate the locked V2.2 RNN seed-42 report into REPORT_CQ_V1.docx."""
from __future__ import annotations

from copy import deepcopy
from pathlib import Path

import pandas as pd
from docx import Document


ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / "CQ" / "docs" / "REPORT_CQ_V1.docx"
RNN = ROOT / "CQ" / "docs" / "REPORT_CQ_RNN_SEED42_V2_2_2026-10-02.docx"
EXPORT = ROOT / "downloads" / "analysis" / "cq_rnn_seed42_v2_2" / "report_id=RNN_seed42_v2_2"


def set_cell(cell, text: str) -> None:
    cell.text = str(text)


def main() -> None:
    base = Document(BASE)
    rnn = Document(RNN)
    inventory = pd.read_parquet(EXPORT / "run_inventory.parquet")
    metrics = pd.read_parquet(EXPORT / "metrics_test_long.parquet")

    # Make the first page identify the report as the current V2.2 release,
    # while retaining its established REPORT_CQ_V1 filename for continuity.
    base.paragraphs[0].text = "BÁO CÁO QUY TRÌNH CQ — RELEASE THỰC NGHIỆM V2.2 VÀ RNN SEED 42"
    base.paragraphs[1].text = (
        "Phạm vi: learner–course enrollment | hybrid split V2.2 | "
        "imputation/augmentation/model RNN seed 42 | cập nhật 03/10/2026"
    )

    for paragraph in base.paragraphs:
        if paragraph.text == "Mục lục nội dung":
            paragraph.add_run(
                "\n9. Data quality, imputation và model contract"
                "\n10. Support, metric và điều kiện báo cáo"
                "\n11. Lộ trình thực nghiệm CQ"
                "\n12. Kết quả RNN seed 42: validation, TEST và AccTEMPO-M3"
            )
            break

    # Replace the formerly registry-only counts by the exact V2.2 model-input
    # counts. TEST is per prefix P1..P4 and is immutable for that window.
    v0 = inventory.loc[inventory.pipeline_id.eq("V0"), ["window_id", "train_rows", "validation_rows"]]
    test = metrics.loc[(metrics.pipeline_id == "V0") & (metrics.phase_id == "P1"), ["window_id", "n_rows"]]
    counts = v0.merge(test, on="window_id").set_index("window_id")
    split_table = next(table for table in base.tables if table.cell(0, 0).text == "W" and table.cell(0, 1).text == "Validation")
    headers = ["W", "TRAIN model input", "VALIDATION", "TEST mỗi prefix", "Temporal blocks", "Run lock"]
    for cell, value in zip(split_table.rows[0].cells, headers):
        set_cell(cell, value)
    for row, window, blocks in zip(split_table.rows[1:], ("W1", "W2", "W3"),
                                   ("A / B / C", "A+B+C / D / E", "A+B+C+D+E / F / G")):
        current = counts.loc[window]
        values = [window, f"{int(current.train_rows):,}", f"{int(current.validation_rows):,}",
                  f"{int(current.n_rows):,}", blocks, "v2_2 / wide_prefix_v2_2"]
        for cell, value in zip(row.cells, values):
            set_cell(cell, value)

    # The final roadmap table must no longer state that current work is pending.
    roadmap = next(table for table in base.tables if table.cell(0, 0).text == "Bước" and table.cell(0, 1).text == "Trạng thái")
    replacements = {
        "Materialize V2.2": ("Đạt", "phase_views_v2_2 và test_prefix_views_v2_2/P1…P4 là input khóa của model."),
        "P0 + S0": ("Đạt", "S0/DQ và registry V2.2 đã được materialize; facts được giữ tách khỏi release V1."),
        "V0 + imputation": ("Đạt", "V0, median, mean, ExtraTrees và MICE đã tạo model inputs V2.2; augmentation chỉ TRAIN."),
    }
    for row in roadmap.rows[1:]:
        if row.cells[0].text in replacements:
            status, detail = replacements[row.cells[0].text]
            set_cell(row.cells[1], status)
            set_cell(row.cells[2], detail)

    # Append all RNN report elements except its duplicate standalone title and
    # subtitle.  This preserves every exact TEST and AccTEMPO table.
    base.add_page_break()
    part_title = base.add_paragraph()
    title_run = part_title.add_run("PHẦN B — KẾT QUẢ THỰC NGHIỆM RNN SEED 42")
    title_run.bold = True
    title_run.font.size = base.styles["Normal"].font.size
    base.add_paragraph(
        "Phần này nối tiếp protocol ở Phần A. Nguồn kết quả là immutable L2 export "
        "RNN_seed42_v2_2: 17 pipeline × 3 window × 4 TEST prefix = 204 cell. "
        "Không dùng output V1 hay test để chọn checkpoint."
    )
    children = list(rnn.element.body)
    # First two children are the standalone report title and subtitle.  Keep
    # its scope paragraph and all subsequent headings, paragraphs and tables.
    for child in children[2:]:
        if child.tag.endswith("sectPr"):
            continue
        base.element.body.append(deepcopy(child))

    base.save(BASE)
    print(BASE)


if __name__ == "__main__":
    main()
