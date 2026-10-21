"""Update existing CQ/LO reports while preserving a pre-audit backup."""
from pathlib import Path
import shutil
from docx import Document
from docx.shared import Pt


def replace_paragraph(document, match, replacement):
    hits = [p for p in document.paragraphs if match in p.text]
    if len(hits) != 1:
        raise ValueError(f"Expected one paragraph containing {match!r}; got {len(hits)}")
    hits[0].text = replacement


def table(doc, headings, rows):
    t = doc.add_table(rows=1, cols=len(headings))
    t.style = "Table Grid"
    for cell, text in zip(t.rows[0].cells, headings):
        cell.text = text
    for row in rows:
        for cell, value in zip(t.add_row().cells, row):
            cell.text = str(value)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10)


def update_cq(doc):
    replace_paragraph(doc, "Cần chạy/đưa vào evidence audit prefix-alignment cho CQ",
        "Đã bổ sung audit CQ W1/P1 cho RNN Modal và LSTM/GRU/BiLSTM VastAI, seed 42; "
        "các kết quả và giới hạn được trình bày ở phụ lục audit ngày 17/10. "
        "Chưa thể kết luận validation selection không thiên lệch cho đến khi kiểm tra tensor validation "
        "và replay checkpoint. PatchTST/MOMENT, seed 43/44 và replication backend cùng revision "
        "vẫn là các pha tiếp theo; không suy ra stability từ một seed.")
    doc.add_heading("Phụ lục: audit prefix và course cohort CQ — 17/10/2026", level=1)
    doc.add_paragraph("Phạm vi: RNN Modal và LSTM, GRU, BiLSTM VastAI; mỗi family 51 cell (17 V × 3 window), seed 42. "
                      "Kiểm tra predictions và phân nhóm tập trung vào W1; không train lại. "
                      "153 manifest VastAI khớp inventory và 51 RNN được chọn theo release và data_release_id.")
    table(doc, ["Kiểm tra", "Kết quả", "Giới hạn"], [
        ("Layout", "204 layout; mỗi family có một signature", "13 cột static; cần duyệt nguồn và thời điểm sẵn có"),
        ("TEST static P1–P4", "Các cột static giống nhau theo enrollment", "Không chứng minh bản thân cột static an toàn khỏi leakage"),
        ("W1 predictions", "RNN 136 cell; ba family VastAI 408 cell", "17 V × 2 split × 4 phase cho mỗi family"),
        ("W1/P1 phân nhóm", "RNN 68; ba family VastAI 204 kết quả", "Ghép theo salted enrollment hash và đối chiếu nhãn"),
    ])
    doc.add_paragraph("Quét toàn bộ file TEST prefix CQ đã tải: mỗi phase có 538.608 dòng, gộp W1–W3. "
                      "Chỉ xét các cột hậu tố phase tương lai, không gồm phase_available. "
                      "P4 không có phase tương lai; đây là kiểm tra giá trị null trong nguồn view, "
                      "không thay thế kiểm tra tensor validation.")
    table(doc, ["Prefix", "Dòng quét", "Ô phase tương lai", "Ô không null"], [
        ("P1", "538.608", "122.802.624", "0"), ("P2", "538.608", "81.868.416", "0"),
        ("P3", "538.608", "40.934.208", "0"), ("P4", "538.608", "0", "0"),
    ])
    doc.add_paragraph("Trong view TRAIN/VALIDATION và TEST P1–P4 đã quét, primary_risk_set luôn bằng 1 "
                      "và split_rule_version chỉ có giá trị offering_temporal_label_balanced_v2_2. "
                      "Đây là bằng chứng ở view nguồn, không tự chứng minh mọi phép biến đổi model input. "
                      "Cần duyệt cả 13 cột static, đặc biệt timeline_source và window_duration_seconds; "
                      "code hiện tại loại split_rule_version, trong khi layout của run lịch sử vẫn chứa cột này.")
    table(doc, ["W1/P1 TEST", "Enrollment", "warning", "average", "good"], [
        ("Seen course", "166.964", "146.031", "9.629", "11.304"),
        ("Unseen course", "11.131", "10.661", "467", "3"),
    ])
    doc.add_paragraph("Chỉ có 3 mẫu good trong nhóm unseen; mọi so sánh macro-F1/per-class ở nhóm này "
                      "cần kèm support và không được diễn giải như ước lượng ổn định. "
                      "STRICT/OVERLAP và SEEN/UNSEEN là lát mô tả, không thay cohort TEST chính.")
    doc.add_paragraph("Việc còn mở: duyệt định nghĩa/thời điểm sẵn có 13 static; kiểm tra validation-prefix "
                      "và replay predictions từ checkpoint trong môi trường phù hợp trước khi chốt selection. "
                      "Chi tiết: CQ/resuilt/cq_full_prefix_static_audit_20261017, "
                      "CQ/resuilt/audit_prefix_validation_alignment_v1_20261017 và "
                      "CQ/resuilt/rnn_prefix_audit_20261017_files.")


def update_lo(doc):
    replace_paragraph(doc, "Audit read-only prefix-validation đã chạy trên release LO",
        "Audit read-only đã kiểm tra 204 layout recurrent LO V3.1: một signature và 15 static field. "
        "Mẫu TEST prefix được kiểm tra không có giá trị future-dynamic không null; "
        "các static field được đối chiếu theo enrollment giữa P1–P4. Không thể từ mẫu này "
        "khẳng định toàn bộ TEST sạch hoặc xác nhận trực tiếp validation-prefix/checkpoint. "
        "Các kết quả seen/unseen và STRICT/OVERLAP trước đây được ghép theo thứ tự nhãn, "
        "chưa chứng minh định danh từng enrollment; cần chạy lại ghép theo hash trước khi dùng làm kết luận.")
    replace_paragraph(doc, "Audit prefix-alignment đã đóng gate W1/P1",
        "Audit prefix-alignment LO cung cấp bằng chứng ban đầu, chưa đóng gate W1/P1. "
        "Cần quét toàn bộ future-mask TEST, duyệt 15 static field, kiểm tra validation-prefix và replay checkpoint; "
        "các phân tích theo course cohort cần đối chiếu lại theo enrollment hash. "
        "Chưa kết luận cần hay không cần train lại hoặc chọn lại checkpoint trước khi hoàn tất kiểm tra. "
        "Các kết quả seed 42 vẫn nên trình bày kèm support và phạm vi evidence; "
        "PatchTST/MOMENT và seed replication là các pha tiếp theo.")
    doc.add_heading("Cập nhật phạm vi bằng chứng audit — 17/10/2026", level=1)
    doc.add_paragraph("Các số liệu LO trước đó vẫn được giữ để tham khảo, nhưng kết luận theo cohort là sơ bộ "
                      "cho đến khi xác minh ghép đúng enrollment. Kiểm tra mẫu TEST không chứng minh toàn bộ "
                      "phase-prefix validation hay tính an toàn về thời điểm của các static field.")


def main():
    for filename, updater in (
        ("CQ/docs/REPORT_CQ_FULL_E2E_SEED42_2026-10-16.docx", update_cq),
        ("LO/docs/REPORT_LO_FULL_E2E_SEED42_2026-10-16.docx", update_lo),
    ):
        path = Path(filename)
        backup = path.with_name(path.stem + "_PRE_AUDIT_20261017.docx")
        if backup.exists():
            raise FileExistsError(f"Backup already exists; refusing to overwrite: {backup}")
        doc = Document(path)
        updater(doc)
        shutil.copy2(path, backup)
        doc.save(path)
        print(f"updated {path}; backup {backup}")


if __name__ == "__main__":
    main()
