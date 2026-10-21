"""Create updated CQ/LO end-to-end reports without overwriting historical copies."""
import json
from pathlib import Path
from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
probe_dir = ROOT / 'LO/resuilt/static_probe_w1_v0_20261019'
probe = json.loads((probe_dir/'probe.json').read_text(encoding='utf-8'))
audit = json.loads((probe_dir/'static_audit.json').read_text(encoding='utf-8'))
scan = json.loads((probe_dir/'full_scan.json').read_text(encoding='utf-8'))

def add_table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = 'Table Grid'
    for c, v in zip(t.rows[0].cells, headers): c.text = str(v)
    for row in rows:
        for c, v in zip(t.add_row().cells, row): c.text = str(v)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs: run.font.size = Pt(11)
    doc.add_paragraph('')

def historical_mark(doc, starts, note):
    for p in doc.paragraphs:
        if any(p.text.startswith(s) for s in starts):
            p.text = note + p.text

lo = Document(ROOT/'LO/docs/REPORT_LO_FULL_E2E_SEED42_2026-10-16.docx')
historical_mark(lo, ('Audit read-only đã kiểm tra 204 layout', 'Audit prefix-alignment LO cung cấp',
                     'Các số liệu LO trước đó vẫn được giữ'),
                '[Lịch sử trước cập nhật 19/10: xem phụ lục mới ở cuối báo cáo] ')
lo.add_page_break(); lo.add_heading('Cập nhật kiểm tra LO — 19/10/2026', 0)
lo.add_paragraph('Áp dụng cho LO V3.1 clean-grain. Những kết quả dưới đây bổ sung báo cáo hiện hữu; full-scan thay thế giới hạn chỉ kiểm tra mẫu đối với mask TEST prefix, nhưng không thay thế kiểm tra validation-prefix, truy vết cohort hoặc ablation.')
lo.add_heading('1. Quét toàn bộ mask TEST prefix', 1)
lo.add_paragraph('Đọc toàn bộ file và row group trong input LO V3.1 trên Modal; mỗi dòng là một enrollment–window trong view test. P4 không có phase tương lai để kiểm tra. Null và NaN đều được coi là vắng dữ liệu trong kiểm tra này.')
add_table(lo, ['Phase','File / row group','Dòng','Ô tương lai','Ô không null','Đạt'], [[x['phase'],f"{x['files']} / {x['row_groups']}",f"{x['rows']:,}",f"{x['future_cells']:,}",f"{x['non_null_future_cells']:,}",'Có' if x['pass'] else 'Không'] for x in scan['results']])
lo.add_paragraph('Nguồn bằng chứng: LO/resuilt/static_probe_w1_v0_20261019/full_scan.json; Modal app ap-szzzh6D513kiHIaTzn9UAK. Phép thử kiểm tra các trường có hậu tố phase trong TEST prefix; không chứng minh mọi feature static sẵn có tại thời điểm dự đoán.')
lo.add_heading('2. Hợp đồng và audit 15 cột static', 1)
lo.add_paragraph('Layout RNN W1/V0 seed 42 ghi 15 cột. Audit dùng mẫu TRAIN 400.000 dòng seed 42; VALIDATION và TEST P1 không lấy mẫu. OFFERING_ATTRIBUTE là cờ thống kê theo course_id_code hiện có, không tự chứng minh không rò rỉ; tính sẵn có theo thời gian vẫn cần xác nhận.')
add_table(lo, ['Cột', 'Quy định registry hiện tại', 'Cờ trong audit'], [[r['column'],r['allowlist_status'],', '.join(r['flags']) or '—'] for r in audit['columns']])
lo.add_paragraph('split_rule_version nằm trong layout run cũ nhưng registry hiện xếp EXCLUDE; audit mẫu thấy cột này hằng. Đây là sai lệch hợp đồng cần truy vết, không đủ để kết luận có leakage hay gây thay đổi dự đoán. Quyết định timeline_source và window_duration_seconds cho lưới mới chưa được khóa; không áp lại vào run đã hoàn tất.')
lo.add_heading('3. Probe chỉ dùng static — LO W1/V0', 1)
lo.add_paragraph('Logistic regression C=1, median imputation và standardization fit trên mẫu TRAIN; dùng cả cấu hình không trọng số và class_weight=balanced. VALIDATION và TEST dùng toàn bộ các dòng trong model input. Macro-F1 chính dùng ba lớp cố định; macro theo lớp hiện diện được giữ riêng trong JSON. Probe không phải RNN và không chứng minh nhân quả.')
rows=[]
for group, b in probe['feature_sets'].items():
    for weighting in ('none','balanced'):
        for split, s in b[f'class_weight_{weighting}'].items():
            rows.append([group,weighting,split,str(s['n_rows']),f"{s['macro_f1']:.4f}",f"{s['balanced_accuracy']:.4f}",f"{s['log_loss']:.4f}",f"{s['per_class']['G']['recall']:.4f}",f"{s['per_class']['E']['recall']:.4f}"])
add_table(lo,['Nhóm','Trọng số','Split','N','Macro-F1','Bal. acc.','Log loss','Recall G','Recall E'],rows)
lo.add_paragraph('Trong cấu hình balanced, COURSE_FINGERPRINT có balanced accuracy 0,6422 trên validation và 0,2067 trên test; chỉ là dấu hiệu phân phối/tín hiệu khác nhau. Cần ablation và đối chiếu cohort trước khi quy nguyên nhân cho nhận diện khóa học. Bảng đầy đủ, support từng lớp và danh sách cột: LO/resuilt/static_probe_w1_v0_20261019/probe.json.')
lo.add_heading('4. Những việc còn mở', 1)
lo.add_paragraph('Chưa có kết quả hai cell ablation bỏ static; chưa chứng minh selection/checkpoint và phép ghép cohort theo enrollment hash ở mọi cell. Đề xuất cho lưới mới: timeline_source làm CONTEXT_ONLY, window_duration_seconds giữ trong tensor nếu được chứng minh đã biết lúc dự đoán. Đây là đề xuất chờ chủ protocol xác nhận, không phải thay đổi đã thực hiện. Các kết luận về seed 43/44, năng lượng toàn hệ thống và CO₂ không được suy từ bộ seed 42 hiện có.')
lo_out=ROOT/'LO/docs/REPORT_LO_FULL_E2E_SEED42_UPDATED_2026-10-19.docx'; lo.save(lo_out)

cq = Document(ROOT/'CQ/docs/REPORT_CQ_FULL_E2E_BATCH2048_SEED42_2026-10-19.docx')
historical_mark(cq, ('Các facts dưới đây là export L2 đã ingest', 'Trong view TRAIN/VALIDATION và TEST',
                    'Việc còn mở: duyệt định nghĩa', 'Đã kiểm tra 153 run_manifest.json trong bundle CQ VastAI',
                    'LSTM/GRU và BiLSTM thuộc các runner provenance khác nhau', 'Các manifest lịch sử mẫu không có energy_null_reason'),
                '[Lịch sử bộ cũ; xem cập nhật batch 2.048 và phụ lục cuối] ')
cq.add_page_break(); cq.add_heading('Tổng kết đối chiếu CQ — 19/10/2026', 0)
cq.add_paragraph('Báo cáo giữ lại dữ liệu lịch sử để truy vết. Kết quả LSTM/GRU/BiLSTM batch 2.048 trong phần cập nhật gần cuối là bộ hiện hành; không gộp với bộ batch 512 cũ. V0–V16 gồm 17 pipeline. RNN và các model khác không nằm trong đợt ingest recurrent này.')
add_table(cq,['Nội dung','Bằng chứng / trạng thái'],[
 ['Recurrent batch 2.048','153 run: 51 LSTM, 51 GRU, 51 BiLSTM; W1–W3, V0–V16, seed 42.'],
 ['Lineage và lưu trữ','Manifest batch mới được đối chiếu; cùng runner hash và GPU NVIDIA L4 trong ingest. Source files đã so SHA-256 với VastAI.'],
 ['L2 chính thức','Ingest vastai_cq_f1_batch2048_seed42_slice_r2_20261019; 66 partition, receipt và số partition đã xác nhận.'],
 ['Sanity và slice r2','Mỗi family 408 sanity, 1.224 per-class, 3.672 bootstrap B=2.000, 1.224 ALL/STRICT/OVERLAP và 51 checkpoint selection. ALL đối chiếu khớp sanity.'],
 ['CQ prefix','Full scan nguồn: 538.608 dòng mỗi phase, không có ô future-phase không null.'],
 ['Năng lượng','Một manifest batch mới có NVML GPU-only estimate; chưa xác nhận đầy đủ mọi run. Không suy CO₂ hay năng lượng toàn hệ thống.'],
 ['Còn mở','Duyệt thời điểm sẵn có static, validation-prefix/checkpoint và các thí nghiệm, seed chưa có bằng chứng; không coi ingest là bằng chứng đóng mọi yêu cầu khoa học.']])
cq.add_paragraph('AccTEMPO là độ đo chính; S_perf và sáu thành phần sanity cần đọc riêng. S_leak của slice dùng canonical real-parent theo toàn eval split/phase; drift và metric dự đoán được tính trên từng slice. STRICT/OVERLAP là phân tích mô tả, không thay tập TEST khóa.')
cq_out=ROOT/'CQ/docs/REPORT_CQ_FULL_E2E_SEED42_UPDATED_2026-10-19.docx'; cq.save(cq_out)
print(lo_out); print(cq_out)
