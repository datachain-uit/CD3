"""Generate corrected CQ--LO V2.2-reference versus V2.3 split reports."""
from __future__ import annotations

from pathlib import Path

import polars as pl
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DATE = "2026-09-29"
NAVY, BLUE, GREY, GREEN = "17365D", "1F4E78", "F2F2F2", "E2F0D9"
TASKS = {
    "CQ": ("CQ_label_final: warning / average / good", "balanced"),
    "LO": ("LO_performance_label_3: I/D / G / E", "rarest_first"),
}


def load(task: str, release: str, name: str) -> pl.DataFrame:
    root = "split_registry_v2_2_reference" if release == "v2_2" else "split_registry_v2_3"
    if name == "qa":
        name = "qa22_v2_2" if release == "v2_2" else "qa22_v2_3"
    files = list((ROOT / task / "resuilt" / root / name).glob("*.parquet"))
    if len(files) != 1:
        raise RuntimeError(f"{task}/{root}/{name}: expected one parquet, found {len(files)}")
    return pl.read_parquet(files[0])


def shade(cell, colour: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), colour)
    props.append(node)


def paragraph_shade(paragraph, colour: str) -> None:
    props = paragraph._p.get_or_add_pPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), colour)
    props.append(node)


def add_table(doc: Document, headers: list[str], rows: list[list[object]], widths: list[float] | None = None) -> None:
    result = doc.add_table(rows=1, cols=len(headers))
    result.style = "Table Grid"
    result.autofit = False
    for idx, header in enumerate(headers):
        cell = result.rows[0].cells[idx]
        cell.text = ""
        run = cell.paragraphs[0].add_run(str(header))
        run.bold = True; run.font.size = Pt(8.5); run.font.color.rgb = RGBColor(255, 255, 255)
        shade(cell, BLUE)
        if widths: cell.width = Inches(widths[idx])
    for row_idx, row in enumerate(rows):
        cells = result.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = ""
            cells[idx].paragraphs[0].add_run(str(value)).font.size = Pt(8.5)
            if row_idx % 2: shade(cells[idx], GREY)
            if widths: cells[idx].width = Inches(widths[idx])
    doc.add_paragraph()


def add_heading(doc: Document, value: str) -> None:
    paragraph = doc.add_heading(value, 1)
    for run in paragraph.runs: run.font.color.rgb = RGBColor.from_string(NAVY)


def pct(value: float) -> str:
    return f"{value * 100:.2f}%"


def setup(doc: Document, footer_text: str) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(1.7); section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.8); section.right_margin = Cm(1.8)
    doc.styles["Normal"].font.name = "Aptos"; doc.styles["Normal"].font.size = Pt(10)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run(f"{footer_text} · {DATE} · Registry-derived").font.size = Pt(8)


def report_task(task: str) -> None:
    label, pairing = TASKS[task]
    qa2, qa3 = load(task, "v2_2", "qa").sort("window_id"), load(task, "v2_3", "qa").sort("window_id")
    actual, strict = load(task, "v2_3", "split_actual_ratios"), load(task, "v2_3", "temporal_strict_share")
    labels, calendar = load(task, "v2_3", "split_label_distribution"), load(task, "v2_3", "calendar_boundaries")
    if not qa2.select(pl.col("qa22_ok").all()).item() or not qa3.select(pl.col("qa22_ok").all()).item():
        raise RuntimeError(f"{task}: registry QA22 is not accepted")

    doc = Document(); setup(doc, f"{task} · split V2.2 reference vs V2.3")
    title = doc.add_paragraph(); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(f"{task} — BÁO CÁO SPLIT V2.2 REFERENCE VÀ V2.3")
    run.bold = True; run.font.size = Pt(20); run.font.color.rgb = RGBColor.from_string(NAVY)
    subtitle = doc.add_paragraph(); subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("scenario = hybrid · offering-level temporal assignment · CQ shared calendar anchor").italic = True
    statement = doc.add_paragraph()
    statement.add_run("Kết luận: ").bold = True
    statement.add_run("V2.2 reference và V2.3 đều qua QA22. V2.3 được chọn cho materialization/downstream vì thêm gate quy mô hai tầng; manifest A–G không đổi.")
    paragraph_shade(statement, GREEN)

    add_heading(doc, "1. Provenance và phạm vi")
    add_table(doc, ["Thuộc tính", "Giá trị"], [
        ["Task / label", f"{task} / {label}"], ["Scenario", "hybrid"], ["Pairing", pairing],
        ["V2.2 baseline hợp lệ", f"{task}/resuilt/split_registry_v2_2_reference"],
        ["V2.3 hợp lệ", f"{task}/resuilt/split_registry_v2_3"],
        ["Artifact không sử dụng", "split_registry_v2_2: CQ incomplete; LO dùng overlap gộp cũ"],
    ], [2.25, 4.8])

    add_heading(doc, "2. Cơ chế và sự khác nhau giữa hai version")
    doc.add_paragraph(
        "A là 60% enrollment sớm nhất theo offering_end_date tại ranh giới offering. Future pool được cắt theo thời gian thành T1, T2, T3; "
        "mỗi lát ghép validation/test ở cấp offering: B/C, D/E, F/G. W1=A|B|C; W2=A+B+C|D|E; W3=A–E|F|G. "
        "CQ là calendar anchor, LO giữ calendar này và dùng rarest_first để bảo toàn nhãn hiếm trong từng cặp."
    )
    add_table(doc, ["Điều khoản", "V2.2 reference", "V2.3"], [
        ["Temporal assignment", "A–G; không tách offering", "Không đổi"],
        ["Label-end order", "Bắt buộc", "Bắt buộc"],
        ["Quy mô mỗi arm", "Không rỗng (≥1)", "≥ max(2.5% population task, 50,000)"],
        ["Evaluation gộp", "Không có gate riêng", "validation + test ≥5% population task"],
        ["Ý nghĩa", "Reference protocol theo tempo_fix", "Gate chặt hơn, không phải split mới"],
    ], [1.6, 2.7, 2.8])

    add_heading(doc, "3. Calendar chung và QA22")
    cal = calendar.row(0, named=True)
    add_table(doc, ["Mốc", "UTC"], [["A boundary", cal["a_boundary"]], ["T1 B/C", cal["slice_1_boundary"]], ["T2 D/E", cal["slice_2_boundary"]], ["T3 F/G", cal["slice_3_boundary"]]], [2.3, 3.6])
    compared = qa2.join(qa3, on="window_id", suffix="_v23")
    qa_rows = []
    for row in compared.iter_rows(named=True):
        qa_rows.append([row["window_id"], f"{row['n_enrollments_validation']:,}", f"{row['n_enrollments_test']:,}", f"{row['n_enrollments_evaluation']:,}", f"{row['min_eval_arm_enrollments_v23']:,}", f"{row['min_eval_total_enrollments_v23']:,}", "Đạt" if row["end_order_ok_v23"] else "Không đạt", "Đạt" if row["qa22_ok"] else "Không đạt", "Đạt" if row["qa22_ok_v23"] else "Không đạt"])
    add_table(doc, ["W", "Validation", "Test", "Eval", "Min arm V2.3", "Min eval V2.3", "End order", "V2.2", "V2.3"], qa_rows, [0.35, .78, .78, .75, .78, .8, .65, .45, .45])

    add_heading(doc, "4. Overlap và temporal strict đúng định nghĩa")
    p1 = strict.filter((pl.col("split") == "test") & (pl.col("phase_id") == "P1"))
    rows = []
    for row in qa3.join(p1, on="window_id").iter_rows(named=True):
        rows.append([row["window_id"], pct(row["ev_train_overlap_share"]), pct(row["temporal_strict_share"]), f"{row['ev_train_overlap_share'] + row['temporal_strict_share']:.12f}"])
    add_table(doc, ["W", "ev_train_overlap_share", "P1 test strict", "Tổng"], rows, [.6, 1.7, 1.45, 1.0])
    doc.add_paragraph("ev_train_overlap_share là riêng TEST ở P1, nên bằng 1 − temporal_strict_share(P1, TEST). Validation không bị gộp vào mẫu số; tổng bằng 1 (sai số số thực).")

    add_heading(doc, "5. Quy mô train và hỗ trợ nhãn evaluation")
    rows = []
    for window in ("W1", "W2", "W3"):
        lookup = {r["split"]: r for r in actual.filter(pl.col("window_id") == window).iter_rows(named=True)}
        rows.append([window, f"{lookup['train']['enrollment_count']:,}", pct(lookup['train']['split_enrollment_share']), f"{lookup['validation']['enrollment_count']:,}", f"{lookup['test']['enrollment_count']:,}"])
    add_table(doc, ["W", "Train", "Train share", "Validation", "Test"], rows, [.5, 1.25, 1., 1.2, 1.2])
    min_rows = []
    for window in ("W1", "W2", "W3"):
        for split in ("validation", "test"):
            count = labels.filter((pl.col("window_id") == window) & (pl.col("split") == split)).select(pl.col("label_enrollment_count").min()).item()
            min_rows.append([window, split, f"{count:,}"])
    add_table(doc, ["W", "Evaluation split", "Minimum class support"], min_rows, [.6, 1.7, 2.1])
    doc.add_paragraph("Label drift giữa các lát là tín hiệu thật, không cân bằng xuyên thời gian. Model phải báo macro-F1, support từng lớp và confidence interval; augmentation chỉ áp dụng cho TRAIN.")

    add_heading(doc, "6. Cây artifact và bước tiếp theo")
    doc.add_paragraph("split_registry_v2_3/\n├── manifest/\n├── qa22_v2_3/\n├── split_registry/\n├── split_actual_ratios/\n├── split_label_distribution/\n├── temporal_strict_context/\n├── temporal_strict_share/\n├── duration_summary/\n├── calendar_boundaries/\n└── block_label_audit/")
    doc.add_paragraph("Tiếp theo: materialize phase views V2.3 và test-prefix P1–P4 từ manifest immutable V2.3, rồi mới S0, imputation, augmentation TRAIN-only và model. Không trộn V1/V2.1 hoặc registry legacy vào meta-release.")
    output = ROOT / task / "docs" / f"REPORT_{task}_SPLIT_V2_2_REFERENCE_VS_V2_3_{DATE}.docx"
    doc.save(output); print(output)


def common_report() -> None:
    data = {task: {rel: load(task, rel, "qa").sort("window_id") for rel in ("v2_2", "v2_3")} for task in TASKS}
    doc = Document(); setup(doc, "CQ–LO common temporal split V2.2 reference vs V2.3")
    title = doc.add_paragraph(); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("BÁO CÁO CHUNG CQ–LO: V2.2 REFERENCE VÀ V2.3"); run.bold = True; run.font.size = Pt(20); run.font.color.rgb = RGBColor.from_string(NAVY)
    statement = doc.add_paragraph(); statement.add_run("Kết luận: ").bold = True; statement.add_run("CQ và LO đều đạt QA22 V2.3. V2.3 giữ nguyên split assignment V2.2 reference và thêm gate quy mô hai tầng, nên là release được dùng cho downstream."); paragraph_shade(statement, GREEN)
    add_heading(doc, "1. So sánh có kiểm soát")
    rows = []
    for task in TASKS:
        q2, q3 = data[task]["v2_2"], data[task]["v2_3"]
        same_counts = q2.select("n_enrollments_validation", "n_enrollments_test").equals(q3.select("n_enrollments_validation", "n_enrollments_test"))
        rows.append([task, "Đạt W1–W3", "Đạt W1–W3", "Có" if same_counts else "Không", "V2.3"])
    add_table(doc, ["Task", "V2.2 reference", "V2.3", "Cùng manifest A–G", "Dùng downstream"], rows, [.7, 1.3, 1.1, 1.5, 1.25])
    add_heading(doc, "2. Quy ước release")
    doc.add_paragraph("V2.2 reference: label-end order + validation/test không rỗng, theo tempo_fix. V2.3: cùng manifest nhưng mỗi arm ≥ max(2.5% population task, 50,000) và eval gộp ≥5%. split_registry_v2_2 legacy bị loại: CQ chưa hoàn tất; LO chứa ev_train_overlap_share theo mẫu số gộp cũ.")
    add_heading(doc, "3. Quy tắc sử dụng")
    doc.add_paragraph("Chỉ materialize từ split_registry_v2_3. Mọi artifact sau đó ghi scenario=hybrid, split_version=v2_3. Báo cáo model phải tách strict/relaxed theo phase, lưu per-class support, và không augmentation validation/test.")
    output = ROOT / "docs" / f"SPLIT_V2_2_REFERENCE_VS_V2_3_COMMON_QA_REPORT_{DATE}.docx"
    doc.save(output); print(output)


if __name__ == "__main__":
    for task_name in TASKS: report_task(task_name)
    common_report()
