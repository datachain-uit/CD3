from docx import Document


document = Document("docs/280926/Mo_ta_meta-dataset_v1.6.docx")
lines = [paragraph.text for paragraph in document.paragraphs]
lines.extend(" / ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows)
for term in ("L0", "L1", "L2", "L3", "attempt_id", "run_id", "manifest", "environment", "probability", "confusion", "checkpoint_selection", "sanity", "resource", "meta_release"):
    print(f"\n### {term}")
    print("\n".join(line for line in lines if term.lower() in line.lower())[:8000])
