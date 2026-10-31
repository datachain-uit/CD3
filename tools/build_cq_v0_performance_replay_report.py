"""Build the CQ PatchTST V0 performance and replay report from audited facts."""
from pathlib import Path
import json
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / 'CQ/resuilt/cq_patchtst_v0_current_seed42_r21_report'
OUT = ROOT / 'CQ/docs/BAO_CAO_CQ_PATCHTST_V0_HIEU_SUAT_REPLAY_2026-10-31.docx'

def table(doc, headers, rows):
    t=doc.add_table(rows=1, cols=len(headers)); t.style='Table Grid'
    for c,v in zip(t.rows[0].cells,headers): c.text=str(v)
    for row in rows:
        for c,v in zip(t.add_row().cells,row): c.text=str(v)

def f(x): return f'{float(x):.4f}'

def main():
    runs=pd.read_csv(DATA/'run_contracts.csv')
    hist=pd.read_csv(DATA/'training_history.csv')
    metrics=pd.read_csv(DATA/'metrics_and_sanity.csv')
    classes=pd.read_csv(DATA/'per_class_recall_ci.csv')
    replay=json.loads((DATA/'current_epoch1_replay_receipt.json').read_text(encoding='utf-8'))
    assert len(runs)==3 and len(metrics)==24 and replay['pass']
    d=Document(); sec=d.sections[0]; sec.left_margin=sec.right_margin=Inches(.7)
    d.styles['Normal'].font.name='Calibri'; d.styles['Normal'].font.size=Pt(11)
    d.add_heading('ĐÁNH GIÁ HIỆU SUẤT VÀ REPLAY',0)
    d.add_paragraph('CQ · PatchTST V0 · seed 42 · W1–W3 · 31/10/2026')
    d.add_heading('1. Kết luận',1)
    d.add_paragraph('Ba run đã hoàn thành và có audit validation/test theo P1–P4. V0 học được tín hiệu nhưng kết quả thay đổi giữa các window; W2 có chênh lệch validation–test lớn ở phase sau và recall lớp average còn thấp. Kết quả chỉ thuộc seed 42; chưa dùng để kết luận độ ổn định giữa các seed.')
    d.add_paragraph('Replay W1/V0 một epoch khớp tuyệt đối ba metric được kiểm tra với run hiện hành. Bằng chứng chỉ bao phủ epoch đầu; chưa so toàn bộ prediction hoặc chứng minh tương đương với pilot cũ.')
    d.add_heading('2. Cấu hình và lịch sử huấn luyện',1)
    d.add_paragraph('NVIDIA L4, BF16; batch hiệu dụng 2.048 gồm 4 micro-batch 512; learning rate 0,001; tối đa 50 epoch, early stopping patience 5. Checkpoint được chọn theo Macro-F1 validation trung bình P1–P4.')
    table(d,['Window','Train','Validation','Test/phase','Epoch chọn / chạy','Validation Macro-F1','Train giờ'],[[r.window,f'{int(r.train_rows):,}',f'{int(r.validation_rows):,}',f'{int(r.test_rows):,}',f'{int(r.best_epoch)} / {int(r.epochs_run)}',f(r.best_validation_macro_f1),f'{r.train_seconds/3600:.2f}'] for r in runs.itertuples(index=False)])
    d.add_paragraph('W1 cải thiện đến epoch 20; W2 dao động và chọn epoch 4; W3 chọn epoch 1, giảm mạnh ở epoch 3 rồi phục hồi một phần. Một run mỗi window không đủ để quy nguyên nhân dao động cho optimizer hoặc khẳng định model kém ổn định trên mọi seed.')
    d.add_heading('3. Hiệu suất validation và test',1)
    d.add_paragraph('AccTEMPO M3 là composite chính; Macro-F1 và balanced accuracy giúp đọc hiệu suất theo lớp. Validation dùng chọn checkpoint; TEST dùng báo cáo, không dùng chọn hyperparameter.')
    table(d,['W','Tập','Phase','N','AccTEMPO','Macro-F1','Balanced acc.'],[[r.window_id,r.eval_split,r.phase_id,f'{int(r.n_eval_samples):,}',f(r.acctempo_m3),f(r.macro_f1),f(r.balanced_acc)] for r in metrics.itertuples(index=False)])
    d.add_paragraph('W2/P4: Macro-F1 validation 0,8334 và test 0,6655. TEST AccTEMPO giảm từ 0,5811 ở P1 xuống 0,4954 ở P4 trong khi Macro-F1 tăng; cần đọc các thành phần sanity cùng hiệu suất, không diễn giải hai xu hướng là một.')
    d.add_heading('4. Recall theo lớp trên TEST',1)
    d.add_paragraph('Khoảng tin cậy là 95% theo audit stratified_binomial_equivalent_v1 với 2.000 lần. Bảng giữ từng window và phase để tránh gộp các cohort khác nhau.')
    table(d,['W','P','Lớp','Support','Recall','CI thấp','CI cao'],[[r.window_id,r.phase_id,r.class_code,f'{int(r.n_support):,}',f(r.recall),f(r.ci_lower),f(r.ci_upper)] for r in classes[classes.eval_split=='TEST'].itertuples(index=False)])
    d.add_heading('5. Replay W1/V0',1)
    d.add_paragraph('Replay epoch 1 cùng runner, data release, input manifest, layout hash, seed, batch, GPU, PyTorch và CUDA theo receipt. Bảng so sánh các giá trị được ghi; không bao gồm đối chiếu prediction toàn run.')
    table(d,['Metric epoch 1','Run hiện hành','Replay','Khớp tuyệt đối'],[[k,repr(v['current']),repr(v['replay']),'Có' if v['equal_bitwise_float64'] else 'Không'] for k,v in replay['metrics'].items()])
    d.add_paragraph('Kết luận: PASS cho ba metric epoch đầu và các trường contract trong receipt. Không suy rộng sang W2/W3, toàn bộ epoch, prediction cuối cùng hay runner/layout của pilot cũ.')
    d.add_heading('6. Nguồn và phạm vi',1)
    d.add_paragraph('Nguồn: CQ/resuilt/cq_patchtst_v0_current_seed42_r21_report gồm run_contracts.csv, training_history.csv, metrics_and_sanity.csv, per_class_recall_ci.csv và current_epoch1_replay_receipt.json. Số test theo phase là số dòng từng tập, không cộng các phase thành enrollment độc lập.')
    OUT.parent.mkdir(parents=True,exist_ok=True); d.save(OUT); print(OUT)

if __name__=='__main__': main()
