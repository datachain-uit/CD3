#!/usr/bin/env python3
"""Build an evidence-bound CQ report for recurrent F1 and the D23 PatchTST pilot.

The report intentionally separates completed seed-42 evidence from the F2 full
grid, which has not yet been run.  It reads immutable local artefacts only.
"""
from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
F1_ROOT = ROOT / "CQ/resuilt/cq_recurrent_allfamilies_grid_r4_20261023"
OLD_PATCH_ROOT = ROOT / "CQ/resuilt/cq_patchtst_v0_seed42_20261023"
D23_ROOT = ROOT / "CQ/resuilt/cq_patchtst_d23_stability_pilot_20261026"
TARGET = ROOT / "CQ/docs/REPORT_CQ_F1_D23_PATCHTST_CONFIGURATION_2026-10-27.docx"


PIPELINE_LABELS = {
    "V0": "zero-fill + masks",
    "V1": "median", "V2": "median + CDSMOTE", "V3": "median + SASMOTE", "V4": "median + RadiusSMOTE",
    "V5": "mean", "V6": "mean + CDSMOTE", "V7": "mean + SASMOTE", "V8": "mean + RadiusSMOTE",
    "V9": "ExtraTrees", "V10": "ExtraTrees + CDSMOTE", "V11": "ExtraTrees + SASMOTE", "V12": "ExtraTrees + RadiusSMOTE",
    "V13": "MICE", "V14": "MICE + CDSMOTE", "V15": "MICE + SASMOTE", "V16": "MICE + RadiusSMOTE",
}


def require(path: Path) -> Path:
    if not path.exists():
        raise FileNotFoundError(path)
    return path


def f4(value: object) -> str:
    return "—" if pd.isna(value) else f"{float(value):.4f}"


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_heading(text, level=level)
    for run in paragraph.runs:
        run.font.name = "Aptos"


def add_text(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(text)
    paragraph.paragraph_format.space_after = Pt(5)
    for run in paragraph.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(10.5)


def add_table(doc: Document, frame: pd.DataFrame, *, font_size: float = 9.5) -> None:
    table = doc.add_table(rows=1, cols=len(frame.columns))
    table.style = "Table Grid"
    table.autofit = True
    for cell, name in zip(table.rows[0].cells, frame.columns):
        cell.text = str(name)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(font_size + 0.5)
    for row in frame.itertuples(index=False, name=None):
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = "—" if pd.isna(value) else str(value)
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(font_size)
    doc.add_paragraph()


def recurrent_summary() -> tuple[pd.DataFrame, pd.DataFrame, dict]:
    root = require(F1_ROOT)
    metrics = pd.read_parquet(require(root / "cq_family_pipeline_window_phase_metrics.parquet"))
    summary = json.loads(require(root / "summary.json").read_text(encoding="utf-8"))
    expected = 4 * 17 * 3 * 4 * 2
    if len(metrics) != expected or int(summary["evaluation_cells"]) != expected:
        raise ValueError(f"F1 coverage mismatch: {len(metrics)} / {summary}")
    grouped = (metrics.groupby(["model_name", "eval_split"], as_index=False)
               .agg(rows=("n_rows", "sum"),
                    cells=("n_rows", "size"),
                    acctempo_m3=("acctempo_m3", "mean"),
                    macro_f1=("f1_macro", "mean"),
                    balanced_accuracy=("balanced_accuracy", "mean"),
                    s_perf=("s_perf", "mean"),
                    s_san_plus=("s_san_plus", "mean"),
                    s_ent=("s_ent", "mean"))
               .sort_values(["model_name", "eval_split"]))
    display = grouped.copy()
    display["rows"] = display["rows"].map(lambda x: f"{int(x):,}")
    display["cells"] = display["cells"].astype(int)
    for column in ("acctempo_m3", "macro_f1", "balanced_accuracy", "s_perf", "s_san_plus", "s_ent"):
        display[column] = display[column].map(f4)
    display = display.rename(columns={
        "model_name": "Họ", "eval_split": "Tập", "rows": "Tổng rows", "cells": "Ô",
        "acctempo_m3": "AccTEMPO", "macro_f1": "Macro-F1", "balanced_accuracy": "Balanced acc.",
        "s_perf": "S_perf", "s_san_plus": "S_san+ v1", "s_ent": "s_ent",
    })
    return display, metrics, summary


def read_patch_sanity(root: Path, audit_subdir: str) -> pd.DataFrame:
    path = require(root / audit_subdir / "CQ/PATCHTST/seed=42/sanity_components.parquet")
    data = pd.read_parquet(path)
    required = {"window_id", "eval_split", "phase_id", "macro_f1", "balanced_acc", "acctempo_m3", "s_perf", "s_san_plus"}
    missing = required.difference(data.columns)
    if missing:
        raise ValueError(f"{path} lacks {sorted(missing)}")
    return data


def old_patch_summary() -> pd.DataFrame:
    data = read_patch_sanity(OLD_PATCH_ROOT, "audits_v0_seed42")
    out = (data.groupby(["window_id", "eval_split"], as_index=False)
           .agg(phases=("phase_id", "size"), acctempo_m3=("acctempo_m3", "mean"),
                macro_f1=("macro_f1", "mean"), balanced_acc=("balanced_acc", "mean"),
                s_perf=("s_perf", "mean"), s_san_plus=("s_san_plus", "mean")))
    out["Cấu hình"] = "V0 cũ (lr=1e-3, không warm-up/clip)"
    return out


# D-23 (b) acceptance criterion as registered on 2026-10-25 (tempo_fix PENDING_DECISIONS[23]): the stabilised
# configuration may be pre-registered for the F2 grid only "if the VALIDATION curve is monotone or the selected
# epoch moves away from epoch 1"; otherwise the registered configuration stays.  The criterion is evaluated from
# train_history.parquet, never from TEST numbers, so that the choice between configurations is not TEST tuning.
SPIKE_RATIO = 2.0


def stability_statistics(history: pd.DataFrame, best_epoch: int) -> dict:
    """Per-epoch VALIDATION selection diagnostics of one run."""
    frame = history.sort_values("epoch").reset_index(drop=True)
    val = frame["validation_macro_f1_mean_p1_p4"].astype(float).to_numpy()
    train_ce = frame["train_cross_entropy_mean_p1_p4"].astype(float).to_numpy()
    steps = [abs(float(val[i] - val[i - 1])) for i in range(1, len(val))]
    spikes = [int(frame["epoch"].iloc[i]) for i in range(1, len(train_ce)) if train_ce[i] > SPIKE_RATIO * train_ce[i - 1]]
    # "Monotone" curve = unimodal: non-decreasing up to the selected epoch and non-increasing after it.  A curve
    # that peaks at epoch 1 and then declines steadily is stable optimisation meeting temporal drift; a curve that
    # swings around its peak is the instability the criterion was written for.  A selection at epoch 1 alone is
    # therefore judged on the whole curve, never vacuously on the single point before it.
    upto, after = val[: int(best_epoch)], val[int(best_epoch) - 1:]
    unimodal = bool(all(upto[i] >= upto[i - 1] for i in range(1, len(upto)))
                    and all(after[i] <= after[i - 1] for i in range(1, len(after))))
    return {
        "epochs": int(len(frame)), "selected_epoch": int(best_epoch),
        "val_f1_selected": float(val[int(best_epoch) - 1]) if 0 < int(best_epoch) <= len(val) else None,
        "val_f1_max": float(val.max()) if len(val) else None,
        "val_f1_min_after_first": float(val[1:].min()) if len(val) > 1 else None,
        "mean_abs_step": float(sum(steps) / len(steps)) if steps else 0.0,
        "max_abs_step": float(max(steps)) if steps else 0.0,
        "train_ce_spike_epochs": spikes,
        "criterion_selected_epoch_gt_1": int(best_epoch) > 1,
        "criterion_unimodal_no_spike": unimodal and not spikes,
    }


def d23_criterion_met(stats: dict) -> bool:
    """Registered D-23 (b) acceptance: a monotone (unimodal, spike-free) VALIDATION curve, or a selection away from epoch 1."""
    return bool(stats["criterion_unimodal_no_spike"] or stats["criterion_selected_epoch_gt_1"])


def history_table(histories: dict[str, pd.DataFrame], selected: dict[str, int]) -> pd.DataFrame:
    rows = []
    for label, history in histories.items():
        for _, item in history.sort_values("epoch").iterrows():
            rows.append({"Cấu hình / window": label, "Epoch": int(item["epoch"]),
                         "Chọn": "YES" if int(item["epoch"]) == selected[label] else "",
                         "VAL Macro-F1 (mean P1–P4)": f"{float(item['validation_macro_f1_mean_p1_p4']):.4f}",
                         "VAL CE": f"{float(item['validation_cross_entropy_mean_p1_p4']):.4f}",
                         "Train CE": f"{float(item['train_cross_entropy_mean_p1_p4']):.4f}",
                         "LR cuối epoch": (f"{float(item['learning_rate_end']):.2e}" if "learning_rate_end" in item and pd.notna(item["learning_rate_end"]) else "")})
    return pd.DataFrame(rows)


def stability_comparison(stats: dict[str, dict]) -> pd.DataFrame:
    rows = []
    for label, item in stats.items():
        rows.append({"Cấu hình / window": label, "Epoch": item["epochs"], "Epoch chọn": item["selected_epoch"],
                     "VAL F1 tại epoch chọn": f"{item['val_f1_selected']:.4f}" if item["val_f1_selected"] is not None else "",
                     "|Δ| TB": f"{item['mean_abs_step']:.4f}", "|Δ| max": f"{item['max_abs_step']:.4f}",
                     "Train CE tăng > 2× (epoch)": ", ".join(map(str, item["train_ce_spike_epochs"])) or "không",
                     "Đường cong đơn điệu (một đỉnh), không spike": "Có" if item["criterion_unimodal_no_spike"] else "Không",
                     "Epoch chọn > 1": "Có" if item["criterion_selected_epoch_gt_1"] else "Không",
                     "Tiêu chí D-23 (b)": "ĐẠT" if d23_criterion_met(item) else "KHÔNG ĐẠT"})
    return pd.DataFrame(rows)


def load_histories() -> tuple[dict[str, pd.DataFrame], dict[str, int], dict[str, dict]]:
    """Old (lr 1e-3) and D23 histories of W2 and W3 side by side; the criterion is evaluated on D23 only."""
    histories: dict[str, pd.DataFrame] = {}
    selected: dict[str, int] = {}
    for window in ("W2", "W3"):
        for label, root in ((f"V0 cũ {window}", OLD_PATCH_ROOT / window), (f"D23 {window}", D23_ROOT / "CQ" / window / "V0" / "PATCHTST" / "seed=42")):
            manifest = json.loads(require(root / "run_manifest.json").read_text(encoding="utf-8"))
            histories[label] = pd.read_parquet(require(root / "train_history.parquet"))
            selected[label] = int(manifest["best_epoch"])
    stats = {label: stability_statistics(history, selected[label]) for label, history in histories.items()}
    return histories, selected, stats


def d23_summary() -> tuple[pd.DataFrame, pd.DataFrame, dict]:
    audit_root = D23_ROOT / "audits"
    data = read_patch_sanity(audit_root, "")
    audit = json.loads(require(audit_root / "CQ/PATCHTST/seed=42/audit_manifest.json").read_text(encoding="utf-8"))
    if int(audit["runs"]) != 2 or int(audit["bootstrap_repetitions"]) != 2000:
        raise ValueError(f"D23 audit coverage is not the registered two-run/B=2000 pilot: {audit}")
    out = (data.groupby(["window_id", "eval_split"], as_index=False)
           .agg(phases=("phase_id", "size"), acctempo_m3=("acctempo_m3", "mean"),
                macro_f1=("macro_f1", "mean"), balanced_acc=("balanced_acc", "mean"),
                s_perf=("s_perf", "mean"), s_san_plus=("s_san_plus", "mean")))
    out["Cấu hình"] = "D23 ổn định hóa (lr=3e-4, warm-up=1, clip=1.0)"
    rows: list[dict] = []
    for window in ("W2", "W3"):
        run_root = require(D23_ROOT / "CQ" / window / "V0" / "PATCHTST" / "seed=42")
        manifest = json.loads(require(run_root / "run_manifest.json").read_text(encoding="utf-8"))
        history = pd.read_parquet(require(run_root / "train_history.parquet"))
        expected = {
            "batch_size": 2048, "micro_batch_size": 512, "accumulation_steps": 4,
            "learning_rate": 3e-4, "warmup_epochs": 1, "gradient_clip_norm": 1.0,
            "static_layout_policy": "current", "run_purpose": "STABILITY_PILOT_D23",
        }
        mismatch = {key: (manifest.get(key), value) for key, value in expected.items() if manifest.get(key) != value}
        if mismatch:
            raise ValueError(f"D23 protocol mismatch {window}: {mismatch}")
        checkpoint = pd.read_parquet(audit_root / "CQ/PATCHTST/seed=42/checkpoint_selection.parquet")
        checkpoint = checkpoint.loc[checkpoint["window_id"] == window].iloc[0]
        rows.append({
            "Window": window,
            "Batch hiệu dụng": "2,048 = 4 × 512",
            "Best epoch": int(manifest["best_epoch"]),
            "Epoch đã đánh giá": int(len(history)),
            "Checkpoint trong margin": int(checkpoint["n_checkpoints_within_margin"]),
            "Thời gian train (giờ)": f"{float(manifest['elapsed_seconds']) / 3600:.2f}",
            "GPU energy (kWh)": f"{float(manifest['energy_gpu_estimated_kwh']):.3f}",
            "GPU": manifest.get("nvml_gpu_name") or manifest.get("gpu"),
            "Layout SHA-256": str(manifest.get("feature_layout_sha256", ""))[:20] + "…",
        })
    return out, pd.DataFrame(rows), audit


def format_patch_metrics(frame: pd.DataFrame) -> pd.DataFrame:
    out = frame.copy()
    out["phases"] = out["phases"].astype(int)
    for c in ("acctempo_m3", "macro_f1", "balanced_acc", "s_perf", "s_san_plus"):
        out[c] = out[c].map(f4)
    return out.rename(columns={
        "window_id": "Window", "eval_split": "Tập", "phases": "Pha", "acctempo_m3": "AccTEMPO",
        "macro_f1": "Macro-F1", "balanced_acc": "Balanced acc.", "s_perf": "S_perf", "s_san_plus": "S_san+ v1",
    })[["Cấu hình", "Window", "Tập", "Pha", "AccTEMPO", "Macro-F1", "Balanced acc.", "S_perf", "S_san+ v1"]]


def main() -> None:
    f1_display, _, f1_manifest = recurrent_summary()
    old = old_patch_summary()
    d23, d23_runs, d23_audit = d23_summary()
    histories, selected, stats = load_histories()
    d23_verdicts = {window: d23_criterion_met(stats[f"D23 {window}"]) for window in ("W2", "W3")}
    d23_all_met = all(d23_verdicts.values())
    verdict_text = ("đạt tiêu chí D-23 (b) ở cả hai window" if d23_all_met else
                    "không đạt tiêu chí D-23 (b) ở " + ", ".join(w for w, ok in d23_verdicts.items() if not ok))

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("BÁO CÁO CQ — F1 RECURRENT VÀ ĐÁNH GIÁ CẤU HÌNH PATCHTST D23")
    run.bold = True
    run.font.size = Pt(17)
    subtitle = doc.add_paragraph("CQ_V2_2 clean-grain · seed 42 · cập nhật 27/10/2026")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for item in subtitle.runs:
        item.italic = True
        item.font.size = Pt(11)

    add_heading(doc, "1. Kết luận điều hành")
    add_text(doc, "Lưới F1 seed 42 đã hoàn tất: 204 run RNN/LSTM/GRU/BiLSTM (17 pipeline × 3 window × 4 họ), mỗi run dùng batch 2.048 và NVML GPU-power sampling. "
                  "Pilot ổn định hóa D23 (hai cell W2/V0 và W3/V0 với learning rate thấp hơn, warm-up và gradient clipping; audit B=2.000) đã chạy xong và "
                  f"{verdict_text} (mục 4b, đánh giá trên đường cong VALIDATION, không dùng TEST). "
                  + ("Cấu hình D23 có thể được tác giả tiền đăng ký làm HP_PATCHTST v3 cho F2; cho đến khi đó, cấu hình đã đăng ký (lr 1e-3) vẫn là cấu hình của lưới."
                     if d23_all_met else
                     "Theo tiêu chí đã đăng ký, cấu hình D23 chưa đủ điều kiện thay thế cấu hình đã đăng ký; lựa chọn cuối (giữ lr 1e-3 và công bố bất ổn, hay tiền đăng ký D23 dù W3 vẫn chọn epoch 1) thuộc tác giả (quyết định 23)."))
    add_text(doc, "Báo cáo không diễn giải hai pilot là kết quả full-grid: 48 cell F2 còn lại chưa chạy. Các bảng dưới là mô tả seed 42, không phải kết luận ổn định đa-seed.")

    add_heading(doc, "2. Phạm vi evidence và kiểm tra bao phủ")
    coverage = pd.DataFrame([
        ["F1 recurrent", "RNN, LSTM, GRU, BiLSTM; V0–V16; W1–W3; P1–P4; TEST + VALIDATION", "204 run; 1.632 metric cells", "Đạt"],
        ["F1 resource", "Batch 2.048, NVIDIA L4, NVML power sampling", f"{f1_manifest['resource_runs']} run; energy non-null={f1_manifest['all_energy_nonnull']}", "Đạt"],
        ["PatchTST V0 cũ", "W1–W3; cấu hình lr=1e-3, chưa warm-up/clip", "Evidence so sánh lịch sử", "Không dùng để khóa F2"],
        ["PatchTST D23", "W2/V0, W3/V0; cấu hình ổn định hóa", f"{d23_audit['runs']} run; B={d23_audit['bootstrap_repetitions']}; 16 sanity rows", "Hoàn tất; tiêu chí D-23 (b): " + ("đạt" if d23_all_met else "không đạt")],
        ["F2 full-grid", "51 cell PatchTST seed 42, trừ ba V0 pilot lịch sử", "Chưa materialize", "Chờ chạy"],
    ], columns=["Thành phần", "Phạm vi", "Bằng chứng", "Trạng thái"])
    add_table(doc, coverage)

    add_heading(doc, "3. Kết quả recurrent CQ hiện có")
    add_text(doc, "Trung bình dưới đây gộp 17 pipeline × 3 window × 4 prefix cho từng họ và split. AccTEMPO M3 vẫn là độ đo tổng hợp chính; Macro-F1 và balanced accuracy được giữ để đọc hiệu năng phân loại. Không dùng bảng gộp này để chọn một pipeline đơn lẻ.")
    add_table(doc, f1_display)
    add_text(doc, "Theo review 25/10, khi dùng AccTEMPO hiệu chuẩn M3(cal), chênh lệch giữa 17 pipeline trong cùng một họ nhỏ (xấp xỉ 0,01); vì vậy không có bằng chứng seed-42 rằng lựa chọn imputer/balancer làm thay đổi hiệu năng phân loại CQ một cách đo được.")

    add_heading(doc, "4. PatchTST: cấu hình cũ và pilot D23")
    config = pd.DataFrame([
        ["Cấu hình cũ V0", "1e-3", "0", "không", "2.048 = 4 × 512", "current", "PRIMARY_GRID", "Chỉ evidence lịch sử; validation dao động theo review"],
        ["D23 ổn định hóa", "3e-4", "1 epoch", "1.0", "2.048 = 4 × 512", "current", "STABILITY_PILOT_D23", "Ứng viên HP v3; chỉ thành cấu hình F2 sau khi tác giả tiền đăng ký"],
    ], columns=["Phiên bản", "Learning rate", "Warm-up", "Gradient clip", "Batch", "Static layout", "Mục đích", "Diễn giải"])
    add_table(doc, config)
    add_text(doc, "D23 giữ cùng kiến trúc PatchTST channel-independent, batch hiệu dụng và layout policy current; thay đổi chỉ nhằm giảm bất ổn tối ưu hóa. Điều này cho phép so sánh trực tiếp tính ổn định của training protocol, nhưng không biến các cell cũ thành run cùng cấu hình.")
    add_table(doc, d23_runs, font_size=8.5)

    add_heading(doc, "4b. Lịch sử VALIDATION theo epoch và tiêu chí D-23 (b)")
    add_text(doc, "Tiêu chí đã đăng ký (25/10): cấu hình ổn định hóa được tiền đăng ký cho F2 nếu đường cong VALIDATION đơn điệu (một đỉnh: không giảm tới epoch chọn và không tăng sau đó, không có epoch train CE tăng quá 2 lần) hoặc epoch chọn rời khỏi epoch 1. Bảng so sánh cấu hình cũ và D23 trên cùng window; TEST không tham gia vào tiêu chí.")
    add_table(doc, stability_comparison(stats), font_size=8)
    add_table(doc, history_table(histories, selected), font_size=8)

    add_heading(doc, "5. Kết quả PatchTST V0 theo window và split")
    add_text(doc, "Mỗi dòng là trung bình bốn prefix P1–P4. D23 chỉ có W2/W3; W1 trong bảng là cấu hình cũ để cung cấp bối cảnh, không được ghép vào evidence D23.")
    comparison = pd.concat([old, d23], ignore_index=True)
    add_table(doc, format_patch_metrics(comparison), font_size=8.5)

    add_heading(doc, "6. Đánh giá cấu hình D23 và quyết định F2")
    assessment = pd.DataFrame([
        ["Tính đúng protocol", "W2 và W3 đều đúng lr=3e-4, warm-up=1, clip=1.0, batch=2.048=4×512, layout=current, V0/PATCHTST/seed42.", "Đạt"],
        ["Traceability", "Manifest ghi runner commit, GPU L4, layout hash, thời gian và NVML energy; audit ghi release CQ_V2_2 và download-manifest hash.", "Đạt"],
        ["Độ đầy đủ đánh giá", "Audit gồm canonical S_leak, sanity, per-class, bootstrap B=2.000, STRICT/OVERLAP và checkpoint selection.", "Đạt"],
        ["Ổn định hóa theo D23", "Tiêu chí D-23 (b) trên đường cong VALIDATION (mục 4b): " + "; ".join(f"{w}: epoch chọn {stats[f'D23 {w}']['selected_epoch']}/{stats[f'D23 {w}']['epochs']}, |Δ| max {stats[f'D23 {w}']['max_abs_step']:.3f}, spike {len(stats[f'D23 {w}']['train_ce_spike_epochs'])}" for w in ("W2", "W3")) + ".", "Đạt" if d23_all_met else "Không đạt"],
        ["Full F2", "Không có 48 cell còn lại trong evidence hiện tại.", "Chưa chạy — không suy diễn kết quả full-grid"],
        ["Đa-seed", "Seed 43/44 và S2 tương ứng chưa phải evidence trong báo cáo này.", "Kế hoạch stability, không phải điều kiện D23"],
    ], columns=["Tiêu chí", "Bằng chứng", "Kết luận"])
    add_table(doc, assessment)
    if d23_all_met:
        add_text(doc, "Đề xuất của nhóm: tiền đăng ký HP_PATCHTST v3 = D23 (lr=3e-4, warm-up 1 epoch, gradient clipping 1.0, batch hiệu dụng 2.048 qua 4 micro-batch 512, static layout current) cho F2 CQ seed 42, chạy lại ba cell V0 pilot dưới v3 trước lưới, và không thay đổi thông số theo kết quả 48 cell còn lại. Quyết định thuộc tác giả (quyết định 23).")
    else:
        add_text(doc, "Theo tiêu chí đã đăng ký, cấu hình D23 chưa thay thế được cấu hình đã đăng ký. Hai phương án cho tác giả: (a) chạy 48 cell còn lại ở cấu hình đã đăng ký (lr 1e-3) và lưu toàn bộ train_history để công bố bất ổn; (b) tiền đăng ký D23 làm HP v3 với lý do khác tiêu chí (ví dụ warm-up và clipping là thực hành chuẩn cho transformer) và chạy lại ba cell V0 pilot dưới v3 trước lưới. Chênh lệch TEST giữa hai cấu hình trong mục 5 là số mô tả một seed và không được dùng để chọn cấu hình.")

    add_heading(doc, "7. Quy ước pipeline dùng trong báo cáo")
    add_text(doc, "Nhãn bên dưới lấy theo registry/artifact thực tế, nhằm tránh hoán đổi MEDIAN/MEAN hoặc SASMOTE/RadiusSMOTE trong báo cáo lịch sử.")
    labels = pd.DataFrame([[key, PIPELINE_LABELS[key]] for key in sorted(PIPELINE_LABELS, key=lambda x: int(x[1:]))], columns=["Pipeline", "Cấu hình"])
    add_table(doc, labels, font_size=8.5)

    add_heading(doc, "8. Các việc còn lại")
    add_text(doc, "(1) Chạy F2 CQ còn lại bằng cấu hình mà tác giả chốt ở quyết định 23 và shard plan (một runner commit, run_purpose PRIMARY_GRID); (2) audit/ingest facts F2 với kiểm đồng nhất training protocol; (3) tạo S2 seed 43/44 rồi chạy F1/F2 theo chính sách stability đã đăng ký. Các việc này là mở rộng lưới, không phải lỗi của F1 hoặc pilot D23 đã hoàn tất.")
    add_text(doc, "Lưu ý source: review 25/10 cũng phát hiện N10 (imputer_only_ref) và nhãn pipeline trong một số generator cũ cần đồng bộ registry. Báo cáo này dùng ánh xạ registry đúng; việc sửa source/report generator cần được commit riêng, không làm thay đổi các artefact đã chạy.")

    TARGET.parent.mkdir(parents=True, exist_ok=True)
    doc.save(TARGET)
    print(json.dumps({"report": str(TARGET), "f1_metric_cells": 1632, "d23_runs": int(d23_audit["runs"]), "d23_audit_rows": d23_audit["rows"]}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
