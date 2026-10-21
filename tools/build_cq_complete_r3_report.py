"""Build the complete CQ seed-42 four-family closure report from immutable facts."""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
FACTS = ROOT / "CQ/resuilt/cq_recurrent_allfamilies_grid_r3_20261021"
AUDIT_ROOT = ROOT / "CQ/resuilt/cq_recurrent_batch2048_slice_r3_partition_20261021"
SOURCE = ROOT / "CQ/docs/REPORT_CQ_TONG_HOP_DAY_DU_SEED42_GRID_R3_2026-10-21.docx"
TARGET = ROOT / "CQ/docs/REPORT_CQ_TONG_HOP_HOAN_CHINH_SEED42_2026-10-21.docx"


def table(document: Document, frame: pd.DataFrame, columns: list[tuple[str, str]], *, size: float = 7.5) -> None:
    output = document.add_table(rows=1, cols=len(columns))
    output.style = "Table Grid"
    for cell, (_, label) in zip(output.rows[0].cells, columns):
        cell.text = label
    for _, record in frame.iterrows():
        cells = output.add_row().cells
        for cell, (name, _) in zip(cells, columns):
            value = record[name]
            if isinstance(value, float):
                cell.text = "" if pd.isna(value) else f"{value:.4f}"
            else:
                cell.text = str(value)
    for row in output.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(size)


def summary_block(document: Document) -> None:
    summary = json.loads((FACTS / "summary.json").read_text(encoding="utf-8"))
    document.add_page_break()
    document.add_paragraph("BẢN TỔNG HỢP HOÀN CHỈNH — CQ CLEAN-GRAIN, SEED 42")
    document.add_paragraph(
        "Bản này hợp nhất báo cáo clean-grain trước đó với evidence r3 mới. Phạm vi đánh giá gồm bốn họ "
        "RNN, LSTM, GRU và BiLSTM; V0–V16; W1–W3; P1–P4; TEST và VALIDATION. "
        "Tổng số fact hiệu năng tại hạt family × pipeline × window × phase × split là "
        f"{summary['evaluation_cells']:,}; số run có resource manifest là {summary['resource_runs']:,}."
    )
    document.add_paragraph(
        "AccTEMPO là độ đo hiệu năng chính. Macro-F1, balanced accuracy, s_ent và recall lớp được giữ song song "
        "để diễn giải tác động của mất cân bằng lớp; accuracy không được dùng để chọn mô hình."
    )


def audit_block(document: Document) -> None:
    document.add_paragraph("1. Kết luận kiểm tra STRICT/OVERLAP")
    rows = []
    for family in ("LSTM", "GRU", "BILSTM"):
        value = json.loads((AUDIT_ROOT / family / "summary.json").read_text(encoding="utf-8"))
        rows.append({"family": family, "runs": value["runs"], "files": value["files"], "rows": value["rows"],
                     "failed": value["failed_cells"], "status": "Đạt" if value["all_passed"] else "Không đạt"})
    document.add_paragraph(
        "Audit r3 đọc prediction thô của ba họ lịch sử. Mỗi cell được kiểm ID hash không null/không trùng, "
        "kiểm cờ STRICT theo phase là nhị phân và kiểm ma trận confusion STRICT + OVERLAP = ALL. "
        "RNN historical-layout đã có audit tương đương trong bundle riêng; do đó coverage toàn CQ là 204/204 run."
    )
    table(document, pd.DataFrame(rows), [("family", "Họ"), ("runs", "Run"), ("files", "File"),
                                          ("rows", "Dòng prediction"), ("failed", "Cell lỗi"), ("status", "Kết quả")], size=9)


def detailed_metrics(document: Document) -> None:
    document.add_page_break()
    document.add_paragraph("2. Bảng đầy đủ hiệu năng theo family × pipeline × window × phase")
    document.add_paragraph(
        "Các bảng dưới là trung bình của các metric theo đúng cell, không gộp window hoặc phase. "
        "Mỗi bảng có 204 dòng = 17 pipeline × 3 window × 4 phase. Cột STRICT W2 chỉ có giá trị ở W2; "
        "tỷ lệ này được tính từ số prediction STRICT/ALL, riêng theo split và phase."
    )
    cells = pd.read_parquet(FACTS / "cq_family_pipeline_window_phase_metrics.parquet")
    order = {f"V{i}": i for i in range(17)}
    columns = [
        ("pipeline_id", "Pipeline"), ("window_id", "W"), ("phase_id", "P"), ("acctempo_m3", "AccTEMPO"),
        ("f1_macro", "Macro-F1"), ("balanced_accuracy", "Bal. acc."), ("s_ent", "s_ent"),
        ("c0_recall", "Recall c0"), ("c1_recall", "Recall c1"), ("c2_recall", "Recall c2"),
        ("w2_strict_share", "STRICT W2"),
    ]
    for split in ("TEST", "VALIDATION"):
        for family in ("RNN", "LSTM", "GRU", "BILSTM"):
            document.add_paragraph(f"{split} — {family}")
            frame = cells[(cells.eval_split == split) & (cells.model_name == family)].copy()
            frame["_order"] = frame.pipeline_id.map(order)
            frame = frame.sort_values(["_order", "window_id", "phase_id"]).drop(columns="_order")
            table(document, frame, columns)


def resource_block(document: Document) -> None:
    document.add_page_break()
    document.add_paragraph("3. Bảng thời gian train và năng lượng GPU")
    document.add_paragraph(
        "train_seconds là tổng thời gian train theo epoch từ train_history, dùng để so sánh công bằng giữa run. "
        "energy_gpu_estimated_kwh là ước lượng GPU-only qua NVML; không phải năng lượng toàn hệ thống và không dùng để suy CO₂. "
        "NVML có dữ liệu hợp lệ cho 204/204 run.")
    frame = pd.read_parquet(FACTS / "cq_train_time_energy_family_pipeline.parquet")
    order = {f"V{i}": i for i in range(17)}
    frame["_order"] = frame.pipeline_id.map(order)
    frame = frame.sort_values(["model_name", "_order"]).drop(columns="_order")
    table(document, frame, [
        ("model_name", "Họ"), ("pipeline_id", "Pipeline"), ("runs", "Số W"),
        ("train_seconds_mean", "Train TB (s)"), ("train_seconds_sum", "Train tổng (s)"),
        ("energy_gpu_estimated_kwh_mean", "GPU TB (kWh)"), ("energy_gpu_estimated_kwh_sum", "GPU tổng (kWh)"),
        ("nvml_samples_min", "NVML min"), ("nvml_missing_rate_max", "Thiếu max"),
    ])
    document.add_paragraph("4. Lineage L2")
    document.add_paragraph(
        "Không trộn hai runner revision: RNN historical-layout có receipt L2 riêng "
        "vastai_cq_rnn_historical_layout_seed42_20261020. Ba họ LSTM/GRU/BiLSTM có receipt "
        "vastai_cq_f1_batch2048_oldfamilies_r2_20261021: 153 contract, 1.224 metric rows, 153 resource rows và 66 partition. "
        "Hai receipt đều khóa CQ_V2_2, seed 42 và NVIDIA L4."
    )
    document.add_paragraph(
        "Máy đọc chi tiết có thể dùng các file CSV/Parquet cùng thư mục facts: "
        "cq_family_pipeline_window_phase_metrics, cq_family_window_phase_mean, "
        "cq_train_time_energy_family_pipeline và cq_resource_per_run."
    )


def main() -> None:
    document = Document(SOURCE)
    summary_block(document)
    audit_block(document)
    detailed_metrics(document)
    resource_block(document)
    document.save(TARGET)
    print(TARGET)


if __name__ == "__main__":
    main()
